"""IC deck builder on the FEP template (docs/IC_DECK_STANDARDS.md).

Every number is bound to a cell of the recalculated workbook with `Deck.bind(sheet, cell, fmt)` (or a workbook
name); `Deck.save()` writes `deck_bindings.csv` (slide, shape, text, cell, value) next to the deck.

    d = Deck("out/model.xlsm", source="Demo model", sections=["Summary", "Returns"])
    d.cover("Project Demo", "Bid strategy", "12 October 2026")
    s = d.content("Summary", "Bid 62.4 EUR/MWh", f"IRR {d.bind('Summary', 'E11', fmt_pct)}")
    s.table(["Output", "Unit", "Base"], [["IRR", "%", d.bind("Out_Project_IRR", fmt=fmt_pct)]])
    d.save("out/deck.pptx")
"""
from __future__ import annotations

import copy
import csv
import datetime as dt
from pathlib import Path

import openpyxl
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "templates" / "fep_ic_template.pptx"
REFERENCE = ROOT / "reference" / "Bora_IC_RtB_strategies.pptx"

TEXT, GREEN, DARK_GREEN, LIME, GREY = "3C3C3B", "00A651", "00563C", "8CC63F", "919191"
LABEL_FILL, KEY_FILL, WHITE = "F2F2F2", "E6F4EC", "FFFFFF"
SERIES = [DARK_GREEN, GREEN, LIME, "A6A6A6", "CFE8D8", "B7DEC6"]
FONT = "Aptos"
MINUS, DASH = "−", "–"


# ---- number formatters (IC_DECK_STANDARDS "Number formats") ----------------------------------
def _na(v):
    return v is None or isinstance(v, str)


def fmt_eur_m(v, dp=1):
    """EUR m, negatives in brackets: (5.3)."""
    if _na(v):
        return DASH
    s = f"{abs(v):,.{dp}f}"
    return f"({s})" if round(v, dp) < 0 else s


def fmt_delta(v, dp=1):
    """Signed delta: +3.3 / (3.3) / 0.0."""
    if _na(v):
        return DASH
    return "+" + f"{v:,.{dp}f}" if round(v, dp) > 0 else fmt_eur_m(v, dp)


def fmt_pct(v, dp=1):
    """Percent with a true minus sign: −7.5%."""
    if _na(v):
        return DASH
    s = f"{abs(v) * 100:,.{dp}f}%"
    return MINUS + s if round(v * 100, dp) < 0 else s


def fmt_mult(v, dp=2):
    return DASH if _na(v) else f"{v:,.{dp}f}x"


def fmt_num(v, dp=1):
    """Plain number with thousands separator (EUR/MWh, EUR k/MW with dp=0)."""
    if _na(v):
        return DASH
    s = f"{abs(v):,.{dp}f}"
    return MINUS + s if round(v, dp) < 0 else s


# ---- template ---------------------------------------------------------------------------------
def make_template(reference=REFERENCE, out=TEMPLATE) -> Path:
    """Copy the reference deck, drop every slide except the cover, and neutralise the cover text."""
    prs = Presentation(reference)
    ids = prs.slides._sldIdLst  # noqa: SLF001
    for sld in list(ids)[1:]:
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)
    _set_cover_text(prs.slides[0], "Project name", "Topic", "Date")
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.last_modified_by, cp.keywords, cp.comments = (
        "FEP IC template", "", "FEP", "FEP", "", "")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def _set_cover_text(slide, project, topic, date):
    box = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    paras = [p for p in box.text_frame.paragraphs if p.runs]
    for p, text in zip(paras, (project, topic, date)):
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)  # noqa: SLF001
        for br in p._p.findall(qn("a:br")):  # noqa: SLF001
            p._p.remove(br)  # noqa: SLF001


# ---- deck -------------------------------------------------------------------------------------
class Bound(str):
    """Formatted text bound to a workbook cell."""


class BoundNum(float):
    """Raw number bound to a workbook cell (chart values)."""


def _rgb(hexstr):
    return RGBColor.from_string(hexstr)


def _font(run_or_font, size, bold=False, color=TEXT):
    f = getattr(run_or_font, "font", run_or_font)
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = _rgb(color)


class Deck:
    def __init__(self, workbook, *, template=TEMPLATE, source="model", date: dt.date | None = None,
                 sections=()):
        self.workbook = Path(workbook)
        self.wb = openpyxl.load_workbook(self.workbook, data_only=True)
        self.prs = Presentation(template)
        self.layout = next(lo for lo in self.prs.slide_masters[0].slide_layouts if lo.name == "Base text")
        self.source, self.date = source, date or dt.date.today()
        self.sections = list(sections)
        self.bindings: list[tuple] = []          # slide, shape, text, cell, value
        self._bound: dict[str, set] = {}         # text -> {(cell, value)}
        self._has_cover = False

    # ---- bindings -------------------------------------------------------------------------
    def _resolve(self, sheet_or_name, cell):
        if cell is None:
            dn = self.wb.defined_names[sheet_or_name]
            sheet, cell = next(dn.destinations)
        else:
            sheet = sheet_or_name
        cell = cell.replace("$", "")
        return f"{sheet}!{cell}", self.wb[sheet][cell].value

    def bind(self, sheet_or_name, cell=None, fmt=None):
        """Value of `Sheet!cell` (or a workbook name) from the recalculated workbook.

        With `fmt` returns formatted text (Bound), without it a number for charts (BoundNum)."""
        ref, value = self._resolve(sheet_or_name, cell)
        if fmt is None:
            b = BoundNum(value if isinstance(value, (int, float)) else 0.0)
        else:
            b = Bound(fmt(value))
            self._bound.setdefault(str(b), set()).add((ref, value))
        b.ref, b.value = ref, value
        return b

    def stored(self, output_key, scenario: int, fmt=None):
        """Stored scenario result (1-based scenario) of output `Out_<output_key>`: on its Summary row, or on the
        Dashboard in the `<output_key>_Stored` row (e.g. `stored("HTM_IRR_2", 3)`)."""
        names = self.wb.defined_names
        if f"{output_key}_Stored" in names:                      # Dashboard: own stored row
            sheet, rng = next(names[f"{output_key}_Stored"].destinations)
            _, first, row = rng.split(":")[0].split("$")
        else:                                                    # Summary: stored columns on the output row
            sheet, cell = next(names[f"Out_{output_key}"].destinations)
            first = next(names["Stored_Results"].destinations)[1].split("$")[1]
            row = cell.split("$")[-1]
        col = openpyxl.utils.column_index_from_string(first) + scenario - 1
        return self.bind(sheet, f"{openpyxl.utils.get_column_letter(col)}{row}", fmt)

    def _record_text(self, slide_no, shape, text):
        for t, refs in self._bound.items():
            if t and t in text:
                for ref, value in sorted(refs, key=str):
                    self.bindings.append((slide_no, shape, t, ref, value))

    # ---- slides ---------------------------------------------------------------------------
    def cover(self, project, topic, date_text):
        _set_cover_text(self.prs.slides[0], project, topic, date_text)
        self._has_cover = True

    def content(self, section, title, subtitle=None) -> "Slide":
        return Slide(self, section, title, subtitle)

    def save(self, path) -> Path:
        path = Path(path)
        if not self._has_cover:
            ids = self.prs.slides._sldIdLst  # noqa: SLF001
            self.prs.part.drop_rel(ids[0].rId)
            ids.remove(ids[0])
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        used = {(b[3], b[2]) for b in self.bindings}
        orphans = [(ref, t) for t, refs in self._bound.items() for ref, _ in refs if (ref, t) not in used]
        if orphans:
            raise ValueError(f"bound values never placed on a slide: {orphans}")
        with open(path.parent / "deck_bindings.csv", "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["slide", "shape", "text", "cell", "value"])
            w.writerows(self.bindings)
        return path


class Slide:
    LEFT, WIDTH = 4.92, 15.0

    def __init__(self, deck: Deck, section, title, subtitle):
        self.d = deck
        self.s = deck.prs.slides.add_slide(deck.layout)
        for ph in list(self.s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)  # noqa: SLF001 - unused body placeholder
        t = self.s.shapes.title
        t.width = Inches(13.8)
        t.text = title
        self._record("Title", title)
        self._page_number()
        self._navigator(section)
        if subtitle:
            self.text("Subtitle", subtitle, self.LEFT, 1.20, self.WIDTH, 0.45, size=18, bold=True)

    @property
    def number(self):
        return self.d.prs.slides.index(self.s) + 1

    def _record(self, shape, text):
        self.d._record_text(self.number, shape, text)  # noqa: SLF001

    def _page_number(self):
        for sp in self.s.slide_layout.placeholders:
            if sp.placeholder_format.type is not None and sp.placeholder_format.type == 13:  # SLIDE_NUMBER
                self.s.shapes._spTree.append(copy.deepcopy(sp._element))  # noqa: SLF001

    def _navigator(self, active):
        secs = self.d.sections
        if not secs:
            return
        shape = self.s.shapes.add_table(len(secs), 1, 0, 0, Inches(4.12), Inches(0.85 * len(secs)))
        shape.name = "Navigator"
        tbl = shape.table
        tbl.first_row = tbl.horz_banding = False
        for i, name in enumerate(secs):
            cell = tbl.cell(i, 0)
            on = name == active
            cell.text = name
            cell.margin_left = Emu(180000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _font(cell.text_frame.paragraphs[0].runs[0], 18, on, WHITE if on else TEXT)
            _borders(cell, top=6350, bottom=6350)
            if on:
                cell.fill.gradient()
                cell.fill.gradient_angle = 0
                stops = cell.fill.gradient_stops
                stops[0].color.rgb, stops[1].color.rgb = _rgb(GREEN), _rgb(DARK_GREEN)
            else:
                cell.fill.background()

    # ---- components -----------------------------------------------------------------------
    def text(self, name, text, left, top, width, height, size=13, bold=False, color=TEXT):
        box = self.s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            _font(p.add_run(), size, bold, color)
            p.runs[0].text = line
        self._record(name, text)
        return box

    def table(self, header, rows, *, key_rows=(), sections=(), left=LEFT, top=1.75, width=WIDTH, height=None,
              col_widths=None, name="Table"):
        """FEP table: green header, grey label + unit columns, key rows bold on light green.

        `rows` hold strings (Bound values carry their cell); `sections` are row indices rendered as sub-headers."""
        n_rows, n_cols = len(rows) + 1, len(header)
        height = height or 0.59 + 0.34 * len(rows)
        shape = self.s.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.name = name
        tbl = shape.table
        tbl.first_row = tbl.horz_banding = False
        if col_widths is None:
            data_w = (width - 4.0 - 1.2) / max(1, n_cols - 2)
            col_widths = [4.0, 1.2] + [data_w] * (n_cols - 2)
        for j, w in enumerate(col_widths):
            tbl.columns[j].width = Inches(w)
        tbl.rows[0].height = Inches(0.59)
        for i in range(1, n_rows):
            tbl.rows[i].height = Inches((height - 0.59) / len(rows))
        for i, values in enumerate([header] + list(rows)):
            body = i - 1
            for j, v in enumerate(values):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = Emu(73152)
                cell.margin_top = cell.margin_bottom = Emu(27432)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.text = str(v)
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if j < 2 else PP_ALIGN.RIGHT
                key = body in key_rows
                if i == 0:
                    fill, color, bold = GREEN, WHITE, True
                elif body in sections:
                    fill, color, bold = KEY_FILL, DARK_GREEN, True
                else:
                    fill = LABEL_FILL if j < 2 else (KEY_FILL if key else None)
                    color, bold = (GREY if j == 1 else TEXT), key
                if p.runs:
                    _font(p.runs[0], 13, bold, color)
                if fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(fill)
                else:
                    cell.fill.background()
                if isinstance(v, Bound):
                    self.d.bindings.append((self.number, name, str(v), v.ref, v.value))
        return shape

    def chart(self, kind, categories, series: dict, *, left=LEFT, top=6.0, width=WIDTH, height=4.0, title=None,
              number_format="#,##0.0", name="Chart"):
        """Editable chart: kind = column | stacked | bar. Series values are BoundNum (from Deck.bind)."""
        ctype = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "stacked": XL_CHART_TYPE.COLUMN_STACKED,
                 "bar": XL_CHART_TYPE.BAR_CLUSTERED}[kind]
        data = CategoryChartData(number_format=number_format)
        data.categories = categories
        for sname, vals in series.items():
            data.add_series(sname, [float(v) for v in vals])
        gf = self.s.shapes.add_chart(ctype, Inches(left), Inches(top), Inches(width), Inches(height), data)
        gf.name = name
        ch = gf.chart
        ch.font.name, ch.font.size = FONT, Pt(12)
        ch.font.color.rgb = _rgb(TEXT)
        ch.has_title = bool(title)
        if title:
            ch.chart_title.text_frame.text = title
            _font(ch.chart_title.text_frame.paragraphs[0].runs[0], 13, True)
        va = ch.value_axis
        va.visible = False
        va.has_major_gridlines = va.has_minor_gridlines = False
        ch.category_axis.has_major_gridlines = False
        ch.category_axis.format.line.color.rgb = _rgb("BFBFBF")
        ch.category_axis.tick_labels.font.size = Pt(12)
        ch.category_axis.reverse_order = kind == "bar"   # bars read top-down in category order
        plot = ch.plots[0]
        plot.gap_width = 55 if kind == "stacked" else 80
        if kind == "stacked":
            plot.overlap = 100
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = number_format, False
        dl.font.size = Pt(12)
        dl.position = XL_LABEL_POSITION.CENTER if kind == "stacked" else XL_LABEL_POSITION.OUTSIDE_END
        if kind == "stacked":
            dl.font.color.rgb = _rgb(WHITE)
        ch.has_legend = len(series) > 1
        if ch.has_legend:
            ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
            ch.legend.font.size = Pt(12)
        for i, s in enumerate(plot.series):
            s.invert_if_negative = False
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = _rgb(SERIES[i % len(SERIES)])
        for sname, vals in series.items():
            for cat, v in zip(categories, vals):
                if hasattr(v, "ref"):
                    self.d.bindings.append((self.number, name, f"{sname} / {cat}", v.ref, v.value))
        return gf

    def card(self, header, bullets, *, left=15.35, top=1.75, width=4.55, height=4.75, name="Card"):
        """Grey side card: header (16 pt dark green) and bullets [(bold lead-in, text)]."""
        box = self.s.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
        box.name = name
        box.fill.solid()
        box.fill.fore_color.rgb = _rgb(LABEL_FILL)
        box.line.fill.background()
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.15)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        _font(p.add_run(), 16, True, DARK_GREEN)
        p.runs[0].text = header
        text = header
        for lead, body in bullets:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(8)
            for t, bold in ((lead, True), (body, False)):
                r = p.add_run()
                r.text = t
                _font(r, 13, bold)
            text += f"\n{lead}{body}"
        self._record(name, text)
        return box

    def footnote(self, text, top=10.55):
        src = f"Source: {self.d.source}, {self.d.date.day} {self.d.date:%b %Y}."
        return self.text("Footnote", f"{text}\n{src}" if text else src, self.LEFT, top, 14.4, 0.7, size=11,
                         color=GREY)


def _borders(cell, top=None, bottom=None, color=TEXT):
    """Set top/bottom borders (EMU widths) and remove left/right borders of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()  # noqa: SLF001
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for i, (tag, w) in enumerate((("a:lnL", None), ("a:lnR", None), ("a:lnT", top), ("a:lnB", bottom))):
        ln = etree.SubElement(tcPr, qn(tag))
        tcPr.remove(ln)
        tcPr.insert(i, ln)
        if w:
            ln.set("w", str(w))
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", color)
        else:
            etree.SubElement(ln, qn("a:noFill"))
