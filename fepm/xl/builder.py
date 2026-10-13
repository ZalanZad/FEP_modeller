"""Workbook builder for the FEP modelling standard (docs/MODELLING_STANDARDS.md).

    m = Model("Project Demo", periods=26, periodicity="annual", start=date(2026, 1, 1), scenarios=["Base", "Low"])
    capex = m.inputs.add_input("Capex", "Capex", "EUR k", 13000, {2: 15000}, source="EPC quote")
    m.calc.section("Capex")
    m.calc.row("capex", "Capex", "EUR k", "=-" + capex + "*{constr}")
    m.save("out/demo.xlsx")

Calc formulas are templates rendered once per timeline column, so every row has one formula by construction.
Tokens (`key` is a Calc row key or a header key: start, end, days, yf, year, period):

    {key}          same column            J15 / J$4
    {key@prev}     previous column        I15 / I$4
    {key.E} {key.F} {key.I}               $E15  (row constants / total)
    {key.all}      whole timeline         $J15:$AI15
    {key.todate}   first column to here   $J15:J15
    {self...}      the row being written

Workbook names (inputs `In_<key>`, constants such as `Thousand`) are written as plain text.
On calc sheets tokens render without a sheet prefix; on Checks / Summary / Dashboard they resolve to one calc sheet.

Layouts:
    layout="standard"   Cover | Summary | Inputs | Curves | Calc | Checks          (default)
    layout="dashboard"  Dashboard | Inputs | Curves | Calc_1..Calc_n | Checks    (no Cover, no Summary)

Several identical calc sheets: `calc_sheets=["Calc_1", "Calc_2"], bids=["BORA 1", "BORA 2"]`. One set of row
templates is written to every sheet with identical formula text; `A1` holds the bid number (input style) and
`m.calc.lever("Dash_X")` gives `INDEX(Dash_X,1,$A$1)`. `m.calc.ref(key, "E", sheet="Calc_2")` addresses one sheet.
A check whose formula has row tokens gets one row per calc sheet (label suffixed with the sheet name);
`sheet="Calc_2"` pins it to one sheet, and a check without tokens is a single row.

Dashboard (`m.dashboard`): section / subheader bars (bid labels over the value columns, optional headers over the
grid columns), `input` (deal levers `Dash_<key>`, one value per bid, optional list validation), `link` (a calc
result per bid, optional scale and `out=` → `Out_<key>_<n>`), `status` (OK/CHECK), `stored_row` / `stored_block`
(purple, `<name>_Stored`, written by a macro), `scenario_results` (`Live_Results` / `Stored_Results`), and the
right-hand panel: `panel_section`, `panel_link` (Live values from Inputs) and `selector` (`Scenario_Name`).
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from . import styles
from .styles import apply

J = 10                       # first timeline column
HEADER = {"start": 3, "end": 4, "days": 5, "yf": 6, "year": 7, "period": 8}
FIRST_ROW = 11               # first content row on Calc / Checks (panes frozen at J10)
UNIT_FMT = {"EUR k": "money", "EUR m": "money1", "%": "pct", "x": "mult", "date": "date", "flag": "flag",
            "MW": "money1", "MWh": "money", "EUR/MWh": "money2", "EUR k/MW": "money", "years": "int",
            "factor": "factor", "text": "text", "#": "general"}
_TOKEN = re.compile(r"\{(\w+)(@prev|\.E|\.F|\.I|\.all|\.todate)?\}")
_LINK = re.compile(r"^=(?:[A-Za-z_][\w.]*|\w+!\$?[A-Z]{1,3}\$?\d+)$")
OK_FMT = [CellIsRule(operator="equal", formula=['"OK"'], font=Font(color=styles.OK_GREEN, bold=True)),
          CellIsRule(operator="equal", formula=['"CHECK"'], font=Font(color=styles.CHECK_RED, bold=True))]


def _fmt(unit, fmt):
    return fmt or UNIT_FMT.get(unit, "money")


def _status(ref):
    return f'=IF({ref}=0,"OK","CHECK")'


def _ok_format(ws, rng):
    for rule in OK_FMT:
        ws.conditional_formatting.add(rng, rule)


@dataclass
class Row:
    key: str
    label: str
    unit: str
    r: int
    formula: str | None = None        # timeline template
    e: str | None = None
    f: str | None = None
    total: str = "sum"                # sum | none | seed
    seed: str | None = None
    fmt: str | None = None
    ef_fmt: str | None = None
    bold: bool = False
    style: str | None = None          # force a style for the timeline (e.g. "macro")
    values: list = field(default_factory=list)   # macro-written timeline values (PASTE rows)


class Model:
    def __init__(self, project: str, purpose: str = "Investment model", *, periods: int, start: dt.date,
                 periodicity: str = "annual", scenarios=("Base",), version: str = "v1", author: str = "FEP",
                 date: dt.date | None = None, curves: bool = False, layout: str = "standard",
                 calc_sheets=("Calc",), bids=None, extra_cols: int = 0, source_col: bool = True):
        if periodicity not in ("annual", "quarterly"):
            raise ValueError("periodicity must be 'annual' or 'quarterly'")
        if not 1 <= len(scenarios) <= 8:
            raise ValueError("1 to 8 scenarios")
        if layout not in ("standard", "dashboard"):
            raise ValueError("layout must be 'standard' or 'dashboard'")
        calc_sheets = list(calc_sheets)
        if calc_sheets != ["Calc"] and calc_sheets != [f"Calc_{i}" for i in range(1, len(calc_sheets) + 1)]:
            raise ValueError("calc_sheets must be ['Calc'] or ['Calc_1', 'Calc_2', ...]")
        self.calc_sheets = calc_sheets
        self.bids = list(bids or [f"Bid {i}" for i in range(1, len(calc_sheets) + 1)])
        if len(self.bids) != len(calc_sheets):
            raise ValueError("one bid label per calc sheet")
        self.layout = layout
        self.home = "Cover" if layout == "standard" else "Dashboard"
        self.wb = Workbook()
        styles.register(self.wb)
        self.n, self.last = periods, J + periods - 1
        self.scenarios = list(scenarios)
        self.names: dict[str, str] = {}
        self.solver_topics: list[str] = []
        self.goal_topics: list[str] = []
        ws = self.wb.active
        ws.title = self.home
        self.cover = self.summary = self.dashboard = None
        if layout == "standard":
            self.cover = Cover(self, ws, project, purpose, version, author, date or dt.date.today())
            self.summary = Summary(self, self.wb.create_sheet("Summary"))
        else:
            self.dashboard = Dashboard(self, ws, project, purpose, version, date or dt.date.today(), extra_cols,
                                       source_col)
        self.inputs = Inputs(self, self.wb.create_sheet("Inputs"), start, periodicity)
        self.curves = Curves(self, self.wb.create_sheet("Curves")) if curves else None
        self.calc = Calc(self, [self.wb.create_sheet(name) for name in calc_sheets])
        self.checks = Checks(self, self.wb.create_sheet("Checks"))
        for w in self.wb.worksheets:
            w.sheet_view.showGridLines = False
            if w.title != self.home:
                w["B2"] = f'={self.home}!$C$3&" – {self._purpose(w.title)}'
                _banner(w, 2, self._width(w.title))

    def _purpose(self, title):
        """Banner text after the project name (closing quote included; bid sheets add their bid label)."""
        if title.startswith("Calc_"):
            bid = "INDEX(Bid_Labels,1,$A$1)" if self.dashboard else '"bid "&$A$1'
            return f'Calculations – "&{bid}'
        return f'{_PURPOSE[title]}"'

    # ---- names and references -------------------------------------------------------------
    def define(self, name: str, ref: str) -> str:
        """Define a workbook name for `Sheet!$A$1` (or a range) and return the name."""
        if name in self.names:
            raise ValueError(f"name {name} already defined")
        self.names[name] = ref
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        return name

    def address(self, name_or_ref: str) -> str:
        return self.names.get(name_or_ref, name_or_ref)

    def render(self, template: str, row: Row | None, col: int | None, sheet: str = "Calc",
               target: str | None = None) -> str:
        """Render a formula template for column `col` (None for row constants / totals).

        On a calc sheet tokens refer to the same sheet; elsewhere to `target` (default: the first calc sheet)."""
        prefix = "" if sheet in self.calc_sheets else f"{target or self.calc_sheets[0]}!"

        def sub(m):
            key, mod = m.group(1), m.group(2) or ""
            if key == "self":
                key = row.key
            if key in HEADER:
                r, rd = HEADER[key], f"${HEADER[key]}"
            elif key in self.calc.rows:
                r = self.calc.rows[key].r
                rd = str(r)
            else:
                raise KeyError(f"unknown row key {{{key}}} in {template!r}")
            if mod in (".E", ".F", ".I"):
                return f"{prefix}${mod[1]}{rd}"
            if mod == ".all":
                return f"{prefix}${L(J)}{rd}:${L(self.last)}{rd}"
            if col is None:
                raise ValueError(f"{{{key}{mod}}} needs a timeline column ({template!r})")
            if mod == ".todate":
                return f"{prefix}${L(J)}{rd}:{L(col)}{rd}"
            return f"{prefix}{L(col - 1 if mod == '@prev' else col)}{rd}"

        return _TOKEN.sub(sub, template)

    # ---- macro-driven values --------------------------------------------------------------
    def goal_seek(self, topic: str, label: str, unit: str, seed: float, achieved: str, target: str,
                  source: str = "Solved by GoalSeek.bas", fmt: str | None = None) -> str:
        """Purple input `<topic>_Input` that GoalSeek.bas solves so that `achieved` = `target`.

        `achieved` / `target` are workbook names or `Sheet!$A$1` refs. Adds the tolerance check."""
        name = self.inputs.add_macro_cell(f"{topic}_Input", label, unit, seed, source, fmt)
        self.define(f"{topic}_Achieved", self.address(achieved))
        self.define(f"{topic}_Target", self.address(target))
        self.goal_topics.append(topic)
        self.checks.add_check(f"Goal seek {topic}: achieved = target",
                              f"=N(ABS({topic}_Achieved-{topic}_Target)>Tolerance)", per_period=False)
        return name

    # ---- output ---------------------------------------------------------------------------
    def _width(self, title):
        if title == "Dashboard":
            return self.dashboard.last
        return {"Cover": 8, "Summary": 6 + len(self.scenarios), "Inputs": 16}.get(title, self.last)

    def save(self, path):
        self.inputs.finish()
        if self.curves:
            self.curves.finish()
        self.calc.finish()
        if self.summary:
            self.summary.write_outputs()     # adds the stored = live check
        self.checks.finish()
        if self.summary:
            self.summary.write_checks()
            self.cover.finish()
        if self.dashboard:
            self.dashboard.finish()
        for ws in self.wb.worksheets:
            ws.column_dimensions.group(L(self._width(ws.title) + 1), "XFD", hidden=True)
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(path)
        return path


_PURPOSE = {"Summary": "Summary", "Inputs": "Inputs and scenarios", "Curves": "Time-series inputs",
            "Calc": "Calculations", "Checks": "Checks"}
STORED_CHECK = ("Stored results = live case",
                "=--(SUMPRODUCT(ABS(Stored_Results-Live_Results)*(Scenario_Names=Scenario_Name))>Tolerance)")


def _banner(ws, r, last_col):
    for c in range(1, last_col + 1):
        apply(ws.cell(r, c), "banner")


def _bar(ws, r, title, last_col, kind="section"):
    ws.cell(r, 1, "x")
    ws.cell(r, 3, title)
    for c in range(1, last_col + 1):
        apply(ws.cell(r, c), kind)


def _end(ws, r, last_col):
    ws.cell(r, 3, "***END***")
    apply(ws.cell(r, 3), "end")


def _columns(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _validate(ws, cells, options):
    """List validation: `options` is a workbook name / range or a list of allowed values."""
    f = options if isinstance(options, str) else '"' + ",".join(str(o) for o in options) + '"'
    dv = DataValidation(type="list", formula1=f, allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(cells)


def _write(ws, r, c, formula, kind="calc", fmt=None):
    """Write a formula; pure links get the green link style."""
    cell = ws.cell(r, c, formula)
    if kind == "calc" and isinstance(formula, str) and _LINK.match(formula):
        kind = "link"
    apply(cell, kind, fmt)
    return cell


# =========================================================================================
class Cover:
    def __init__(self, model, ws, project, purpose, version, author, date):
        self.m, self.ws = model, ws
        ws["B2"] = '=C3&" – Cover"'
        _banner(ws, 2, 8)
        rows = [("Project", project, None), ("Purpose", purpose, None), ("Version", version, None),
                ("Date", date, "date"), ("Author", author, None)]
        for i, (label, value, fmt) in enumerate(rows):
            ws.cell(3 + i, 2, label)
            apply(ws.cell(3 + i, 3, value), "input", fmt)
        ws["C3"].font = Font(name=styles.FONT, size=16, bold=True, color=styles.BLUE)
        _columns(ws, {"A": 3, "B": 16, "C": 34, "D": 3, "E": 16, "F": 30, "G": 3, "H": 3})

    def finish(self):
        ws = self.ws
        _bar(ws, 9, "Colour legend", 8)
        legend = [("input", "Hard-coded input", "Inputs and Curves only, with unit and source"),
                  ("link", "Link from another sheet", ""),
                  ("calc", "Calculation", ""),
                  ("macro", "Macro-written value", "Written by vba/*.bas only (PASTE rows, solved inputs, "
                                                   "stored scenario results)"),
                  ("seed", "Opening seed", "Corkscrew opening balance")]
        for i, (kind, text, note) in enumerate(legend):
            apply(ws.cell(10 + i, 3, text), kind)
            ws.cell(10 + i, 3).alignment = styles.Alignment(horizontal="left")
            ws.cell(10 + i, 5, note)
        _bar(ws, 16, "Model status", 8)
        ws.cell(17, 2, "Master check")
        _write(ws, 17, 3, _status("Master_Check"))
        ws["C17"].alignment = styles.Alignment(horizontal="left")
        _ok_format(ws, "C17")
        ws.cell(18, 2, "Macros")
        ws.cell(18, 3, "RunAllScenarios (Scenarios.bas) stores all scenarios and runs the solvers")
        _end(ws, 20, 8)


# =========================================================================================
class Inputs:
    """B label, C unit, D source, F Live, H:O scenarios 1-8 (MODELLING_STANDARDS §4)."""

    def __init__(self, model, ws, start, periodicity):
        self.m, self.ws, self.r = model, ws, 4
        n = len(model.scenarios)
        _columns(ws, {"A": 3, "B": 40, "C": 10, "D": 36, "E": 2, "F": 12, "G": 2,
                      **{L(8 + i): 12 for i in range(8)}, "P": 2})
        ws["B4"], ws["B5"] = "Scenario selected", "Scenario number"
        if model.dashboard:              # the selector lives on the Dashboard
            _write(ws, 4, 6, "=Scenario_Name")
            ws["D4"] = "Set on the Dashboard"
        else:
            apply(ws.cell(4, 6, model.scenarios[0]), "input")
            model.define("Scenario_Name", "Inputs!$F$4")
            _validate(ws, "F4", "Scenario_Names")
        names_ref = f"Inputs!$H$7:${L(7 + n)}$7"
        _write(ws, 5, 6, "=MATCH(Scenario_Name,Scenario_Names,0)", fmt="int")
        model.define("Scenario_Selected", "Inputs!$F$5")
        for c, h in zip((2, 3, 4, 6), ("Input", "Unit", "Source / note", "Live")):
            apply(ws.cell(7, c, h), "header")
        ws["B7"].alignment = ws["C7"].alignment = ws["D7"].alignment = styles.Alignment(horizontal="left")
        for i in range(8):
            cell = ws.cell(7, 8 + i, model.scenarios[i] if i < n else None)
            apply(cell, "input" if i < n else "header")
        model.define("Scenario_Names", names_ref)
        ws.freeze_panes = "E8"
        self.r = 9
        self.section("Constants")
        for name, value, unit, note in [
                ("Thousand", 1000, "#", "Scale"), ("Million", 1000000, "#", "Scale"),
                ("Days_In_Year", 365, "days", "Day count for year fractions and XIRR"),
                ("Months_In_Year", 12, "months", ""),
                ("Months_Per_Period", 12 if periodicity == "annual" else 3, "months",
                 f"Timeline periodicity ({periodicity})"),
                ("Tolerance", 0.000001, "#", "Check and solver tolerance")]:
            self.add_constant(name, name.replace("_", " "), unit, value, note)
        self.section("Timing")
        self.model_start = self.add_input("Model_Start", "Model start date", "date", start, source="Timeline start")

    def section(self, title):
        self.r += 1
        _bar(self.ws, self.r, title, 16)
        self.r += 1

    add_section = section

    def _label(self, label, unit, source):
        ws = self.ws
        ws.cell(self.r, 2, label)
        apply(ws.cell(self.r, 3, unit), "unit")
        ws.cell(self.r, 4, source)

    def add_constant(self, name, label, unit, value, source=""):
        self._label(label, unit, source)
        apply(self.ws.cell(self.r, 6, value), "input", "general")
        self.m.define(name, f"Inputs!$F${self.r}")
        self.r += 1
        return name

    def add_input(self, key, label, unit, base, scenarios: dict | None = None, source: str = "",
                  fmt: str | None = None) -> str:
        """Scenario input: base in H, overrides by scenario index (1-based) in I:O, Live in F. Returns `In_<key>`."""
        ws, r, fmt = self.ws, self.r, _fmt(unit, fmt)
        self._label(label, unit, source)
        for i in range(len(self.m.scenarios)):
            v = base if i == 0 else (scenarios or {}).get(i + 1)
            apply(ws.cell(r, 8 + i, v), "input", fmt)
        live = f"INDEX($H{r}:$O{r},1,Scenario_Selected)"
        _write(ws, r, 6, f'=IF({live}="",$H{r},{live})', fmt=fmt)     # blank scenario cell = base case
        self.r += 1
        return self.m.define(f"In_{key}", f"Inputs!$F${r}")

    def add_macro_cell(self, name, label, unit, seed, source, fmt=None) -> str:
        """Purple cell that only a macro writes (solved input). `seed` is the starting guess."""
        self._label(label, unit, source)
        apply(self.ws.cell(self.r, 6, seed), "macro", _fmt(unit, fmt))
        self.r += 1
        return self.m.define(name, f"Inputs!$F${self.r - 1}")

    def finish(self):
        if self.m.solver_topics or self.m.goal_topics:      # only models with a circularity or a goal seek
            self._macro_control()
        _end(self.ws, self.r + 1, 16)

    def _macro_control(self):
        """Topic lists read by Solver.bas / GoalSeek.bas; their presence also makes package() embed those modules."""
        self.section("Macro control")
        for name, topics, note in [("Solver_Topics", self.m.solver_topics, "COPY/PASTE topics for Solver.bas"),
                                   ("GoalSeek_Topics", self.m.goal_topics, "Topics for GoalSeek.bas")]:
            self._label(name.replace("_", " "), "text", note)
            apply(self.ws.cell(self.r, 6, ",".join(topics) or None), "input", "text")
            self.m.define(name, f"Inputs!$F${self.r}")
            self.r += 1


# =========================================================================================
class Curves:
    """Annual curve blocks picked by name (INDEX/MATCH), two spare rows per block."""

    def __init__(self, model, ws):
        self.m, self.ws, self.r = model, ws, 4
        _columns(ws, {"A": 3, "B": 3, "C": 40, "D": 10, "E": 12, "F": 12, "G": 2, "H": 2, "I": 12,
                      **{L(c): 11 for c in range(J, model.last + 1)}})
        ws.freeze_panes = "J4"

    def add_block(self, name, label, unit, years: list[int], curves: dict[str, list], source="", fmt=None):
        """Defines `<name>_Years`, `<name>_Labels`, `<name>_Data`. Returns the block name."""
        ws, fmt = self.ws, _fmt(unit, fmt)
        if len(years) > self.m.n:
            raise ValueError("curve longer than the timeline")
        self.r += 1
        _bar(ws, self.r, label, self.m.last)
        ws.cell(self.r, 5, source).font = Font(name=styles.FONT, size=styles.SIZE, color=styles.WHITE, italic=True)
        self.r += 1
        ws.cell(self.r, 3, "Year")
        for i, y in enumerate(years):
            apply(ws.cell(self.r, J + i, y), "input", "int")
        last = L(J + len(years) - 1)
        self.m.define(f"{name}_Years", f"Curves!${L(J)}${self.r}:${last}${self.r}")
        first = self.r + 1
        for cname, vals in list(curves.items()) + [(None, None), (None, None)]:
            self.r += 1
            apply(ws.cell(self.r, 3, cname), "input" if cname else "label")
            apply(ws.cell(self.r, 4, unit), "unit")
            for i in range(len(years)):
                apply(ws.cell(self.r, J + i, vals[i] if vals else None), "input", fmt)
        self.m.define(f"{name}_Labels", f"Curves!$C${first}:$C${self.r}")
        self.m.define(f"{name}_Data", f"Curves!${L(J)}${first}:${last}${self.r}")
        self.r += 1
        return name

    def finish(self):
        _end(self.ws, self.r + 1, self.m.last)


# =========================================================================================
class Calc:
    """Row templates written to every calc sheet (`Calc`, or `Calc_1..Calc_n` with the bid number in A1)."""

    def __init__(self, model, sheets):
        self.m, self.sheets, self.ws = model, sheets, sheets[0]
        self.rows: dict[str, Row] = {}
        self.r = FIRST_ROW
        for i, ws in enumerate(sheets, 1):
            _columns(ws, {"A": 3, "B": 3, "C": 40, "D": 10, "E": 12, "F": 12, "G": 2, "H": 2, "I": 12,
                          **{L(c): 11 for c in range(J, model.last + 1)}})
            ws.freeze_panes = "J10"
            if ws.title != "Calc":
                apply(ws.cell(1, 1, i), "input", "int")
                ws.cell(1, 3, "Bid number")

    # ---- layout ---------------------------------------------------------------------------
    def section(self, title):
        if self.r > FIRST_ROW:
            self.r += 1
        for ws in self.sheets:
            _bar(ws, self.r, title, self.m.last)
        self.r += 1

    def subheader(self, title):
        for ws in self.sheets:
            _bar(ws, self.r, title, self.m.last, "subheader")
            ws.cell(self.r, 1).value = None
        self.r += 1

    def blank(self):
        self.r += 1

    add_section = section

    # ---- rows -----------------------------------------------------------------------------
    def row(self, key, label, unit, formula=None, *, e=None, f=None, total="sum", seed=None, fmt=None,
            ef_fmt=None, bold=False, style=None) -> str:
        """Add a row. `formula` is the timeline template; `e`/`f` row constants or results; `total`
        sum (I = lifetime total) | none | seed (I = opening seed, grey). Returns the key."""
        if key in self.rows or key in HEADER or key == "self":
            raise ValueError(f"duplicate or reserved row key {key}")
        self.rows[key] = Row(key, label, unit, self.r, formula, e, f, total, seed, fmt, ef_fmt, bold, style)
        self.r += 1
        return key

    def ref(self, key, col="I", sheet=None) -> str:
        """Absolute reference to a row's E/F/I cell on one calc sheet (default: the first)."""
        return f"{sheet or self.m.calc_sheets[0]}!${col}${self.rows[key].r}"

    def refs(self, key, col="I") -> list[str]:
        """The same cell on every calc sheet, in bid order."""
        return [self.ref(key, col, s) for s in self.m.calc_sheets]

    def name(self, name, key, col="E", sheet=None) -> str:
        return self.m.define(name, self.ref(key, col, sheet))

    @staticmethod
    def lever(name) -> str:
        """This sheet's value of a 1 x n Dashboard range (`Dash_<key>`): `INDEX(name,1,$A$1)`."""
        return f"INDEX({name},1,$A$1)"

    # ---- flags (MODELLING_STANDARDS §6.5) -----------------------------------------------
    def window(self, key, label, start, end):
        """Share of each period inside [start, end] (dates or names)."""
        return self.row(key, label, "flag", "=MAX(0,(MIN({self.F},{end})-MAX({self.E},{start})+1)/{days})",
                        e=start, f=end, total="none", fmt="factor", ef_fmt="date")

    def event(self, key, label, date):
        return self.row(key, label, "flag", "=N(AND({start}<={self.E},{end}>={self.E}))", e=date,
                        total="sum", fmt="flag", ef_fmt="date")

    def cumulative(self, key, label, of):
        return self.row(key, label, "flag", f"=MIN(1,SUM({{{of}.todate}}))", total="none", fmt="flag")

    def index(self, key, label, base_year, rate_key):
        """Indexation: 1 in the base year, then compounding with the row `rate_key`."""
        return self.row(key, label, "factor", f"=IF({{year}}={{self.E}},1,{{self@prev}}*(1+{{{rate_key}}}))",
                        e=base_year, total="seed", seed="=1", fmt="factor", ef_fmt="int")

    def curve(self, key, label, unit, block, selector, fmt=None):
        """Pick a curve by name from a Curves block for each period's year."""
        return self.row(key, label, unit,
                        f"=INDEX({block}_Data,MATCH({{self.E}},{block}_Labels,0),MATCH({{year}},{block}_Years,0))",
                        e=f"={selector}", total="none", fmt=fmt, ef_fmt="text")

    def corkscrew(self, key, label, unit, movements: list[tuple[str, str, str]], seed=None):
        """Opening / movements / closing. `movements` = [(key, label, template)]. Returns the closing key."""
        self.subheader(label)
        self.row(f"{key}_open", "Opening balance", unit, f"={{{key}_close@prev}}", total="none")
        for mk, ml, mf in movements:
            self.row(mk, ml, unit, mf)
        last = movements[-1][0]
        return self.row(f"{key}_close", "Closing balance", unit, f"=SUM({{{key}_open}}:{{{last}}})",
                        total="seed", seed=seed, bold=True)

    add_row = row

    # ---- circularity (COPY / PASTE / Delta, §7) -------------------------------------------
    def _topics(self, topic):
        """(sheet, topic name) per calc sheet: `Fee` on a single sheet, `Fee_1`, `Fee_2`... otherwise."""
        sheets = self.m.calc_sheets
        return [(s, topic if len(sheets) == 1 else f"{topic}_{i}") for i, s in enumerate(sheets, 1)]

    def paste(self, topic, label, unit):
        """Macro-written PASTE row (place it above the rows that use it). Key: `<topic>_paste`."""
        key = self.row(f"{topic.lower()}_paste", f"{label} – paste (macro)", unit, None, total="none",
                       style="macro")
        self.rows[key].values = [0] * self.m.n
        for sheet, t in self._topics(topic):
            self.m.define(f"{t}_Paste", self._span(key, sheet))
        return key

    def copy(self, topic, label, unit, formula):
        """COPY row with the live formula and its Delta row; registers the topic(s) for Solver.bas."""
        ck = self.row(f"{topic.lower()}_copy", f"{label} – copy", unit, formula, total="none")
        dk = self.row(f"{topic.lower()}_delta", f"{label} – delta", unit,
                      f"=ABS({{{ck}}}-{{{topic.lower()}_paste}})", total="sum")
        for sheet, t in self._topics(topic):
            self.m.define(f"{t}_Copy", self._span(ck, sheet))
            self.m.define(f"{t}_Delta", self.ref(dk, "I", sheet))
            self.m.solver_topics.append(t)
            self.m.checks.add_check(f"Solver {t}: delta = 0", f"=N({t}_Delta>Tolerance)", per_period=False)
        return ck

    def _span(self, key, sheet):
        r = self.rows[key].r
        return f"{sheet}!${L(J)}${r}:${L(self.m.last)}${r}"

    # ---- returns (column E) -------------------------------------------------------------
    def xirr(self, key, label, cf, iferror=False):
        f = f"XIRR({{{cf}.all}},{{end.all}})"
        return self.row(key, label, "%", e=f'=IFERROR({f},"n/a")' if iferror else "=" + f, total="none",
                        ef_fmt="pct", bold=True)

    def xnpv(self, key, label, rate, cf, unit="EUR k"):
        return self.row(key, label, unit, e=f"=XNPV({rate},{{{cf}.all}},{{end.all}})", total="none",
                        ef_fmt="money", bold=True)

    def moic(self, key, label, cf):
        return self.row(key, label, "x", e=f'=SUMIF({{{cf}.all}},">0")/-SUMIF({{{cf}.all}},"<0")',
                        total="none", ef_fmt="mult", bold=True)

    # ---- write ----------------------------------------------------------------------------
    def finish(self):
        for ws in self.sheets:
            self._finish(ws)

    def _finish(self, ws):
        m = self.m
        for r, (label, unit) in zip(range(3, 9), [("Period start", "date"), ("Period end", "date"),
                                                  ("Days in period", "days"), ("Year fraction", "factor"),
                                                  ("Year", "year"), ("Period counter", "#")]):
            ws.cell(r, 3, label)
            apply(ws.cell(r, 4, unit), "unit")
        _write(ws, 4, 9, f"={m.inputs.model_start}-1", fmt="date")
        tmpl = {3: "={end@prev}+1", 4: "=EOMONTH({start},Months_Per_Period-1)", 5: "={end}-{start}+1",
                6: "={days}/Days_In_Year", 7: "=YEAR({end})", 8: "={period@prev}+1"}
        fmts = {3: "date", 4: "date", 5: "int", 6: "factor", 7: "int", 8: "int"}
        for c in range(J, m.last + 1):
            for r, t in tmpl.items():
                cell = _write(ws, r, c, m.render(t, None, c, ws.title), fmt=fmts[r])
                if r == 4:
                    cell.font = Font(name=styles.FONT, size=styles.SIZE, bold=True)
        for row in self.rows.values():
            _write_row(ws, m, row, ws.title)
        _end(ws, self.r + 1, m.last)


def _write_row(ws, m, row: Row, sheet):
    r, fmt = row.r, _fmt(row.unit, row.fmt)
    ws.cell(r, 3, row.label)
    apply(ws.cell(r, 4, row.unit), "unit")
    if row.bold:
        ws.cell(r, 3).font = Font(name=styles.FONT, size=styles.SIZE, bold=True)
    for c, t in ((5, row.e), (6, row.f)):
        if t is not None:
            t = t if str(t).startswith("=") else f"={t}"
            _write(ws, r, c, m.render(t, row, None, sheet), "total" if row.bold else "calc",
                   row.ef_fmt or fmt)
    if row.total == "sum" and (row.formula or row.values):
        _write(ws, r, 9, f"=SUM({L(J)}{r}:{L(m.last)}{r})", "total" if row.bold else "calc", fmt)
    elif row.total == "seed":
        cell = ws.cell(r, 9, m.render(row.seed, row, None, sheet) if row.seed else None)
        apply(cell, "seed", fmt)
    for i, c in enumerate(range(J, m.last + 1)):
        if row.values:
            apply(ws.cell(r, c, row.values[i]), "macro", fmt)
        elif row.formula:
            _write(ws, r, c, m.render(row.formula, row, c, sheet), row.style or ("total" if row.bold else "calc"),
                   fmt)


# =========================================================================================
class Checks:
    """One row per check; per-period tests in J→ summed in F, status in E, Master_Check at the bottom."""

    def __init__(self, model, ws):
        self.m, self.ws = model, ws
        self.items: list[tuple[str, str, bool, str | None]] = []
        _columns(ws, {"A": 3, "B": 3, "C": 48, "D": 10, "E": 12, "F": 12, "G": 2, "H": 2, "I": 2,
                      **{L(c): 11 for c in range(J, model.last + 1)}})
        ws.freeze_panes = "J10"

    def add_check(self, label, formula, per_period=True, sheet=None):
        """`formula` returns 0 when OK; row tokens resolve to a calc sheet ({key} → Calc_1!J15).

        With several calc sheets a formula with row tokens gives one row per sheet (label suffixed with the
        sheet name); `sheet` pins it to one sheet. A formula without tokens is a single row."""
        self.items.append((label, formula, per_period, sheet))

    def equal(self, label, a, b):
        """Per-period check that two Calc rows (templates) agree within Tolerance."""
        self.add_check(label, f"=N(ABS({a}-{b})>Tolerance)")

    def _expand(self):
        multi = len(self.m.calc_sheets) > 1
        for label, formula, per_period, sheet in self.items:
            sheets = [sheet] if sheet else self.m.calc_sheets if _TOKEN.search(formula) else [None]
            for s in sheets:
                yield (f"{label} – {s}" if multi and s else label), formula, per_period, s

    def finish(self):
        ws, m = self.ws, self.m
        ws.cell(4, 3, "Period end")
        for c in range(J, m.last + 1):
            _write(ws, 4, c, f"={m.calc_sheets[0]}!{L(c)}$4", fmt="date")
        apply(ws.cell(9, 5, "Status"), "header")
        apply(ws.cell(9, 6, "Fails"), "header")
        r = FIRST_ROW
        _bar(ws, r, "Checks (0 = OK)", m.last)
        first = r + 1
        self.rows = []
        for label, formula, per_period, sheet in self._expand():
            r += 1
            ws.cell(r, 3, label)
            apply(ws.cell(r, 4, "flag"), "unit")
            if per_period:
                for c in range(J, m.last + 1):
                    _write(ws, r, c, m.render(formula, None, c, "Checks", sheet), fmt="flag")
                _write(ws, r, 6, f"=SUM({L(J)}{r}:{L(m.last)}{r})", fmt="flag")
            else:
                _write(ws, r, 6, m.render(formula, None, None, "Checks", sheet), fmt="flag")
            _write(ws, r, 5, _status(f"F{r}"))
            self.rows.append((label, r))
        if len(self.rows) > 1 and self.rows[-1][0] == STORED_CHECK[0]:
            r += 2      # every check but stored = live: a scenario macro can store it per set without a circularity
            ws.cell(r, 3, "Checks excluding stored = live")
            _write(ws, r, 6, f"=SUMPRODUCT(ABS(F{first}:F{r - 3}))", "total", "flag")
            m.define("Model_Checks", f"Checks!$F${r}")
        r += 2
        ws.cell(r, 3, "Master check")
        ws.cell(r, 3).font = Font(name=styles.FONT, size=styles.SIZE, bold=True)
        last = self.rows[-1][1]
        _write(ws, r, 6, f"=SUMPRODUCT(ABS(F{first}:F{last}))", "total", "flag")
        _write(ws, r, 5, _status(f"F{r}"), "total")
        m.define("Master_Check", f"Checks!$F${r}")
        _ok_format(ws, f"E{first}:E{r}")
        _end(ws, r + 2, m.last)


# =========================================================================================
class Summary:
    """Panel 1 live outputs (E) next to Panel 2 stored scenario results (G→, macro), Panel 3 checks."""

    def __init__(self, model, ws):
        self.m, self.ws = model, ws
        self.outputs: list[tuple] = []
        n = len(model.scenarios)
        _columns(ws, {"A": 3, "B": 3, "C": 40, "D": 10, "E": 12, "F": 2, **{L(7 + i): 12 for i in range(n)}})
        ws.freeze_panes = "E6"

    def add_output(self, key, label, unit, ref, scale=None, fmt=None, bold=False):
        """Live output `=ref[/scale]` named `Out_<key>`; the stored-results row is written by RunAllScenarios."""
        self.outputs.append((key, label, unit, ref, scale, fmt, bold))

    def write_outputs(self):
        ws, m = self.ws, self.m
        n, last = len(m.scenarios), 6 + len(m.scenarios)
        _bar(ws, 4, "Headline outputs – live case and stored scenario results", last)
        apply(ws.cell(5, 5, "Live"), "header")
        for i in range(n):
            _write(ws, 5, 7 + i, f"=Inputs!{L(8 + i)}$7", "header")
        r = 6
        for key, label, unit, ref, scale, fmt, bold in self.outputs:
            fmt = _fmt(unit, fmt)
            ws.cell(r, 3, label)
            apply(ws.cell(r, 4, unit), "unit")
            ref = m.address(ref) if not ref.startswith("=") else ref[1:]
            f = f"={ref}/{scale}" if scale else f"={ref}"
            _write(ws, r, 5, f, "total" if bold else "calc", fmt)
            if bold:
                ws.cell(r, 3).font = Font(name=styles.FONT, size=styles.SIZE, bold=True)
            for i in range(n):
                apply(ws.cell(r, 7 + i), "macro", fmt)
            m.define(f"Out_{key}", f"Summary!$E${r}")
            r += 1
        if self.outputs:
            m.define("Live_Results", f"Summary!$E$6:$E${r - 1}")
            m.define("Stored_Results", f"Summary!$G$6:${L(last)}${r - 1}")
            m.checks.add_check(*STORED_CHECK, per_period=False)
        self.r = r

    def write_checks(self):
        ws, m = self.ws, self.m
        last, r = 6 + len(m.scenarios), self.r + 1
        _bar(ws, r, "Checks", last)
        r += 1
        ws.cell(r, 3, "Master check")
        ws.cell(r, 3).font = Font(name=styles.FONT, size=styles.SIZE, bold=True)
        _write(ws, r, 5, _status("Master_Check"), "total")
        _ok_format(ws, f"E{r}:E{r + len(m.checks.rows)}")
        for label, cr in m.checks.rows:
            r += 1
            ws.cell(r, 3, label)
            _write(ws, r, 5, f"=Checks!$E${cr}")
        _end(ws, r + 2, last)


# =========================================================================================
class Dashboard:
    """Front sheet of the dashboard layout: header, deal levers, per-bid results and a right-hand panel.

    Columns: C label | D unit | E.. one value column per bid | grid (per-item info, stored scenario results)
    | Source | panel (label, unit, live value). Rows 3-8 header, row 9 column headers, content from row 11."""

    TOP = 11

    def __init__(self, model, ws, project, purpose, version, date, extra_cols, source_col=True):
        self.m, self.ws = model, ws
        self.nb = len(model.calc_sheets)
        self.nx = max(extra_cols, len(model.scenarios))     # grid also holds the stored scenario results
        self.V = 5                                           # first value column (E)
        self.X = self.V + self.nb + 1                        # first grid column
        self.SRC = self.X + self.nx + 1 if source_col else None   # source column (optional)
        self.P = (self.SRC or self.X + self.nx - 1) + 2           # panel: label, unit, value
        self.left, self.last = self.SRC or self.X + self.nx - 1, self.P + 2
        self.r = self.pr = self.TOP
        self.rows: dict[str, int] = {}
        self.outs: list[tuple] = []                          # (out key, name, label, unit, fmt, formula)
        self.selector_cell = None
        _columns(ws, {"A": 3, "B": 3, "C": 44, "D": 10, **{L(self.V + i): 12 for i in range(self.nb)},
                      L(self.X - 1): 2, **{L(self.X + i): 12 for i in range(self.nx)}, L(self.P - 1): 3,
                      L(self.P): 34, L(self.P + 1): 10, L(self.P + 2): 12,
                      **({L(self.SRC - 1): 2, L(self.SRC): 28} if self.SRC else {})})
        ws["B2"] = '=C3&" – Dashboard"'
        _banner(ws, 2, self.last)
        apply(ws.cell(3, 3, project), "input")
        ws["C3"].font = Font(name=styles.FONT, size=16, bold=True, color=styles.BLUE)
        ws["C3"].alignment = styles.Alignment(horizontal="left")
        ws.cell(4, 3, purpose)
        for r, label, value, fmt in ((5, "Version", version, "text"), (6, "Model date", date, "date")):
            ws.cell(r, 3, label)
            apply(ws.cell(r, self.V, value), "input", fmt)
        ws.cell(7, 3, "Key-assumption set")
        ws.cell(8, 3, "Master check")
        for c, h in ((3, "Item"), (4, "Unit"), (self.SRC, "Source"), (self.P, "Key assumptions"),
                     (self.P + 1, "Unit"), (self.P + 2, "Live")):
            if c is None:
                continue
            apply(ws.cell(9, c, h), "header")
            if c in (3, 4, self.SRC, self.P, self.P + 1):
                ws.cell(9, c).alignment = styles.Alignment(horizontal="left")
        for i, b in enumerate(model.bids):
            apply(ws.cell(9, self.V + i, b), "header")
        model.define("Bid_Labels", f"Dashboard!${L(self.V)}$9:${L(self.V + self.nb - 1)}$9")
        ws.freeze_panes = "A10"

    # ---- layout ---------------------------------------------------------------------------
    def section(self, title, headers=()):
        """Section bar with the bid labels over the value columns and `headers` over the grid columns."""
        if self.r > self.TOP:
            self.r += 1
        self._bar(title, "section", headers)

    def subheader(self, title, headers=()):
        self._bar(title, "subheader", headers)
        self.ws.cell(self.r - 1, 1).value = None

    def blank(self):
        self.r += 1

    def _bar(self, title, kind, headers, values=None):
        ws, r = self.ws, self.r
        if len(headers) > self.nx:
            raise ValueError(f"{len(headers)} grid headers, only {self.nx} grid columns (extra_cols)")
        _bar(ws, r, title, self.left, kind)
        for c, h in [(self.V + i, b) for i, b in enumerate(values or self.m.bids)] + \
                    [(self.X + i, h) for i, h in enumerate(headers)]:
            ws.cell(r, c, h).alignment = styles.RIGHT
        self.r += 1

    def _row(self, key, label, unit, source=None, extra=(), bold=False):
        ws, r = self.ws, self.r
        if key is not None:
            if key in self.rows:
                raise ValueError(f"duplicate Dashboard row {key}")
            self.rows[key] = r
        ws.cell(r, 3, label)
        apply(ws.cell(r, 4, unit), "unit")
        if source and self.SRC:
            ws.cell(r, self.SRC, source)
        if bold:
            ws.cell(r, 3).font = Font(name=styles.FONT, size=styles.SIZE, bold=True)
        if len(extra) > self.nx:
            raise ValueError(f"{len(extra)} grid values, only {self.nx} grid columns (extra_cols)")
        for i, v in enumerate(extra):
            if isinstance(v, str) and v.startswith("="):
                _write(ws, r, self.X + i, v)
            elif v is not None:
                ws.cell(r, self.X + i, v)
        self.r += 1
        return r

    def _values(self, values):
        vals = list(values) if isinstance(values, (list, tuple)) else [values] * self.nb
        if len(vals) != self.nb:
            raise ValueError(f"{self.nb} values needed (one per bid)")
        return vals

    def _span(self, r, first=None, n=None):
        first = first or self.V
        return f"Dashboard!${L(first)}${r}:${L(first + (n or self.nb) - 1)}${r}"

    def ref(self, key, bid=1, grid=None) -> str:
        """Address of a row's value for `bid` (1-based), or of grid column `grid` (1-based)."""
        c = self.X + grid - 1 if grid else self.V + bid - 1
        return f"Dashboard!${L(c)}${self.rows[key]}"

    # ---- rows -----------------------------------------------------------------------------
    def input(self, key, label, unit, values, *, source="", fmt=None, options=None, extra=()) -> str:
        """Deal lever: one input per bid (blue on yellow), named `Dash_<key>` over the 1 x n range.

        `options` adds list validation (a list of values or a workbook name). Calc reads it with
        `m.calc.lever("Dash_<key>")`."""
        fmt = _fmt(unit, fmt)
        r = self._row(key, label, unit, source, extra)
        for i, v in enumerate(self._values(values)):
            apply(self.ws.cell(r, self.V + i, v), "input", fmt)
        if options:
            _validate(self.ws, f"{L(self.V)}{r}:{L(self.V + self.nb - 1)}{r}", options)
        return self.m.define(f"Dash_{key}", self._span(r))

    def link(self, label, unit, ref, *, key=None, scale=None, fmt=None, out=None, bold=False, extra=()):
        """A calc result per bid. `ref` is a row template (`"{irr.E}"`, `"-{capex.I}"`) rendered once per calc
        sheet; `scale` divides it (e.g. "Thousand" for EUR m). `out` names the cells `Out_<out>_<bid>`
        (`Out_<out>` with one bid). Returns the Dashboard row."""
        fmt = _fmt(unit, fmt)
        r = self._row(key, label, unit, None, extra, bold)
        t = ref[1:] if ref.startswith("=") else ref
        for i, (sheet, bid) in enumerate(zip(self.m.calc_sheets, self.m.bids)):
            x = self.m.render(t, None, None, "Dashboard", sheet)
            if scale:
                x = f"{x}/{scale}" if re.fullmatch(r"[\w$!.]+", x) else f"({x})/{scale}"
            f = "=" + x
            _write(self.ws, r, self.V + i, f, "total" if bold else "calc", fmt)
            if out:
                name = f"Out_{out}_{i + 1}" if self.nb > 1 else f"Out_{out}"
                self.m.define(name, f"Dashboard!${L(self.V + i)}${r}")
                self.outs.append((out, name, f"{label} – {bid}" if self.nb > 1 else label, unit, fmt, f))
        return r

    def status(self, label, check=None, *, ref=None, key=None, extra=()):
        """OK/CHECK cells: `check` is a calc template that is 0 when OK (one cell per bid), or `ref` a single
        workbook name / cell (first value column)."""
        r = self._row(key, label, "flag", None, extra)
        if check is not None:
            t = check[1:] if check.startswith("=") else check
            for i, sheet in enumerate(self.m.calc_sheets):
                _write(self.ws, r, self.V + i, _status(self.m.render(t, None, None, "Dashboard", sheet)))
        else:
            _write(self.ws, r, self.V, _status(ref))
        _ok_format(self.ws, f"{L(self.V)}{r}:{L(self.V + self.nb - 1)}{r}")
        return r

    def stored_row(self, key, label, unit, fmt=None, extra=()) -> str:
        """Purple macro-written cells, one per bid, named `<key>_Stored`."""
        r = self._row(key, label, unit, None, extra)
        for i in range(self.nb):
            apply(self.ws.cell(r, self.V + i), "macro", _fmt(unit, fmt))
        return self.m.define(f"{key}_Stored", self._span(r))

    def stored_block(self, name, first, last, cols, fmt="general") -> str:
        """Purple macro-written block over grid columns `cols` (1-based, contiguous) of Dashboard rows
        `first`..`last` (row keys). Named `<name>_Stored`."""
        cols = list(cols)
        if cols != list(range(cols[0], cols[-1] + 1)) or cols[-1] > self.nx:
            raise ValueError("cols must be contiguous grid columns")
        r1, r2 = self.rows[first], self.rows[last]
        for r in range(r1, r2 + 1):
            for c in cols:
                apply(self.ws.cell(r, self.X + c - 1), "macro", fmt)
        c1, c2 = L(self.X + cols[0] - 1), L(self.X + cols[-1] - 1)
        return self.m.define(f"{name}_Stored", f"Dashboard!${c1}${r1}:${c2}${r2}")

    def scenario_results(self, title="Stored results by key-assumption set", outs=None):
        """Live (E) and stored (grid, purple) value of the `out=` links (all, or the `outs` keys), one row per
        output and bid, written by RunAllScenarios. Defines `Live_Results`, `Stored_Results`, `<Out key>_Stored`
        and the stored = live check. Leave out results that can be text ("n/a"): the check needs numbers."""
        ws, m, n = self.ws, self.m, len(self.m.scenarios)
        rows = [o[1:] for o in self.outs if outs is None or o[0] in outs]
        if not rows:
            raise ValueError("no outputs: use link(..., out=...) first")
        self.section(title, [f"=Inputs!{L(8 + i)}$7" for i in range(n)])
        ws.cell(self.r - 1, self.V, "Live")
        for i in range(1, self.nb):
            ws.cell(self.r - 1, self.V + i).value = None
        first = self.r
        for name, label, unit, fmt, formula in rows:
            r = self._row(None, label, unit)
            _write(ws, r, self.V, formula, "calc", fmt)        # same link as the output: no chain
            for i in range(n):
                apply(ws.cell(r, self.X + i), "macro", fmt)
            m.define(f"{name[4:]}_Stored", self._span(r, self.X, n))
        last = self.r - 1
        m.define("Live_Results", f"Dashboard!${L(self.V)}${first}:${L(self.V)}${last}")
        m.define("Stored_Results", f"Dashboard!${L(self.X)}${first}:${L(self.X + n - 1)}${last}")
        m.checks.add_check(*STORED_CHECK, per_period=False)

    # ---- right-hand panel ---------------------------------------------------------------
    def panel_section(self, title):
        if self.pr > self.TOP:
            self.pr += 1
        for c in range(self.P, self.P + 3):
            apply(self.ws.cell(self.pr, c), "section")
        self.ws.cell(self.pr, self.P, title)
        self.pr += 1

    def panel_link(self, label, unit, ref, fmt=None):
        """Live value from Inputs (a workbook name such as `In_Premium`, or `Sheet!$A$1`)."""
        ws, r = self.ws, self.pr
        ws.cell(r, self.P, label)
        apply(ws.cell(r, self.P + 1, unit), "unit")
        _write(ws, r, self.P + 2, f"={ref}", fmt=_fmt(unit, fmt))
        self.pr += 1
        return r

    def selector(self, label="Key-assumption set"):
        """Put the `Scenario_Name` input (list of `Scenario_Names`) in the panel; the header then links to it.
        Without this call the selector sits in the header."""
        ws, r = self.ws, self.pr
        ws.cell(r, self.P, label)
        apply(ws.cell(r, self.P + 1, "text"), "unit")
        self.selector_cell = (r, self.P + 2)
        self.pr += 1
        return "Scenario_Name"

    # ---- write ----------------------------------------------------------------------------
    def finish(self):
        ws, m = self.ws, self.m
        r, c = self.selector_cell or (7, self.V)
        apply(ws.cell(r, c, m.scenarios[0]), "input", "text")
        ws.cell(r, c).alignment = styles.Alignment(horizontal="left")
        _validate(ws, f"{L(c)}{r}", "Scenario_Names")
        m.define("Scenario_Name", f"Dashboard!${L(c)}${r}")
        if self.selector_cell:
            _write(ws, 7, self.V, "=Scenario_Name").alignment = styles.Alignment(horizontal="left")
        _write(ws, 8, self.V, _status("Master_Check")).alignment = styles.Alignment(horizontal="left")
        _ok_format(ws, f"{L(self.V)}8")
        _end(ws, max(self.r, self.pr) + 1, self.last)
