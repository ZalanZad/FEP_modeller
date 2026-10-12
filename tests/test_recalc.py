import datetime as dt

import openpyxl

from fepm.recalc import recalc, values


def xirr(cfs, dates):
    def npv(r):
        return sum(c / (1 + r) ** ((d - dates[0]).days / 365) for c, d in zip(cfs, dates))
    lo, hi = -0.99, 10.0
    for _ in range(300):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if npv(mid) > 0 else (lo, mid)
    return (lo + hi) / 2


def test_recalc_keeps_formulas_and_computes_xirr(tmp_path):
    dates = [dt.date(2026, 12, 31), dt.date(2027, 12, 31), dt.date(2028, 12, 31), dt.date(2029, 12, 31)]
    cfs = [-100, 30, 45, 60]
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, (d, c) in enumerate(zip(dates, cfs), start=2):
        ws.cell(1, i, d)
        ws.cell(2, i, c)
    ws["A3"] = "=XIRR(B2:E2,B1:E1)"
    ws["A4"] = "=XNPV(0.08,B2:E2,B1:E1)"
    ws["A5"] = "=SUM(B2:E2)"
    src = tmp_path / "x.xlsx"
    wb.save(src)                                   # openpyxl writes no cached values
    out = recalc(src, tmp_path / "x_recalc.xlsx")
    v = values(out).active
    assert abs(v["A3"].value - xirr(cfs, dates)) < 1e-9
    xnpv = sum(c / 1.08 ** ((d - dates[0]).days / 365) for c, d in zip(cfs, dates))
    assert abs(v["A4"].value - xnpv) < 1e-9
    assert v["A5"].value == 35
    assert openpyxl.load_workbook(out).active["A3"].value == "=XIRR(B2:E2,B1:E1)"


def test_recalc_with_inputs(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"], ws["A2"] = 2, "=A1*3"
    wb.save(tmp_path / "y.xlsx")
    out = recalc(tmp_path / "y.xlsx", tmp_path / "y2.xlsx", inputs={"S!A1": 5})
    assert values(out)["S"]["A2"].value == 15
