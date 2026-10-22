"""Dashboard layout: Dashboard | Inputs | Calc_1 | Calc_2 | Checks, built, macro-run and recalculated in LibreOffice."""
import datetime as dt
import shutil

import openpyxl
import pytest

from fepm.audit.lint import lint, lint_values
from fepm.recalc import name_value, recalc, values
from fepm.xl.builder import Model
from fepm.xl.macros import VBA_DIR, _calls, embedded_modules, package

TOL = 1e-9
VOLUMES = [1000, 2000]
STORE_BAS = '''Attribute VB_Name = "TestStore"
Option Explicit
Public Sub StoreBase()
    Dim i As Long
    Application.Calculate
    For i = 1 To Range("Base_Revenue_Stored").Columns.Count
        Range("Base_Revenue_Stored").Cells(1, i).Value = Range("Out_Revenue_" & i).Value
    Next i
End Sub
'''


def build_twin(path):
    m = Model("Twin", "Two-bid test", periods=4, start=dt.date(2026, 1, 1), scenarios=["Base", "High price"],
              layout="dashboard", calc_sheets=["Calc_1", "Calc_2"], bids=["BID 1", "BID 2"], extra_cols=1)
    i, d, c, k = m.inputs, m.dashboard, m.calc, m.checks
    price = i.add_input("Price", "Power price", "EUR/MWh", 50, {2: 60}, source="test")
    capex = i.add_input("Capex", "Capex", "EUR k", 300, source="test")
    rate = i.add_input("Fee_Rate", "Fee rate on gross revenue", "%", 0.2, source="test")
    d.section("Deal levers", headers=["Range"])
    vol = d.input("Volume", "Annual volume", "MWh", VOLUMES, source="test", extra=["0–5,000"])
    d.input("Form", "Security form", "text", ["cash", "guarantee"], source="test", options=["cash", "guarantee"])
    c.section("Revenue")
    c.row("vol", "Volume", "MWh", "={self.E}*{yf}", e="=" + c.lever(vol), ef_fmt="money")
    c.row("rev", "Revenue", "EUR k", f"={{vol}}*{price}/Thousand")
    c.paste("Fee", "Fee", "EUR k")
    c.row("gross", "Gross revenue", "EUR k", "={rev}+{fee_paste}")
    c.copy("Fee", "Fee", "EUR k", f"={{gross}}*{rate}")
    c.row("capex", "Capex", "EUR k", f"=-{capex}*N({{period}}=1)")
    c.row("cf", "Cash flow", "EUR k", "={gross}+{capex}", bold=True)
    c.xirr("irr", "IRR", "cf", iferror=True)
    k.add_check("Volume not negative", "=N({vol}<0)")
    k.add_check("Cash flow adds up", "=N(ABS({cf.I}-{gross.I}-{capex.I})>Tolerance)", per_period=False)
    d.section("Results")
    d.link("Revenue", "EUR m", "{rev.I}", key="rev", scale="Thousand", out="Revenue")
    d.link("IRR", "%", "{irr.E}", out="IRR", bold=True)
    d.status("Volume valid", "N({vol.E}<0)")
    d.stored_row("Base_Revenue", "Revenue, base set (macro)", "EUR m")
    span = f"{d.ref('rev', 1)}:{d.ref('rev', 2).split('!')[1]}"
    k.add_check("Stored base revenue = live", f"=--(SUMPRODUCT(ABS(Base_Revenue_Stored-{span}))>Tolerance)",
                per_period=False)
    d.scenario_results(outs=["Revenue", "IRR"])
    d.panel_section("Key assumptions")
    d.selector()
    d.panel_link("Power price", "EUR/MWh", price)
    return m, m.save(path)


@pytest.fixture(scope="module")
def twin(tmp_path_factory):
    out = tmp_path_factory.mktemp("twin")
    m, xlsx = build_twin(out / "twin.xlsx")
    vba = out / "vba"
    vba.mkdir()
    for f in VBA_DIR.glob("*.bas"):
        shutil.copy(f, vba)
    (vba / "TestStore.bas").write_text(STORE_BAS)
    xlsm = package(xlsx, out / "twin.xlsm", run=["TestStore.StoreBase", ("Scenarios", "RunAllScenarios")],
                   vba_dir=vba)
    rc = recalc(xlsm, out / "twin_recalc.xlsx")
    return {"m": m, "xlsx": xlsx, "xlsm": xlsm, "rc": rc, "out": out}


def test_layout_and_lint(twin):
    wb = openpyxl.load_workbook(twin["xlsx"])
    assert wb.sheetnames == ["Dashboard", "Inputs", "Calc_1", "Calc_2", "Checks"]
    assert (wb["Calc_1"]["A1"].value, wb["Calc_2"]["A1"].value) == (1, 2)
    assert wb["Calc_2"]["B2"].value == '=Dashboard!$C$3&" – Calculations – "&INDEX(Bid_Labels,1,$A$1)'
    assert wb["Inputs"]["F4"].value == "=Scenario_Name"
    assert lint(twin["xlsx"]) == []
    assert lint(twin["xlsm"]) == []


def test_calc_sheets_have_identical_formula_text(twin):
    wb = openpyxl.load_workbook(twin["xlsx"])
    a, b = wb["Calc_1"], wb["Calc_2"]
    cells = lambda ws: {c.coordinate: c.value for r in ws.iter_rows() for c in r  # noqa: E731
                        if c.value is not None and c.coordinate != "A1"}
    assert cells(a) == cells(b)
    assert any(str(v).startswith("=INDEX(Dash_Volume,1,$A$1)") for v in cells(a).values())


def test_checks_one_row_per_calc_sheet(twin):
    labels = [label for label, _ in twin["m"].checks.rows]
    assert "Volume not negative – Calc_1" in labels and "Volume not negative – Calc_2" in labels
    assert "Solver Fee_1: delta = 0" in labels and "Solver Fee_2: delta = 0" in labels
    assert labels.count("Stored results = live case") == 1


def test_macros_run_in_order_and_master_check(twin):
    assert {"Scenarios", "Solver", "GoalSeek", "TestStore"} <= set(embedded_modules(twin["xlsm"]))
    assert lint_values(twin["rc"]) == []                          # Master_Check = 0, no errors
    wb = values(twin["rc"])
    v = lambda n: name_value(wb, n)  # noqa: E731
    stored = [wb["Dashboard"].cell(int(next(wb.defined_names["Base_Revenue_Stored"].destinations)[1]
                                        .split("$")[2].split(":")[0]), 5 + i).value for i in range(2)]
    assert stored == [v("Out_Revenue_1"), v("Out_Revenue_2")]    # TestStore.StoreBase ran
    sheet, rng = next(wb.defined_names["Revenue_1_Stored"].destinations)
    r = int(rng.split("$")[2].split(":")[0])
    base, high = wb[sheet].cell(r, 8).value, wb[sheet].cell(r, 9).value
    assert abs(base - v("Out_Revenue_1")) < TOL and abs(high / base - 60 / 50) < TOL    # RunAllScenarios ran
    assert v("Fee_1_Delta") < 1e-6 and v("Fee_2_Delta") < 1e-6     # Solver closed both sheets


def test_dash_input_drives_its_bid(twin):
    m = twin["m"]
    wb = values(twin["rc"])
    yf = [wb["Calc_1"].cell(6, c).value for c in range(10, 14)]
    for n, vol in enumerate(VOLUMES, 1):
        assert abs(name_value(wb, f"Out_Revenue_{n}") - sum(yf) * vol * 50 / 1000 / 1000) < TOL
    rc2 = recalc(twin["xlsx"], twin["out"] / "bump.xlsx", inputs={m.dashboard.ref("Volume", 2): 3000})
    wb2 = values(rc2)
    r = m.calc.rows["vol"].r
    assert wb2["Calc_1"].cell(r, 5).value == 1000 and wb2["Calc_2"].cell(r, 5).value == 3000
    assert abs(name_value(wb2, "Out_Revenue_2") / name_value(wb2, "Out_Revenue_1") - 3) < TOL


def test_lint_dashboard_rules(twin, tmp_path):
    path = tmp_path / "planted.xlsx"
    wb = openpyxl.load_workbook(twin["xlsx"])
    m = twin["m"]
    r = m.calc.rows["rev"].r
    wb["Dashboard"]["C40"] = 7                                   # number outside a Dash_ / _Stored name
    wb["Calc_2"].cell(r, 11).value = wb["Calc_2"].cell(r, 11).value + "+0"      # Calc_2 drifts from Calc_1
    wb["Calc_2"]["A1"] = 3
    wb.save(path)
    f = lint(path)
    where = {(x.rule, x.where) for x in f}
    assert ("a-hardcode", "Dashboard!C40") in where
    assert ("c-sheet", f"Calc_2!row {r}") in where
    assert ("a-bid", "Calc_2!A1") in where
    assert not any(x.where.startswith(("Dashboard!E", "Dashboard!F")) for x in f)   # deal levers allowed


def test_run_list_forms():
    assert _calls(("Scenarios", "RunAllScenarios")) == [("Scenarios", "RunAllScenarios")]
    assert _calls("Scenarios.RunAllScenarios") == [("Scenarios", "RunAllScenarios")]
    assert _calls(["A.B", ("C", "D")]) == [("A", "B"), ("C", "D")]
    assert _calls(None) == []


def test_section_values_and_across(tmp_path):
    """`section(values=)` relabels the value columns; `across()` sums one calc cell over every bid sheet."""
    m, _ = build_twin(tmp_path / "twin.xlsx")
    d = m.dashboard
    d.section("All bids together", values=["Total"])
    r = d.r - 1
    assert d.ws.cell(r, d.V).value == "Total" and d.ws.cell(r, d.V + 1).value is None
    rev = m.calc.rows["rev"].r
    assert d.across("{rev.I}") == f"Calc_1!$I{rev}+Calc_2!$I{rev}"
    assert d.across("={capex.I}") == d.across("{capex.I}")
    d.section("Per bid")
    assert [d.ws.cell(d.r - 1, d.V + i).value for i in range(2)] == ["BID 1", "BID 2"]
