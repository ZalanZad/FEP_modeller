import datetime as dt

import openpyxl
import pytest

from fepm.audit.lint import lint, to_r1c1
from fepm.xl.builder import Model


@pytest.fixture
def tiny(tmp_path):
    m = Model("Tiny", periods=5, start=dt.date(2026, 1, 1), scenarios=["Base", "Low"])
    price = m.inputs.add_input("Price", "Price", "EUR/MWh", 50, {2: 40}, source="test")
    m.calc.section("Revenue")
    m.calc.row("vol", "Volume", "MWh", "=Thousand*{yf}")
    m.calc.row("rev", "Revenue", "EUR k", f"={{vol}}*{price}/Thousand")
    m.calc.corkscrew("cash", "Cash", "EUR k", [("cash_in", "Cash in", "={rev}")])
    m.calc.xirr("irr", "IRR", "rev", iferror=True)
    m.checks.add_check("Volume not negative", "=N({vol}<0)")
    m.summary.add_output("Rev", "Revenue", "EUR m", m.calc.ref("rev"), scale="Thousand")
    path = m.save(tmp_path / "tiny.xlsx")
    return path, m


def _rules(findings):
    return {f.rule for f in findings}


def test_clean_model(tiny):
    assert lint(tiny[0]) == []


def test_catches_planted_hardcode(tiny):
    path, m = tiny
    wb = openpyxl.load_workbook(path)
    r = m.calc.rows["rev"].r
    wb["Calc"].cell(r, 12).value = 123          # L: a number pasted over a formula
    wb.save(path)
    f = lint(path)
    assert "a-hardcode" in _rules(f)
    assert any(x.rule == "a-hardcode" and x.where == f"Calc!L{r}" for x in f)


def test_catches_inconsistent_row(tiny):
    path, m = tiny
    wb = openpyxl.load_workbook(path)
    r = m.calc.rows["vol"].r
    wb["Calc"].cell(r, 11).value = f"=Thousand*K6+J{r}"
    wb.save(path)
    assert "c-row" in _rules(lint(path))


def test_catches_literal_direction_and_forbidden(tiny):
    path, m = tiny
    wb = openpyxl.load_workbook(path)
    ws = wb["Calc"]
    r = m.calc.rows["vol"].r
    for c in range(10, 15):
        col = openpyxl.utils.get_column_letter(c)
        ws.cell(r, c).value = f"={col}{r + 1}*4+OFFSET({col}6,0,0)"
    ws.merge_cells("C1:D1")
    wb.save(path)
    rules = _rules(lint(path))
    assert {"b-literal", "d-direction", "e-function", "e-merged"} <= rules


def test_r1c1_relative():
    assert to_r1c1("=J5+$E5*J$4", 5, 10) == to_r1c1("=K5+$E5*K$4", 5, 11)
    assert to_r1c1("=J5", 5, 10) != to_r1c1("=J5", 5, 11)
