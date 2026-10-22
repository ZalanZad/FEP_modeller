"""Zalahaláp (ZHLP) – 63 MWp PV + 50 MW / 200 MWh BESS: investment model. Built to SPEC.md v1 (+ user decisions of
22-Oct-2026: share 60 % at exactly 90 k€/MW; early works provisional).

Dashboard | Inputs | Curves | Calc_1 | Calc_2 | Calc_3 | Checks, quarterly from 01-Oct-2025 to 30-Jun-2064.
Each Calc sheet runs one revenue case picked on the Dashboard (default A = reconstruction of the original model
Zalahalap_model_10112025_v3, B = Aurora merchant, the decision case, C = offtake offer + Aurora merchant tail) through
the same rows: project (PV, BESS, opex, capex, debt, tax), buyer view and implied prices, then the JV/FEP layer for the
three strategies (sale at RtB, sale at COD, hold to maturity). Solver.bas closes the tax circularity, GoalSeek.bas
solves the asset-deal implied prices, Scenarios.bas stores the 18 key-assumption sets.
"""
import csv
import datetime as dt
from pathlib import Path

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L

from fepm.xl.builder import Model
from fepm.xl.macros import package
from fepm.xl.styles import FMT, RIGHT, apply

HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
OLD = DATA / "old_engine"
NAME = "Zalahalap_Model"
D = dt.date
START, PERIODS = D(2025, 10, 1), 155                       # 01-Oct-2025 → 30-Jun-2064
YEARS = list(range(2025, 2065))
SETS = ["Base", "Reconciliation (v3 rules)", "PV price −10%", "PV price +10%", "PV yield −10%", "BESS revenue −20%",
        "Capex +10%", "Opex +10%", "Buyer IRR +1 pp", "RtB and COD price −10%", "Gearing −10 pp", "Debt rate +1 pp",
        "COD +12 months", "BESS 2 h", "No BESS", "PV and BESS prices −20%", "C: E.ON PPA + option B",
        "C: MVM 45 + MET floor"]
BLOCK = "2027 solar DC 1.0cy 4h"                            # default Aurora battery block (SPEC §15 item 5)
BLOCK_2H = "2027 solar DC 1.0cy 2h"
CAPTURE = "Aurora implied capture – "                        # PV curve names: CAPTURE + block
PH = "placeholder – pending user (SPEC §16 Q{})"
EW_NOTE = "provisional – early works to be revisited (user 22-Oct-26)"
EXIT = 'SUMIF({start.all},"<="&{exit_ev.E},{self.all})'    # amount to the end of the exit period
TO_COD = 'SUMIF({end.all},"<="&{cod_ev.E},{self.all})'     # amount to COD
SPREAD = "=MAX(0,MIN({self.F},{end})-MAX({self.E},{start})+1)/({self.F}-{self.E}+1)"   # amount spread over E → F
WINDOW = "=MAX(0,MIN({self.F},{end})-MAX({self.E},{start})+1)/{days}"                 # share of the period in E → F
PATCH = {}                                                  # (calc key, column) -> number format


def fmt(key, col, f):
    PATCH[(key, col)] = f


# =========================================================================================== inputs
def inputs(m):
    i = m.inputs
    a = i.add_input
    i.add_constant("Sens_Step", "Sensitivity step (price, yield, capex, opex sets)", "%", 0.1, "SPEC §10")
    i.add_constant("Sens_Step_Large", "Large sensitivity step (BESS revenue, downside sets)", "%", 0.2, "SPEC §10")
    i.add_constant("Sens_PP", "Rate step (buyer IRR, debt rate sets)", "%", 0.01, "SPEC §10")
    i.add_constant("Hours_In_Day", "Hours in a day", "h", 24, "Grid export limit")
    i.add_constant("Recon_Set", "Key-assumption set of the reconciliation", "#", 2, "SPEC §10 set 2")
    a("Acq", "Acquisition, JV start, deposits paid", "date", D(2025, 11, 30), source="ZHLP v3 Stat_inp E39")
    a("RtB", "Ready to build (cases B, C)", "date", D(2026, 12, 31), source="Teaser p6; " + PH.format(5))
    a("RtB_Exit", "RtB sale closing (exit), cases B, C", "date", D(2027, 3, 31), source=PH.format(5))
    a("NTP", "Notice to proceed (cases B, C)", "date", D(2027, 6, 30), source=PH.format(5))
    a("COD", "Commercial operation, PV and BESS (cases B, C)", "date", D(2028, 3, 31),
      source="Teaser p6; " + PH.format(5))
    a("COD_Delay", "COD delay (construction longer)", "months", 0, {13: 12}, source="SPEC §10 set 13", fmt="int")
    a("Life", "Operating life", "years", 35, source="ZHLP v3 Control D30", fmt="int")
    a("PreRtB_First", "Pre-RtB development costs: first month", "date", D(2026, 10, 31),
      source="ZHLP v3 Control D116")
    a("PreRtB_Last", "Pre-RtB development costs: last month", "date", D(2027, 3, 31), source="ZHLP v3 Control D117")

    i.section("PV and grid (cases B, C; case A on Curves)")
    a("MWp", "PV capacity", "MWp", 63, source="Teaser p1", fmt="money1")
    a("Grid_MVA", "Grid export capacity", "MVA", 49.9, source="Teaser p1", fmt="money1")
    a("Yield", "Specific yield", "MWh/MWp", 1265, source="Teaser p1 (v3: 1,250)", fmt="money")
    a("Yield_Adj", "Yield change (P90 proxy)", "%", 0, {5: "=-Sens_Step"}, source="SPEC §10 set 5")
    a("PV_Avail", "PV availability", "%", 0.995, source="ZHLP v3 Control D56")
    a("Degradation", "PV degradation per year", "%", 0.003, source="ZHLP v3 Control D57", fmt="0.00%")
    a("Clip_Loss", "Clipping at the grid limit, share of DC energy", "%", 0.02, source=PH.format(13))

    i.section("BESS")
    a("BESS_On", "BESS built (1 = yes)", "flag", 1, {15: 0}, source="Teaser p1; SPEC §10 set 15")
    a("BESS_MW", "BESS power", "MW", 50, source="Teaser p1; ZHLP v3 Control D91", fmt="money1")
    a("BESS_Hours", "BESS duration", "h", 4, {14: 2}, source="Teaser; " + PH.format(6), fmt="money1")
    a("BESS_Block", "Aurora battery block (Degraded batteries)", "text", BLOCK, {14: BLOCK_2H},
      source="Aurora Q2-26; " + PH.format(6))
    a("BESS_Avail", "BESS availability, project", "%", 0.98, source=PH.format(6))
    a("Aurora_Avail", "BESS availability inside Aurora's dispatch", "%", 0.98, source="Aurora Q2-26 Inputs C18")
    a("BESS_RTE", "Round-trip efficiency (energy check only)", "%", 0.86, source="Aurora Q2-26 Inputs C17")
    a("BESS_Haircut", "Haircut on merchant BESS revenue", "%", 0, {6: "=Sens_Step_Large", 16: "=Sens_Step_Large"},
      source="SPEC §6.2; sets 6, 16")
    a("BESS_Grid_Charge", "Grid charge on charging, offers that deduct it", "EUR k/MW", 0, source=PH.format(12))
    a("BESS_Capex", "BESS capex", "EUR k/MWh", 110, source="v3 Control D94; " + PH.format(6))
    a("Augment_Cost", "BESS augmentation cost", "EUR k/MWh", 110, source="v3 Control D100; " + PH.format(6))
    a("Augment_Year", "Augmentation year (0 = Aurora block repower year)", "year", 0, source="SPEC §6.2", fmt="int")
    a("Augment_Year_v3", "Augmentation year, v3 rule (fix 7 off)", "year", 2048, source="ZHLP v3 Control D101",
      fmt="int")
    a("BESS_Opex", "BESS fixed opex", "EUR k/MW", 30.8, source="ZHLP v3 Control D96", fmt="money1")

    i.section("Prices and offers")
    a("PV_Curve", "PV price curve (Curves, real 2025)", "text", CAPTURE + BLOCK, {14: CAPTURE + BLOCK_2H},
      source="SPEC §0.2, §6.2 (user: Aurora file only)")
    a("PV_Price_Adj", "PV merchant price change", "%", 0, {3: "=-Sens_Step", 4: "=Sens_Step", 16: "=-Sens_Step_Large"},
      source="SPEC §10 sets 3, 4, 16")
    a("GoO", "Guarantee of origin price, real 2025", "EUR/MWh", 1.75, source="v3 Control D58; " + PH.format(12),
      fmt="money2")
    a("Offer_PV", "PV offer in case C (C1, C2a, C2b, C3)", "text", "C2a", {17: "C1", 18: "C3"},
      source="SPEC §6.3; " + PH.format(12))
    a("Offer_BESS", "BESS offer in case C (C1A, C1B, C2c, C2d, none)", "text", "C2c", {17: "C1B", 18: "C2d"},
      source="SPEC §6.3; " + PH.format(12))

    i.section("Capex")
    a("PV_Capex", "PV equipment and balance of plant", "EUR k/MWp", 350, source="ZHLP v3 Control D82:D83")
    a("Cable_Capex", "MV cable to the grid connection point", "EUR k", 3700, source="v3 Control D104; " + PH.format(13))
    a("Substation", "HV substation", "EUR k", 9500, source="v3 Control D105; " + PH.format(13))
    a("Land_Capex", "Land to COD (not depreciated)", "EUR k", 620, source="ZHLP v3 Control D106")
    a("Capex_Adj", "Capex change", "%", 0, {7: "=Sens_Step"}, source="SPEC §10 set 7")
    a("PreRtB_Devex", "Pre-RtB development costs", "EUR k", 3182, source="ZHLP v3 Control D115")
    a("PreRtB_Reimbursed", "Buyer reimburses pre-RtB costs (1 = yes)", "flag", 1, source="ZHLP v3 Control D118")

    i.section("Early works (provisional – to be revisited)")
    a("Milestones_Scope", "Early-works lines paid in development (4 or 6)", "#", 4, {2: 6}, source=EW_NOTE,
      fmt="int")
    a("Milestones_Late_To_Capex", "Lines outside the scope paid at NTP (1 = yes)", "flag", 1, source=EW_NOTE)
    a("EW_Reimbursed", "Buyer reimburses early works at the exit (1 = yes)", "flag", 1, source=EW_NOTE)
    a("EW_Interest", "Early works earn the FEP return (1 = yes)", "flag", 1, source=EW_NOTE)
    a("EW_To_PSH", "Early works are PSH milestone income (v3)", "flag", 0, {2: 1}, source="ZHLP v3 Calculation F221")

    i.section("Opex and levies (real 2025, Aurora euro index)")
    a("OM", "O&M and fixed opex", "EUR k/MWp", 14.96, source="ZHLP v3 Control D59", fmt="money2")
    a("Cable_OM", "MV cable O&M", "EUR k/MWp", 5, source="ZHLP v3 Control D62", fmt="money1")
    a("Balancing", "Balancing cost", "EUR/MWh", 5, source="ZHLP v3 Control D60", fmt="money1")
    a("Lease", "Land lease from COD", "EUR k", 96, source="ZHLP v3 Control D61")
    a("RTM", "Route-to-market fee on PV revenue", "%", 0, source="ZHLP v3 Control D77")
    a("Opex_Adj", "Opex change", "%", 0, {8: "=Sens_Step"}, source="SPEC §10 set 8")
    a("HIPA", "Local business tax (HIPA), on revenue", "%", 0.02, {2: 0}, source="SPEC §15 item 9 (Bana)")
    a("Innovation", "Innovation contribution, on revenue", "%", 0.003, {2: 0}, source="Act LXXVI of 2014, 15. §",
      fmt="0.00%")

    i.section("Deposits and guarantees")
    a("FX", "FX rate", "HUF/EUR", 400, source="ZHLP v3 Stat_inp E23 (Aurora 397.77)", fmt="money")
    a("BG_Base", "Bank guarantee base rate (1M EURIBOR)", "%", 0.0187, source="ZHLP v3 Dyn_Inp G50:Y50 average",
      fmt="0.00%")
    a("BG_Margin", "Bank guarantee margin", "%", 0.0175, source="ZHLP v3 Stat_inp E110", fmt="0.00%")

    i.section("JV, RtB exit and COD sale")
    a("RtB_Price", "RtB price (0 = implied at the buyer target IRR)", "EUR k/MWp", 90,
      source="v3 Stat_inp E24; " + PH.format(3))
    a("COD_Price", "COD price (0 = implied at the buyer target IRR)", "EUR k/MWp", 0, source=PH.format(8))
    a("Price_Step", "RtB and COD price change", "%", 0, {10: "=-Sens_Step"}, source="SPEC §10 set 10")
    a("Share_Rule", "Share rule (1 = step, 0 = v3 slide)", "flag", 1, {2: 0}, source="SPEC §0.1; set 2 = v3")
    a("FEP_Share_Low", "Futureal share at or below the threshold", "%", 0.6, source="User 22-Oct-26 (SPEC §0.1)")
    a("FEP_Share_High", "Futureal share above the threshold", "%", 0.5, source="User 22-Oct-26 (SPEC §0.1)")
    a("Share_Threshold", "Share threshold (60 % at or below)", "EUR k/MWp", 90, source="User 22-Oct-26: 60 % at 90")
    a("Share_Slide_End", "v3 slide: 50 % reached at", "EUR k/MWp", 100, source="ZHLP v3 Stat_inp J36")
    a("Transaction", "Transaction costs (FEP)", "EUR k", 200, source="ZHLP v3 Stat_inp E37")
    a("SPV_Price", "FEP SPV acquisition price", "EUR k", 549, {2: 150},
      source="JV sheet G19; " + PH.format(3))
    a("SPV_Interest", "SPV price earns the FEP return (1 = yes)", "flag", 0, {2: 1}, source="ZHLP v3 Stat_inp E66")
    a("DSA_Per_MWp", "DSA loan to PSH", "EUR k/MWp", 15, {2: 0}, source="JV sheet F55; " + PH.format(3))
    a("Loan_Rate", "FEP return on loans and costs", "%", 0.12, source="JV sheet G30; v3 Stat_inp E69")
    a("PSH_Land", "PSH land apport (in kind)", "EUR k", 150, source="ZHLP v3 Stat_inp E34")
    a("PSH_Land_Step", "PSH land and its return repaid before the split", "flag", 0, {2: 1},
      source="ZHLP v3 Calculation F155:F156")
    a("PSH_Devex", "PSH own development costs", "EUR k/MWp", 15, source="ZHLP v3 Stat_inp E79")
    a("Sale_Costs_RtB", "Sale costs at the RtB exit", "EUR k", 0, source=PH.format(7))
    a("Sale_Costs_COD", "Sale costs at COD", "EUR k", 0, source=PH.format(7))
    a("PSH_Stays", "HTM: PSH keeps its share (0 = FEP buys it out)", "flag", 0, source=PH.format(4))

    i.section("Deliberate fixes vs v3 (SPEC §13; 0 = v3 method)")
    for key, label, src in [("Fix_Reimb", "Fix 2: pre-RtB costs reimbursed as spent", "v3 RTB_Model H69"),
                            ("Fix_BG", "Fix 3: no bank-guarantee phantom", "v3 Calculation H118:H120"),
                            ("Fix_Clip", "Fix 5: PV export clipped at the grid", "v3 Control D16"),
                            ("Fix_BESS", "Fix 6: Aurora block total for case A", "v3 Control D97:D98"),
                            ("Fix_Augment", "Fix 7: augmentation indexed, at repower year", "v3 Control D101"),
                            ("Fix_Interest", "Fix 8: dated act/365 return, 31-Dec capitalised", "v3 Stat_inp E69")]:
        a(key, label, "flag", 1, {2: 0}, source=src)

    i.section("Valuation")
    a("Discount_Rate", "FEP discount rate (NPV)", "%", 0.08, source="ZHLP v3 Control D9")
    a("Buyer_IRR_RtB", "Buyer target IRR at RtB (unlevered, post-tax)", "%", 0.08, {9: "=H{r}+Sens_PP"},
      source="v3 Control D9; " + PH.format(8))
    a("Buyer_IRR_COD", "Buyer target IRR at COD (unlevered, post-tax)", "%", 0.08, {9: "=H{r}+Sens_PP"},
      source="v3 Control D9; " + PH.format(8))
    a("Sale_Form", "Deal form (0 = share deal, 1 = asset deal)", "flag", 0, source=PH.format(7))
    a("Sale_Gain_Taxed", "FEP gain on the sale taxed (0 = exempt)", "flag", 0, source=PH.format(7))

    i.section("Senior debt (Bana indications, not ZHLP terms)")
    a("Gearing", "Debt share of construction cost (cap)", "%", 0.55, {2: 0, 11: "=H{r}-Sens_Step"}, source=PH.format(9))
    a("Debt_Per_MW", "Debt per MWp (0 = gearing only)", "EUR k/MWp", 0, source=PH.format(9))
    a("DSCR_Contracted", "Sizing DSCR on contracted revenue", "x", 1.20, source=PH.format(9), fmt="mult")
    a("DSCR_Merchant", "Sizing DSCR on merchant and BESS revenue", "x", 1.35, source=PH.format(9), fmt="mult")
    a("Debt_Tenor", "Annuity tenor from COD", "years", 15, source=PH.format(9), fmt="int")
    a("Debt_Rate", "All-in fixed interest rate", "%", 0.05, {12: "=H{r}+Sens_PP"}, source=PH.format(9))
    a("Upfront_Fee", "Upfront fee on the facility", "%", 0.0175, source=PH.format(9), fmt="0.00%")
    a("Commitment_Fee", "Commitment fee on undrawn debt", "%", 0.009, source=PH.format(9), fmt="0.00%")
    a("DSRA_Months", "DSRA: months of debt service, funded at COD", "months", 6, source=PH.format(9), fmt="int")
    a("Sweep", "Cash sweep, share of cash after debt service", "%", 0, source=PH.format(9))

    i.section("Tax (Hungary)")
    a("CIT", "Corporate income tax", "%", 0.09, source="v3 Control D66; " + PH.format(10))
    a("RH_On", "Robin Hood tax applies (1 = yes)", "flag", 0, source=PH.format(10))
    a("RH_Rate", "Robin Hood tax on positive pre-tax profit", "%", 0.31, source="Thtv. (Bana)")
    a("Dep_Years", "Tax depreciation: PV, grid, development", "years", 30, source="ZHLP v3 Control D65", fmt="int")
    a("Dep_BESS_Years", "Tax depreciation: BESS and augmentation", "years", 10, {2: 30}, source=PH.format(10),
      fmt="int")
    a("Loss_Cap", "Tax losses usable, share of the tax base", "%", 0.5, source="ZHLP v3 Control D102")
    a("Loss_Years", "Tax losses lapse after NTP +", "years", 5, {2: 40}, source="Tao. 17. § (Bana); set 2: none",
      fmt="int")


def _fix_set_formulas(m):
    """`{r}` in a set override is the input's own row (resolved after the row is written)."""
    ws = m.inputs.ws
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            if isinstance(c.value, str) and "{r}" in c.value:
                c.value = c.value.replace("{r}", str(c.row))


# =========================================================================================== curves
def _old_curves():
    rows = {int(r["year"]): r for r in csv.DictReader(open(OLD / "old_curves.csv"))}

    def col(key):
        out = []
        for y in YEARS:
            try:
                out.append(float(rows[y][key]))
            except ValueError:
                out.append(0.0)                 # 2025: Aurora_Curve!C8 = "no " (SPEC §13 item 1); no operations
        return out
    return col


def _aurora():
    data = {}
    for r in csv.DictReader(open(DATA / "aurora_curves.csv")):
        vals = [float(r[str(y)]) for y in range(2026, 2061)]
        data[(r["series"], r["block"])] = (vals, r["source"])
    return data


def _extend(vals, first=0.0, grow=None):
    """Aurora 2026-2060 → model years 2025-2064: 2025 = `first`; after 2060 flat real (or grown by `grow`)."""
    out = [first] + list(vals)
    while len(out) < len(YEARS):
        out.append(out[-1] * (1 + grow) if grow else out[-1])
    return out


def curves(m):
    cv, ws = m.curves, m.curves.ws
    old, au = _old_curves(), _aurora()
    blocks = sorted({b for s, b in au if s == "storage_total"}, key=lambda b: (b[:4], "4h" not in b, b))

    def bar(title, source=""):
        cv.r += 1
        for c in range(1, m.last + 1):
            apply(ws.cell(cv.r, c), "section")
        ws.cell(cv.r, 1, "x")
        ws.cell(cv.r, 3, title)
        ws.cell(cv.r, 5, source)
        cv.r += 1

    def header(cells):
        for c, h in cells.items():
            apply(ws.cell(cv.r, c, h), "header")
        cv.r += 1

    def line(label, unit, values, name=None, f=None, e=None, ef=None, kind="input"):
        r = cv.r
        ws.cell(r, 3, label)
        apply(ws.cell(r, 4, unit), "unit")
        for k, v in enumerate(values):
            if v is not None:
                apply(ws.cell(r, 10 + k, v), "calc" if isinstance(v, str) and v.startswith("=") else kind, f)
        if name:
            m.define(name, f"Curves!$J${r}:${L(9 + len(values))}${r}")
        for col, item in ((5, e), (6, ef)):
            if item:
                apply(ws.cell(r, col, item[1]), "input", item[2] if len(item) > 2 else f)
                if item[0]:
                    m.define(item[0], f"Curves!${L(col)}${r}")
        cv.r += 1
        return r

    def note(text):
        ws.cell(cv.r, 3, text)
        cv.r += 1

    bar("Notes")
    note("Aurora battery revenue blocks (Degraded batteries) are fixed dispatch outputs of Aurora's model: they do "
         "NOT re-optimise when inputs change.")
    note("BESS MW scales them linearly; hours and cycles switch blocks. RTE, degradation, grid cap, PV:BESS ratio "
         "and charging rules act only through the haircut, availability and augmentation inputs.")
    note("Aurora data: data/aurora_extract.py → data/aurora_curves.csv (source rows cited per curve). Case A curves: "
         "data/old_engine/old_curves.csv (ZHLP v3 Aurora_Curve).")

    # ---- revenue-case table (SPEC §6.4) ----------------------------------------------------------------
    bar("Revenue-case table (one column per case; Calc picks the case on the Dashboard)", "SPEC §6.4")
    header({3: "Item", 10: "A", 11: "B", 12: "C"})
    first = cv.r
    v3 = "ZHLP v3 "
    rows = [
        ("Case_Names", "Case", "text", ["A", "B", "C"], "text"),
        ("Case_Labels", "Case name", "text", ["v3 reconstruction", "Aurora merchant", "Offer + merchant tail"], "text"),
        ("Case_PV_Curve", "PV price curve", "text", [v3 + "capture", "=In_PV_Curve", "=In_PV_Curve"], "text"),
        ("Case_BESS_Curve", "BESS revenue curve", "text",
         [f'=IF(In_Fix_BESS=1,In_BESS_Block,"{v3}BESS net")', "=In_BESS_Block", "=In_BESS_Block"], "text"),
        ("Case_Offer", "Offtake offer applies (1 = yes)", "flag", [0, 0, 1], "flag"),
        ("Case_MWp", "PV capacity (MWp)", "MWp", [61, "=In_MWp", "=In_MWp"], "money1"),
        ("Case_MVA", "Grid export capacity (MVA)", "MVA", [50, "=In_Grid_MVA", "=In_Grid_MVA"], "money1"),
        ("Case_Yield", "Specific yield (MWh/MWp)", "MWh/MWp", [1250, "=In_Yield", "=In_Yield"], "money"),
        ("Case_RtB", "Ready to build", "date", [D(2026, 7, 31), "=In_RtB", "=In_RtB"], "date"),
        ("Case_Exit", "RtB sale closing (exit)", "date", [D(2026, 10, 31), "=In_RtB_Exit", "=In_RtB_Exit"], "date"),
        ("Case_NTP", "NTP (construction start − 1 day)", "date", [D(2027, 3, 31), "=In_NTP", "=In_NTP"], "date"),
        ("Case_COD", "COD (operations from the next day)", "date", [D(2028, 3, 31), "=In_COD", "=In_COD"], "date"),
    ]
    for name, label, unit, vals, f in rows:
        r = line(label, unit, vals, name, f)
        ws.cell(r, 5, "ZHLP v3 Control D15:D29" if name not in ("Case_Names", "Case_Labels", "Case_Offer") else "")
    m.define("Case_Table", f"Curves!$J${first}:$L${cv.r - 1}")
    for r in range(first, cv.r):
        for c in (10, 11, 12):
            ws.cell(r, c).alignment = RIGHT

    # ---- curves picked by name -------------------------------------------------------------------------
    price = {v3 + "capture": old("pv_capture_selected_real2025")}
    price[v3 + "capture Low"] = [x * 0.85 for x in price[v3 + "capture"]]
    price[v3 + "capture High"] = [x * 1.15 for x in price[v3 + "capture"]]
    for b in blocks:
        if ("implied_capture", b) in au:
            price[CAPTURE + b] = _extend(au[("implied_capture", b)][0])
    for b in blocks:
        if ("da_to_grid", b) in au:
            price["Aurora DA ÷ to-grid – " + b] = _extend(au[("da_to_grid", b)][0])
    price["Aurora baseload (information)"] = _extend(au[("baseload", "Central")][0])
    cv.add_block("Price", "PV price curves, real 2025 (picked by name; v3 Low/High = Central ∓ 15 %)", "EUR/MWh",
                 YEARS, price, source="v3 Aurora_Curve L; Aurora Degraded batteries s+37/(s+31+s+32)",
                 fmt="money2")
    r = line("Aurora euro inflation index (2025 = 1; +2 % a year after 2060)", "index",
             _extend(au[("euro_index", "Central")][0], first=1.0, grow=0.02), "Euro_Index", "factor")
    ws.cell(r, 5, "Aurora Common Assumptions row 14")
    stack = {v3 + "BESS net": old("bess_net_real2025_eur_kw")}
    for b in blocks:
        stack[b] = _extend(au[("storage_total", b)][0])
    cv.add_block("Stack", "BESS storage total cashflow, real 2025 (Aurora block row s+27; v3: Aurora_Curve AB)",
                 "EUR k/MW", YEARS, stack, source="Fixed Aurora dispatch: no re-optimisation", fmt="money2")
    vols = [("Discharge", "BESS discharge volume (s+11; v3: U)", "discharge", "bess_discharge_mwh_per_mw", 1),
            ("ChargeRES", "BESS charge from PV (s+13, negative; v3: Z)", "charge_res",
             "bess_charge_from_pv_mwh_per_mw", -1),
            ("ChargeGrid", "BESS charge from the grid (s+12, negative)", "charge_grid", None, 1)]
    for name, label, key, oldkey, sign in vols:
        data = {v3 + "BESS net": [sign * x for x in old(oldkey)] if oldkey else [0.0] * len(YEARS)}
        for b in blocks:
            data[b] = _extend(au[(key, b)][0])
        cv.add_block(name, label, "MWh/MW", YEARS, data, source="Aurora Q2-26 Degraded batteries", fmt="money")

    bar("Aurora battery blocks: repower year (first year with repower count ≥ 1, block row s+1)",
        "Aurora Degraded batteries")
    header({3: "Block", 5: "Repower"})
    first = cv.r
    for b in [v3 + "BESS net"] + blocks:
        if b.startswith(v3):
            year = 2048
        else:
            rc = au[("repower_count", b)][0]
            year = next((2026 + k for k, x in enumerate(rc) if x >= 1), 0)
        apply(ws.cell(cv.r, 3, b), "input")
        apply(ws.cell(cv.r, 5, year), "input", "int")
        cv.r += 1
    m.define("Block_Names", f"Curves!$C${first}:$C${cv.r - 1}")
    m.define("Block_Repower", f"Curves!$E${first}:$E${cv.r - 1}")

    bar(f"Information only: components of the storage total, block {BLOCK} (never added to the total)",
        "Aurora Degraded batteries s+14…s+26")
    for k in range(14, 27):
        vals, src = au[(f"component_{k}", BLOCK)]
        line(f"s+{k} ({src})", "EUR k/MW", _extend(vals), None, "money2")

    # ---- offers (SPEC §6.3) -----------------------------------------------------------------------------
    bar("PV offers (case C): pay-as-produced; indexed = merchant price − fee", "Offer PDFs, SPEC §6.3")
    header({3: "Item", 10: "C1 E.ON", 11: "C2a MET", 12: "C2b MET", 13: "C3 MVM"})
    pv = [("OPV_Names", "Offer", "text", ["C1", "C2a", "C2b", "C3"], "text"),
          ("OPV_Type", "Fixed price (1 = yes, 0 = indexed)", "flag", [0, 1, 0, 1], "flag"),
          ("OPV_Price", "Fixed price (nominal, 3-year tenor)", "EUR/MWh", [0, 43.96, 0, 45], "money2"),
          ("OPV_Fee", "Handling fee, share of the price", "%", [0.076, 0, 0.05, 0], "0.0%"),
          ("OPV_Fee_Min", "Handling fee floor", "EUR/MWh", [5, 0, 5.75, 0], "money2"),
          ("OPV_Fee_Max", "Handling fee cap", "EUR/MWh", [1000, 0, 10, 0], "money2"),
          ("OPV_Start", "Start (blank = day after COD)", "date", [None, D(2028, 1, 1), D(2028, 1, 1), D(2028, 1, 1)],
           "date"),
          ("OPV_Tenor", "Tenor", "years", [3, 3, 3, 5], "int"),
          ("OPV_Cost", "Fixed offer fee (E.ON balance group 250 €/month)", "EUR k", [3, 0, 0, 0], "money1")]
    for name, label, unit, vals, f in pv:
        line(label, unit, vals, name, f)
    note("E.ON p4 §5.1 (fee 7.6 %, min 5); MET p3 (43.96 €/MWh 2028-30) and p4 (fee 5 %, 5.75–10); MVM e-mail (45 "
         "€/MWh 2028-32). All indicative; fees applied to the annual price.")
    bar("BESS offers (case C): owner revenue = max(floor, fixed + share × (stack − grid charge − fixed) − E.ON fee)",
        "Offer PDFs, SPEC §6.3")
    header({3: "Item", 10: "C1A E.ON", 11: "C1B E.ON", 12: "C2c MET", 13: "C2d MET", 14: "none"})
    bess = [("OB_Names", "Offer", "text", ["C1A", "C1B", "C2c", "C2d", "none"], "text"),
            ("OB_Fixed", "Fixed fee (tolling or floor), per MW-year", "EUR k/MW", [0, 0, 49.7, 45.6, 0], "money1"),
            ("OB_Hours", "Duration in the offer (scales the fixed fee)", "h", [2, 2, 1.8, 1.8, 1], "money1"),
            ("OB_Share", "Owner share of the stack above the fixed fee", "%", [0.9, 0.9, 0, 0.5, 1], "0.0%"),
            ("OB_Grid", "Grid charge deducted before the share (1 = yes)", "flag", [1, 1, 0, 0, 0], "flag"),
            ("OB_Bands", "E.ON option A band table applies (1 = yes)", "flag", [1, 0, 0, 0, 0], "flag"),
            ("OB_Start", "Start (blank = day after COD)", "date", [None, None, D(2028, 1, 1), D(2028, 1, 1), None],
             "date"),
            ("OB_Tenor", "Tenor", "years", [3, 3, 3, 3, 0], "int")]
    for name, label, unit, vals, f in bess:
        line(label, unit, vals, name, f)
    note("E.ON p3-5 (50 MW / 100 MWh; option B 90 %); MET p5-6 (tolling 49.7, floor 45.6 + 50/50, 3 years, 50 MW / "
         "90 MWh, scaled to the built hours). MVM: no BESS instrument.")
    bar("E.ON option A: bands on the stack (EUR k/MW-yr) and E.ON fee per band; outside these years option B applies",
        "E.ON offer p5")
    header({3: "Item", 5: "Year", 6: "Floor", 10: "Band 1", 11: "Band 2", 12: "Band 3", 13: "Band 4", 14: "Band 5"})
    r1 = line("Lower bound, 2027", "EUR k/MW", [0, 47, 52, 80, 109], None, "money1", e=(None, 2027, "int"))
    r2 = line("Lower bound, 2028", "EUR k/MW", [0, 38, 43, 67, 91], None, "money1", e=(None, 2028, "int"))
    for r in (r1, r2):
        apply(ws.cell(r, 6, f"=K{r}"), "calc", "money1")
    m.define("EB_Years", f"Curves!$E${r1}:$E${r2}")
    m.define("EB_Floor", f"Curves!$F${r1}:$F${r2}")
    m.define("EB_LB", f"Curves!$J${r1}:$N${r2}")
    f1 = line("E.ON fee in the band, 2027", "%", [0, 1, 0.1, 0.15, 0.18], None, "0.0%")
    f2 = line("E.ON fee in the band, 2028", "%", [0, 1, 0.1, 0.15, 0.18], None, "0.0%")
    d1 = cv.r
    for fr in (f1, f2):
        line(f"Fee step at the band (marginal), row {fr}", "%",
             [f"=J{fr}"] + [f"={L(10 + k)}{fr}-{L(9 + k)}{fr}" for k in range(1, 5)], None, "0.0%")
    m.define("EB_dR", f"Curves!$J${d1}:$N${cv.r - 1}")

    # ---- dated lines --------------------------------------------------------------------------------
    bar("Early works lines (provisional – early works to be revisited, user 22-Oct-2026)", "ZHLP v3 Stat_inp F44:F58")
    header({3: "Line (rank = order in Stat_inp)", 5: "HUF m", 6: "EUR k", 9: "EUR k", 10: "Paid", 11: "Rank"})
    first = cv.r
    ew = [("Land 1st payment", 60, 0, D(2025, 11, 30), 1), ("Land 2nd payment", 60, 0, D(2026, 2, 28), 2),
          ("ICS execution design 1st payment", 0, 25, D(2026, 2, 28), 3),
          ("Critical equipment 1st (HV/ICS/substation, 132 kV transformer)", 0, 460, D(2026, 4, 30), 4),
          ("ICS design 2nd + 132 kV transformer 2nd (Stat_inp E48)", 0, 650, D(2026, 7, 31), 5),
          ("Land 3rd payment (Stat_inp E49)", 30, 0, D(2026, 7, 31), 6)]
    for label, huf, eur, date, rank in ew:
        r = cv.r
        apply(ws.cell(r, 3, label), "input")
        for c, v, f in ((5, huf, "money1"), (6, eur, "money1"), (10, date, "date"), (11, rank, "int")):
            apply(ws.cell(r, c, v), "input", f)
        apply(ws.cell(r, 9, f"=E{r}*Thousand/In_FX+F{r}"), "calc", "money1")
        cv.r += 1
    for name, c in (("EW_Keur", "I"), ("EW_Date", "J"), ("EW_Rank", "K")):
        m.define(name, f"Curves!${c}${first}:${c}${cv.r - 1}")

    bar("Grid deposits and guarantees (released blank = at the RtB exit, or at COD on the hold routes)",
        "ZHLP v3 Stat_inp E84:E106, Dyn_Inp rows 29-33")
    header({3: "Line", 5: "HUF m", 6: "Phantom", 9: "EUR k", 10: "Paid", 11: "Released", 12: "At exit", 13: "Cash"})
    first = cv.r
    dep = [("Connection fee (reimbursement of seller)", 45, 0, D(2025, 11, 30), None, 1),
           ("Capacity guarantee", 180, 0, D(2025, 11, 30), None, 1),
           ("Cash additional guarantee (3 instalments paid by seller)", 485.55, 0, D(2025, 11, 30), D(2026, 1, 31), 1),
           ("Additional guarantee, first half", 242.775, 0, D(2026, 1, 31), D(2026, 3, 31), 0),
           ("Additional guarantee, second half", 242.775, 0, D(2026, 1, 31), None, 0),
           ("Additional guarantee, final instalments", 80.925, 0, D(2026, 4, 30), None, 0),
           ("v3 phantom guarantee (Calculation H119 one row high; fix 3)", 485.55, 1, D(2025, 11, 30), None, 0)]
    for label, huf, ph, paid, rel, cash in dep:
        r = cv.r
        apply(ws.cell(r, 3, label), "input")
        for c, v, f in ((5, huf, "money2"), (6, ph, "flag"), (10, paid, "date"), (11, rel, "date"), (13, cash, "flag")):
            apply(ws.cell(r, c, v), "input", f)
        apply(ws.cell(r, 9, f"=E{r}*Thousand/In_FX*(1-F{r}*In_Fix_BG)"), "calc", "money2")
        apply(ws.cell(r, 12, f"=N(ISBLANK(K{r}))"), "calc", "flag")
        cv.r += 1
    for name, c in (("Dep_Keur", "I"), ("Dep_Paid", "J"), ("Dep_Rel", "K"), ("Dep_Exit", "L"), ("Dep_Cash", "M")):
        m.define(name, f"Curves!${c}${first}:${c}${cv.r - 1}")

    recon_targets(m, line, bar, header, ws, cv)


RECON = [  # key in data/recon_targets.csv, label, calc template (Calc_1, set 2), unit
    ("rtb_profit", "FEP profit – sale at RtB", "{rtb_cf.I}", "EUR k"),
    ("rtb_irr", "FEP IRR – sale at RtB", "{rtb_irr.E}", "%"),
    ("htm_irr", "FEP IRR – hold to maturity", "{htm_irr.E}", "%"),
    ("proj_irr", "Project IRR, unlevered, after actual tax", "{proj_irr.E}", "%"),
    ("buyer_irr", "Buyer IRR at the RtB price, unlevered, after tax", "{buy_r_irr.E}", "%"),
    ("life_revenue", "Lifetime revenue", "{rev.I}", "EUR k"),
    ("life_ebitda", "Lifetime EBITDA", "{ebitda.I}", "EUR k"),
]


def recon_targets(m, line, bar, header, ws, cv):
    """Scenario A in the reconciliation set (2) vs the independent Python shadow (reconcile.py, SPEC §11)."""
    path = DATA / "recon_targets.csv"
    stored = {r["key"]: (float(r["shadow"]), float(r["v3"]) if r["v3"] else None) for r in csv.DictReader(open(path))} \
        if path.exists() else {}
    bar("Scenario A, reconciliation set: independent Python shadow (reconcile.py) and the original v3 values",
        "out/reconciliation.md")
    header({3: "Headline (Calc_1)", 5: "Shadow", 6: "Tolerance", 9: "v3 original"})
    for key, label, _, unit in RECON:
        sh, v3 = stored.get(key, (0.0, None))
        f = "pct" if unit == "%" else "money"
        apply(ws.cell(cv.r, 3, label), "input")
        apply(ws.cell(cv.r, 4, unit), "unit")
        apply(ws.cell(cv.r, 5, sh), "input", FMT[f])
        apply(ws.cell(cv.r, 6, f"=Recon_Tol*MAX(1,ABS(E{cv.r}))"), "calc", FMT[f])
        m.define(f"Recon_{key}", f"Curves!$E${cv.r}")
        m.define(f"Recon_Tol_{key}", f"Curves!$F${cv.r}")
        if v3 is not None:
            apply(ws.cell(cv.r, 9, v3), "input", FMT[f])
        cv.r += 1
    line("Tolerance: 0.1 % of the value (0.1 pp on rates)", "%", [], None, "pct", e=("Recon_Tol", 0.001, "0.0%"))


# =========================================================================================== calc
def calc(m):
    c = m.calc
    case = c.lever("Dash_Case")

    c.section("Case and timing")
    c.row("case", "Revenue case (E: lever, F: column of the case table)", "text", e=case,
          f="MATCH({self.E},Case_Names,0)", total="none", ef_fmt="text")
    fmt("case", "F", "int")
    c.row("mwp", "PV capacity (E: MWp) and grid export (F: MVA)", "MW", e="INDEX(Case_MWp,1,{case.F})",
          f="INDEX(Case_MVA,1,{case.F})", total="none", ef_fmt="money1")
    c.row("yld", "Specific yield (E: case, F: after the yield set)", "MWh/MWp", e="INDEX(Case_Yield,1,{case.F})",
          f="{self.E}*(1+In_Yield_Adj)", total="none", ef_fmt="money")
    c.row("rtbd", "RtB (E) and RtB exit (F)", "date", e="INDEX(Case_RtB,1,{case.F})",
          f="INDEX(Case_Exit,1,{case.F})", total="none")
    c.row("codd", "NTP (E) and COD (F)", "date", e="INDEX(Case_NTP,1,{case.F})",
          f="EOMONTH(INDEX(Case_COD,1,{case.F}),In_COD_Delay)", total="none")
    c.event("acq_ev", "Acquisition and JV start", "In_Acq")
    c.event("exit_ev", "RtB exit (sale closing)", "{rtbd.F}")
    c.event("ntp_ev", "NTP", "{codd.E}")
    c.event("cod_ev", "COD", "{codd.F}")
    c.row("col_x", "Exit column on the timeline (buyer view at RtB)", "#", e="MATCH({exit_ev.E}-1,{end.all},1)+1",
          total="none", ef_fmt="int")
    c.row("col_c", "COD column on the timeline (buyer view at COD)", "#", e="MATCH({cod_ev.E},{end.all},0)",
          total="none", ef_fmt="int")
    c.row("post_x", "After the exit period", "flag", "=N({start}>{exit_ev.E})", total="none", fmt="flag")
    c.row("post_c", "After COD", "flag", "=N({end}>{cod_ev.E})", total="none", fmt="flag")
    c.row("constr", "Construction spread (NTP → COD)", "factor", SPREAD, e="{codd.E}+1", f="{codd.F}",
          ef_fmt="date")
    c.window("ops", "Operations (COD → end of life)", "{codd.F}+1", "EDATE({codd.F},In_Life*Months_In_Year)")
    c.row("elapsed", "Operating years elapsed at the period start", "years",
          "=MAX(0,{start}-1-{codd.F})/Days_In_Year", total="none", fmt="factor")
    c.row("dsa_w", "DSA loan spread (acquisition → RtB)", "factor", SPREAD, e="In_Acq+1", f="{rtbd.E}",
          ef_fmt="date")
    c.row("pre_w", "Pre-RtB development cost spread", "factor", SPREAD, e="EOMONTH(In_PreRtB_First,-1)+1",
          f="In_PreRtB_Last", ef_fmt="date")
    c.window("tenor", "Debt repayment period", "{codd.F}+1", "EDATE({codd.F},In_Debt_Tenor*Months_In_Year)")
    c.event("tenor_end", "Last debt repayment", "EDATE({codd.F},In_Debt_Tenor*Months_In_Year)")
    c.row("dep", "Tax depreciation spread: PV, grid, development", "factor", SPREAD, e="{codd.F}+1",
          f="EDATE({codd.F},In_Dep_Years*Months_In_Year)", ef_fmt="date")
    c.row("dep_b", "Tax depreciation spread: BESS", "factor", SPREAD, e="{codd.F}+1",
          f="EDATE({codd.F},In_Dep_BESS_Years*Months_In_Year)", ef_fmt="date")
    c.row("idx", "Debt compounding index (E: at COD, F: rate per period)", "factor", "=(1+{self.F})^{period}",
          e="SUMIF({end.all},{codd.F},{self.all})", f="In_Debt_Rate*Months_Per_Period/Months_In_Year",
          total="none", ef_fmt="factor")
    fmt("idx", "F", "0.000%")
    c.row("q4", "31-Dec capitalisation", "flag", "=N(MONTH({end})=Months_In_Year)", total="none", fmt="flag")
    c.row("q1", "First quarter of the year", "flag", "=N(MONTH({end})<=Months_Per_Period)", total="none", fmt="flag")

    c.section("Prices (real 2025 curves × Aurora euro index)")
    c.row("infl", "Aurora euro index", "factor", "=INDEX(Euro_Index,MATCH({year},Price_Years,0))", total="none")
    c.curve("pv_real", "PV price, real 2025 (E: curve)", "EUR/MWh", "Price", "INDEX(Case_PV_Curve,1,{case.F})",
            fmt="money2")
    c.row("pv_mkt", "PV merchant price, nominal (price set applied)", "EUR/MWh",
          "={pv_real}*{infl}*(1+In_PV_Price_Adj)", total="none", fmt="money2")
    c.row("goo", "Guarantee of origin price, nominal", "EUR/MWh", "=In_GoO*{infl}", total="none", fmt="money2")
    c.curve("stk_real", "BESS storage total cashflow, real 2025 (E: curve)", "EUR k/MW", "Stack",
            "INDEX(Case_BESS_Curve,1,{case.F})", fmt="money2")
    c.row("stack", "BESS merchant revenue per MW-year, nominal (haircut, availability)", "EUR k/MW",
          "={stk_real}*{infl}*(1-In_BESS_Haircut)*In_BESS_Avail/In_Aurora_Avail", total="none", fmt="money2")

    c.section("Offtake offers (case C; merchant tail after the tenor)")
    c.row("offer", "Offer applies on this sheet (case table)", "flag", e="INDEX(Case_Offer,1,{case.F})",
          total="none", ef_fmt="flag")
    c.row("op", "PV offer: column (E), fixed price (F: 1 = yes)", "#", e="MATCH(In_Offer_PV,OPV_Names,0)",
          f="INDEX(OPV_Type,{self.E})", total="none", ef_fmt="int")
    c.row("op_fee", "PV offer handling fee: share (E), floor (F)", "EUR/MWh", e="INDEX(OPV_Fee,{op.E})",
          f="INDEX(OPV_Fee_Min,{op.E})", total="none", ef_fmt="money2")
    fmt("op_fee", "E", "pct")
    c.row("op_px", "PV offer: fee cap (E), fixed price (F)", "EUR/MWh", e="INDEX(OPV_Fee_Max,{op.E})",
          f="INDEX(OPV_Price,{op.E})", total="none", ef_fmt="money2")
    c.row("op_win", "PV offer in force (E: start, F: end)", "factor", "={offer.E}*" + WINDOW[1:],
          e="IF(INDEX(OPV_Start,{op.E})>0,INDEX(OPV_Start,{op.E}),{codd.F}+1)",
          f="EDATE({self.E},INDEX(OPV_Tenor,{op.E})*Months_In_Year)-1", total="none", ef_fmt="date")
    c.row("pv_fee", "PV offer handling fee", "EUR/MWh", "=MIN({op_px.E},MAX({op_fee.E}*{pv_mkt},{op_fee.F}))",
          total="none", fmt="money2")
    c.row("pv_price", "PV price received (offer in force, else merchant)", "EUR/MWh",
          "={pv_mkt}+{op_win}*(IF({op.F}=1,{op_px.F},{pv_mkt}-{pv_fee})-{pv_mkt})", total="none", fmt="money2")
    c.row("ob", "BESS offer: column (E), fixed fee per MW-year at the built hours (F)", "EUR k/MW",
          e="MATCH(In_Offer_BESS,OB_Names,0)", f="INDEX(OB_Fixed,{self.E})*In_BESS_Hours/INDEX(OB_Hours,{self.E})",
          total="none", ef_fmt="money1")
    fmt("ob", "E", "int")
    c.row("ob_p", "BESS offer: owner share (E), grid charge deducted (F)", "EUR k/MW", e="INDEX(OB_Share,{ob.E})",
          f="INDEX(OB_Grid,{ob.E})*In_BESS_Grid_Charge", total="none", ef_fmt="money1")
    fmt("ob_p", "E", "pct")
    c.row("ob_win", "BESS offer in force (E: start, F: end)", "factor", "={offer.E}*" + WINDOW[1:],
          e="IF(INDEX(OB_Start,{ob.E})>0,INDEX(OB_Start,{ob.E}),{codd.F}+1)",
          f="EDATE({self.E},INDEX(OB_Tenor,{ob.E})*Months_In_Year)-1", total="none", ef_fmt="date")
    c.row("b_band", "E.ON band table row this year (0 = none; E: bands apply)", "#",
          "={self.E}*IFERROR(MATCH({year},EB_Years,0),0)", e="INDEX(OB_Bands,{ob.E})", total="none", fmt="int",
          ef_fmt="flag")
    c.row("b_row", "Band table row looked up", "#", "=MAX(1,{b_band})", total="none", fmt="int")
    c.row("b_fee", "E.ON fee on the stack, by band", "EUR k/MW",
          "=({b_band}>0)*SUMPRODUCT(({stack}>INDEX(EB_LB,{b_row},0))*({stack}-INDEX(EB_LB,{b_row},0)),"
          "INDEX(EB_dR,{b_row},0))", total="none", fmt="money2")
    c.row("b_floor", "Guaranteed floor in band years", "EUR k/MW", "=({b_band}>0)*INDEX(EB_Floor,{b_row})",
          total="none", fmt="money2")
    c.row("b_share", "Owner share of the stack above the fixed fee", "%", "=IF({b_band}>0,1,{ob_p.E})",
          total="none", fmt="pct")
    c.row("b_own", "Owner BESS revenue per MW-year under the offer", "EUR k/MW",
          "=MAX({b_floor},{ob.F}+{b_share}*MAX(0,{stack}-{ob_p.F}-{ob.F})-{b_fee})", total="none", fmt="money2")
    c.row("bess_px", "BESS revenue per MW-year received", "EUR k/MW", "={stack}+{ob_win}*({b_own}-{stack})",
          total="none", fmt="money2")

    c.section("Capacity and energy")
    c.row("bess_mw", "BESS built (E: MW, F: MWh)", "MW", e="In_BESS_MW*In_BESS_On", f="{self.E}*In_BESS_Hours",
          total="none", ef_fmt="money1")
    c.row("gen", "PV generation after availability and degradation (E: first full year)", "MWh",
          "={mwp.E}*{yld.F}*In_PV_Avail*(1-In_Degradation)^{elapsed}*{ops}*{yf}",
          e="SUMIF({year.all},YEAR({codd.F})+1,{self.all})")
    c.row("clip", "Clipped at the grid limit (fix 5)", "MWh", "=-{gen}*In_Clip_Loss*In_Fix_Clip")
    c.curve("v_res", "Aurora block: charge from PV per MW (E: curve)", "MWh", "ChargeRES", "{stk_real.E}", fmt="money")
    c.curve("v_grid", "Aurora block: charge from the grid per MW", "MWh", "ChargeGrid", "{stk_real.E}", fmt="money")
    c.curve("v_dis", "Aurora block: discharge per MW", "MWh", "Discharge", "{stk_real.E}", fmt="money")
    c.row("to_bess", "PV energy routed to the BESS", "MWh", "=MIN({gen}+{clip},-{v_res}*{bess_mw.E}*{ops}*{yf})")
    c.row("pv_exp", "PV energy exported directly", "MWh", "={gen}+{clip}-{to_bess}")
    c.row("dis", "BESS discharge to the grid", "MWh", "={v_dis}*{bess_mw.E}*{ops}*{yf}")
    c.row("chg", "BESS charge from the grid", "MWh", "=-{v_grid}*{bess_mw.E}*{ops}*{yf}")
    c.row("export", "Energy exported (E: first full year)", "MWh", "={pv_exp}+{dis}",
          e="SUMIF({year.all},YEAR({codd.F})+1,{self.all})")
    c.row("grid_h", "Grid export limit (MVA × hours)", "MWh", "={mwp.F}*Hours_In_Day*{days}*{ops}")

    c.section("Revenue, opex and EBITDA")
    c.row("pv_rev", "PV revenue (E: first full year)", "EUR k", "=({gen}+{clip})*{pv_price}/Thousand",
          e="SUMIF({year.all},YEAR({codd.F})+1,{self.all})")
    c.row("goo_rev", "Guarantee of origin revenue", "EUR k", "=({gen}+{clip})*{goo}/Thousand")
    c.row("bess_rev", "BESS revenue (E: first full year)", "EUR k", "={bess_px}*{bess_mw.E}*{ops}*{yf}",
          e="SUMIF({year.all},YEAR({codd.F})+1,{self.all})")
    c.row("rev", "Revenue (E: first full year)", "EUR k", "={pv_rev}+{goo_rev}+{bess_rev}",
          e="SUMIF({year.all},YEAR({codd.F})+1,{self.all})", bold=True)
    c.row("contr", "Of which contracted (fixed-price PV offer, BESS fixed fee)", "EUR k",
          "={op_win}*{op.F}*{pv_rev}+{ob_win}*{ob.F}*{bess_mw.E}*{ops}*{yf}")
    c.row("opex_om", "O&M and MV cable O&M", "EUR k",
          "=-(In_OM+In_Cable_OM)*{mwp.E}*{infl}*(1+In_Opex_Adj)*{ops}*{yf}")
    c.row("opex_bal", "Balancing", "EUR k", "=-In_Balancing*({gen}+{clip})/Thousand*{infl}*(1+In_Opex_Adj)")
    c.row("opex_lease", "Land lease", "EUR k", "=-In_Lease*{infl}*(1+In_Opex_Adj)*{ops}*{yf}")
    c.row("opex_bess", "BESS fixed opex", "EUR k", "=-In_BESS_Opex*{bess_mw.E}*{infl}*(1+In_Opex_Adj)*{ops}*{yf}")
    c.row("opex_rtm", "Route-to-market and offer fees", "EUR k", "=-In_RTM*{pv_rev}-INDEX(OPV_Cost,{op.E})*{op_win}*{yf}")
    c.row("levies", "HIPA and innovation contribution", "EUR k", "=-{rev}*(In_HIPA+In_Innovation)")
    c.row("ebitda", "EBITDA (E: first full year)", "EUR k", "={rev}+SUM({opex_om}:{levies})",
          e="SUMIF({year.all},YEAR({codd.F})+1,{self.all})", bold=True)

    c.section("Development, early works, deposits and guarantees (E: to the exit, F: to COD)")
    c.row("spv", "FEP SPV acquisition price / connection rights", "EUR k", "=-In_SPV_Price*{acq_ev}")
    c.row("trans", "Transaction costs (FEP)", "EUR k", "=-In_Transaction*{acq_ev}")
    c.row("dsa", "DSA loan to PSH", "EUR k", "=-In_DSA_Per_MWp*{mwp.E}*{dsa_w}")
    c.row("ew_dev", "Early works paid in development (lines in scope)", "EUR k",
          '=-SUMIFS(EW_Keur,EW_Date,">="&{start},EW_Date,"<="&{end},EW_Rank,"<="&In_Milestones_Scope)', e=EXIT)
    c.row("ew_ntp", "Early works outside the scope, paid at NTP", "EUR k",
          '=-In_Milestones_Late_To_Capex*{ntp_ev}*SUMIF(EW_Rank,">"&In_Milestones_Scope,EW_Keur)')
    c.row("pre_dev", "Pre-RtB development costs", "EUR k", "=-In_PreRtB_Devex*{pre_w}")
    c.row("dep_cash", "Cash deposits paid", "EUR k",
          '=-SUMIFS(Dep_Keur,Dep_Paid,">="&{start},Dep_Paid,"<="&{end},Dep_Cash,1)')
    c.row("dep_back", "Cash deposits released on their date (E: outstanding until exit or COD)", "EUR k",
          '=SUMIFS(Dep_Keur,Dep_Rel,">="&{start},Dep_Rel,"<="&{end},Dep_Cash,1)',
          e="SUMIFS(Dep_Keur,Dep_Exit,1,Dep_Cash,1)")
    c.row("dep_cod", "Cash deposits released at COD (hold routes)", "EUR k", "={dep_back.E}*{cod_ev}")
    c.row("bg_d", "Guarantees outstanding, dated release (E: all-in rate)", "EUR k",
          '=SUMIFS(Dep_Keur,Dep_Paid,"<="&{end},Dep_Rel,">"&{end},Dep_Cash,0)', e="In_BG_Base+In_BG_Margin",
          total="none", ef_fmt="0.00%")
    c.row("bg_x", "Guarantees outstanding, released at the exit or COD", "EUR k",
          '=SUMIFS(Dep_Keur,Dep_Paid,"<="&{end},Dep_Exit,1,Dep_Cash,0)', total="none")
    c.row("bg_hold", "Guarantee fees, hold routes (to COD)", "EUR k",
          "=-({bg_d}+{bg_x}*(1-{post_c}))*{bg_d.E}*{days}/Days_In_Year", e=EXIT, f=TO_COD)
    c.row("bg_rtb", "Guarantee fees, RtB route (to the exit period)", "EUR k", "={bg_hold}*(1-{post_x})")

    c.section("Construction")
    c.row("capex", "Remaining construction capex (E: total incl. early works, F: of which BESS)", "EUR k",
          "=-({self.E}+{ew_dev.I}+{ew_ntp.I})*{constr}",
          e="(In_PV_Capex*{mwp.E}+In_Cable_Capex+In_Substation)*(1+In_Capex_Adj)+In_Land_Capex+{self.F}",
          f="In_BESS_Capex*{bess_mw.F}*(1+In_Capex_Adj)")
    c.row("aug_y", "BESS augmentation year (E: used, F: repower year of the curve)", "year",
          e="IF(In_Fix_Augment=0,In_Augment_Year_v3,IF(In_Augment_Year>0,In_Augment_Year,{self.F}))",
          f="INDEX(Block_Repower,MATCH({stk_real.E},Block_Names,0))", total="none", ef_fmt="int")
    c.row("aug", "BESS augmentation (Q1 of the year; indexed with fix 7)", "EUR k",
          "=-In_Augment_Cost*{bess_mw.F}*IF(In_Fix_Augment=1,{infl},1)*({year}={aug_y.E})*{q1}")
    c.row("cost", "Construction cost (E: compounded to COD at the debt rate)", "EUR k", "={capex}+{ew_ntp}",
          e="SUMPRODUCT({self.all}/{idx.all})*{idx.E}", bold=True)

    c.section("Debt sizing on CFADS")
    c.paste("Tax", "Corporate income tax", "EUR k")
    c.row("cfads", "CFADS (EBITDA less tax, after COD)", "EUR k", "=({ebitda}+{tax_paste})*{post_c}")
    c.row("w_dscr", "1 ÷ target DSCR, weighted by contracted and merchant revenue", "factor",
          "=IF({rev}>0,({contr}/In_DSCR_Contracted+({rev}-{contr})/In_DSCR_Merchant)/{rev},0)", total="none")
    c.row("ds_cap", "CFADS ÷ target DSCR (E: annuity debt capacity, F: lowest quarter)", "EUR k",
          '=IF({tenor}=1,{cfads}*{w_dscr},"")', e="PV({idx.F},In_Debt_Tenor*Months_In_Year/Months_Per_Period,-{self.F})",
          f="MIN({self.all})", total="none")
    c.row("dtgt", "Debt quantum: per MWp (off: whole cost)", "EUR k",
          e="IF(In_Debt_Per_MW>0,In_Debt_Per_MW*{mwp.E},-{cost.E})", total="none")
    c.row("k", "Debt share of construction cost (E), share of the capacity used (F)", "%",
          e="MAX(0,MIN(In_Gearing,{ds_cap.E}/-{cost.E},{dtgt.E}/-{cost.E}))", f="{self.E}*-{cost.E}/{ds_cap.E}",
          total="none", ef_fmt="pct", bold=True)
    c.row("d_cod", "Debt at COD (E) and level annuity per quarter (F)", "EUR k", e="{k.E}*-{cost.E}",
          f="PMT({idx.F},In_Debt_Tenor*Months_In_Year/Months_Per_Period,-{self.E})", total="none")

    c.section("Senior debt")
    c.row("d_int", "Interest paid", "EUR k", "=-{d_close@prev}*{idx.F}*{post_c}")
    c.row("d_sched", "Scheduled repayment (annuity)", "EUR k", "=-MIN({d_close@prev},{d_cod.F}+{d_int})*{tenor}")
    c.row("d_service", "Scheduled debt service", "EUR k", "={d_int}+{d_sched}")
    c.row("d_sweep", "Cash sweep", "EUR k",
          "=-MIN({d_close@prev}+{d_sched},In_Sweep*MAX(0,{cfads}+{d_service}))*{post_c}")
    c.row("d_ds", "Debt service incl. sweep", "EUR k", "={d_service}+{d_sweep}")
    c.row("d_open", "Opening balance", "EUR k", "={d_close@prev}", total="none")
    c.row("d_draw", "Drawdown", "EUR k", "=-{k.E}*{cost}")
    c.row("d_idc", "Interest capitalised", "EUR k", "={d_open}*{idx.F}*(1-{post_c})")
    c.row("d_rep", "Repayment", "EUR k", "={d_sched}+{d_sweep}")
    c.row("d_close", "Closing balance (E: at COD)", "EUR k", "=SUM({d_open}:{d_rep})",
          e="SUMIF({end.all},{codd.F},{self.all})", total="seed", bold=True)
    c.row("d_up", "Upfront fee", "EUR k", "=-In_Upfront_Fee*{d_close.E}*({d_open}=0)*({d_draw}>0)")
    c.row("d_cf", "Commitment fee", "EUR k",
          "=-In_Commitment_Fee*({d_draw.I}-SUM({d_draw.todate})+{d_draw})*{yf}*({constr}>0)")
    c.row("dsra", "DSRA funded at COD, released at the end (E: balance)", "EUR k", "={self.E}*({tenor_end}-{cod_ev})",
          e='-SUMIFS({d_service.all},{end.all},">"&{codd.F},{end.all},"<="&EDATE({codd.F},In_DSRA_Months))')
    c.row("dscr", "DSCR on scheduled debt service (E: minimum)", "x", '=IF({d_service}<-Tolerance,{cfads}/-{d_service},"")',
          e="MIN({self.all})", total="none", fmt="mult", ef_fmt="mult")

    c.section("Tax")
    c.row("tax_dep", "Tax depreciation: PV, grid, development (E: base)", "EUR k", "=-{self.E}*{dep}",
          e="{capex.E}-{capex.F}-In_Land_Capex+{d_idc.I}-{pre_dev.I}-{dsa.I}-{spv.I}")
    c.row("tax_dep_b", "Tax depreciation: BESS (E: base)", "EUR k", "=-{self.E}*{dep_b}", e="{capex.F}")
    c.row("tax_dep_a", "Tax depreciation: augmentation (E: start, F: end)", "EUR k", "={aug.I}*" + SPREAD[1:],
          e='SUMIF({aug.all},"<0",{end.all})+1', f="EDATE({self.E}-1,In_Dep_BESS_Years*Months_In_Year)",
          ef_fmt="date")
    c.row("taxable", "Tax base before losses", "EUR k",
          "={ebitda}+{tax_dep}+{tax_dep_b}+{tax_dep_a}+{d_int}+{d_up}+{d_cf}+{bg_hold}")
    c.row("lcf_open", "Tax losses: opening", "EUR k", "={lcf_close@prev}", total="none")
    c.row("lcf_add", "Losses added", "EUR k", "=MAX(0,-{taxable})")
    c.row("lcf_use", "Losses used (capped)", "EUR k", "=-MIN({lcf_open},In_Loss_Cap*MAX(0,{taxable}))")
    c.row("lcf_lapse", "Losses lapsed (E: date, NTP + carry-forward period)", "EUR k",
          "=-({lcf_open}+{lcf_add}+{lcf_use})*N(AND({start}<={self.E},{end}>={self.E}))",
          e="EDATE({codd.E},In_Loss_Years*Months_In_Year)")
    fmt("lcf_lapse", "E", "date")
    c.row("lcf_close", "Tax losses: closing", "EUR k", "=SUM({lcf_open}:{lcf_lapse})", total="seed")
    c.copy("Tax", "Corporate income tax and Robin Hood tax", "EUR k",
           "=-In_CIT*(MAX(0,{taxable})+{lcf_use})-In_RH_On*In_RH_Rate*MAX(0,{taxable})")

    c.section("Project and equity cash flows (100 %)")
    c.row("dev_cf", "Development, deposits and guarantees", "EUR k",
          "={spv}+{trans}+{dsa}+{ew_dev}+{pre_dev}+{dep_cash}+{dep_back}+{dep_cod}+{bg_hold}")
    c.row("proj_pre", "Project cash flow before financing and tax", "EUR k",
          "={dev_cf}+{capex}+{ew_ntp}+{aug}+{ebitda}")
    c.row("proj_cf", "Project cash flow before financing, after tax", "EUR k", "={proj_pre}+{tax_copy}", bold=True)
    c.xirr("proj_irr", "Project IRR, unlevered, after the SPV's actual tax", "proj_cf")
    c.xirr("proj_irr_pre", "Project IRR, unlevered, before tax", "proj_pre")
    c.row("debt_cf", "Debt flows to equity (draws, fees, service, DSRA)", "EUR k",
          "={d_draw}+{d_up}+{d_cf}+{d_ds}+{dsra}")
    c.row("eq_cf", "Equity cash flow, levered, after tax", "EUR k", "={proj_cf}+{debt_cf}", bold=True)
    c.subheader("Funding to COD (E: uses and sources)")
    c.row("f_dev", "Development, deposits and guarantees (net)", "EUR k", e='-SUMIF({end.all},"<="&{codd.F},{dev_cf.all})')
    c.row("f_capex", "Construction capex and early works at NTP", "EUR k", e="-({capex.I}+{ew_ntp.I})")
    c.row("f_fin", "Capitalised interest and debt fees", "EUR k", e="{d_idc.I}-{d_up.I}-{d_cf.I}")
    c.row("f_uses", "Total uses", "EUR k", e="SUM({f_dev.E}:{f_fin.E})", bold=True)
    c.row("f_equity", "Equity to COD (E), peak including the DSRA (F)", "EUR k", e="{f_uses.E}-{d_close.E}",
          f="{self.E}+{dsra.E}")
    c.row("f_gear", "Debt share of total uses (E), check: equity + equity flows to COD (F)", "%",
          e="{d_close.E}/{f_uses.E}", f='{f_equity.F}+SUMIF({end.all},"<="&{cod_ev.E},{eq_cf.all})', ef_fmt="pct")
    fmt("f_gear", "F", "money")

    c.section("Buyer view and implied prices (share deal: SPV tax base unchanged)")
    c.subheader("Buyer at the RtB exit")
    c.row("reimb_x", "Buyer reimburses at the exit: early works, cash deposits", "EUR k",
          "={exit_ev}*(In_EW_Reimbursed*-{ew_dev.E}+{dep_back.E})")
    c.row("pre_reimb", "Buyer reimburses pre-RtB costs (E: at the exit; fix 2: as spent after it)", "EUR k",
          "=In_PreRtB_Reimbursed*({exit_ev}*{self.E}-In_Fix_Reimb*{post_x}*{pre_dev})",
          e='IF(In_Fix_Reimb=1,-SUMIF({start.all},"<="&{exit_ev.E},{pre_dev.all}),-{pre_dev.I})')
    c.row("bu_rtb", "Buyer flows after the exit before the price (unlevered, after tax)", "EUR k",
          "={post_x}*({proj_cf}-{dev_cf}+{bg_hold}+{dep_cod}+{ew_dev})-{reimb_x}-{pre_reimb}")
    last = L(m.last)

    def npv_from(rate, key, col_key):
        """XNPV of a row from a start column (E of `col_key`) to the end of the timeline."""
        r = c.rows[key].r
        return (f"XNPV({rate},INDEX({{{key}.all}},{{{col_key}.E}}):${last}${r},"
                f"INDEX({{end.all}},{{{col_key}.E}}):${last}$4)")
    c.row("imp_r", "Implied RtB price at the target IRR (E: share deal, closed form; F: asset deal, goal seek)",
          "EUR k/MW", e=npv_from("In_Buyer_IRR_RtB", "bu_rtb", "col_x") + "/{mwp.E}", f="INDEX(Asset_RtB_Inputs,$A$1)", total="none", ef_fmt="money1")
    c.row("px_rtb", "RtB price used (E: k€/MWp, F: consideration)", "EUR k",
          e="IF(In_RtB_Price>0,In_RtB_Price,IF(In_Sale_Form=1,{imp_r.F},{imp_r.E}))*(1+In_Price_Step)",
          f="{self.E}*{mwp.E}", total="none", bold=True)
    fmt("px_rtb", "E", "money1")
    c.row("buy_r", "Buyer cash flow at the RtB price, unlevered, after tax", "EUR k",
          "={bu_rtb}-{px_rtb.F}*{exit_ev}")
    c.xirr("buy_r_irr", "Buyer IRR at RtB, unlevered, after tax", "buy_r", start="col_x")
    c.row("buy_r_pre", "Buyer cash flow at the RtB price, unlevered, before tax", "EUR k",
          "={buy_r}-{post_x}*{tax_copy}")
    c.xirr("buy_r_irr_pre", "Buyer IRR at RtB, unlevered, before tax", "buy_r_pre", start="col_x")
    c.row("buy_r_l", "Buyer cash flow at the RtB price, levered, after tax", "EUR k", "={buy_r}+{post_x}*{debt_cf}")
    c.xirr("buy_r_irr_l", "Buyer IRR at RtB, levered, after tax", "buy_r_l", start="col_x")
    c.row("as_r_dep", "Asset deal: buyer's depreciation of the solved price", "EUR k",
          "=-{imp_r.F}*{mwp.E}*{dep}")
    c.row("as_r_cf", "Asset deal: buyer flows at the solved price, tax saving capped at the SPV's tax", "EUR k",
          "={bu_rtb}-{imp_r.F}*{mwp.E}*{exit_ev}+MIN(-{tax_copy},-In_CIT*{as_r_dep})")
    c.xirr("as_r_irr", "Asset deal: buyer IRR at the solved RtB price (goal seek → target)", "as_r_cf",
           start="col_x")
    c.subheader("Buyer at COD")
    c.row("bu_cod", "Buyer flows after COD before the price (unlevered, after tax)", "EUR k",
          "={post_c}*({ebitda}+{tax_copy}+{aug})")
    c.row("imp_c", "Implied COD price at the target IRR (E: share deal, closed form; F: asset deal, goal seek)",
          "EUR k/MW", e=npv_from("In_Buyer_IRR_COD", "bu_cod", "col_c") + "/{mwp.E}", f="INDEX(Asset_COD_Inputs,$A$1)", total="none", ef_fmt="money")
    c.row("px_cod", "COD price used (E: k€/MWp, F: enterprise value)", "EUR k",
          e="IF(In_COD_Price>0,In_COD_Price,IF(In_Sale_Form=1,{imp_c.F},{imp_c.E}))*(1+In_Price_Step)",
          f="{self.E}*{mwp.E}", total="none", bold=True)
    c.row("buy_c", "Buyer cash flow at the COD price, unlevered, after tax", "EUR k", "={bu_cod}-{px_cod.F}*{cod_ev}")
    c.xirr("buy_c_irr", "Buyer IRR at COD, unlevered, after tax", "buy_c", start="col_c")
    c.row("buy_c_pre", "Buyer cash flow at the COD price, unlevered, before tax", "EUR k",
          "={buy_c}-{post_c}*{tax_copy}")
    c.xirr("buy_c_irr_pre", "Buyer IRR at COD, unlevered, before tax", "buy_c_pre", start="col_c")
    c.row("buy_c_l", "Buyer cash flow at COD, levered: pays price less debt, funds the DSRA", "EUR k",
          "={buy_c}+{d_ds}+{dsra}+{d_close.E}*{cod_ev}")
    c.xirr("buy_c_irr_l", "Buyer IRR at COD, levered, after tax", "buy_c_l", start="col_c")
    c.row("as_c_dep", "Asset deal: buyer's depreciation of the solved price", "EUR k", "=-{imp_c.F}*{mwp.E}*{dep}")
    c.row("as_c_cf", "Asset deal: buyer flows at the solved price, tax saving capped at the SPV's tax", "EUR k",
          "={bu_cod}-{imp_c.F}*{mwp.E}*{cod_ev}+MIN(-{tax_copy},-In_CIT*{as_c_dep})")
    c.xirr("as_c_irr", "Asset deal: buyer IRR at the solved COD price (goal seek → target)", "as_c_cf", start="col_c")

    c.section("FEP return on loans and costs (12 %; E: to the exit, F: to COD)")
    c.row("cl_f", "Return per period (E: v3 monthly compounding; fix 8: act/365 simple)", "factor",
          "=IF(In_Fix_Interest=1,In_Loan_Rate*{days}/Days_In_Year,{self.E})",
          e="(1+In_Loan_Rate/Months_In_Year)^Months_Per_Period-1", total="none", fmt="0.000%", ef_fmt="0.000%")
    c.row("cl_add", "Principal: DSA, early works, deposits, guarantee fees, SPV price", "EUR k",
          "=-({dsa}+In_EW_Interest*{ew_dev}+{dep_cash}+{dep_back}+{bg_hold}+In_SPV_Interest*{spv})")
    c.row("cl_open", "Opening balance (principal and capitalised return)", "EUR k", "={cl_close@prev}", total="none")
    c.row("cl_ret", "Return accrued", "EUR k", "={cl_open}*{cl_f}*(1-{post_c})", e=EXIT, f=TO_COD)
    c.row("cl_cap", "Return capitalised (31-Dec with fix 8, else every period)", "EUR k",
          "=IF(In_Fix_Interest=1,{q4}*SUMIF({year.todate},{year},{cl_ret.todate}),{cl_ret})")
    c.row("cl_close", "Closing balance", "EUR k", "={cl_open}+{cl_add}+{cl_cap}", total="seed")
    c.row("ps_open", "PSH land: opening balance (v3 rule)", "EUR k", "={ps_close@prev}", total="none")
    c.row("ps_ret", "PSH land: return accrued", "EUR k", "={ps_open}*{cl_f}*(1-{post_c})", e=EXIT, f=TO_COD)
    c.row("ps_cap", "PSH land: return capitalised", "EUR k",
          "=IF(In_Fix_Interest=1,{q4}*SUMIF({year.todate},{year},{ps_ret.todate}),{ps_ret})")
    c.row("ps_close", "PSH land: closing balance", "EUR k", "={ps_open}+In_PSH_Land*{acq_ev}+{ps_cap}", total="seed")

    c.section("Sale at RtB (JV split: 221025 sheet, step share rule)")
    c.row("shr_d", "Share rule inputs (E: low − high share, F: v3 slide width)", "%",
          e="In_FEP_Share_Low-In_FEP_Share_High", f="In_Share_Slide_End-In_Share_Threshold", total="none",
          ef_fmt="pct")
    fmt("shr_d", "F", "money1")
    c.row("shr_rule", "Futureal share by rule (E: step, F: v3 slide)", "%",
          e="IF({px_rtb.E}<=In_Share_Threshold,In_FEP_Share_Low,In_FEP_Share_High)",
          f="In_FEP_Share_Low-{shr_d.E}*MIN(1,MAX(0,({px_rtb.E}-In_Share_Threshold)/{shr_d.F}))", total="none",
          ef_fmt="pct")
    c.row("shr", "Share of the distributable amount (E: Futureal, F: PSH)", "%",
          e="IF(In_Share_Rule=1,{shr_rule.E},{shr_rule.F})", f="1-{self.E}", total="none", ef_fmt="pct", bold=True)
    c.subheader("RtB waterfall (E: due, F: paid)")
    c.row("w_avail", "Cash available: consideration (E and F)", "EUR k", e="{px_rtb.E}*{mwp.E}",
          f="{self.E}", total="none")
    waterfall(c, "w", [("cost", "Sale costs", "In_Sale_Costs_RtB"),
                       ("spv", "FEP SPV acquisition price / connection rights", "In_SPV_Price"),
                       ("dsa", "DSA loan principal", "-{dsa.I}"),
                       ("ret", "FEP 12 % return", "{cl_ret.E}"),
                       ("bg", "Guarantee fees paid by FEP", "-{bg_rtb.I}"),
                       ("psh", "PSH land apport and return (v3 rule)", "In_PSH_Land_Step*(In_PSH_Land+{ps_ret.E})")])
    c.row("w_chk", "Check: paid + distributable − available (E), steps paid above due (F)", "EUR k",
          e="SUM({w_cost.F}:{w_psh.F})+{w_dist.F}-{w_avail.F}",
          f="SUMPRODUCT(--({w_cost.F}:{w_psh.F}>{w_cost.E}:{w_psh.E}+Tolerance))", total="none")
    c.row("w_fep", "FEP receipts at the exit", "EUR k", f="SUM({w_spv.F}:{w_bg.F})+{shr.E}*{w_dist.F}",
          total="none", bold=True)
    c.row("w_psh_r", "PSH receipts at the exit", "EUR k", f="{w_psh.F}+{shr.F}*{w_dist.F}", total="none")
    c.subheader("FEP returns – sale at RtB")
    c.row("rtb_out", "FEP outlays – RtB route", "EUR k",
          "=({spv}+{trans}+{dsa}+{ew_dev}+{dep_cash})*(1-{post_x})+{bg_rtb}+{pre_dev}")
    c.row("rtb_in", "FEP receipts – RtB route", "EUR k",
          "={dep_back}*(1-{post_x})+{reimb_x}+{pre_reimb}+{w_fep.F}*{exit_ev}")
    c.row("rtb_pre", "FEP cash flow before tax on the gain", "EUR k", "={rtb_out}+{rtb_in}")
    c.row("rtb_tax", "Tax on the RtB gain (E: amount)", "EUR k", "={self.E}*{exit_ev}",
          e="-In_CIT*In_Sale_Gain_Taxed*MAX(0,{rtb_pre.I})")
    c.row("rtb_cf", "FEP cash flow – sale at RtB", "EUR k", "={rtb_pre}+{rtb_tax}", bold=True)
    c.xirr("rtb_irr", "FEP IRR – sale at RtB, after tax", "rtb_cf")
    c.xirr("rtb_irr_pre", "FEP IRR – sale at RtB, before tax", "rtb_pre")
    c.row("rtb_moic", "FEP MOIC – sale at RtB (receipts after tax ÷ outlays)", "x",
          e="({rtb_in.I}+{rtb_tax.I})/-{rtb_out.I}", total="none", ef_fmt="mult", bold=True)
    c.xnpv("rtb_npv", "FEP NPV – sale at RtB", "In_Discount_Rate", "rtb_cf")
    c.row("psh_r", "PSH cash flow – sale at RtB", "EUR k",
          "=-In_PSH_Land*{acq_ev}-In_PSH_Devex*{mwp.E}*{dsa_w}-{dsa}-In_EW_To_PSH*{ew_dev}+{w_psh_r.F}*{exit_ev}")
    c.xirr("psh_r_irr", "PSH IRR – sale at RtB", "psh_r", iferror=True)
    c.moic("psh_r_moic", "PSH MOIC – sale at RtB", "psh_r")

    c.section("Sale at COD (same JV split of the net proceeds)")
    c.row("cod_contrib", "FEP contributions to COD (equity, levered)", "EUR k", "={eq_cf}*(1-{post_c})-{dep_cod}")
    c.row("v_avail", "Cash available: debt repaid (E), price less debt (F)", "EUR k", e="{d_close.E}",
          f="MAX(0,{px_cod.F}-{self.E})", total="none")
    waterfall(c, "v", [("cost", "Sale costs", "In_Sale_Costs_COD"),
                       ("fund", "FEP funding returned", "-{cod_contrib.I}"),
                       ("ret", "FEP 12 % return to COD", "{cl_ret.F}"),
                       ("psh", "PSH land apport and return (v3 rule)", "In_PSH_Land_Step*(In_PSH_Land+{ps_ret.F})")])
    c.row("v_chk", "Check: paid + distributable − available (E), steps paid above due (F)", "EUR k",
          e="SUM({v_cost.F}:{v_psh.F})+{v_dist.F}-{v_avail.F}",
          f="SUMPRODUCT(--({v_cost.F}:{v_psh.F}>{v_cost.E}:{v_psh.E}+Tolerance))", total="none")
    c.row("v_fep", "FEP proceeds at COD (E: tax on the gain)", "EUR k", f="{v_fund.F}+{v_ret.F}+{shr.E}*{v_dist.F}",
          e="-In_CIT*In_Sale_Gain_Taxed*MAX(0,{self.F}+{cod_contrib.I})", total="none", bold=True)
    c.row("v_psh_r", "PSH receipts at COD", "EUR k", f="{v_psh.F}+{shr.F}*{v_dist.F}", total="none")
    c.row("cod_pre", "FEP cash flow – sale at COD, before tax on the gain", "EUR k",
          "={cod_contrib}+{v_fep.F}*{cod_ev}+{dep_cod}")
    c.row("cod_cf", "FEP cash flow – sale at COD", "EUR k", "={cod_pre}+{v_fep.E}*{cod_ev}", bold=True)
    c.xirr("cod_irr", "FEP IRR – sale at COD, levered, after tax", "cod_cf")
    c.xirr("cod_irr_pre", "FEP IRR – sale at COD, levered, before tax", "cod_pre")
    c.moic("cod_moic", "FEP MOIC – sale at COD", "cod_cf")
    c.xnpv("cod_npv", "FEP NPV – sale at COD", "In_Discount_Rate", "cod_cf")
    c.row("jv_c", "JV (100 %) cash flow to the COD sale, unlevered, after tax", "EUR k",
          "={proj_cf}*(1-{post_c})+({px_cod.F}-In_Sale_Costs_COD)*{cod_ev}")
    c.xirr("jv_c_irr", "JV IRR – sale at COD, unlevered, after tax", "jv_c")
    c.row("jv_c_pre", "JV (100 %) cash flow to the COD sale, unlevered, before tax", "EUR k",
          "={jv_c}-{tax_copy}*(1-{post_c})")
    c.xirr("jv_c_irr_pre", "JV IRR – sale at COD, unlevered, before tax", "jv_c_pre")
    c.row("psh_c", "PSH cash flow – sale at COD", "EUR k",
          "=-In_PSH_Land*{acq_ev}-In_PSH_Devex*{mwp.E}*{dsa_w}-{dsa}-In_EW_To_PSH*{ew_dev}+{v_psh_r.F}*{cod_ev}")
    c.xirr("psh_c_irr", "PSH IRR – sale at COD", "psh_c", iferror=True)

    c.section("Hold to maturity (FEP buys PSH out at the exit)")
    c.row("htm_k", "FEP share of equity flows after the exit (E), PSH buy-out at the exit (F)", "%",
          e="1-In_PSH_Stays*{shr.F}", f="(1-In_PSH_Stays)*{w_psh_r.F}", total="none", ef_fmt="pct")
    fmt("htm_k", "F", "money")
    c.row("htm_w", "FEP share of the equity flows", "%", "=1-{post_x}*(1-{htm_k.E})", total="none", fmt="pct")
    c.row("htm_cf", "FEP cash flow – hold to maturity", "EUR k", "={eq_cf}*{htm_w}-{htm_k.F}*{exit_ev}", bold=True)
    c.xirr("htm_irr", "FEP IRR – hold to maturity, levered, after tax", "htm_cf")
    c.row("htm_pre", "FEP cash flow – hold to maturity, before tax", "EUR k", "={htm_cf}-{tax_copy}*{htm_w}")
    c.xirr("htm_irr_pre", "FEP IRR – hold to maturity, levered, before tax", "htm_pre")
    c.moic("htm_moic", "FEP MOIC – hold to maturity", "htm_cf")
    c.xnpv("htm_npv", "FEP NPV – hold to maturity", "In_Discount_Rate", "htm_cf")
    c.row("htm_cap", "FEP capital invested – hold to maturity (outflows)", "EUR k", e='-SUMIF({htm_cf.all},"<0")',
          total="none")
    c.row("psh_h", "PSH cash flow – hold to maturity", "EUR k",
          "={psh_r}-{w_psh_r.F}*{exit_ev}+{htm_k.F}*{exit_ev}+(1-{htm_w})*{eq_cf}")
    c.xirr("psh_h_irr", "PSH IRR – hold to maturity", "psh_h", iferror=True)


def waterfall(c, p, steps):
    """One row per step: E = amount due, F = MIN(due, cash still available); distributable = cash above all dues
    (the Bora helper)."""
    first = f"{p}_{steps[0][0]}"
    prev = None
    for k, label, due in steps:
        rest = "{%s_avail.F}" % p + ("" if prev is None else "-SUM({%s.F}:{%s.F})" % (first, prev))
        c.row(f"{p}_{k}", label, "EUR k", e=due, f="MIN({self.E}," + rest + ")", total="none")
        prev = f"{p}_{k}"
    c.row(f"{p}_dist", "Distributable", "EUR k", f="MAX(0,{%s_avail.F}-SUM({%s.F}:{%s.F}))" % (p, first, prev),
          total="none", bold=True)


# =========================================================================================== macros: goal seek
def goal_seeks(m):
    """Asset deal (SPEC §9, task 2.5): the price feeds the buyer's depreciation, so GoalSeek.bas solves it per sheet."""
    i = m.inputs
    i.section("Solved by GoalSeek.bas (asset deal implied prices)")
    for topic, key, target, seed, label in [("Asset_RtB", "as_r_irr", "In_Buyer_IRR_RtB", 100, "RtB"),
                                            ("Asset_COD", "as_c_irr", "In_Buyer_IRR_COD", 1000, "COD")]:
        first = i.r
        for n, sheet in enumerate(m.calc_sheets, 1):
            m.goal_seek(f"{topic}_{n}", f"Asset-deal implied {label} price, {sheet}", "EUR k/MWp", seed,
                        m.calc.ref(key, "E", sheet), target, source="Solved by GoalSeek.bas", fmt="money1")
        m.define(f"{topic}_Inputs", f"Inputs!$F${first}:$F${i.r - 1}")


# =========================================================================================== checks
def checks(m):
    k = m.checks
    k.add_check("Revenue case, offers and BESS curve found",
                "=N(ISNA({case.F}))+N(ISNA({op.E}))+N(ISNA({ob.E}))+N(ISNA({aug_y.F}))", per_period=False)
    k.add_check("Revenue-case table complete for the selected case",
                "=--(COUNTBLANK(INDEX(Case_Table,0,{case.F}))>0)", per_period=False)
    k.add_check("COD falls on a quarter end", "=N(COUNTIF({end.all},{cod_ev.E})=0)", per_period=False)
    k.add_check("Debt at COD within the DSCR capacity", "=N({k.F}>1+Tolerance)", per_period=False)
    k.add_check("Debt repaid by the end of the tenor", "=N(({end}>={tenor.F})*ABS({d_close})>Tolerance)")
    k.add_check("DSCR not below the weighted target", "=N(({tenor}=1)*({cfads}*{w_dscr}+{d_service})<-Tolerance)")
    k.add_check("Sources = uses to COD", "=N(ABS({f_gear.F})>Tolerance)", per_period=False)
    k.add_check("PV energy: generation = direct export + BESS charging + clipped",
                "=N(ABS({gen}-{pv_exp}-{to_bess}+{clip})>Tolerance)+N({pv_exp}<-Tolerance)")
    k.add_check("BESS discharge ≤ charge × RTE", "=N({dis}>({to_bess}+{chg})*In_BESS_RTE+Tolerance)")
    k.add_check("Export ≤ grid capacity × hours", "=N({export}>{grid_h}+Tolerance)")
    k.add_check("RtB waterfall: paid + distributable = available, no step above its due",
                "=N(ABS({w_chk.E})>Tolerance)+N({w_chk.F}>0)", per_period=False)
    k.add_check("RtB: FEP + PSH + sale costs = consideration",
                "=N(ABS({w_fep.F}+{w_psh_r.F}+{w_cost.F}-{px_rtb.F})>Tolerance)", per_period=False)
    k.add_check("COD waterfall: paid + distributable = available, no step above its due",
                "=N(ABS({v_chk.E})>Tolerance)+N({v_chk.F}>0)", per_period=False)
    k.add_check("COD: FEP + PSH + sale costs = price less debt",
                "=N(ABS({v_fep.F}+{v_psh_r.F}+{v_cost.F}-{v_avail.F})>Tolerance)", per_period=False)
    k.add_check("XIRR returned a number: FEP routes",
                "=N(ISERROR({rtb_irr.E}))+N(ISERROR({cod_irr.E}))+N(ISERROR({htm_irr.E}))", per_period=False)
    k.add_check("XIRR returned a number: project and buyers",
                "=N(ISERROR({proj_irr.E}))+N(ISERROR({buy_r_irr.E}))+N(ISERROR({buy_c_irr.E}))", per_period=False)
    for key, label, t, _ in RECON:
        k.add_check(f"Scenario A = shadow, reconciliation set: {label}",
                    f"=(Scenario_Selected=Recon_Set)*N(ABS({t}-Recon_{key})>Recon_Tol_{key})",
                    per_period=False, sheet="Calc_1")


# =========================================================================================== dashboard
GLANCE = [  # subheader or (label, unit, template, scale, out, bold)
    "Sale at RtB",
    ("RtB price used", "EUR k/MW", "{px_rtb.E}", None, "RtB_Price", False),
    ("FEP profit", "EUR m", "{rtb_cf.I}", "Thousand", "RtB_Profit", False),
    ("FEP IRR, after tax (unlevered: no debt before NTP)", "%", "{rtb_irr.E}", None, "RtB_IRR", True),
    ("FEP IRR, before tax", "%", "{rtb_irr_pre.E}", None, "RtB_IRR_Pre", False),
    ("FEP MOIC", "x", "{rtb_moic.E}", None, "RtB_MOIC", False),
    ("FEP NPV", "EUR m", "{rtb_npv.E}", "Thousand", "RtB_NPV", False),
    ("PSH IRR", "%", "{psh_r_irr.E}", None, None, False),
    ("Buyer IRR at our price, unlevered, after tax", "%", "{buy_r_irr.E}", None, "Buyer_RtB_IRR", False),
    ("Buyer IRR at our price, unlevered, before tax", "%", "{buy_r_irr_pre.E}", None, "Buyer_RtB_IRR_Pre", False),
    ("Buyer IRR at our price, levered, after tax", "%", "{buy_r_irr_l.E}", None, "Buyer_RtB_IRR_L", False),
    ("Implied RtB price at the buyer target IRR (share deal)", "EUR k/MW", "{imp_r.E}", None, "Implied_RtB", True),
    ("Implied RtB price at the buyer target IRR (asset deal)", "EUR k/MW", "{imp_r.F}", None, "Implied_RtB_Asset",
     False),
    "Sale at COD",
    ("COD price used (enterprise value per MWp)", "EUR k/MW", "{px_cod.E}", None, "COD_Price", False),
    ("FEP profit", "EUR m", "{cod_cf.I}", "Thousand", "COD_Profit", False),
    ("FEP IRR, levered, after tax", "%", "{cod_irr.E}", None, "COD_IRR", True),
    ("FEP IRR, levered, before tax", "%", "{cod_irr_pre.E}", None, "COD_IRR_Pre", False),
    ("FEP MOIC", "x", "{cod_moic.E}", None, "COD_MOIC", False),
    ("FEP NPV", "EUR m", "{cod_npv.E}", "Thousand", "COD_NPV", False),
    ("JV IRR, unlevered, after tax", "%", "{jv_c_irr.E}", None, "COD_JV_IRR", False),
    ("JV IRR, unlevered, before tax", "%", "{jv_c_irr_pre.E}", None, "COD_JV_IRR_Pre", False),
    ("PSH IRR", "%", "{psh_c_irr.E}", None, None, False),
    ("Buyer IRR at our price, unlevered, after tax", "%", "{buy_c_irr.E}", None, "Buyer_COD_IRR", False),
    ("Buyer IRR at our price, unlevered, before tax", "%", "{buy_c_irr_pre.E}", None, "Buyer_COD_IRR_Pre", False),
    ("Buyer IRR at our price, levered, after tax", "%", "{buy_c_irr_l.E}", None, "Buyer_COD_IRR_L", False),
    ("Implied COD price at the buyer target IRR (share deal)", "EUR k/MW", "{imp_c.E}", None, "Implied_COD", True),
    ("Implied COD price at the buyer target IRR (asset deal)", "EUR k/MW", "{imp_c.F}", None, "Implied_COD_Asset",
     False),
    "Hold to maturity",
    ("FEP profit", "EUR m", "{htm_cf.I}", "Thousand", "HTM_Profit", False),
    ("FEP IRR, levered, after tax", "%", "{htm_irr.E}", None, "HTM_IRR", True),
    ("FEP IRR, levered, before tax", "%", "{htm_irr_pre.E}", None, "HTM_IRR_Pre", False),
    ("FEP MOIC", "x", "{htm_moic.E}", None, "HTM_MOIC", False),
    ("FEP NPV", "EUR m", "{htm_npv.E}", "Thousand", "HTM_NPV", False),
    ("Project IRR, unlevered, after the SPV's actual tax", "%", "{proj_irr.E}", None, "Project_IRR", False),
    ("Project IRR, unlevered, before tax", "%", "{proj_irr_pre.E}", None, "Project_IRR_Pre", False),
    ("PSH IRR", "%", "{psh_h_irr.E}", None, None, False),
]


def dashboard(m):
    d, S = m.dashboard, m.calc_sheets
    m2 = "money2"

    d.section("Levers (revenue case per sheet; everything else from Inputs, set by the key-assumption set)")
    d.input("Case", "Revenue case: A v3 reconstruction, B Aurora merchant, C offer", "text", ["A", "B", "C"],
            options="Case_Names")
    d.link("Case name", "text", "INDEX(Case_Labels,1,{case.F})", fmt="text")
    d.link("PV price curve", "text", "{pv_real.E}", fmt="text")
    d.link("BESS revenue curve (Aurora: fixed dispatch, no re-optimisation)", "text", "{stk_real.E}", fmt="text")
    d.link("PV offer / BESS offer (case C only)", "text", 'IF({offer.E}=1,In_Offer_PV&" / "&In_Offer_BESS,"–")',
           fmt="text")
    d.link("BESS built", "MW", "{bess_mw.E}", fmt="money1")
    d.link("BESS energy", "MWh", "{bess_mw.F}", fmt="money")
    d.link("Buyer target IRR at RtB (unlevered, after tax)", "%", "In_Buyer_IRR_RtB", fmt="pct")
    d.link("Futureal share of the distributable amount", "%", "{shr.E}", out="FEP_Share", fmt="pct")

    d.section("Returns at a glance")
    for g in GLANCE:
        if isinstance(g, str):
            d.subheader(g)
            continue
        label, unit, t, scale, out, bold = g
        f = "pct" if unit == "%" else ("mult" if unit == "x" else (m2 if scale else "money1"))
        d.link(label, unit, t, scale=scale, out=out, fmt=f, bold=bold)

    d.section("Capacity and energy")
    d.link("PV capacity", "MWp", "{mwp.E}", out="MWp", fmt="money1")
    d.link("Grid export capacity", "MVA", "{mwp.F}", fmt="money1")
    d.link("PV generation, first full year", "GWh", "{gen.E}", scale="Thousand", out="Gen_Year1", fmt="money1")
    d.link("PV clipped, lifetime", "GWh", "-{clip.I}", scale="Thousand", fmt="money1")
    d.link("Energy exported, first full year", "GWh", "{export.E}", scale="Thousand", fmt="money1")
    d.link("BESS discharge, lifetime", "GWh", "{dis.I}", scale="Thousand", fmt="money1")
    d.link("BESS full cycles per year (lifetime average)", "#", "{dis.I}/MAX(1,{bess_mw.F})/In_Life", fmt="money1")
    d.link("Revenue, first full year", "EUR m", "{rev.E}", scale="Thousand", out="Revenue_Year1", fmt=m2)
    d.link("of which BESS", "EUR m", "{bess_rev.E}", scale="Thousand", fmt=m2)
    d.link("EBITDA, first full year", "EUR m", "{ebitda.E}", scale="Thousand", out="EBITDA_Year1", fmt=m2)
    d.link("Lifetime revenue", "EUR m", "{rev.I}", scale="Thousand", out="Revenue_Life", fmt=m2)
    d.link("Lifetime EBITDA", "EUR m", "{ebitda.I}", scale="Thousand", out="EBITDA_Life", fmt=m2)

    d.section("Funding to COD (100 %)")
    for t, label, out, bold in [("{f_dev.E}", "Development, deposits and guarantees (net)", "Uses_Development", False),
                                ("{f_capex.E}", "Construction capex and early works at NTP", "Uses_Capex", False),
                                ("{f_fin.E}", "Capitalised interest and debt fees", "Uses_Financing", False),
                                ("{f_uses.E}", "Total uses", "Uses_Total", True),
                                ("{d_close.E}", "Senior debt at COD (incl. capitalised interest)", "Debt_COD", False),
                                ("{f_equity.E}", "Equity to COD", "Equity_COD", True),
                                ("{f_equity.F}", "Peak equity, hold (incl. DSRA)", "Equity_Peak", False)]:
        d.link(label, "EUR m", t, scale="Thousand", out=out, fmt=m2, bold=bold)
    d.link("Debt share of total uses", "%", "{f_gear.E}", out="Gearing_Uses", fmt="pct")

    d.section("Debt and tax")
    d.link("Debt capacity at the weighted target DSCR (PV at COD)", "EUR m", "{ds_cap.E}", scale="Thousand", fmt=m2)
    d.link("Debt share of construction cost, sized", "%", "{k.E}", out="Debt_Share", fmt="pct")
    d.link("Share of the DSCR capacity used", "%", "{k.F}", fmt="pct")
    d.link("DSRA funded at COD", "EUR m", "{dsra.E}", scale="Thousand", fmt=m2)
    d.link("Minimum DSCR", "x", "{dscr.E}", out="Min_DSCR", fmt="mult")
    d.link("Lifetime corporate income tax", "EUR m", "SUM({tax_copy.all})", scale="Thousand", out="Tax_Life", fmt=m2)
    d.link("Tax losses lapsed", "EUR m", "-{lcf_lapse.I}", scale="Thousand", fmt=m2)

    d.section("JV waterfall at the RtB exit (paid)")
    for t, label, bold in [("{w_avail.E}", "Consideration (price × MWp)", True),
                           ("-{w_cost.F}", "Less sale costs", False),
                           ("-{w_spv.F}", "Less FEP SPV acquisition price", False),
                           ("-{w_dsa.F}", "Less DSA loan principal", False),
                           ("-{w_ret.F}", "Less FEP 12 % return", False),
                           ("-{w_bg.F}", "Less guarantee fees paid by FEP", False),
                           ("-{w_psh.F}", "Less PSH land apport and return (v3 rule)", False),
                           ("{w_dist.F}", "Distributable", True),
                           ("{w_fep.F}", "FEP receipts at the exit", True),
                           ("{w_psh_r.F}", "PSH receipts at the exit", True),
                           ("{reimb_x.I}+{pre_reimb.I}", "Memo: buyer reimbursements to FEP", False)]:
        d.link(label, "EUR m", t, scale="Thousand", fmt=m2, bold=bold)
    d.section("JV waterfall at COD (paid)")
    for t, label, bold in [("{px_cod.F}", "Price (enterprise value)", True),
                           ("-{v_avail.E}", "Less debt repaid", False),
                           ("-{v_cost.F}", "Less sale costs", False),
                           ("-{v_fund.F}", "Less FEP funding returned", False),
                           ("-{v_ret.F}", "Less FEP 12 % return to COD", False),
                           ("-{v_psh.F}", "Less PSH land apport and return (v3 rule)", False),
                           ("{v_dist.F}", "Distributable", True),
                           ("{v_fep.F}", "FEP proceeds at COD", True),
                           ("{v_psh_r.F}", "PSH receipts at COD", True)]:
        d.link(label, "EUR m", t, scale="Thousand", fmt=m2, bold=bold)
    d.link("Checks failed in this set (0 = OK)", "#", "Model_Checks", out="Checks", fmt="int")

    d.scenario_results(outs=["RtB_Profit", "RtB_IRR", "RtB_MOIC", "Implied_RtB", "COD_IRR", "COD_MOIC", "Implied_COD",
                             "HTM_IRR", "HTM_MOIC", "HTM_NPV", "Project_IRR", "Buyer_RtB_IRR", "Debt_COD",
                             "Min_DSCR", "Revenue_Life", "Checks"])

    d.panel_section("Key-assumption set")
    d.selector()
    d.panel_section("Key assumptions (live)")
    for label, unit, ref, f in [
            ("RtB price (0 = implied)", "EUR k/MW", "In_RtB_Price", "money1"),
            ("COD price (0 = implied)", "EUR k/MW", "In_COD_Price", "money1"),
            ("Buyer target IRR, RtB", "%", "In_Buyer_IRR_RtB", "pct"),
            ("Buyer target IRR, COD", "%", "In_Buyer_IRR_COD", "pct"),
            ("Deal form (0 share, 1 asset)", "flag", "In_Sale_Form", "flag"),
            ("PV merchant price change", "%", "In_PV_Price_Adj", "pct"),
            ("Yield, cases B and C", "MWh/MWp", "In_Yield", "money"),
            ("BESS power", "MW", "In_BESS_MW", "money1"), ("BESS duration", "h", "In_BESS_Hours", "money1"),
            ("BESS built", "flag", "In_BESS_On", "flag"), ("BESS haircut", "%", "In_BESS_Haircut", "pct"),
            ("PV offer (case C)", "text", "In_Offer_PV", "text"), ("BESS offer (case C)", "text", "In_Offer_BESS", "text"),
            ("Gearing cap", "%", "In_Gearing", "pct"), ("Debt rate", "%", "In_Debt_Rate", "pct"),
            ("DSCR contracted / merchant", "x", "In_DSCR_Merchant", "mult"), ("CIT", "%", "In_CIT", "pct"),
            ("Discount rate", "%", "In_Discount_Rate", "pct"), ("FX", "HUF/EUR", "In_FX", "money"),
            ("COD, cases B and C", "date", "In_COD", "date"), ("RtB exit, cases B and C", "date", "In_RtB_Exit", "date"),
            ("Share rule (1 step, 0 v3)", "flag", "In_Share_Rule", "flag"),
            ("Early works lines in development", "#", "In_Milestones_Scope", "int")]:
        d.panel_link(label, unit, ref, fmt=f)
    d.panel_section("Not re-optimised in Excel")
    for t in ["Aurora's BESS revenue is a fixed annual series per MW",
              "for one duration, cycle cap, RTE 0.86, availability",
              "0.98, DC co-location 1:1:1 and entry year. MW",
              "scales it; hours/cycles switch blocks; RTE, grid",
              "cap or charging rules act only via the haircut,",
              "availability and augmentation inputs."]:
        d.ws.cell(d.pr, d.P, t)
        d.pr += 1


# =========================================================================================== build
def model():
    m = Model("Project Zalahaláp", "63 MWp PV + 50 MW / 200 MWh BESS – sale at RtB, sale at COD or hold",
              periods=PERIODS, periodicity="quarterly", start=START, scenarios=SETS, version="v1",
              date=D(2026, 10, 22), curves=True, layout="dashboard", calc_sheets=["Calc_1", "Calc_2", "Calc_3"],
              bids=["A Original", "B Aurora", "C Offer"], extra_cols=len(SETS), source_col=False)
    inputs(m)
    curves(m)
    calc(m)
    goal_seeks(m)
    checks(m)
    dashboard(m)
    return m


def build_model(out: Path) -> Path:
    m = model()
    _fix_set_formulas(m)
    xlsx = m.save(out / f"{NAME}.xlsx")
    for (key, col), f in PATCH.items():          # second format for rows whose E and F hold different units
        for ws in m.calc.sheets:
            ws[f"{col}{m.calc.rows[key].r}"].number_format = FMT.get(f, f)
    dash = m.wb["Dashboard"]                     # long text results (case and curve names) shrink to the column
    for row in dash.iter_rows(min_col=5, max_col=7):
        for c in row:
            if c.number_format == "@":
                c.alignment = Alignment(horizontal="right", shrink_to_fit=True)
    m.wb.save(xlsx)
    xlsm = package(xlsx, out / f"{NAME}.xlsm", run=["Scenarios.RunAllScenarios"], include=["Scenarios"])
    xlsx.unlink()
    return xlsm


def build(out: Path):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    return {"model": build_model(out)}
