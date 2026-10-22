"""Extract the Aurora Q2-2026 Hungary flexible-energy curves that the Zalahaláp model uses (SPEC §6.2).

    python3 projects/zalahalap/data/aurora_extract.py  →  projects/zalahalap/data/aurora_curves.csv

Source: reference/zalahalap/Aurora_Q2_26_HUN_Flexible_Energy_Data_AG.xlsm (release 6-May-2026), EUR real 2025,
calendar-year averages 2026–2060, Central scenario only. One CSV row per series, with the sheet and row it comes from:
  - Central prices row 20 (baseload) and Common Assumptions row 14 (euro inflation index, 2025 = 1);
  - per battery block of sheet "Degraded batteries" (block start s, rows 53/96/139/182/225/268):
      storage total cashflow (s+27, EUR/kW/yr = k€/MW/yr; already nets charging, PV energy stored and network charges),
      its components (s+14 … s+26, information only), total discharge (s+11), charge from grid (s+12),
      charge from RES (s+13), repower count (s+1), RES-to-grid and RES-to-storage volumes (s+31, s+32), RES
      day-ahead and total cashflow (s+34, s+37);
  - "implied PV capture" per co-located block = RES total cashflow (s+37: day-ahead + intraday + the transfer for
    energy stored) × 1000 ÷ RES output (to grid s+31 + to storage s+32), EUR/MWh: the value of every MWh the PV
    produces, consistent with the storage total, which already deducts the PV energy stored (s+20). The SPEC's
    literal ratio, RES day-ahead cashflow (s+34) ÷ RES-to-grid volume (s+31), is kept as "da_to_grid": applied to
    all PV output it would count the stored energy twice.
Aurora's battery blocks are fixed dispatch outputs: nothing here re-optimises when model inputs change.
Stdlib + openpyxl.
"""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
SRC = HERE.parents[2] / "reference" / "zalahalap" / "Aurora_Q2_26_HUN_Flexible_Energy_Data_AG.xlsm"
OUT = HERE / "aurora_curves.csv"
SHEET = "Degraded batteries"
BLOCKS = {  # block start row -> model name
    53: "2027 solar DC 1.5cy 2h", 96: "2027 solar DC 1.0cy 2h", 139: "2027 solar DC 1.0cy 4h",
    182: "2027 standalone 1.0cy 2h", 225: "2030 solar DC 1.5cy 2h", 268: "2030 solar DC 1.0cy 4h"}
SERIES = [  # offset, key, unit
    (27, "storage_total", "k€/MW/yr"), (11, "discharge", "MWh/MW"), (12, "charge_grid", "MWh/MW"),
    (13, "charge_res", "MWh/MW"), (1, "repower_count", "#"), (31, "res_to_grid", "MWh/MW RES"),
    (32, "res_to_storage", "MWh/MW RES"), (34, "res_da_cashflow", "k€/MW RES/yr"),
    (37, "res_total_cashflow", "k€/MW RES/yr")]
COMPONENTS = range(14, 27)                       # s+14 … s+26: components of the storage total (information only)
FIRST_COL, YEARS = 6, list(range(2026, 2061))    # F = 2026 … AN = 2060


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True, keep_vba=False)
    rows = []

    def take(sheet, r, key, block, unit, c0=FIRST_COL):
        ws = wb[sheet]
        vals = [ws.cell(r, c0 + k).value or 0.0 for k in range(len(YEARS))]
        rows.append([key, block, unit, f"{sheet}!{r}", ws.cell(r, 4).value or ws.cell(r, 3).value] + vals)
        return vals

    ws = wb["Central prices"]
    assert [ws.cell(19, 6 + k).value for k in range(2)] == [2026, 2027]
    take("Central prices", 20, "baseload", "Central", "EUR/MWh")
    ws = wb["Common Assumptions"]                 # E = 2025, F = 2026 …
    assert ws.cell(7, 6).value == 2026
    take("Common Assumptions", 14, "euro_index", "Central", "index 2025=1")
    deg = wb[SHEET]
    for s, name in BLOCKS.items():
        assert str(deg.cell(s, 2).value).startswith(name[:4]), (s, deg.cell(s, 2).value)
        got = {}
        for off, key, unit in SERIES:
            got[key] = take(SHEET, s + off, key, name, unit)
        for off in COMPONENTS:
            take(SHEET, s + off, f"component_{off}", name, "k€/MW/yr")
        if "standalone" not in name:
            out = [g + t for g, t in zip(got["res_to_grid"], got["res_to_storage"])]
            cap = [c * 1000 / vol if vol else 0.0 for c, vol in zip(got["res_total_cashflow"], out)]
            rows.append(["implied_capture", name, "EUR/MWh", f"{SHEET}!{s + 37}/({s + 31}+{s + 32})",
                         "RES total cashflow x 1000 / RES output (to grid + to storage)"] + cap)
            da = [c * 1000 / vol if vol else 0.0 for c, vol in zip(got["res_da_cashflow"], got["res_to_grid"])]
            rows.append(["da_to_grid", name, "EUR/MWh", f"{SHEET}!{s + 34}/{s + 31}",
                         "RES Day-Ahead cashflow x 1000 / RES-to-grid volume"] + da)
    with open(OUT, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["series", "block", "unit", "source", "label"] + YEARS)
        w.writerows(rows)
    print(f"wrote {OUT} ({len(rows)} series)")


if __name__ == "__main__":
    main()
