"""Dump the original Zalahalap model (Zalahalap_model_10112025_v3.xlsb) into the old-engine extracts (SPEC §3).

Writes, next to this script:
  old_inputs.json   every §3.1 input: sheet!cell, value, unit, label, formula (links show their source)
  old_curves.csv    Aurora_Curve (annual 2025-2070, the curves the original model uses) + IMF inflation
  old_targets.csv   the §3.2 reconciliation targets: cached value, value in the C8-patched recalc, the target used
                    ([R] = from the patched copy) and the SPEC figure with a tie flag

Steps:
  1. soffice converts the xlsb to xlsx in --work (skipped when the xlsx is already there);
  2. fepm.recalc writes the patched copy with Aurora_Curve!C8 = 0 (the cell holds the text "no ", which breaks the
     HTM chain; SPEC §3.2) and recalculates it in LibreOffice;
  3. values are read from the conversion (Excel's cached values) and, when pyxlsb is importable, every dumped cell
     is cross-checked against the xlsb itself.

Usage: python3 projects/zalahalap/data/old_engine/dump_old.py [--work DIR] [--xlsb PATH]
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as col_idx, get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
sys.path.insert(0, str(ROOT))
from fepm.recalc import recalc, soffice  # noqa: E402

XLSB = ROOT / "reference" / "zalahalap" / "Zalahalap_model_10112025_v3.xlsb"
PATCH = {"Aurora_Curve!C8": 0}
XLSB_ERRORS = {"0x0": "#NULL!", "0x7": "#DIV/0!", "0xf": "#VALUE!", "0x17": "#REF!", "0x1d": "#NAME?",
               "0x24": "#NUM!", "0x2a": "#N/A"}   # pyxlsb returns error cells as hex codes
LABELS = {"Control": "BC", "Stat_inp": "CD", "HTM_Model": "BC", "Calculation": "CD", "RTB_Model": "BC",
          "Dyn_Inp": "CD", "Aurora_Curve": "BC"}            # sheet -> label column, unit column

# ---- SPEC §3.1 inputs: group -> [(key, cell, overrides)] -----------------------------------------------------
INPUTS = {
    "switches": [
        ("strategy", "Control!D5", {}),
        ("aurora_scenario", "Control!D6", {"note": "1 Low / 2 Central / 3 High; Low/High = Central x (1 -/+ D2:D3)"}),
        ("capture_basis", "Control!D7", {}),
        ("psh_buyout_in_htm", "Control!D8", {}),
        ("inflation_index", "Control!D13", {}),
        ("price_sensitivity", "Control!D10", {}), ("capex_sensitivity", "Control!D11", {}),
        ("yield_sensitivity", "Control!D12", {}),
    ],
    "capacity": [
        ("pv_mwp", "Control!D15", {}), ("grid_mva", "Control!D16", {}), ("conversion_rate", "Control!D17", {}),
        ("mw_at_rtb", "Control!D18", {}), ("grid_intake", "Control!D90", {}),
    ],
    "bess": [
        ("switch", "Control!D89", {}), ("mw", "Control!D91", {}), ("hours", "Control!D92", {}),
        ("mwh", "Control!D93", {}), ("capex_per_mwh", "Control!D94", {}), ("capex_total", "Control!D95", {}),
        ("fixed_opex", "Control!D96", {}), ("network_charge_retention", "Control!D97", {}),
        ("no_grid_discharge_scaling", "Control!D98", {"note": "SPEC §13 fix 6: uses the 2 h rows 62-63"}),
        ("pv_energy_valued", "Control!D99", {}), ("augment_cost_per_mwh", "Control!D100", {}),
        ("augment_year", "Control!D101", {}),
        ("block_row_offset", "Aurora_Curve!G4", {"label": "BESS block row offset in Aurora_Raw (26 = 4 h)",
                                                 "unit": "rows"}),
    ],
    "yield": [
        ("base_yield", "Control!D54", {}), ("yield_in_use", "Control!D55", {}), ("availability", "Control!D56", {}),
        ("degradation", "Control!D57", {}),
    ],
    "dates": [
        ("acquisition", "Control!D19", {}), ("dev_months", "Control!D20", {}), ("rtb", "Control!D21", {}),
        ("rtb_to_sale_months", "Control!D22", {}), ("acq_to_sale_months", "Control!D23", {}),
        ("sale", "Control!D24", {}), ("construction_start", "Control!D25", {}),
        ("construction_start_override", "Control!D26", {}), ("construction_months", "Control!D27", {}),
        ("construction_end", "Control!D28", {}), ("cod", "Control!D29", {}), ("life_years", "Control!D30", {}),
        ("operations_end", "Control!D31", {}), ("htm_timeline_start", "Control!D32", {}),
    ],
    "capex": [
        ("pv_total_per_mwp", "Control!D52", {}), ("early_works_in_capex", "Control!D53", {}),
        ("pv_equipment_per_mwp", "Control!D82", {}), ("bop_per_mwp", "Control!D83", {}),
        ("mv_cable", "Control!D104", {}), ("hv_substation", "Control!D105", {}), ("land_to_cod", "Control!D106", {}),
        ("decommissioning_per_mwp", "Control!D103", {}),
        ("total_construction", "HTM_Model!D147", {}),
        ("early_works_total", "Calculation!F60", {"label": "Successful DEVEX (early works) total"}),
        ("milestones_per_mw", "Calculation!F30", {"unit": "EUR k/MW"}),
        ("pre_rtb_certain_1", "Control!D109", {}), ("pre_rtb_certain_2", "Control!D110", {}),
        ("pre_rtb_uncertain", "Control!D111", {}), ("pre_rtb_uncertain_switch", "Control!D112", {}),
        ("pre_rtb_registry_fee", "Control!D113", {}), ("pre_rtb_fence_road", "Control!D114", {}),
        ("pre_rtb_total", "Control!D115", {}), ("pre_rtb_first_month", "Control!D116", {}),
        ("pre_rtb_last_month", "Control!D117", {}), ("pre_rtb_reimbursed", "Control!D118", {}),
    ],
    "milestones": [(f"m{i - 43}", f"Stat_inp!E{i}", {}) for i in range(44, 50)]
                  + [(f"m{i - 52}_month", f"Stat_inp!E{i}", {}) for i in range(53, 59)],
    "opex": [
        ("om_per_mw", "Control!D59", {}), ("balancing_per_mwh", "Control!D60", {}), ("land_lease", "Control!D61", {}),
        ("mv_cable_om_per_mw", "Control!D62", {}), ("route_to_market", "Control!D77", {}),
        ("days_receivable", "Control!D63", {}), ("days_payable", "Control!D64", {}),
        ("real_base_year", "Control!D68", {}),
    ],
    "revenue": [
        ("flat_capture_rate", "Aurora_Curve!D4", {}), ("low_adjustment", "Aurora_Curve!D2", {}),
        ("high_adjustment", "Aurora_Curve!D3", {}), ("goo_real", "Control!D58", {}),
    ],
    "rtb_deal": [
        ("active_case", "Stat_inp!E17", {}), ("active_case_name", "Stat_inp!E18", {}),
        ("fx", "Stat_inp!E23", {}), ("rtb_premium", "Stat_inp!E24", {}),
        ("futureal_share", "Stat_inp!E36", {"note": "case 4 (J36) is price-driven: 60 % below 90, "
                                                    "50 % + (100 - price)/100 up to 100, else 50 %"}),
        ("connection_rights", "Stat_inp!E33", {}), ("transaction_costs", "Stat_inp!E37", {}),
        ("psh_land_apport", "Stat_inp!E34", {}), ("investment_date", "Stat_inp!E39", {}),
        ("acq_cost_accruing", "Stat_inp!E66", {}), ("milestones_accruing", "Stat_inp!E67", {}),
        ("pref_rate", "Stat_inp!E69", {}), ("euribor_curve", "Stat_inp!E109", {}), ("bg_margin", "Stat_inp!E110", {}),
        ("deposit_profile", "Stat_inp!E99", {}),
    ] + [(f"deposit_{i - 83}_hufk_per_mva", f"Stat_inp!E{i}", {}) for i in range(84, 89)]
      + [(f"deposit_{i - 91}", f"Stat_inp!E{i}", {}) for i in range(92, 97)]
      + [("deposits_total", "Stat_inp!E97", {})]
      + [(f"deposit_{i - 101}_cash", f"Stat_inp!E{i}", {}) for i in range(102, 107)]
      + [(f"developer_m{i - 73}", f"Stat_inp!E{i}", {}) for i in range(74, 79)]
      + [("developer_total", "Stat_inp!E79", {})],
    "htm": [
        ("psh_buyout", "Control!D79", {}), ("buyer_price", "Control!D70", {}),
        ("buyer_price_override", "Control!D71", {}), ("buyer_reimburses_early_works", "Control!D72", {}),
        ("deposit_release", "Control!D73", {}), ("bg_rate_after_sale", "Control!D74", {}),
        ("guaranteed_deposits_after_sale", "Control!D75", {"note": "SPEC §13 fix 3: Calculation row 122"}),
        ("cash_deposits_at_sale", "Control!D76", {}),
    ],
    "tax": [("cit", "Control!D66", {}), ("depreciation_years", "Control!D65", {}), ("loss_cap", "Control!D102", {})],
    "valuation": [
        ("npv_rate", "Control!D9", {}),
    ] + [(f"buyer_irr_axis_{i - 70}", f"Sensitivities!D{i}", {"label": "Target buyer IRR (sensitivity axis)",
                                                               "unit": "%"}) for i in range(71, 76)],
}
SERIES = [   # monthly / quarterly rows kept with their dates: key, label, values row, dates row (None = index)
    ("euribor", "1M EURIBOR curve in use (Nov-25), monthly", "Dyn_Inp!G50:Y50", "Dyn_Inp!G49:Y49"),
    ("rtb_flag", "RTB flag by month", "Dyn_Inp!G7:Y7", "Dyn_Inp!G6:Y6"),
    ("sale_flag", "Project sale flag by month", "Dyn_Inp!G8:Y8", "Dyn_Inp!G6:Y6"),
] + [(f"deposit_{r - 21}_profile", f"Deposit profile (selected), Dyn_Inp row {r}", f"Dyn_Inp!G{r}:Y{r}",
      "Dyn_Inp!G6:Y6") for r in range(22, 27)] + [
    ("capex_phasing", "Construction capex phasing by construction month (buyer model)", "Buyer_Inp!E6:AB6",
     "Buyer_Inp!E5:AB5"),
]

# ---- Aurora_Curve columns -> curve names (rows 8:53 = 2025-2070) -------------------------------------------
CURVES = {"C": "baseload_central_real2025", "D": "baseload_low_real2025", "E": "baseload_central_copy_real2025",
          "F": "baseload_high_real2025", "G": "baseload_selected_real2025", "H": "hybrid_solar_to_grid_mwh_per_mw",
          "I": "hybrid_capture_price_real2025", "J": "hybrid_capture_rate", "K": "capture_rate_applied",
          "L": "pv_capture_selected_real2025", "M": "aurora_euro_index", "N": "imf_buyer_index",
          "O": "inflation_index_selected", "P": "pv_capture_selected_nominal", "Q": "goo_nominal",
          "R": "baseload_selected_nominal", "S": "aurora_data_year_flag", "T": "bess_retention",
          "U": "bess_discharge_mwh_per_mw", "V": "bess_gross_discharge_eur_kw", "W": "bess_ancillary_eur_kw",
          "X": "bess_grid_charging_eur_kw", "Y": "bess_network_charges_eur_kw",
          "Z": "bess_charge_from_pv_mwh_per_mw", "AA": "bess_cost_pv_stored_eur_kw",
          "AB": "bess_net_real2025_eur_kw", "AC": "bess_net_nominal_eur_kw"}

# ---- SPEC §3.2 targets: key, label, unit, cell, SPEC figure in the cell's unit, decimals shown in the SPEC -----
TARGETS = [
    ("rtb_xirr_incl", "RtB Futureal XIRR incl. pre-RtB reimbursement", "%", "Summary!D6", 2.245, 3),
    ("rtb_xirr_excl", "RtB Futureal XIRR excl. pre-RtB reimbursement", "%", "Calculation!F144", 1.818, 3),
    ("rtb_moic_incl_dep", "RtB MOIC incl. deposits", "x", "Summary!D7", 1.427, 3),
    ("rtb_moic_excl_dep", "RtB MOIC excl. deposits", "x", "Summary!D25", 1.576, 3),
    ("calc_moic_incl_dep", "RtB MOIC incl. deposits, Calculation basis", "x", "Calculation!G179", 1.795, 3),
    ("calc_moic_excl_dep", "RtB MOIC excl. deposits, Calculation basis", "x", "Calculation!G180", 2.529, 3),
    ("rtb_net_profit", "RtB Futureal net profit", "EUR m", "Summary!D8", 2.941, 3),
    ("rtb_capital", "RtB Futureal capital", "EUR m", "Summary!D9", 6.882, 3),
    ("rtb_proceeds", "RtB Futureal cash proceeds", "EUR m", "Summary!D23", 9.824, 3),
    ("developer_profit", "RtB developer net profit", "EUR k", "Calculation!F226", 2430, 0),
    ("rtb_breakeven", "RtB break-even price for 1.0x MOIC (Summary!D28 points at D75: bug)", "EUR k/MW",
     "RTB_Model!D78", 9.63, 2),
    ("htm_xirr", "HTM Futureal XIRR", "%", "Summary!H6", 0.0550, 4),
    ("htm_moic", "HTM Futureal MOIC", "x", "Summary!H7", 1.959, 3),
    ("htm_npv", "HTM Futureal NPV at 8 %", "EUR m", "Summary!H8", -13.01, 2),
    ("htm_capital", "HTM Futureal capital", "EUR m", "Summary!H9", 85.42, 2),
    ("htm_payback", "HTM payback from acquisition", "years", "Summary!H10", 16.8, 1),
    ("project_irr", "Project IRR (unlevered, post-CIT)", "%", "Summary!D31", 0.0637, 4),
    ("buyer_irr", "Buyer unlevered IRR at 90 k/MW", "%", "Summary!D33", 0.0557, 4),
    ("buyer_moic", "Buyer MOIC", "x", "Summary!D34", 1.986, 3),
    ("buyer_npv", "Buyer NPV at 8 %", "EUR m", "Summary!D36", -12.52, 2),
    ("ffy_gwh", "First full year (2029) net generation", "GWh", "Summary!H21", 75.61, 2),
    ("ffy_price", "First full year PV capture price, nominal", "EUR/MWh", "Summary!H22", 72.56, 2),
    ("ffy_revenue", "First full year revenue", "EUR m", "Summary!H24", 8.849, 3),
    ("life_pv_revenue", "Lifetime PV power revenue", "EUR m", "Summary!H25", 233.19, 2),
    ("life_goo", "Lifetime GoO revenue", "EUR m", "Summary!H26", 6.71, 2),
    ("life_revenue", "Total lifetime revenue", "EUR m", "Summary!H27", 349.82, 2),
    ("life_ebitda", "Lifetime EBITDA", "EUR m", "Summary!H30", 177.87, 2),
    ("life_margin", "Lifetime EBITDA margin", "%", "Summary!H31", 0.508, 3),
    ("life_capex", "Total construction capex", "EUR m", "Summary!H33", 57.17, 2),
    ("life_opex", "Total lifetime opex", "EUR m", "Summary!H53", 171.95, 2),
    ("life_bess_net", "Lifetime BESS net revenue", "EUR m", "Summary!H59", 109.92, 2),
]


# ---- workbook access ---------------------------------------------------------------------------------------
def split(ref):
    sheet, cell = ref.split("!")
    return sheet, cell


def cells(rng):
    """'Dyn_Inp!G50:Y50' -> ['Dyn_Inp!G50', ...] (one row)."""
    sheet, span = split(rng)
    a, b = span.split(":")
    r = int("".join(ch for ch in a if ch.isdigit()))
    c1, c2 = col_idx(a.rstrip("0123456789")), col_idx(b.rstrip("0123456789"))
    return [f"{sheet}!{L(c)}{r}" for c in range(c1, c2 + 1)]


class Book:
    """Formulas and cached values of one xlsx."""

    def __init__(self, path, formulas=True):
        self.v = openpyxl.load_workbook(path, data_only=True)
        self.f = openpyxl.load_workbook(path) if formulas else None

    def value(self, ref):
        sheet, cell = split(ref)
        return self.v[sheet][cell].value

    def formula(self, ref):
        sheet, cell = split(ref)
        x = self.f[sheet][cell].value
        if isinstance(x, ArrayFormula):
            return x.text
        return x if isinstance(x, str) and x.startswith("=") else None


def xlsb_values(path, refs):
    """Cached values of `refs` read from the xlsb itself (None if pyxlsb is not installed)."""
    try:
        from pyxlsb import open_workbook
    except ImportError:
        return None
    want = {}
    for ref in refs:
        sheet, cell = split(ref)
        col = col_idx(cell.rstrip("0123456789"))
        want.setdefault(sheet, {})[(int(cell[len(cell.rstrip("0123456789")):]) - 1, col - 1)] = ref
    out = {}
    with open_workbook(str(path)) as wb:
        for sheet, pos in want.items():
            with wb.get_sheet(sheet) as ws:
                for row in ws.rows(sparse=True):
                    for c in row:
                        if (c.r, c.c) in pos:
                            out[pos[(c.r, c.c)]] = XLSB_ERRORS.get(c.v, c.v)
    return out


def plain(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    return v


def serial(v):
    """Comparable number for a date (Excel serial) or the value itself."""
    if isinstance(v, dt.datetime):
        return (v - dt.datetime(1899, 12, 30)).days + (v - dt.datetime(1899, 12, 30)).seconds / 86400
    return v


def same(a, b):
    a, b = serial(a), serial(b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))
    return (a in (None, "") and b in (None, "")) or str(a).strip() == str(b).strip()


def label_unit(book, ref):
    sheet, cell = split(ref)
    cols = LABELS.get(sheet)
    if not cols:
        return None, None
    r = cell[len(cell.rstrip("0123456789")):]
    lab, unit = (book.value(f"{sheet}!{c}{r}") for c in cols)
    return (lab.strip() if isinstance(lab, str) else lab), unit


def case_values(book, ref):
    """Stat_inp column E picks the active case from G:Q; keep the non-blank case cells."""
    sheet, cell = split(ref)
    if sheet != "Stat_inp" or not cell.startswith("E"):
        return None
    r = cell[1:]
    out = {}
    for c in range(col_idx("G"), col_idx("Q") + 1):
        name = book.value(f"Stat_inp!{L(c)}18")
        v = book.value(f"Stat_inp!{L(c)}{r}")
        if v is not None:
            f = book.formula(f"Stat_inp!{L(c)}{r}")
            out[f"{L(c)} {name}"] = {"value": plain(v), **({"formula": f} if f else {})}
    return out or None


# ---- extracts ----------------------------------------------------------------------------------------------
def dump_inputs(book, xlsb_name):
    data = {"meta": {
        "source": f"reference/zalahalap/{xlsb_name} (cached values via soffice xlsx conversion)",
        "units": "as in the workbook: EUR k unless stated; dates ISO; % as fractions",
        "active_case": f"Stat_inp!E17 = {book.value('Stat_inp!E17')} ({book.value('Stat_inp!E18')}); "
                       "Stat_inp column E = IF(ISBLANK(INDEX(G:Q,E17)),G,INDEX(G:Q,E17))",
        "script": "projects/zalahalap/data/old_engine/dump_old.py",
    }}
    refs = []
    for group, items in INPUTS.items():
        data[group] = {}
        for key, ref, extra in items:
            lab, unit = label_unit(book, ref)
            e = {"value": plain(book.value(ref)), "unit": extra.get("unit", unit),
                 "cell": ref, "label": extra.get("label", lab)}
            f = book.formula(ref)
            if f:
                e["formula"] = f
            cases = case_values(book, ref)
            if cases:
                e["cases"] = cases
            if "note" in extra:
                e["note"] = extra["note"]
            data[group][key] = e
            refs.append(ref)
    data["series"] = {}
    for key, label, vals, dates in SERIES:
        vs, ds = cells(vals), cells(dates)
        data["series"][key] = {"label": label, "cells": vals, "dates_cells": dates,
                               "dates": [plain(book.value(d)) for d in ds],
                               "values": [book.value(v) for v in vs]}
        refs += vs
    (HERE / "old_inputs.json").write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n")
    return refs


def dump_curves(book):
    refs = []
    with open(HERE / "old_curves.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["year", *CURVES.values(), "source_cells"])
        for r in range(8, 54):
            row = [book.value(f"Aurora_Curve!B{r}")]
            for c in CURVES:
                row.append(book.value(f"Aurora_Curve!{c}{r}"))
                refs.append(f"Aurora_Curve!{c}{r}")
            note = f"Aurora_Curve!C{r}:AC{r}"
            if r == 8:
                note += " (C8 holds the text 'no '; the patched copy sets it to 0)"
            w.writerow(row + [note])
    return refs


def dump_targets(book, patched):
    rows, refs = [], []
    for key, label, unit, ref, spec, dec in TARGETS:
        cached, fixed = book.value(ref), patched.value(ref)
        use_patch = not same(cached, fixed)
        target = fixed if use_patch else cached
        ok = isinstance(target, (int, float)) and abs(target - spec) <= 0.5 * 10 ** -dec + 1e-12
        rows.append([key, label, unit, ref, book.formula(ref) or "", cached, fixed, "R" if use_patch else "cached",
                     target, spec, "yes" if ok else "NO"])
        refs.append(ref)
    with open(HERE / "old_targets.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["key", "output", "unit", "cell", "formula", "cached_value", "patched_value", "basis",
                    "target", "spec_value", "ties_to_spec"])
        w.writerows(rows)
    return refs, rows


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--xlsb", type=Path, default=XLSB)
    ap.add_argument("--work", type=Path, default=Path(tempfile.gettempdir()) / "zhlp_old_engine")
    a = ap.parse_args()
    a.work.mkdir(parents=True, exist_ok=True)
    xlsx = a.work / (a.xlsb.stem + ".xlsx")
    if not xlsx.exists():
        subprocess.run([soffice(), "--headless", f"-env:UserInstallation=file://{a.work}/lo_profile",
                        "--convert-to", "xlsx", "--outdir", str(a.work), str(a.xlsb)], check=True,
                       capture_output=True, timeout=600)
    patched_path = a.work / (a.xlsb.stem + "_C8patched.xlsx")
    if not patched_path.exists():
        recalc(xlsx, patched_path, inputs=PATCH)
    book, patched = Book(xlsx), Book(patched_path, formulas=False)

    refs = dump_inputs(book, a.xlsb.name) + dump_curves(book)
    trefs, rows = dump_targets(book, patched)
    refs += trefs

    raw = xlsb_values(a.xlsb, refs)
    if raw is None:
        print("pyxlsb not installed: xlsb cross-check skipped")
    else:
        bad = [r for r in refs if not same(raw.get(r), book.value(r))]
        print(f"xlsb cross-check: {len(refs) - len(bad)} of {len(refs)} cells equal")
        for r in bad:
            print(f"  differs {r}: xlsb {raw.get(r)!r} vs xlsx {book.value(r)!r}")
    for key, *_, basis, target, spec, ok in rows:
        if ok != "yes":
            print(f"target {key}: {target!r} does not tie to the SPEC figure {spec}")
    print(f"wrote {HERE / 'old_inputs.json'}, old_curves.csv, old_targets.csv "
          f"({sum(r[7] == 'R' for r in rows)} of {len(rows)} targets from the patched copy)")


if __name__ == "__main__":
    main()
