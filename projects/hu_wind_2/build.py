"""Iváncsa and Dunaújváros: RtB exit, quick hold-to-maturity and COD exit on the user's sheet (SPEC bora_v5_summary §21).

The user's `Bid_Cash` sheet (reference/HU_Wind_bid_deposits_fees_20261027.xlsx) is loaded as is and nothing in rows
1-78 or the band tables is changed. Four blocks are added below it, in the sheet's own layout
(C label | D unit | E Iváncsa | F Dunaújváros | G Total | H Date):
  (a) inputs for the exit cases (blue), (b) RtB exit: claims per line and the waterfall at the sale,
  (c) build and operate to maturity, EBITDA only, with an annual timeline in X:BG and the COD exit, (d) checks.
No macros. Writes out/HU_Wind_bid_deposits_fees.xlsx and out/rows.json (row map for check.py).

    python3 projects/hu_wind_2/build.py
"""
import json
import sys
from datetime import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
from fepm.xl.styles import (BLACK, BLUE, CHECK_RED, DARK, FEP_DARK_GREEN, FMT, FONT, GREY, INPUT_FILL,  # noqa: E402
                            OK_GREEN, SUB_FILL, WHITE)

SRC = ROOT / "reference" / "HU_Wind_bid_deposits_fees_20261027.xlsx"
OUT = HERE / "out" / "HU_Wind_bid_deposits_fees.xlsx"
ROWS = HERE / "out" / "rows.json"

START = 81                                    # first new row (the user's sheet ends at row 78)
TL0, NYEARS = 24, 36                          # timeline X:BG, 2026-2061
TL = [L(TL0 + i) for i in range(NYEARS)]
HELP = ["X", "Y", "Z", "AA", "AB", "AC"]      # claim helper columns (I:W are the user's hidden band tables)
NUM, NUM1, PCT, MULT, DATE, INT = FMT["money"], FMT["money1"], FMT["pct"], FMT["mult"], FMT["date"], "0"
FAC = FMT["factor"]
CHECK_FMT = '"CHECK" 0;"CHECK" -0;"OK"'
THIN = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color=BLACK))

ws = None
R: dict[str, int] = {}


def font(color=BLACK, bold=False, italic=False):
    return Font(name=FONT, size=10, color=color, bold=bold, italic=italic)


def put(ref, value, kind="calc", fmt=None, bold=False, top=False):
    c = ws[ref]
    c.value = value
    if kind == "input":
        c.font, c.fill = font(BLUE), PatternFill("solid", fgColor=INPUT_FILL)
        c.border = Border(THIN, THIN, THIN, THIN)
    elif kind == "note":
        c.font = font(GREY, italic=True)
    else:
        c.font = font(bold=bold)
        if top:
            c.border = TOP
    if fmt:
        c.number_format = fmt
    if kind != "note":
        c.alignment = Alignment(horizontal="right")
    return c


def label(r, text, unit=None, bold=False):
    ws[f"C{r}"].value, ws[f"C{r}"].font = text, font(bold=bold)
    if unit:
        ws[f"D{r}"].value, ws[f"D{r}"].font = unit, font(GREY, italic=True)


def bar(r, title, heads, extra=()):
    """Section bar in the sheet's style (white bold on FEP green, B:H), optionally over extra columns."""
    for col in ["B", "C", "D", "E", "F", "G", "H", *extra]:
        ws[f"{col}{r}"].font = font(WHITE, bold=True)
        ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=FEP_DARK_GREEN)
        ws[f"{col}{r}"].alignment = Alignment(horizontal="right" if col not in "BCD" else None)
    ws[f"C{r}"].value = title
    for col, h in heads.items():
        ws[f"{col}{r}"].value = h


def sub(r, title, extra=()):
    for col in ["C", "D", "E", "F", "G", "H", *extra]:
        ws[f"{col}{r}"].font = font(DARK, bold=True)
        ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=SUB_FILL)
    ws[f"C{r}"].value = title


def fx(template, **kw):
    """Fill {key} placeholders with row numbers (R) plus extra keyword values."""
    return template.format(**{**R, **kw})


# ---------------------------------------------------------------------------------------------------------------------
# (a) Inputs for the exit cases
# key, label, unit, format, value, source note.  COMMON: one value for both projects (column E, like the key dates).
COMMON = [
    ("rtb", "RtB date", "date", DATE, dt(2029, 3, 31), "SPEC §21.1"),
    ("sale", "Sale date (RtB exit)", "date", DATE, dt(2029, 9, 30), "SPEC §21.1"),
    ("freeze", "Interest freeze date", "date", DATE, dt(2029, 6, 30), "Devex file V24"),
    ("devd", "Devex average payment date (spent continuously)", "date", DATE, dt(2027, 12, 31),
     "Bid_Cash H77 'continuous'; SPEC §21.1"),
    ("r12", "Claim rate, cash lines (incl. cash-posted securities)", "%", PCT, 0.12, "Devex file V18; Bora: bonds 6 %"),
    ("r20", "Claim rate, devex", "%", PCT, 0.20, "Devex file V19"),
    ("floor", "Devex minimum return multiple", "x", MULT, 2.0, "Devex file V20"),
    ("life", "Operating life from COD", "years", INT, 30, "Bora Inputs G157"),
    ("miy", "Months in a year", "months", INT, 12, "Constant (Months_In_Year)"),
    ("tol", "Check tolerance", "EUR k", "0.000", 0.001, "Constant (Tolerance)"),
]
PROJECT = [
    ("Exit and valuation", None),
    ("prem", "RtB premium", "EUR k/MVA", NUM, 150, "SPEC §21.1, placeholder"),
    ("sc", "Sale costs", "EUR k", NUM, 250, "SPEC §21.1"),
    ("share", "FEP share of distributable proceeds", "%", PCT, 1.0, "SPEC §21.1, placeholder"),
    ("dsa", "Developer success fee (DSA)", "EUR k", NUM, 0, "SPEC §21.1"),
    ("dr", "Discount rate (buyer's pre-tax return at COD)", "%", PCT, 0.08, "Bora Inputs G226"),
    ("Capex (per MW of wind built unless stated)", None),
    ("sf", "Build shortfall (wind built = committed less this)", "%", PCT, 0.10, "Bora Dash_Build_Shortfall; GYIK Q28/30/76"),
    ("turb", "Turbine supply incl. transport", "EUR k/MW", NUM, 850, "Bora Inputs G162 (Base)"),
    ("inst", "Turbine installation and erection", "EUR k/MW", NUM, 80, "Bora Inputs G163"),
    ("fdn", "Foundations", "EUR k/MW", NUM, 90, "Bora Inputs G164"),
    ("civ", "Balance of plant – civil", "EUR k/MW", NUM, 70, "Bora Inputs G165"),
    ("elec", "Balance of plant – electrical", "EUR k/MW", NUM, 60, "Bora Inputs G166"),
    ("own", "Owner's costs, engineering, PM, construction insurance", "EUR k/MW", NUM, 40, "Bora Inputs G169"),
    ("land", "Land purchase and easements", "EUR k/MW", NUM, 15, "Bora Inputs G170"),
    ("subs", "On-site substation and 132 kV bay (lump sum)", "EUR k", NUM, 12000, "Bora Inputs G167, at the source basis"),
    ("basis", "Source basis of the lump sums", "MVA", NUM, 200, "Bora Inputs E305"),
    ("ccont", "Capex contingency (not on BESS)", "%", PCT, 0.05, "Bora Inputs G173"),
    ("bess", "BESS capex", "EUR k/MWh", NUM, 100, "CLAUDE.md standing assumption, 15-Oct-26"),
    ("Generation and revenue", None),
    ("hrs", "Net P50 hours", "h/yr", NUM, 2500, "Bora Inputs G181 (Base)"),
    ("curt", "Curtailment", "%", PCT, 0.02, "Bora Inputs G183"),
    ("clip", "Clipping", "%", PCT, 0.04, "Bora Dash_Clip_Override (user, 26-Oct-26)"),
    ("deg", "Degradation", "% p.a.", PCT, 0.003, "Bora Inputs G184"),
    ("ppa", "PPA price in the COD year", "EUR/MWh", NUM1, 90, "Bora Inputs G193 (Base)"),
    ("esc", "PPA escalation", "% p.a.", PCT, 0.01, "Bora Inputs G194"),
    ("Opex (COD-year prices)", None),
    ("om", "O&M full-service contract", "EUR k/MW/yr", NUM1, 12, "Bora Inputs G203 (Base)"),
    ("bop", "Balance of plant maintenance", "EUR k/MW/yr", NUM1, 3, "Bora Inputs G204"),
    ("lease", "Land lease", "EUR k/MW/yr", NUM1, 3, "Bora Inputs G205"),
    ("ins", "Operational insurance", "EUR k/MW/yr", NUM1, 2.5, "Bora Inputs G206"),
    ("mgmt", "Asset and commercial management", "EUR k/MW/yr", NUM1, 1.5, "Bora Inputs G207"),
    ("ptax", "Property tax", "EUR k/MW/yr", NUM1, 3, "Bora Inputs G208"),
    ("adm", "Administration, audit, compliance (at the source basis)", "EUR k/yr", NUM, 150, "Bora Inputs G209"),
    ("bal", "Balancing cost", "EUR/MWh", NUM1, 4, "Bora Inputs G196"),
    ("grid", "Grid use-of-system charges", "EUR/MWh", NUM1, 1, "Bora Inputs G197"),
    ("mekh", "MEKH supervisory fee", "% of rev.", PCT, 0.001, "Bora Inputs G198"),
    ("hipa", "HIPA local business tax", "% of rev.", PCT, 0.02, "Bora Inputs G212"),
    ("ocont", "Opex contingency (not on community, HIPA)", "%", PCT, 0.05, "Bora Inputs G210"),
    ("oesc", "Opex escalation", "% p.a.", PCT, 0.025, "Bora Inputs G211"),
]


def block_inputs(r):
    bar(r, "Inputs for the exit cases", {"E": "Iváncsa", "F": "Dunaújváros", "G": "Source"})
    r += 1
    sub(r, "Dates and claim terms (both projects)")
    r += 1
    for key, text, unit, fmt, v, note in COMMON:
        R[key] = r
        label(r, text, unit)
        put(f"E{r}", v, "input", fmt)
        put(f"G{r}", note, "note")
        r += 1
    for item in PROJECT:
        if item[1] is None:
            sub(r, item[0])
            r += 1
            continue
        key, text, unit, fmt, v, note = item
        R[key] = r
        label(r, text, unit)
        put(f"E{r}", v, "input", fmt)
        put(f"F{r}", v, "input", fmt)
        put(f"G{r}", note, "note")
        r += 1
    return r + 1


# ---------------------------------------------------------------------------------------------------------------------
# (b) RtB exit
# Claim lines: key, label, Bid_Cash principal row (EUR k), payment date source, rate key.
LINES = [("c_pf", "Participation fee (not reimbursed)", 71, "$E$35", "r12"),
         ("c_bs", "Bid security (cash)", 72, "$E$35", "r12"),
         ("c_pb", "Performance security payable (cash)", 73, "$E$36", "r12"),
         ("c_tso", "TSO connection-request security", 74, "$E$37", "r12"),
         ("c_dir", "Direct connection fee", 75, "$E$39", "r12"),
         ("c_ind", "Indirect connection fee, payable", 76, "$E$39", "r12"),
         ("c_dev", "Devex: MAX(20 % accrued, 1× principal), not reimbursed", 77, "$E${devd}", "r20")]
FLOWS = [("f_sdd", "Signing fee and DD", 55, "$E$34"), ("f_pf", "Participation fee", 71, "$E$35"),
         ("f_bs", "Bid security", 72, "$E$35"), ("f_pb", "Performance security payable", 73, "$E$36"),
         ("f_tso", "TSO connection-request security", 74, "$E$37"), ("f_dir", "Direct connection fee", 75, "$E$39"),
         ("f_ind", "Indirect connection fee, payable", 76, "$E$39"), ("f_dev", "Devex", 77, "$E${devd}")]


def block_rtb(r):
    bar(r, "RtB exit: FEP claims at the sale, interest to the freeze (act/365, capitalised 31-Dec)",
        {"E": "Iváncsa", "F": "Dunaújváros", "G": "Total", "H": "Paid", "X": "Rate", "Y": "Year paid",
         "Z": "Days, 1st yr", "AA": "Full years", "AB": "Days, last yr", "AC": "Factor"}, HELP)
    r += 1
    R["end"] = r
    label(r, "Interest accrual end (earlier of freeze and sale)", "date")
    put(f"E{r}", fx("=MIN($E${freeze},$E${sale})"), fmt=DATE)
    r += 1
    first = r
    for key, text, src, date, rate in LINES:
        R[key] = r
        label(r, text, "EUR k")
        for c in "EF":
            f = f"={c}{src}*MAX($AC{r}-1,$E${R['floor']}-1)" if key == "c_dev" else f"={c}{src}*($AC{r}-1)"
            put(f"{c}{r}", f, fmt=NUM)
        put(f"G{r}", f"=E{r}+F{r}", fmt=NUM)
        put(f"H{r}", fx("=" + date), fmt=DATE)
        put(f"X{r}", fx("=$E${%s}" % rate), fmt=PCT)
        put(f"Y{r}", f"=YEAR(H{r})", fmt=INT)
        put(f"Z{r}", fx("=MIN(DATE(Y{r}+1,1,1),$E${end}+1)-H{r}", r=r), fmt=INT)
        put(f"AA{r}", fx("=MAX(0,YEAR($E${end})-Y{r}-1)", r=r), fmt=INT)
        put(f"AB{r}", fx("=(YEAR($E${end})>Y{r})*($E${end}-DATE(YEAR($E${end}),1,1)+1)", r=r), fmt=INT)
        put(f"AC{r}", f"=(1+X{r}*Z{r}/Days_In_Year)*(1+X{r})^AA{r}*(1+X{r}*AB{r}/Days_In_Year)", fmt="0.0000")
        r += 1
    R["c_tot"] = r
    label(r, "Total FEP claims", "EUR k", bold=True)
    for c in "EFG":
        put(f"{c}{r}", f"=SUM({c}{first}:{c}{r - 1})", fmt=NUM, bold=True, top=True)
    r += 2

    bar(r, "RtB exit: waterfall at the sale", {"E": "Iváncsa", "F": "Dunaújváros", "G": "Total", "H": "Date"})
    r += 1
    rows = [("w_prem", "RtB premium", "={c}{prem}*{c}5", False),
            ("w_reimb", "Buyer reimburses deposits and fees at cost (not fee, devex)", "={c}78-{c}71-{c}77", False),
            ("w_proc", "Total proceeds", "={c}{w_prem}+{c}{w_reimb}", True),
            ("w_sc", "Less sale costs", "=-MIN({c}{w_proc},{c}{sc})", False),
            ("w_rp", "Less deposits and fees repaid to FEP", "=-MIN(SUM({c}{w_proc}:{c}{w_sc}),{c}{w_reimb})", False),
            ("w_c12", "Less FEP claims at 12 %", "=-MIN(SUM({c}{w_proc}:{c}{w_rp}),SUM({c}{c_pf}:{c}{c_ind}))", False),
            ("w_dev", "Less FEP devex return", "=-MIN(SUM({c}{w_proc}:{c}{w_c12}),{c}{c_dev})", False),
            ("w_dsa", "Less developer success fee (DSA)", "=-MIN(SUM({c}{w_proc}:{c}{w_dev}),{c}{dsa})", False),
            ("w_dist", "Distributable proceeds", "=SUM({c}{w_proc}:{c}{w_dsa})", True),
            ("w_fep", "FEP share of distributable", "={c}{w_dist}*{c}{share}", False),
            ("w_devsh", "Developer share of distributable", "={c}{w_dist}-{c}{w_fep}", False)]
    for key, text, f, bold in rows:
        R[key] = r
        label(r, text, "EUR k", bold=bold)
        for c in "EF":
            put(f"{c}{r}", fx(f, c=c), fmt=NUM, bold=bold, top=key == "w_dist")
        put(f"G{r}", f"=E{r}+F{r}", fmt=NUM, bold=bold, top=key == "w_dist")
        r += 1
    r += 1
    sub(r, "FEP cash flows, dated (outlays, then receipts at the sale)")
    r += 1
    for key, text, src, date in FLOWS:
        R[key] = r
        label(r, text, "EUR k")
        for c in "EF":
            put(f"{c}{r}", f"=-{c}{src}", fmt=NUM1 if key == "f_pf" else NUM)
        put(f"G{r}", f"=E{r}+F{r}", fmt=NUM)
        put(f"H{r}", fx("=" + date), fmt=DATE)
        r += 1
    R["f_rec"] = r
    label(r, "FEP receipts at the sale (repaid, claims, share)", "EUR k", bold=True)
    for c in "EF":
        put(f"{c}{r}", fx("={c}{w_fep}-SUM({c}{w_rp}:{c}{w_dev})", c=c), fmt=NUM, bold=True)
    put(f"G{r}", f"=E{r}+F{r}", fmt=NUM, bold=True)
    put(f"H{r}", fx("=$E${sale}"), fmt=DATE)
    r += 1
    rng = fx("{c}{f_sdd}:{c}{f_rec}", c="{c}")
    res = [("r_profit", "FEP profit, pre-tax", "EUR k", NUM, "=SUM(%s)" % rng, True),
           ("r_moic", "FEP MOIC", "x", MULT, '=SUMIF(%s,">0")/-SUMIF(%s,"<0")' % (rng, rng), False),
           ("r_irr", "FEP IRR, pre-tax (XIRR on the dated flows)", "%", PCT,
            "=XIRR(%s,$H$%d:$H$%d)" % (rng, R["f_sdd"], R["f_rec"]), False),
           ("r_mva", "FEP profit per MVA", "EUR k/MVA", NUM1, None, False)]
    for key, text, unit, fmt, f, bold in res:
        R[key] = r
        label(r, text, unit, bold=bold)
        for c in "EFG":
            if key == "r_mva":
                f2 = f"=E{R['r_profit']}/E5" if c == "E" else (f"=F{R['r_profit']}/F5" if c == "F"
                                                               else f"=G{R['r_profit']}/(E5+F5)")
            else:
                f2 = f.format(c=c)
            put(f"{c}{r}", f2, fmt=fmt, bold=bold, top=key == "r_profit")
        r += 1
    return r + 1


# ---------------------------------------------------------------------------------------------------------------------
# (c) Build and operate, hold to maturity; COD exit
TIMING = [  # key, label, unit, fmt, E constant, timeline formula ({x} column, {w} previous column)
    ("t_yr", "Year", "year", INT, "=YEAR($E$34)", "=$E{t_yr}+COLUMN()-COLUMN($X$1)"),
    ("t_end", "Year end", "date", DATE, None, "=DATE({x}{t_yr}+1,1,1)-1"),
    ("t_n", "Years from the COD year", "years", INT, None, "={x}{t_yr}-YEAR($E$41)"),
    ("t_cm", "Construction months by year end (E: construction period)", "months", INT,
     '=DATEDIF($E$40,$E$41,"m")', "=({x}{t_yr}+1-YEAR($E$40))*Months_In_Year-MONTH($E$40)"),
    ("t_cs", "Construction share of the year (NTP to COD)", "%", PCT, None,
     "=MAX(0,MIN({x}{t_cm},$E{t_cm})-MAX({x}{t_cm}-Months_In_Year,0))/$E{t_cm}"),
    ("t_om", "Months operated by year end (E: operating life)", "months", INT, "=$E${life}*Months_In_Year",
     "=({x}{t_yr}+1-YEAR($E$41))*Months_In_Year-MONTH($E$41)"),
    ("t_op", "Operating share of the year (COD to end of life)", "%", PCT, None,
     "=MAX(0,MIN({x}{t_om},$E{t_om})-MAX({x}{t_om}-Months_In_Year,0))/Months_In_Year"),
    ("t_cod", "COD year flag", "flag", FMT["flag"], None, "=N({x}{t_n}=0)"),
]
# Project rows: key, label, unit, fmt, timeline formula ({p} project input column, {x} column), total in G, bold.
PROJ = [
    ("p_dev", "FEP outlays at their dates (signing, fee, securities, fees, devex)", "EUR k", NUM,
     "=SUMPRODUCT((YEAR($H${f_sdd}:$H${f_dev})={x}${t_yr})*${p}${f_sdd}:${p}${f_dev})", True, False),
    ("p_pbr", "Performance security released (bid security included)", "EUR k", NUM,
     "=(${p}$72+${p}$73)*({x}${t_yr}=YEAR($E$42))", True, False),
    ("p_capex", "Construction capex (straight-line NTP to COD)", "EUR k", NUM, "=-${p}${k_capex}*{x}${t_cs}", True, False),
    ("p_muni", "Municipal one-off at COD", "EUR k", NUM, "=-${p}${k_muni}*{x}${t_cod}", True, False),
    ("p_inv", "Investment cash flow", "EUR k", NUM, "=SUM({x}{p_dev}:{x}{p_muni})", True, True),
    ("p_deg", "Degradation factor (from the first full year)", "x", FAC, "=(1-${p}${deg})^MAX({x}${t_n}-1,0)", False, False),
    ("p_gen", "Net generation", "GWh", NUM1,
     "=${p}${k_built}*${p}${hrs}*(1-${p}${curt})*(1-${p}${clip})*{x}${t_op}*{x}{p_deg}/Thousand", True, False),
    ("p_px", "PPA price", "EUR/MWh", NUM1, "=${p}${ppa}*(1+${p}${esc})^{x}${t_n}*({x}${t_op}>0)", False, False),
    ("p_rev", "Revenue", "EUR k", NUM, "={x}{p_gen}*{x}{p_px}", True, True),
    ("p_idx", "Opex index (COD year = 1)", "x", FAC, "=(1+${p}${oesc})^{x}${t_n}", False, False),
    ("p_fix", "O&M, BoP, lease, insurance, management, property tax", "EUR k", NUM,
     "=-SUM(${p}${om}:${p}${ptax})*${p}${k_built}*{x}{p_idx}*{x}${t_op}", True, False),
    ("p_adm", "Administration, audit, compliance", "EUR k", NUM,
     "=-${p}${adm}*${p}$5/${p}${basis}*{x}{p_idx}*{x}${t_op}", True, False),
    ("p_bg", "Balancing and grid charges", "EUR k", NUM, "=-{x}{p_gen}*(${p}${bal}+${p}${grid})*{x}{p_idx}", True, False),
    ("p_mekh", "MEKH supervisory fee", "EUR k", NUM, "=-{x}{p_rev}*${p}${mekh}", True, False),
    ("p_oc", "Opex contingency", "EUR k", NUM, "=SUM({x}{p_fix}:{x}{p_mekh})*${p}${ocont}", True, False),
    ("p_comm", "Community contribution", "EUR k", NUM, "=-{x}{p_gen}*${p}$10", True, False),
    ("p_hipa", "HIPA local business tax", "EUR k", NUM, "=-{x}{p_rev}*${p}${hipa}", True, False),
    ("p_ebitda", "EBITDA", "EUR k", NUM, "={x}{p_rev}+SUM({x}{p_fix}:{x}{p_hipa})", True, True),
    ("p_cf", "Unlevered pre-tax cash flow", "EUR k", NUM, "={x}{p_inv}+{x}{p_ebitda}", True, True),
    ("p_df", "Discount factor to COD (buyer's rate)", "x", FAC, "=(1+${p}${dr})^(($E$41-{x}${t_end})/Days_In_Year)",
     False, False),
    ("p_pv", "EBITDA valued at COD (total = COD price)", "EUR k", NUM, "={x}{p_ebitda}*{x}{p_df}", True, False),
    ("p_cod", "COD exit: FEP cash flow (to COD, price and security at COD)", "EUR k", NUM,
     "={x}{p_inv}*({x}${t_n}<=0)+{x}${t_cod}*($G{p_pv}+${p}$72+${p}$73)", True, True),
]
CAPEX = [("k_built", "Wind built (committed less the build shortfall)", "MW", NUM1, "={c}17*(1-{c}{sf})", False),
         ("k_turb", "Turbine supply incl. transport", "EUR k", NUM, "={c}{turb}*{c}{k_built}", False),
         ("k_inst", "Turbine installation and erection", "EUR k", NUM, "={c}{inst}*{c}{k_built}", False),
         ("k_fdn", "Foundations", "EUR k", NUM, "={c}{fdn}*{c}{k_built}", False),
         ("k_civ", "Balance of plant – civil", "EUR k", NUM, "={c}{civ}*{c}{k_built}", False),
         ("k_elec", "Balance of plant – electrical", "EUR k", NUM, "={c}{elec}*{c}{k_built}", False),
         ("k_own", "Owner's costs, engineering, PM, construction insurance", "EUR k", NUM, "={c}{own}*{c}{k_built}",
          False),
         ("k_land", "Land purchase and easements", "EUR k", NUM, "={c}{land}*{c}{k_built}", False),
         ("k_subs", "On-site substation and 132 kV bay (scaled by MVA)", "EUR k", NUM, "={c}{subs}*{c}5/{c}{basis}",
          False),
         ("k_cont", "Capex contingency", "EUR k", NUM, "=SUM({c}{k_turb}:{c}{k_subs})*{c}{ccont}", False),
         ("k_bess", "BESS (on BESS energy)", "EUR k", NUM, "={c}{bess}*{c}19", False),
         ("k_capex", "Construction capex", "EUR k", NUM, "=SUM({c}{k_turb}:{c}{k_bess})", True),
         ("k_conn", "Connection fees (TSO security, direct, indirect)", "EUR k", NUM, "={c}74+{c}75+{c}76", False),
         ("k_muni", "Municipal one-off", "EUR k", NUM, "={c}9*{c}5", False),
         ("k_dev", "Development (participation fee, signing and DD, devex)", "EUR k", NUM, "={c}71+{c}55+{c}77",
          False),
         ("k_allin", "All-in cost to COD (securities excluded: returned)", "EUR k", NUM,
          "={c}{k_capex}+SUM({c}{k_conn}:{c}{k_dev})", True)]


def block_htm(r):
    bar(r, "Build and operate, hold to maturity – EBITDA only (no debt, no tax)",
        {"E": "Iváncsa", "F": "Dunaújváros", "G": "Total"})
    r += 1
    # Reserve the summary rows (filled once the rows below exist).
    summary = [("h_irr", "Project IRR, pre-tax (unlevered)", "%", PCT),
               ("h_npv", "NPV at the discount rate, at the first year end", "EUR k", NUM),
               ("h_cmw", "All-in cost to COD per MW built", "EUR k/MW", NUM1),
               ("h_cmva", "All-in cost to COD per MVA", "EUR k/MVA", NUM1),
               ("h_gwh", "First full operating year: net generation", "GWh", NUM1),
               ("h_ebitda", "First full operating year: EBITDA", "EUR k", NUM),
               ("h_ebmw", "First full operating year: EBITDA per MW built", "EUR k/MW", NUM1),
               None,
               ("x_price", "Price at COD (EBITDA after COD at the buyer's rate)", "EUR k", NUM),
               ("x_pmw", "Price per MW built", "EUR k/MW", NUM1),
               ("x_pmva", "Price per MVA", "EUR k/MVA", NUM1),
               ("x_profit", "FEP profit, pre-tax", "EUR k", NUM),
               ("x_irr", "FEP IRR, pre-tax", "%", PCT),
               ("x_moic", "FEP MOIC", "x", MULT)]
    for item in summary:
        if item is None:
            sub(r, "COD exit: sale at COD to a buyer earning the discount rate, pre-tax")
            r += 1
            continue
        key, text, unit, fmt = item
        R[key] = r
        label(r, text, unit, bold=key in ("h_irr", "x_price", "x_profit"))
        r += 1
    R["summary"] = (R["h_irr"], R["x_moic"])
    r += 1
    sub(r, "Capex and all-in cost to COD")
    r += 1
    for key, text, unit, fmt, f, bold in CAPEX:
        R[key] = r
        label(r, text, unit, bold=bold)
        for c in "EF":
            put(f"{c}{r}", fx(f, c=c), fmt=fmt, bold=bold, top=bold)
        put(f"G{r}", f"=E{r}+F{r}", fmt=fmt, bold=bold, top=bold)
        r += 1
    r += 1

    # Timeline (shared by both projects: the key dates on Bid_Cash are common).
    bar(r, "Annual timeline (key dates common to both projects)", {"E": "Value", "G": "Total"}, TL)
    r += 1
    for key, text, unit, fmt, econst, f in TIMING:
        R[key] = r
        label(r, text, unit)
        if econst:
            put(f"E{r}", fx(econst), fmt=INT)
        for x in TL:
            put(f"{x}{r}", fx(f, x=x), fmt=fmt)
        if key in ("t_cs", "t_op"):
            put(f"G{r}", f"=SUM(X{r}:BG{r})", fmt=FMT["money2"])
        r += 1
    r += 1

    proj_rows = {}
    for p, name in (("E", "Iváncsa"), ("F", "Dunaújváros")):
        bar(r, f"{name}: annual cash flow (EUR k)", {"G": "Total"}, TL)
        r += 1
        rows = {}
        for key, text, unit, fmt, f, total, bold in PROJ:
            rows[key] = r
            R[key] = r                     # current project's rows for the templates
            label(r, text, unit, bold=bold)
            for x in TL:
                put(f"{x}{r}", fx(f, x=x, p=p), fmt=fmt, bold=bold, top=bold and key != "p_inv")
            if total:
                put(f"G{r}", f"=SUM(X{r}:BG{r})", fmt=fmt, bold=bold, top=bold and key != "p_inv")
            r += 1
        proj_rows[p] = rows
        r += 1

    # Summary rows at the top of the block, per project.
    yr, end = R["t_yr"], R["t_end"]
    for p in "EF":
        q = proj_rows[p]
        tl = lambda k: f"X{q[k]}:BG{q[k]}"  # noqa: E731
        f = {"h_irr": f"=XIRR({tl('p_cf')},X${end}:BG${end})",
             "h_npv": f"=XNPV({p}{R['dr']},{tl('p_cf')},X${end}:BG${end})",
             "h_cmw": f"={p}{R['k_allin']}/{p}{R['k_built']}",
             "h_cmva": f"={p}{R['k_allin']}/{p}5",
             "h_gwh": f"=SUMIF(X${yr}:BG${yr},YEAR($E$41)+1,{tl('p_gen')})",
             "h_ebitda": f"=SUMIF(X${yr}:BG${yr},YEAR($E$41)+1,{tl('p_ebitda')})",
             "h_ebmw": f"={p}{R['h_ebitda']}/{p}{R['k_built']}",
             "x_price": f"=G{q['p_pv']}",
             "x_pmw": f"={p}{R['x_price']}/{p}{R['k_built']}",
             "x_pmva": f"={p}{R['x_price']}/{p}5",
             "x_profit": f"=G{q['p_cod']}",
             "x_irr": f"=XIRR({tl('p_cod')},X${end}:BG${end})",
             "x_moic": f'=SUMIF({tl("p_cod")},">0")/-SUMIF({tl("p_cod")},"<0")'}
        for key, form in f.items():
            fmt = {"h_irr": PCT, "x_irr": PCT, "x_moic": MULT, "h_npv": NUM, "h_ebitda": NUM, "x_price": NUM,
                   "x_profit": NUM}.get(key, NUM1)
            put(f"{p}{R[key]}", form, fmt=fmt, bold=key in ("h_irr", "x_price", "x_profit"))
    for key in ("h_npv", "h_gwh", "h_ebitda", "x_price", "x_profit"):
        put(f"G{R[key]}", f"=E{R[key]}+F{R[key]}", fmt=NUM if key != "h_gwh" else NUM1,
            bold=key in ("x_price", "x_profit"))
    for key, num in (("h_cmw", "k_allin"), ("h_ebmw", "h_ebitda"), ("x_pmw", "x_price")):
        put(f"G{R[key]}", f"=G{R[num]}/G{R['k_built']}", fmt=NUM1)
    for key, num in (("h_cmva", "k_allin"), ("x_pmva", "x_price")):
        put(f"G{R[key]}", f"=G{R[num]}/(E5+F5)", fmt=NUM1)
    R["proj"] = proj_rows
    return r


# ---------------------------------------------------------------------------------------------------------------------
# (d) Checks
def block_checks(r):
    bar(r, "Checks (0 = OK)", {"E": "Iváncsa", "F": "Dunaújváros", "G": "Total"})
    r += 1
    pi, pd = R["proj"]["E"], R["proj"]["F"]
    checks = [
        ("ch_wf", "RtB waterfall closes (proceeds = all amounts paid out)",
         {c: fx("=N(ABS({c}{w_proc}+{c}{w_sc}+{c}{w_dsa}-{c}{w_devsh}-{c}{f_rec})>Tolerance)", c=c) for c in "EF"}),
        ("ch_end", "Claim lines paid before the interest accrual end",
         {"E": fx('=COUNTIF($H${c_pf}:$H${c_dev},">"&$E${end})')}),
        ("ch_rtb", "RtB date not after the sale date", {"E": fx("=N($E${rtb}>$E${sale})")}),
        ("ch_capex", "Construction capex spread in full (timeline = capex)",
         {"E": fx("=N(ABS(E{k_capex}+$G$%d)>Tolerance)" % pi["p_capex"]),
          "F": fx("=N(ABS(F{k_capex}+$G$%d)>Tolerance)" % pd["p_capex"])}),
        ("ch_allin", "Investment cash flows = all-in cost to COD",
         {"E": fx("=N(ABS(E{k_allin}+$G$%d)>Tolerance)" % pi["p_inv"]),
          "F": fx("=N(ABS(F{k_allin}+$G$%d)>Tolerance)" % pd["p_inv"])}),
        ("ch_life", "Operating life inside the timeline", {"E": fx("=N(ABS($G${t_op}-$E${life})>Tolerance)")}),
    ]
    first = r
    for key, text, cells in checks:
        R[key] = r
        label(r, text, "flag")
        for c, f in cells.items():
            put(f"{c}{r}", f, fmt=FMT["flag"])
        put(f"G{r}", f"=E{r}+F{r}", fmt=FMT["flag"])
        r += 1
    R["master"] = r
    label(r, "Master check", "flag", bold=True)
    put(f"G{r}", f"=SUM(G{first}:G{r - 1})", fmt=CHECK_FMT, bold=True, top=True)
    ws.conditional_formatting.add(f"G{r}", CellIsRule(operator="equal", formula=["0"], font=Font(color=OK_GREEN, bold=True)))
    ws.conditional_formatting.add(f"G{r}", CellIsRule(operator="notEqual", formula=["0"],
                                                      font=Font(color=CHECK_RED, bold=True)))
    r += 2
    ws[f"C{r}"].value, ws[f"C{r}"].font = "***END***", font(GREY, bold=True)
    return r


def build():
    global ws
    wb = openpyxl.load_workbook(SRC)
    ws = wb["Bid_Cash"]
    r = block_inputs(START)
    r = block_rtb(r)
    r = block_htm(r)
    r = block_checks(r)
    for name, key in (("Months_In_Year", "miy"), ("Tolerance", "tol")):
        wb.defined_names[name] = DefinedName(name, attr_text=f"Bid_Cash!$E${R[key]}")
    wb.defined_names["Master_Check"] = DefinedName("Master_Check", attr_text=f"Bid_Cash!$G${R['master']}")
    # Visible columns to the right of the hidden band tables: claim helpers and the timeline.
    if "X" in ws.column_dimensions:
        del ws.column_dimensions["X"]
    for col in TL:
        ws.column_dimensions[col].width = 10
    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    ROWS.write_text(json.dumps(R, indent=1))
    print(f"wrote {OUT} (rows {START}-{r})")


if __name__ == "__main__":
    build()
