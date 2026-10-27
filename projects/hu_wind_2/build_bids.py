"""HU wind bids: Bora (low bid), Iváncsa, Dunaújváros – RtB exit, build and hold (EBITDA), COD exit (SPEC §22).

Two sheets, no macros:
  Dashboard  key inputs (blue, Dash_* names) and executive values on top, then the calculation blocks of §21 per
             project column (Bora | Iváncsa | Dunaújváros | Total) and an annual 2026–2061 timeline in J:AS.
  Inputs     every other input with its source (common inputs named; per-project inputs in E:G) and the band tables.
Base: projects/hu_wind_2/build.py (§21). Writes out/HU_Wind_bids.xlsx and out/rows_bids.json (row map for check.py).

    python3 projects/hu_wind_2/build_bids.py
"""
import json
import sys
from datetime import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
from fepm.xl.styles import (BLACK, BLUE, CHECK_RED, DARK, FEP_DARK_GREEN, FMT, FONT, GREY, INPUT_FILL,  # noqa: E402
                            OK_GREEN, SUB_FILL, WHITE)

OUT = HERE / "out" / "HU_Wind_bids.xlsx"
ROWS = HERE / "out" / "rows_bids.json"
BID_CASH = ROOT / "reference" / "HU_Wind_bid_deposits_fees_20261027.xlsx"     # band tables (user's sheet)

P = ["E", "F", "G"]                                   # Bora, Iváncsa, Dunaújváros (same columns on both sheets)
NAMES = ["Bora", "Iváncsa", "Dunaújváros"]
DCOL = {"E": "J", "F": "K", "G": "L"}                 # payment-date columns
HCOL = {"E": ("K", "L", "M"), "F": ("N", "O", "P"), "G": ("Q", "R", "S")}   # claim helpers: days 1st yr, years, factor
TL = [L(10 + i) for i in range(36)]                   # J:AS, 2026-2061
NUM, NUM1, PCT, MULT, DATE, INT = FMT["money"], FMT["money1"], FMT["pct"], FMT["mult"], FMT["date"], "0"
NUM2, FAC, FLAG = FMT["money2"], FMT["factor"], FMT["flag"]
CHECK_FMT = '"CHECK" 0;"CHECK" -0;"OK"'
THIN = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color=BLACK))

ws = None
R: dict = {}          # Dashboard rows
I: dict = {}          # Inputs rows


def font(color=BLACK, bold=False, italic=False):
    return Font(name=FONT, size=10, color=color, bold=bold, italic=italic)


def put(ref, value, kind="calc", fmt=None, bold=False, top=False, sheet=None):
    c = (sheet or ws)[ref]
    c.value = value
    if kind == "input":
        c.font, c.fill = font(BLUE), PatternFill("solid", fgColor=INPUT_FILL)
        c.border = Border(THIN, THIN, THIN, THIN)
    elif kind == "note":
        c.font = font(GREY, italic=True)
    else:
        c.font = font(bold=bold)
        if top:
            c.border = TOP
    if fmt:
        c.number_format = fmt
    if kind != "note":
        c.alignment = Alignment(horizontal="right")
    return c


def label(r, text, unit=None, bold=False, sheet=None):
    s = sheet or ws
    s[f"C{r}"].value, s[f"C{r}"].font = text, font(bold=bold)
    if unit:
        s[f"D{r}"].value, s[f"D{r}"].font = unit, font(GREY, italic=True)


def bar(r, title, heads, extra=(), sheet=None):
    s = sheet or ws
    for col in ["B", "C", "D", "E", "F", "G", "H", *extra]:
        s[f"{col}{r}"].font = font(WHITE, bold=True)
        s[f"{col}{r}"].fill = PatternFill("solid", fgColor=FEP_DARK_GREEN)
        s[f"{col}{r}"].alignment = Alignment(horizontal=None if col in ("B", "C", "D") else "right")
    s[f"B{r}"].value = "x"
    s[f"C{r}"].value = title
    for col, h in heads.items():
        s[f"{col}{r}"].value = h


def sub(r, title, extra=(), sheet=None):
    s = sheet or ws
    for col in ["C", "D", "E", "F", "G", "H", *extra]:
        s[f"{col}{r}"].font = font(DARK, bold=True)
        s[f"{col}{r}"].fill = PatternFill("solid", fgColor=SUB_FILL)
    s[f"C{r}"].value = title


def fx(t, **kw):
    return t.format(**{**R, **kw})


def name(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


PROJ_HEADS = {"E": "Bora", "F": "Iváncsa", "G": "Dunaújváros", "H": "Total"}

# ---------------------------------------------------------------------------------------------------------------------
# Inputs sheet.  Common inputs: (name, label, unit, fmt, value, source) in column E, workbook name.
# Per-project inputs: (key, label, unit, fmt, [Bora, Iváncsa, Dunaújváros], source).
INPUTS = [
    ("Tender and grid (all projects)", None),
    ("FX", "FX rate", "HUF/EUR", NUM, 365, "Bora Inputs E11; Bid_Cash E45"),
    ("Bench_Wind", "Security base: benchmark investment, wind", "HUF/MW", NUM, 406e6, "Bora Inputs E75; Bid_Cash E46"),
    ("Bench_BESS", "Security base: benchmark investment, BESS", "HUF/MWh", NUM, 200e6, "Bora Inputs E89; Bid_Cash E47"),
    ("Bid_Sec_Rate", "Bid security, share of base", "%", PCT, 0.005, "Bora Inputs E76; Bid_Cash E49"),
    ("PB_Rate", "Performance security, minimum share of base", "%", PCT, 0.025, "Bora Inputs E77; Bid_Cash E50"),
    ("Part_Fee", "Participation fee per bid", "HUF", NUM, 3e6, "Bora Inputs E81; Bid_Cash E48"),
    ("TSO_Rate", "TSO connection-request security", "HUF/MVA", NUM, 4.5e6, "Bora Inputs E83; Bid_Cash E51"),
    ("Per project: fees, devex, dates", "P"),
    ("ind", "Indirect connection fee", "HUF/MVA", NUM, [73071306, 19826680, 19826680], "Bora Inputs G290; Bid_Cash E52"),
    ("dir", "Direct connection fee per connection point", "HUF", NUM, [508345804, 460e6, 460e6],
     "Bora Inputs G291; Bid_Cash E53"),
    ("share_pt", "Bids sharing the connection point", "count", INT, [2, 1, 1], "Bora: 2 bids per point (SPEC §22.2)"),
    ("sdd", "Signing fee and DD", "EUR k", NUM, [200, 150, 150], "Bora Inputs E24+E25; Bid_Cash E55"),
    ("devex", "Devex incl. contingency", "EUR k", NUM1, [1568.9, 1000, 1000], "Bora Devex_FEP E39 (100 MVA); Bid_Cash E56"),
    ("trafo", "EU transformers ≥ 80 % (ÖSZ-9)", "flag", INT, [0, 0, 0], "Bid_Cash E16; Bora low bid: no"),
    ("d_sign", "Signing fee and DD paid", "date", DATE, [dt(2026, 7, 31)] * 3, "Bora Inputs E27; Bid_Cash E34"),
    ("d_bid", "Bid submission: fee paid, bid security posted", "date", DATE, [dt(2026, 10, 30)] * 3,
     "Bora Inputs E95; Bid_Cash E35"),
    ("d_pb", "Performance security posted", "date", DATE, [dt(2026, 12, 30), dt(2027, 1, 15), dt(2027, 1, 15)],
     "Bora Deposits K10 (E96 + 17 d); Bid_Cash E36"),
    ("d_tso", "TSO security paid", "date", DATE, [dt(2026, 12, 30), dt(2027, 1, 15), dt(2027, 1, 15)],
     "Bora Inputs E97; Bid_Cash E37"),
    ("d_hc", "Connection agreement (HCSSZ): fees paid", "date", DATE, [dt(2028, 10, 1)] * 3,
     "Bora Inputs E99; Bid_Cash E39"),
    ("d_dev", "Devex average payment date", "date", DATE, [dt(2028, 1, 16), dt(2027, 12, 31), dt(2027, 12, 31)],
     "Bora Devex_FEP, amount-weighted; Bid_Cash 'continuous'"),
    ("d_ntp", "Notice to proceed", "date", DATE, [dt(2029, 6, 30)] * 3, "Bora Inputs E154; Bid_Cash E40"),
    ("d_cod", "COD", "date", DATE, [dt(2031, 6, 30)] * 3, "Bora Inputs E155; Bid_Cash E41"),
    ("d_rel", "Performance security released", "date", DATE, [dt(2033, 3, 31)] * 3,
     "Bid_Cash E42 (Kiírás XII.13, estimate)"),
    ("d_dsa", "DSA accrual start", "date", DATE, [dt(2026, 10, 26), dt(2026, 7, 31), dt(2026, 7, 31)],
     "Bora Inputs E70; others: signing (DSA off)"),
    ("RtB exit terms (all projects)", None),
    ("Rate_Cash", "Claim rate, cash lines (incl. cash-posted securities)", "%", PCT, 0.12, "Devex file V18; SPEC §21.1"),
    ("Rate_Devex", "Claim rate, devex", "%", PCT, 0.20, "Devex file V19"),
    ("Devex_Floor", "Devex minimum return multiple", "x", MULT, 2.0, "Devex file V20"),
    ("Freeze_Date", "Interest freeze date", "date", DATE, dt(2029, 6, 30), "Devex file V24"),
    ("Sale_Date", "Sale date (RtB exit)", "date", DATE, dt(2029, 9, 30), "Bora Inputs E58; SPEC §21.1"),
    ("Sale_Costs", "Sale costs per project", "EUR k", NUM, 250, "SPEC §21.1"),
    ("DSA_Rate", "Developer success fee (on committed MW, 30/360)", "EUR k/MW/yr", NUM1, 2.5, "Bora Inputs E36"),
    ("DSA_Cap", "Developer success fee cap", "EUR k/MW", NUM1, 10, "Bora Inputs E37"),
    ("CIT_Rate", "Corporate income tax on exit gains", "%", PCT, 0.09, "SPEC §22.2 (exit gains only)"),
    ("Capex (per MW of wind built unless stated)", None),
    ("Turbine_EU", "Turbine supply incl. transport, EU route", "EUR k/MW", NUM, 850, "Bora Inputs G162 (Base)"),
    ("Turbine_CN", "Turbine supply incl. transport, CN route", "EUR k/MW", NUM, 500, "Bora Inputs H162 (High)"),
    ("Capex_Install", "Turbine installation and erection", "EUR k/MW", NUM, 80, "Bora Inputs G163"),
    ("Capex_Fdn", "Foundations", "EUR k/MW", NUM, 90, "Bora Inputs G164"),
    ("Capex_Civil", "Balance of plant – civil", "EUR k/MW", NUM, 70, "Bora Inputs G165"),
    ("Capex_Elec", "Balance of plant – electrical", "EUR k/MW", NUM, 60, "Bora Inputs G166"),
    ("Capex_Owner", "Owner's costs, engineering, PM, construction insurance", "EUR k/MW", NUM, 40, "Bora Inputs G169"),
    ("Capex_Land", "Land purchase and easements", "EUR k/MW", NUM, 15, "Bora Inputs G170"),
    ("Subs_Lump", "On-site substation and 132 kV bay (lump sum)", "EUR k", NUM, 12000, "Bora Inputs G167"),
    ("Lump_Basis", "Source basis of the lump sums", "MVA", NUM, 200, "Bora Inputs E305"),
    ("Capex_Cont", "Capex contingency (not on BESS)", "%", PCT, 0.05, "Bora Inputs G173"),
    ("BESS_Capex", "BESS capex", "EUR k/MWh", NUM, 100, "CLAUDE.md standing assumption, 15-Oct-26"),
    ("Generation and revenue", None),
    ("Hours_P50", "Net P50 hours", "h/yr", NUM, 2500, "Bora Inputs G181 (Base)"),
    ("Curtailment", "Curtailment", "%", PCT, 0.02, "Bora Inputs G183"),
    ("Clipping", "Clipping", "%", PCT, 0.04, "Bora Dash_Clip_Override (user, 26-Oct-26)"),
    ("Degradation", "Degradation", "% p.a.", PCT, 0.003, "Bora Inputs G184"),
    ("PPA_Price", "PPA price in the COD year", "EUR/MWh", NUM1, 90, "Bora Inputs G193 (Base)"),
    ("PPA_Esc", "PPA escalation", "% p.a.", PCT, 0.01, "Bora Inputs G194"),
    ("Opex (COD-year prices)", None),
    ("Opex_OM", "O&M full-service contract", "EUR k/MW/yr", NUM1, 12, "Bora Inputs G203 (Base)"),
    ("Opex_BoP", "Balance of plant maintenance", "EUR k/MW/yr", NUM1, 3, "Bora Inputs G204"),
    ("Opex_Lease", "Land lease", "EUR k/MW/yr", NUM1, 3, "Bora Inputs G205"),
    ("Opex_Ins", "Operational insurance", "EUR k/MW/yr", NUM1, 2.5, "Bora Inputs G206"),
    ("Opex_Mgmt", "Asset and commercial management", "EUR k/MW/yr", NUM1, 1.5, "Bora Inputs G207"),
    ("Opex_PropTax", "Property tax", "EUR k/MW/yr", NUM1, 3, "Bora Inputs G208"),
    ("Opex_Admin", "Administration, audit, compliance (at the source basis)", "EUR k/yr", NUM, 150, "Bora Inputs G209"),
    ("Balancing", "Balancing cost", "EUR/MWh", NUM1, 4, "Bora Inputs G196"),
    ("Grid_Charges", "Grid use-of-system charges", "EUR/MWh", NUM1, 1, "Bora Inputs G197"),
    ("MEKH_Fee", "MEKH supervisory fee", "% of rev.", PCT, 0.001, "Bora Inputs G198"),
    ("HIPA_Rate", "HIPA local business tax", "% of rev.", PCT, 0.02, "Bora Inputs G212"),
    ("Opex_Cont", "Opex contingency (not on community, HIPA)", "%", PCT, 0.05, "Bora Inputs G210"),
    ("Opex_Esc", "Opex escalation", "% p.a.", PCT, 0.025, "Bora Inputs G211"),
    ("Valuation, timeline and constants", None),
    ("Life", "Operating life from COD", "years", INT, 30, "Bora Inputs G157"),
    ("Discount_Rate", "Discount rate (buyer's pre-tax return at COD)", "%", PCT, 0.08, "Bora Inputs G226"),
    ("First_Year", "First year of the annual timeline", "year", INT, 2026, "Year of the first signing"),
    ("Thousand", "Thousand", "x", NUM, 1000, "Constant"),
    ("Days_In_Year", "Days in a year (act/365)", "days", INT, 365, "Constant"),
    ("Months_In_Year", "Months in a year", "months", INT, 12, "Constant"),
    ("Tolerance", "Check tolerance", "EUR k", "0.000", 0.001, "Constant"),
]


def build_inputs(wb):
    s = wb.create_sheet("Inputs")
    s.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", (2, 2, 52, 11, 13, 13, 13, 44)):
        s.column_dimensions[col].width = w
    s["C2"].value, s["C2"].font = "HU wind bids – other inputs (blue) with sources", font(bold=True)
    r = 4
    bar(r, "Inputs", {"E": "Bora", "F": "Iváncsa", "G": "Dunaújváros", "H": "Source"}, sheet=s)
    s["H4"].alignment = Alignment(horizontal=None)
    r += 1
    for item in INPUTS:
        if item[1] in (None, "P"):
            sub(r, item[0] + ("" if item[1] else " (value in column E)"), sheet=s)
            r += 1
            continue
        key, text, unit, fmt, v, note = item
        I[key] = r
        label(r, text, unit, sheet=s)
        if isinstance(v, list):
            for c, x in zip(P, v):
                put(f"{c}{r}", x, "input", fmt, sheet=s)
        else:
            put(f"E{r}", v, "input", fmt, sheet=s)
            name(wb, key, f"Inputs!$E${r}")
        put(f"H{r}", note, "note", sheet=s)
        r += 1
    name(wb, "Opex_Per_MW", f"Inputs!$E${I['Opex_OM']}:$E${I['Opex_PropTax']}")
    # Band tables: the user's Bid_Cash J3:V28 (values and names), Kiírás Annex 4.
    r += 1
    bar(r, "Band tables for the points (ÖSZ, Kiírás Annex 4)", {"E": "Values →"}, sheet=s)
    r += 1
    src = openpyxl.load_workbook(BID_CASH)
    b = src["Bid_Cash"]
    for rr in range(3, 29):
        n = b[f"J{rr}"].value
        vals = [b.cell(rr, c).value for c in range(11, 23) if b.cell(rr, c).value is not None]
        label(r, n, None, sheet=s)
        for j, x in enumerate(vals):
            put(f"{L(5 + j)}{r}", x, "input", "General", sheet=s)
        name(wb, n, f"Inputs!$E${r}:${L(4 + len(vals))}${r}" if len(vals) > 1 else f"Inputs!$E${r}")
        r += 1
    r += 1
    s[f"C{r}"].value, s[f"C{r}"].font = "***END***", font(GREY, bold=True)
    s.freeze_panes = "E5"
    return s


# ---------------------------------------------------------------------------------------------------------------------
# Dashboard: key inputs (blue, Dash_* names): key, label, unit, fmt, [Bora, Iváncsa, Dunaújváros].
KEY = [("Dash_MVA", "Net connection capacity", "MVA", NUM, [100, 50, 50]),
       ("Dash_Overbuild", "Overbuild (ÖSZ-3)", "%", PCT, [0.20, 0.41, 0.41]),
       ("Dash_BESS_Share", "BESS power, share of MVA (ÖSZ-4)", "%", PCT, [0.10, 0.40, 0.40]),
       ("Dash_Hours", "BESS duration", "h", NUM, [4, 4, 4]),
       ("Dash_Shortfall", "Wind built below the bid", "%", PCT, [0.10, 0.10, 0.10]),
       ("Dash_Municipal", "Municipal one-off (ÖSZ-1), at COD", "EUR k/MVA", NUM, [20, 50, 50]),
       ("Dash_Community", "Community contribution (ÖSZ-2)", "EUR/MWh", NUM1, [0, 3, 3]),
       ("Dash_Partial", "Partial award accepted (ÖSZ-5)", "flag", INT, [1, 0, 0]),
       ("Dash_PB_Mult", "Performance security multiple (ÖSZ-6)", "x", MULT, [1, 1, 1]),
       ("Dash_Brownfield", "Brownfield site (ÖSZ-7)", "flag", INT, [0, 0, 0]),
       ("Dash_Shared", "Shared connection (ÖSZ-8)", "flag", INT, [1, 0, 0]),
       ("Dash_Turbine", "Turbine route (EU or CN; EU scores ÖSZ-9)", "text", "General", ["CN", "EU", "EU"]),
       ("turb_px", None, None, None, None),
       ("Dash_Premium", "RtB premium", "EUR k/MVA", NUM, [150, 150, 150]),
       ("Dash_FEP_Share", "FEP share of distributable proceeds", "%", PCT, [0.70, 1.0, 1.0]),
       ("Dash_DSA", "Developer success fee (DSA) on", "flag", INT, [1, 0, 0])]


def key_inputs(wb, r):
    bar(r, "Key inputs", PROJ_HEADS)
    r += 1
    for key, text, unit, fmt, vals in KEY:
        R[key] = r
        if key == "turb_px":
            label(r, "Turbine supply price for the route", "EUR k/MW")
            for c in P:
                put(f"{c}{r}", f'=IF({c}{R["Dash_Turbine"]}="EU",Turbine_EU,Turbine_CN)', fmt=NUM)
        else:
            label(r, text, unit)
            for c, x in zip(P, vals):
                put(f"{c}{r}", x, "input", fmt)
            name(wb, key, f"Dashboard!$E${r}:$G${r}")
        r += 1
    put(f"H{R['Dash_MVA']}", f"=SUM(E{R['Dash_MVA']}:G{R['Dash_MVA']})", fmt=NUM)
    return r + 1


EXEC = [("ex_pts", "Points (ÖSZ, of 100)", "points", NUM1),
        ("ex_cash", "Cash to RtB (deposits, fees, devex, signing)", "EUR k", NUM),
        "RtB exit at the sale (30-Sep-2029)",
        ("ex_rp", "FEP profit, pre-tax", "EUR k", NUM), ("ex_rcit", "CIT on the gain", "EUR k", NUM),
        ("ex_rpost", "FEP profit, post-tax", "EUR k", NUM), ("ex_rirr", "FEP IRR, pre-tax", "%", PCT),
        ("ex_rmoic", "FEP MOIC, pre-tax", "x", MULT), ("ex_rmva", "FEP profit per MVA, pre-tax", "EUR k/MVA", NUM1),
        "Build and operate, hold to maturity (EBITDA only, no debt, no tax)",
        ("ex_built", "Wind built", "MW", NUM1), ("ex_allin", "All-in cost to COD", "EUR m", NUM1),
        ("ex_cmw", "All-in cost per MW built", "EUR k/MW", NUM1), ("ex_cmva", "All-in cost per MVA", "EUR k/MVA", NUM1),
        ("ex_eb", "First full operating year: EBITDA", "EUR k", NUM),
        ("ex_gwh", "First full operating year: net generation", "GWh", NUM1),
        ("ex_irr", "Project IRR, pre-tax (unlevered)", "%", PCT),
        ("ex_npv", "NPV at the discount rate, at the first year end", "EUR k", NUM),
        "COD exit (buyer earns the discount rate, pre-tax)",
        ("ex_price", "Price at COD", "EUR m", NUM1), ("ex_pmw", "Price per MW built", "EUR k/MW", NUM1),
        ("ex_pmva", "Price per MVA", "EUR k/MVA", NUM1), ("ex_xp", "FEP profit, pre-tax", "EUR k", NUM),
        ("ex_xcit", "CIT on the gain", "EUR k", NUM), ("ex_xpost", "FEP profit, post-tax", "EUR k", NUM),
        ("ex_xirr", "FEP IRR, pre-tax", "%", PCT), ("ex_xmoic", "FEP MOIC, pre-tax", "x", MULT)]


def exec_reserve(r):
    bar(r, "Executive values (links to the calculations below)", PROJ_HEADS)
    r += 1
    for item in EXEC:
        if isinstance(item, str):
            sub(r, item)
        else:
            R[item[0]] = r
            label(r, item[1], item[2], bold=item[0] in ("ex_rpost", "ex_irr", "ex_xpost"))
        r += 1
    R["exec"] = (R["ex_pts"], R["ex_xmoic"])
    return r + 1


# ---------------------------------------------------------------------------------------------------------------------
def rows3(r, key, text, unit, fmt, f, total=True, bold=False, top=False):
    """One row per project column (E:G) from a template ({c} = Dashboard/Inputs column), total in H."""
    R[key] = r
    label(r, text, unit, bold=bold)
    for c in P:
        put(f"{c}{r}", fx(f, c=c, d=DCOL[c]), fmt=fmt, bold=bold, top=top)
    if total:
        put(f"H{r}", f"=SUM(E{r}:G{r})", fmt=fmt, bold=bold, top=top)
    return r + 1


POINTS = [("pt_mun", "Points – municipal one-off",
           "=INDEX(Osz1_Pts,1+SUMPRODUCT((Osz1_LB<{c}{Dash_Municipal})+(Osz1_LB={c}{Dash_Municipal})*Osz1_Incl))*Osz1_W"),
          ("pt_com", "Points – community contribution",
           "=INDEX(Osz2_Pts,1+SUMPRODUCT((Osz2_LB<{c}{Dash_Community})+(Osz2_LB={c}{Dash_Community})*Osz2_Incl))*Osz2_W"),
          ("pt_ob", "Points – overbuild",
           "=INDEX(Osz3_Pts,1+SUMPRODUCT((Osz3_LB<{c}{Dash_Overbuild})+(Osz3_LB={c}{Dash_Overbuild})*Osz3_Incl))*Osz3_W"),
          ("pt_bess", "Points – BESS",
           "=INDEX(Osz3_Pts,1+SUMPRODUCT((Osz3_LB<{c}{Dash_BESS_Share})+(Osz3_LB={c}{Dash_BESS_Share})*Osz3_Incl))*Osz4_W"),
          ("pt_part", "Points – partial award", "={c}{Dash_Partial}*Osz5_Pts*Osz5_W"),
          ("pt_pb", "Points – performance security",
           "=INDEX(Osz6_Pts,1+SUMPRODUCT((Osz6_LB<{c}{Dash_PB_Mult})+(Osz6_LB={c}{Dash_PB_Mult})*Osz6_Incl))*Osz6_W"),
          ("pt_brown", "Points – brownfield", "={c}{Dash_Brownfield}*Osz7_Pts*Osz7_W"),
          ("pt_shared", "Points – shared connection", "={c}{Dash_Shared}*Osz8_Pts*Osz8_W"),
          ("pt_eut", "Points – EU turbines", '=N({c}{Dash_Turbine}="EU")*Osz9a_Pts*Osz9_W'),
          ("pt_eutr", "Points – EU transformers", "=Inputs!{c}%d*Osz9b_Pts*Osz9_W")]


def calc_blocks(r):
    # Points and capacity
    bar(r, "Bid points and capacity", PROJ_HEADS)
    r += 1
    for key, text, f in POINTS:
        r = rows3(r, key, text, "points", NUM1, f.replace("%d", str(I["trafo"])), total=False)
    r = rows3(r, "pt_tot", "Points total", "points", NUM1, "=SUM({c}{pt_mun}:{c}{pt_eutr})", total=False, bold=True,
              top=True)
    r = rows3(r, "wind", "Wind committed in the bid", "MW", NUM1, "={c}{Dash_MVA}*(1+{c}{Dash_Overbuild})")
    r = rows3(r, "mwh", "BESS energy", "MWh", NUM1, "={c}{Dash_MVA}*{c}{Dash_BESS_Share}*{c}{Dash_Hours}")
    r = rows3(r, "built", "Wind built (committed less the build shortfall)", "MW", NUM1,
              "={c}{wind}*(1-{c}{Dash_Shortfall})")
    r += 1

    # Deposits and fees to RtB, with payment dates in J:L
    bar(r, "Deposits, fees and devex to RtB (HUF rates at the FX rate)", {**PROJ_HEADS, "J": "Paid: Bora",
                                                                          "K": "Iváncsa", "L": "Dunaújváros"},
        ["J", "K", "L"])
    r += 1
    r = rows3(r, "d_base", "Security base (benchmark investment), not cash", "EUR k", NUM,
              "=({c}{wind}*Bench_Wind+{c}{mwh}*Bench_BESS)/FX/Thousand")
    dep = [("d_pf", "Participation fee", NUM1, "=Part_Fee/FX/Thousand", "d_bid"),
           ("d_bs", "Bid security (cash)", NUM, "=Bid_Sec_Rate*{c}{d_base}", "d_bid"),
           ("d_pb", "Performance security payable (bid security deducted)", NUM,
            "=PB_Rate*{c}{Dash_PB_Mult}*{c}{d_base}-{c}{d_bs}", "d_pb"),
           ("d_tso", "TSO connection-request security", NUM, "=TSO_Rate*{c}{Dash_MVA}/FX/Thousand", "d_tso"),
           ("d_dir", "Direct connection fee (per point ÷ bids sharing it)", NUM,
            "=Inputs!{c}$%d/Inputs!{c}$%d/FX/Thousand" % (I["dir"], I["share_pt"]), "d_hc"),
           ("d_ind", "Indirect connection fee, payable (TSO security deducted)", NUM,
            "=Inputs!{c}$%d*{c}{Dash_MVA}/FX/Thousand-{c}{d_tso}" % I["ind"], "d_hc"),
           ("d_dev", "Devex incl. contingency", NUM, "=Inputs!{c}$%d" % I["devex"], "d_dev"),
           ("d_sdd", "Signing fee and DD", NUM, "=Inputs!{c}$%d" % I["sdd"], "d_sign")]
    for key, text, fmt, f, dkey in dep:
        r = rows3(r, key, text, "EUR k", fmt, f)
        for c in P:
            put(f"{DCOL[c]}{r - 1}", f"=Inputs!{c}${I[dkey]}", fmt=DATE)
    r = rows3(r, "d_tot", "Total cash to RtB", "EUR k", NUM, "=SUM({c}{d_pf}:{c}{d_sdd})", bold=True, top=True)
    r += 1

    # RtB claims
    heads = {**PROJ_HEADS, "J": "Rate"}
    for c, nm in zip(P, NAMES):
        k1, k2, k3 = HCOL[c]
        heads.update({k1: f"{nm}: days 1st yr", k2: "full years", k3: "factor"})
    bar(r, "RtB exit: FEP claims at the sale (interest to the freeze, act/365, capitalised 31-Dec)", heads,
        [L(10 + i) for i in range(10)])
    r += 1
    r = rows3(r, "c_end", "Interest accrual end (earlier of freeze and sale)", "date", DATE,
              "=MIN(Freeze_Date,Sale_Date)", total=False)
    r = rows3(r, "c_last", "Days in the accrual end year", "days", INT, "={c}{c_end}-DATE(YEAR({c}{c_end}),1,1)+1",
              total=False)
    lines = [("c_pf", "Participation fee (not reimbursed)", "d_pf", "Rate_Cash"),
             ("c_bs", "Bid security (cash)", "d_bs", "Rate_Cash"),
             ("c_pb", "Performance security payable (cash)", "d_pb", "Rate_Cash"),
             ("c_tso", "TSO connection-request security", "d_tso", "Rate_Cash"),
             ("c_dir", "Direct connection fee", "d_dir", "Rate_Cash"),
             ("c_ind", "Indirect connection fee, payable", "d_ind", "Rate_Cash"),
             ("c_dev", "Devex: MAX(20 % accrued, 1× principal), not reimbursed", "d_dev", "Rate_Devex")]
    for key, text, dkey, rate in lines:
        R[key] = r
        label(r, text, "EUR k")
        put(f"J{r}", f"={rate}", fmt=PCT)
        for c in P:
            d1, n, fac = HCOL[c]
            dc = f"{DCOL[c]}{R[dkey]}"
            put(f"{d1}{r}", f"=DATE(YEAR({dc})+1,1,1)-{dc}", fmt=INT)
            put(f"{n}{r}", f"=YEAR({c}${R['c_end']})-YEAR({dc})-1", fmt=INT)
            put(f"{fac}{r}", f"=(1+$J{r}*{d1}{r}/Days_In_Year)*(1+$J{r})^{n}{r}*(1+$J{r}*{c}${R['c_last']}/Days_In_Year)",
                fmt="0.0000")
            f = (f"={c}{R[dkey]}*MAX({fac}{r}-1,Devex_Floor-1)" if key == "c_dev" else f"={c}{R[dkey]}*({fac}{r}-1)")
            put(f"{c}{r}", f, fmt=NUM1 if key == "c_pf" else NUM)
        put(f"H{r}", f"=SUM(E{r}:G{r})", fmt=NUM)
        r += 1
    r = rows3(r, "c_tot", "Total FEP claims", "EUR k", NUM, "=SUM({c}{c_pf}:{c}{c_dev})", bold=True, top=True)
    r = rows3(r, "c_dsa", "Developer success fee due (DSA, capped)", "EUR k", NUM,
              "={c}{Dash_DSA}*MIN(DSA_Rate*YEARFRAC(Inputs!{c}$%d,Sale_Date),DSA_Cap)*{c}{wind}" % I["d_dsa"])
    r += 1

    # Waterfall and FEP flows
    bar(r, "RtB exit: waterfall at the sale", {**PROJ_HEADS, "J": "Date: Bora", "K": "Iváncsa", "L": "Dunaújváros"},
        ["J", "K", "L"])
    r += 1
    wf = [("w_prem", "RtB premium", "={c}{Dash_Premium}*{c}{Dash_MVA}", False),
          ("w_reimb", "Buyer reimburses securities and connection fees at cost", "=SUM({c}{d_bs}:{c}{d_ind})", False),
          ("w_proc", "Total proceeds", "={c}{w_prem}+{c}{w_reimb}", True),
          ("w_sc", "Less sale costs", "=-MIN({c}{w_proc},Sale_Costs)", False),
          ("w_rp", "Less securities and fees repaid to FEP", "=-MIN(SUM({c}{w_proc}:{c}{w_sc}),{c}{w_reimb})", False),
          ("w_c12", "Less FEP claims at 12 %", "=-MIN(SUM({c}{w_proc}:{c}{w_rp}),SUM({c}{c_pf}:{c}{c_ind}))", False),
          ("w_dev", "Less FEP devex return", "=-MIN(SUM({c}{w_proc}:{c}{w_c12}),{c}{c_dev})", False),
          ("w_dsa", "Less developer success fee (DSA)", "=-MIN(SUM({c}{w_proc}:{c}{w_dev}),{c}{c_dsa})", False),
          ("w_dist", "Distributable proceeds", "=SUM({c}{w_proc}:{c}{w_dsa})", True),
          ("w_fep", "FEP share of distributable", "={c}{w_dist}*{c}{Dash_FEP_Share}", False),
          ("w_devsh", "Developer share of distributable", "={c}{w_dist}-{c}{w_fep}", False)]
    for key, text, f, bold in wf:
        r = rows3(r, key, text, "EUR k", NUM, f, bold=bold, top=key == "w_dist")
    r += 1
    sub(r, "FEP cash flows, dated (outlays, then receipts at the sale)", ["J", "K", "L"])
    r += 1
    flows = [("f_sdd", "Signing fee and DD", "d_sdd", "d_sign"), ("f_pf", "Participation fee", "d_pf", "d_bid"),
             ("f_bs", "Bid security", "d_bs", "d_bid"), ("f_pb", "Performance security payable", "d_pb", "d_pb"),
             ("f_tso", "TSO connection-request security", "d_tso", "d_tso"),
             ("f_dir", "Direct connection fee", "d_dir", "d_hc"), ("f_ind", "Indirect connection fee, payable", "d_ind",
                                                                    "d_hc"),
             ("f_dev", "Devex", "d_dev", "d_dev")]
    for key, text, src, dkey in flows:
        r = rows3(r, key, text, "EUR k", NUM1 if key == "f_pf" else NUM, "=-{c}{%s}" % src)
        for c in P:
            put(f"{DCOL[c]}{r - 1}", f"=Inputs!{c}${I[dkey]}", fmt=DATE)
    r = rows3(r, "f_rec", "FEP receipts at the sale (repaid, claims, share)", "EUR k", NUM,
              "={c}{w_fep}-SUM({c}{w_rp}:{c}{w_dev})", bold=True)
    for c in P:
        put(f"{DCOL[c]}{r - 1}", "=Sale_Date", fmt=DATE)
    r = rows3(r, "r_profit", "FEP profit, pre-tax", "EUR k", NUM, "=SUM({c}{f_sdd}:{c}{f_rec})", bold=True, top=True)
    r = rows3(r, "r_cit", "CIT on the gain", "EUR k", NUM, "=-MAX(0,{c}{r_profit})*CIT_Rate")
    r = rows3(r, "r_post", "FEP profit, post-tax", "EUR k", NUM, "={c}{r_profit}+{c}{r_cit}", bold=True)
    r += 1

    # Capex and all-in cost to COD
    bar(r, "Capex and all-in cost to COD", PROJ_HEADS)
    r += 1
    capex = [("k_turb", "Turbine supply incl. transport (route price)", "={c}{turb_px}*{c}{built}"),
             ("k_inst", "Turbine installation and erection", "=Capex_Install*{c}{built}"),
             ("k_fdn", "Foundations", "=Capex_Fdn*{c}{built}"),
             ("k_civ", "Balance of plant – civil", "=Capex_Civil*{c}{built}"),
             ("k_elec", "Balance of plant – electrical", "=Capex_Elec*{c}{built}"),
             ("k_own", "Owner's costs, engineering, PM, construction insurance", "=Capex_Owner*{c}{built}"),
             ("k_land", "Land purchase and easements", "=Capex_Land*{c}{built}"),
             ("k_subs", "On-site substation and 132 kV bay (scaled by MVA)", "=Subs_Lump*{c}{Dash_MVA}/Lump_Basis"),
             ("k_cont", "Capex contingency", "=SUM({c}{k_turb}:{c}{k_subs})*Capex_Cont"),
             ("k_bess", "BESS (on BESS energy)", "=BESS_Capex*{c}{mwh}")]
    for key, text, f in capex:
        r = rows3(r, key, text, "EUR k", NUM, f)
    r = rows3(r, "k_capex", "Construction capex", "EUR k", NUM, "=SUM({c}{k_turb}:{c}{k_bess})", bold=True, top=True)
    r = rows3(r, "k_conn", "Connection fees (TSO security, direct, indirect)", "EUR k", NUM,
              "={c}{d_tso}+{c}{d_dir}+{c}{d_ind}")
    r = rows3(r, "k_muni", "Municipal one-off", "EUR k", NUM, "={c}{Dash_Municipal}*{c}{Dash_MVA}")
    r = rows3(r, "k_dev", "Development (participation fee, devex, signing and DD)", "EUR k", NUM,
              "={c}{d_pf}+{c}{d_dev}+{c}{d_sdd}")
    r = rows3(r, "k_allin", "All-in cost to COD (securities excluded: returned)", "EUR k", NUM,
              "={c}{k_capex}+SUM({c}{k_conn}:{c}{k_dev})", bold=True, top=True)
    return r + 1


# Annual timeline: shared calendar rows, then one sub-block per project ({p} project column, {d} date column).
SUBROWS = [
    ("t_n", "Years from the COD year", "years", INT, None, "={x}${t_yr}-YEAR(Inputs!${p}${d_cod})"),
    ("t_cm", "Construction months by year end (value: construction period)", "months", INT,
     '=DATEDIF(Inputs!${p}${d_ntp},Inputs!${p}${d_cod},"m")',
     "=({x}${t_yr}+1-YEAR(Inputs!${p}${d_ntp}))*Months_In_Year-MONTH(Inputs!${p}${d_ntp})"),
    ("t_cs", "Construction share of the year (NTP to COD)", "%", PCT, None,
     "=MAX(0,MIN({x}{t_cm},${p}{t_cm})-MAX({x}{t_cm}-Months_In_Year,0))/${p}{t_cm}"),
    ("t_om", "Months operated by year end (value: operating life)", "months", INT, "=Life*Months_In_Year",
     "=({x}${t_yr}+1-YEAR(Inputs!${p}${d_cod}))*Months_In_Year-MONTH(Inputs!${p}${d_cod})"),
    ("t_op", "Operating share of the year (COD to end of life)", "%", PCT, None,
     "=MAX(0,MIN({x}{t_om},${p}{t_om})-MAX({x}{t_om}-Months_In_Year,0))/Months_In_Year"),
    ("t_cod", "COD year flag", "flag", FLAG, None, "=N({x}{t_n}=0)"),
    ("p_dev", "FEP outlays at their dates (signing, fee, securities, fees, devex)", "EUR k", NUM, None,
     "=SUMPRODUCT((YEAR(${d}${f_sdd}:${d}${f_dev})={x}${t_yr})*${p}${f_sdd}:${p}${f_dev})"),
    ("p_pbr", "Performance security released (bid security included)", "EUR k", NUM, None,
     "=(${p}${d_bs}+${p}${d_pb})*({x}${t_yr}=YEAR(Inputs!${p}${d_rel}))"),
    ("p_capex", "Construction capex (straight-line NTP to COD)", "EUR k", NUM, None, "=-${p}${k_capex}*{x}{t_cs}"),
    ("p_muni", "Municipal one-off at COD", "EUR k", NUM, None, "=-${p}${k_muni}*{x}{t_cod}"),
    ("p_inv", "Investment cash flow", "EUR k", NUM, None, "=SUM({x}{p_dev}:{x}{p_muni})"),
    ("p_deg", "Degradation factor (from the first full year)", "x", FAC, None, "=(1-Degradation)^MAX({x}{t_n}-1,0)"),
    ("p_gen", "Net generation", "GWh", NUM1, None,
     "=${p}${built}*Hours_P50*(1-Curtailment)*(1-Clipping)*{x}{t_op}*{x}{p_deg}/Thousand"),
    ("p_px", "PPA price", "EUR/MWh", NUM1, None, "=PPA_Price*(1+PPA_Esc)^{x}{t_n}*({x}{t_op}>0)"),
    ("p_rev", "Revenue", "EUR k", NUM, None, "={x}{p_gen}*{x}{p_px}"),
    ("p_idx", "Opex index (COD year = 1)", "x", FAC, None, "=(1+Opex_Esc)^{x}{t_n}"),
    ("p_fix", "O&M, BoP, lease, insurance, management, property tax", "EUR k", NUM, None,
     "=-SUM(Opex_Per_MW)*${p}${built}*{x}{p_idx}*{x}{t_op}"),
    ("p_adm", "Administration, audit, compliance", "EUR k", NUM, None,
     "=-Opex_Admin*${p}${Dash_MVA}/Lump_Basis*{x}{p_idx}*{x}{t_op}"),
    ("p_bg", "Balancing and grid charges", "EUR k", NUM, None, "=-{x}{p_gen}*(Balancing+Grid_Charges)*{x}{p_idx}"),
    ("p_mekh", "MEKH supervisory fee", "EUR k", NUM, None, "=-{x}{p_rev}*MEKH_Fee"),
    ("p_oc", "Opex contingency", "EUR k", NUM, None, "=SUM({x}{p_fix}:{x}{p_mekh})*Opex_Cont"),
    ("p_comm", "Community contribution", "EUR k", NUM, None, "=-{x}{p_gen}*${p}${Dash_Community}"),
    ("p_hipa", "HIPA local business tax", "EUR k", NUM, None, "=-{x}{p_rev}*HIPA_Rate"),
    ("p_ebitda", "EBITDA", "EUR k", NUM, None, "={x}{p_rev}+SUM({x}{p_fix}:{x}{p_hipa})"),
    ("p_cf", "Unlevered pre-tax cash flow", "EUR k", NUM, None, "={x}{p_inv}+{x}{p_ebitda}"),
    ("p_df", "Discount factor to COD (buyer's rate)", "x", FAC, None,
     "=(1+Discount_Rate)^((Inputs!${p}${d_cod}-{x}${t_end})/Days_In_Year)"),
    ("p_pv", "EBITDA valued at COD (total = COD price)", "EUR k", NUM, None, "={x}{p_ebitda}*{x}{p_df}"),
    ("p_cod", "COD exit: FEP cash flow (to COD, price and securities at COD)", "EUR k", NUM, None,
     "={x}{p_inv}*({x}{t_n}<=0)+{x}{t_cod}*($H{p_pv}+${p}${d_bs}+${p}${d_pb})"),
]
NO_TOTAL = {"t_n", "t_cm", "t_om", "t_cod", "p_deg", "p_px", "p_idx", "p_df"}
BOLD = {"p_inv", "p_rev", "p_ebitda", "p_cf", "p_cod"}


def timeline(r):
    bar(r, "Annual timeline", {"H": "Total"}, TL)
    r += 1
    R["t_yr"] = r
    label(r, "Year", "year")
    for x in TL:
        put(f"{x}{r}", "=First_Year+COLUMN()-COLUMN($J$1)", fmt=INT)
    r += 1
    R["t_end"] = r
    label(r, "Year end", "date")
    for x in TL:
        put(f"{x}{r}", f"=DATE({x}{R['t_yr']}+1,1,1)-1", fmt=DATE)
    r += 2
    proj = {}
    for c, nm in zip(P, NAMES):
        bar(r, f"{nm}: annual cash flow (EUR k)", {c: "Value", "H": "Total"}, TL)
        r += 1
        rows = {}
        for key, text, unit, fmt, const, f in SUBROWS:
            rows[key] = R[key] = r
            bold = key in BOLD
            label(r, text, unit, bold=bold)
            kw = dict(p=c, d=DCOL[c], d_cod=I["d_cod"], d_ntp=I["d_ntp"], d_rel=I["d_rel"])
            if const:
                put(f"{c}{r}", fx(const, **kw), fmt=INT)
            for x in TL:
                put(f"{x}{r}", fx(f, x=x, **kw),
                    fmt=fmt, bold=bold, top=bold and key != "p_inv")
            if key not in NO_TOTAL:
                put(f"H{r}", f"=SUM(J{r}:AS{r})", fmt=NUM2 if fmt == PCT else fmt, bold=bold,
                    top=bold and key != "p_inv")
            r += 1
        proj[c] = rows
        r += 1
    R["proj"] = proj
    return r


def cod_exit(r):
    bar(r, "COD exit: tax on the gain (price and profit: totals of the cash-flow blocks above)", PROJ_HEADS)
    r += 1
    pr = R["proj"]
    for key, text, f, bold in (("x_cit", "CIT on the gain", "=-MAX(0,$H${cod})*CIT_Rate", False),
                               ("x_post", "FEP profit, post-tax", "=$H${cod}+{c}{x_cit}", True)):
        R[key] = r
        label(r, text, "EUR k", bold=bold)
        for c in P:
            put(f"{c}{r}", fx(f, c=c, cod=pr[c]["p_cod"]), fmt=NUM, bold=bold)
        put(f"H{r}", f"=SUM(E{r}:G{r})", fmt=NUM, bold=bold)
        r += 1
    return r + 1


def checks(r):
    bar(r, "Checks (0 = OK)", PROJ_HEADS)
    r += 1
    pr = R["proj"]
    first = r
    defs = [("ch_wf", "RtB waterfall closes (proceeds = all amounts paid out)",
             lambda c: fx("=N(ABS({c}{w_proc}+{c}{w_sc}+{c}{w_dsa}-{c}{w_devsh}-{c}{f_rec})>Tolerance)", c=c)),
            ("ch_end", "Claim lines paid before the accrual end year",
             lambda c: fx("=SUMPRODUCT(--(YEAR({d}{d_pf}:{d}{d_dev})>=YEAR({c}{c_end})))", c=c, d=DCOL[c])),
            ("ch_route", "Turbine route is EU or CN",
             lambda c: fx('=N(AND({c}{Dash_Turbine}<>"EU",{c}{Dash_Turbine}<>"CN"))', c=c)),
            ("ch_capex", "Construction capex spread in full (timeline = capex)",
             lambda c: fx("=N(ABS({c}{k_capex}+$H$%d)>Tolerance)" % pr[c]["p_capex"], c=c)),
            ("ch_allin", "Investment cash flows = all-in cost to COD",
             lambda c: fx("=N(ABS({c}{k_allin}+$H$%d)>Tolerance)" % pr[c]["p_inv"], c=c)),
            ("ch_life", "Operating life inside the timeline",
             lambda c: "=N(ABS($H$%d-Life)>Tolerance)" % pr[c]["t_op"])]
    for key, text, f in defs:
        R[key] = r
        label(r, text, "flag")
        for c in P:
            put(f"{c}{r}", f(c), fmt=FLAG)
        put(f"H{r}", f"=SUM(E{r}:G{r})", fmt=FLAG)
        r += 1
    R["master"] = r
    label(r, "Master check", "flag", bold=True)
    put(f"H{r}", f"=SUM(H{first}:H{r - 1})", fmt=CHECK_FMT, bold=True, top=True)
    for op, colr in (("equal", OK_GREEN), ("notEqual", CHECK_RED)):
        ws.conditional_formatting.add(f"H{r}", CellIsRule(operator=op, formula=["0"], font=Font(color=colr, bold=True)))
    r += 2
    ws[f"C{r}"].value, ws[f"C{r}"].font = "***END***", font(GREY, bold=True)
    return r


def exec_fill():
    pr = R["proj"]
    ye = f"$J${R['t_end']}:$AS${R['t_end']}"
    yr = f"$J${R['t_yr']}:$AS${R['t_yr']}"
    for c in P:
        q = pr[c]
        tl = lambda k: f"$J${q[k]}:$AS${q[k]}"  # noqa: E731
        fl = f"{c}{R['f_sdd']}:{c}{R['f_rec']}"
        cod = f"YEAR(Inputs!{c}${I['d_cod']})+1"
        f = {"ex_pts": f"={c}{R['pt_tot']}", "ex_cash": f"={c}{R['d_tot']}",
             "ex_rp": f"={c}{R['r_profit']}", "ex_rcit": f"={c}{R['r_cit']}", "ex_rpost": f"={c}{R['r_post']}",
             "ex_rirr": f"=XIRR({fl},{DCOL[c]}{R['f_sdd']}:{DCOL[c]}{R['f_rec']})",
             "ex_rmoic": f'=SUMIF({fl},">0")/-SUMIF({fl},"<0")',
             "ex_rmva": f"={c}{R['r_profit']}/{c}{R['Dash_MVA']}",
             "ex_built": f"={c}{R['built']}", "ex_allin": f"={c}{R['k_allin']}/Thousand",
             "ex_cmw": f"={c}{R['k_allin']}/{c}{R['built']}", "ex_cmva": f"={c}{R['k_allin']}/{c}{R['Dash_MVA']}",
             "ex_eb": f"=SUMIF({yr},{cod},{tl('p_ebitda')})", "ex_gwh": f"=SUMIF({yr},{cod},{tl('p_gen')})",
             "ex_irr": f"=XIRR({tl('p_cf')},{ye})", "ex_npv": f"=XNPV(Discount_Rate,{tl('p_cf')},{ye})",
             "ex_price": f"=$H${q['p_pv']}/Thousand", "ex_pmw": f"=$H${q['p_pv']}/{c}{R['built']}",
             "ex_pmva": f"=$H${q['p_pv']}/{c}{R['Dash_MVA']}", "ex_xp": f"=$H${q['p_cod']}",
             "ex_xcit": f"={c}{R['x_cit']}", "ex_xpost": f"={c}{R['x_post']}",
             "ex_xirr": f"=XIRR({tl('p_cod')},{ye})",
             "ex_xmoic": f'=SUMIF({tl("p_cod")},">0")/-SUMIF({tl("p_cod")},"<0")'}
        fmts = {k[0]: k[3] for k in EXEC if not isinstance(k, str)}
        for key, form in f.items():
            put(f"{c}{R[key]}", form, fmt=fmts[key], bold=key in ("ex_rpost", "ex_irr", "ex_xpost"))
    # Totals (H): sums and ratios on the totals; IRRs per project only.
    for key, src in (("ex_cash", "d_tot"), ("ex_rp", "r_profit"), ("ex_rcit", "r_cit"), ("ex_rpost", "r_post"),
                     ("ex_built", "built"), ("ex_xcit", "x_cit"), ("ex_xpost", "x_post")):
        put(f"H{R[key]}", f"=H{R[src]}", fmt=NUM1 if key == "ex_built" else NUM,
            bold=key in ("ex_rpost", "ex_xpost"))
    for key in ("ex_eb", "ex_gwh", "ex_npv"):
        put(f"H{R[key]}", f"=SUM(E{R[key]}:G{R[key]})", fmt=NUM1 if key == "ex_gwh" else NUM)
    hf = f"H{R['f_sdd']}:H{R['f_rec']}"
    put(f"H{R['ex_rmoic']}", f'=SUMIF({hf},">0")/-SUMIF({hf},"<0")', fmt=MULT)
    put(f"H{R['ex_rmva']}", f"=H{R['r_profit']}/H{R['Dash_MVA']}", fmt=NUM1)
    put(f"H{R['ex_allin']}", f"=H{R['k_allin']}/Thousand", fmt=NUM1)
    put(f"H{R['ex_cmw']}", f"=H{R['k_allin']}/H{R['built']}", fmt=NUM1)
    put(f"H{R['ex_cmva']}", f"=H{R['k_allin']}/H{R['Dash_MVA']}", fmt=NUM1)
    for key in ("ex_xp",):
        put(f"H{R[key]}", f"=SUM(E{R[key]}:G{R[key]})", fmt=NUM)
    put(f"H{R['ex_price']}", f"=SUM(E{R['ex_price']}:G{R['ex_price']})", fmt=NUM1)
    put(f"H{R['ex_pmw']}", f"=H{R['ex_price']}*Thousand/H{R['built']}", fmt=NUM1)
    put(f"H{R['ex_pmva']}", f"=H{R['ex_price']}*Thousand/H{R['Dash_MVA']}", fmt=NUM1)


def build():
    global ws
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False
    build_inputs(wb)
    for col, w in zip("ABCDEFGHI", (2, 2, 54, 10, 13, 13, 13, 13, 2)):
        ws.column_dimensions[col].width = w
    for x in TL:
        ws.column_dimensions[x].width = 10
    ws["C2"].value = "HU wind bids: Bora (low bid), Iváncsa, Dunaújváros – RtB exit, build and hold, COD exit (EUR k)"
    ws["C2"].font = font(bold=True)
    r = key_inputs(wb, 5)
    r = exec_reserve(r)
    r = calc_blocks(r)
    r = timeline(r)
    r = cod_exit(r)
    checks(r)
    exec_fill()
    label(3, "Master check (0 = OK)", "flag", bold=True)
    put("H3", "=Master_Check", fmt=CHECK_FMT, bold=True)
    for op, colr in (("equal", OK_GREEN), ("notEqual", CHECK_RED)):
        ws.conditional_formatting.add("H3", CellIsRule(operator=op, formula=["0"], font=Font(color=colr, bold=True)))
    name(wb, "Master_Check", f"Dashboard!$H${R['master']}")
    ws.freeze_panes = "E5"
    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    ROWS.write_text(json.dumps({"R": R, "I": I}, indent=1))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
