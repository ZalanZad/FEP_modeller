"""HU wind bids v2: unlevered, exits tax-exempt, devex itemised, transparent COD cost stack (SPEC bora_v5_summary §26).

Base: reference/HU_Wind_bids_271026_flagged.xlsx (the user's latest copy). Edited in LibreOffice through UNO (row
inserts / deletes keep references intact). out/HU_Wind_bids.xlsx (§25) is not touched.
  §25.1  key input "JV with a developer" (Bora 1, others 0); every claim line × the flag; claim rows relabelled.
  §25.2  Inputs "Sale date, Case A (before HCSSZ)" 30-Sep-2028; Case A gets its own claims block (accrual end
         MIN(freeze, sale A); robust accrual rows, since land falls in the end year), DSA to sale A, receipts at sale A,
         labels A / B as in the summary, check "Sale date A before HCSSZ", the user's two comments removed.
  §26.1  RtB and COD exits fully exempt: the CIT rows are 0 ("exempt (participation exemption, total)").
  §26.2  Devex itemised on Inputs (BORA_Devex 02-Oct-26 rows 14–48 without the two connection-fee rows), per-project
         scale factor (0.25), project devex = land + development (amount-weighted dates); the category rows, the
         1,000 / 784.5 placeholders and the devex average date are removed; land and development each get a line in
         both RtB cases (claims with the 2× floor, dated flows) and in the COD stack.
  §26.3  COD exit cost stack in the summary; the timeline row "Developer JV buyout at RtB" becomes "RtB exit price
         forgone" (premium × MVA at the Case B sale date).
  §26.4  All row grouping and hiding removed on both sheets.
Then verify(): recalc, no errors, Master_Check 0, lint, independent Python recomputation (both RtB cases for all three
projects, the Iváncsa COD exit), out/dump_v2.txt.

    python3 projects/hu_wind_2/build_v2.py
"""
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[1]))
import build_lev  # noqa: E402
from build_lev import LAST, TL, add_name, col, letter, xirr  # noqa: E402
from build_lev2 import Sheet  # noqa: E402  (formula() converts Inputs! refs to the API grammar)
from fepm.audit.lint import lint, lint_values  # noqa: E402
from fepm.recalc import Office  # noqa: E402

ROOT = HERE.parents[1]
SRC = ROOT / "reference" / "HU_Wind_bids_271026_flagged.xlsx"
DEVEX_FILE = ROOT / "reference" / "BORA_Devex_20261002.xlsx"
OUT = HERE / "out" / "HU_Wind_bids_v2.xlsx"
RC = HERE / "out" / "HU_Wind_bids_v2_recalc.xlsx"
BASE_RC = HERE / "out" / "HU_Wind_bids_v2_base_recalc.xlsx"
MAPS = HERE / "out" / "rows_v2.json"
DUMP = HERE / "out" / "dump_v2.txt"
P = ["E", "F", "G"]
NAMES = ["Bora", "Iváncsa", "Dunaújváros"]
JV = [1, 0, 0]
SALE_A = date(2028, 9, 30)
SCALE = 0.25
CONTINUOUS = date(2027, 12, 31)                         # "Continuous" items: as the Bora model's Devex_FEP
JV_TAG = " on FEP's shareholder loan (JV only)"
EXEMPT = "CIT: exempt (participation exemption, total)"


def serial(d):
    return (d - date(1899, 12, 30)).days


class Rows:
    """Old row → new row after inserts (old position p: rows ≥ p move down) and deletes (rows p..p+n-1 removed)."""
    def __init__(self, ins=(), dels=()):
        self.ins, self.dels = list(ins), list(dels)

    def __call__(self, r):
        for p, n in self.dels:
            if p <= r < p + n:
                return None
        return r + sum(n for p, n in self.ins if p <= r) - sum(n for p, n in self.dels if p + n <= r)

    def block(self, pos):
        n = dict(self.ins)[pos]
        return list(range(self(pos) - n, self(pos)))


def devex_items():
    """The 33 items of BORA_Devex rows 14-48 without the connection-fee rows 18-19 (task, item, date, HUF, class)."""
    v = openpyxl.load_workbook(DEVEX_FILE, data_only=True).active
    out, task = [], None
    for r in range(14, 49):
        if r in (18, 19):
            continue
        task = (v.cell(r, 2).value or task).strip()
        d = v.cell(r, 5).value
        out.append((r, task, str(v.cell(r, 8).value).strip(), d.date() if isinstance(d, datetime) else None,
                    v.cell(r, 9).value, v.cell(r, 12).value))
    return out


# ---------------------------------------------------------------------------------------------------------------------
def inputs_inserts(doc):
    """Inputs, before the Dashboard edits: scale row, Sale date A, devex items block (names). Deletes come last."""
    inp = Sheet(doc, "Inputs")
    m = Rows()
    r_scale = inp.find("Devex: grid connection plan and agreement")
    r_sale = inp.find("Sale date (RtB exit)")
    r_band = inp.find("Band tables for the points", exact=False)
    r_sub = inp.find("RtB exit terms (all projects", exact=False)
    r_date = inp.find("Signing fee and DD paid")
    r_fee = inp.find("Indirect connection fee")
    r_floor = inp.find("Devex minimum return multiple")
    items = devex_items()
    n_blk = len(items) + 2                                   # sub-header, items, blank
    inp.insert(r_band, n_blk)
    m.ins.append((r_band, n_blk))
    inp.insert(r_sale + 1, 1)
    m.ins.append((r_sale + 1, 1))
    idx = [doc.Sheets.getByIndex(i).Name for i in range(doc.Sheets.Count)].index("Inputs")
    # Scale factor (replaces the devex category row; rows 19-24 are deleted at the end).
    inp.text(f"C{r_scale}", "Devex scale (bid MVA ÷ 200 MVA of the devex file)")
    inp.text(f"D{r_scale}", "x")
    for c in P:
        if c != "E":                                          # blue input style, as column E
            inp.s.copyRange(inp.cell(f"{c}{r_scale}").getCellAddress(), inp.cell(f"E{r_scale}").getRangeAddress())
        inp.number(f"{c}{r_scale}", SCALE)
        inp.fmt(f"{c}{r_scale}", f"E{r_floor}")
    inp.text(f"H{r_scale}", "SPEC §26.2: 50/200; 1 = the full devex file")
    # Sale date A.
    r = m.block(r_sale + 1)[0]
    inp.copy_format(r_sale, r, last="H")
    inp.text(f"C{r}", "Sale date, Case A (before HCSSZ)")
    inp.text(f"D{r}", "date")
    inp.number(f"E{r}", serial(SALE_A))
    inp.text(f"H{r}", "SPEC §25.2 (user, 28-Oct-2026)")
    add_name(doc, "Sale_Date_A", idx, col("E"), r)
    inp.text(f"C{r_sale}", "Sale date, Case B (RtB exit after HCSSZ)")
    # Devex items.
    rows = m.block(r_band)
    head, first = rows[0], rows[1]
    inp.copy_format(r_sub, head, last="H")
    inp.text(f"C{head}", "Devex items (BORA_Devex 02-Oct-26; all projects × the scale): HUF | date | class")
    for k, (src, task, item, d, huf, cls) in enumerate(items):
        r = first + k
        inp.copy_format(r_date, r, last="H")
        inp.text(f"C{r}", f"{task} – {item}")
        inp.text(f"D{r}", "HUF")
        inp.number(f"E{r}", huf)
        inp.fmt(f"E{r}", f"E{r_fee}")
        inp.number(f"F{r}", serial(d or CONTINUOUS))
        inp.text(f"G{r}", cls)
        inp.fmt(f"G{r}", f"C{r_sub}")
        note = f"BORA_Devex 02-Oct-26 row {src}" + ("; continuous: 31-Dec-27" if d is None else "")
        inp.text(f"H{r}", note)
    last = first + len(items) - 1
    for name, c in (("Devex_HUF", "E"), ("Devex_Date", "F"), ("Devex_Class", "G")):
        doc.NamedRanges.addNewByName(name, f"$Inputs.${c}${first}:${c}${last}",
                                     _addr(idx, col(c), first), 0)
    return m, {"scale": r_scale, "items": (first, last)}


def _addr(sheet, c, r):
    from com.sun.star.table import CellAddress
    return CellAddress(Sheet=sheet, Column=c, Row=r - 1)


def inputs_deletes(doc, m):
    """Remove the devex category rows, the 1,000 placeholder and the devex average date (no references left)."""
    inp = Sheet(doc, "Inputs")
    d = Sheet(doc, "Dashboard")
    r_cat = inp.find("Devex: land (reservation, purchase, contracting)")
    r_avg = inp.find("Devex average payment date")
    gone = set(range(r_cat, r_cat + 6)) | {r_avg}
    used = d.s.getCellRangeByName(f"A1:{LAST}700").getFormulaArray()
    pat = re.compile(r"\$?Inputs\.\$?[A-Z]+\$?(\d+)")
    for row in used:
        for f in row:
            for hit in pat.findall(f):
                if int(hit) in gone:
                    raise RuntimeError(f"Dashboard still refers to deleted Inputs row {hit}: {f}")
    inp.s.Rows.removeByIndex(r_avg - 1, 1)
    inp.s.Rows.removeByIndex(r_cat - 1, 6)
    m.dels += [(r_avg, 1), (r_cat, 6)]
    return m


# ---------------------------------------------------------------------------------------------------------------------
def dashboard(doc):
    d = Sheet(doc, "Dashboard")
    f = lambda t, s=1: d.find(t, start=s, exact=False)  # noqa: E731
    a = {"dsa_key": d.find("Developer success fee (DSA) on"), "wind": d.find("Wind committed in the bid"),
         "mva": d.find("Net connection capacity"), "prem": d.find("RtB premium")}
    huf = f("Guarantees, securities and connection fees, DEVEX (HUF)")
    eur = f("Guarantees, securities and connection fees, DEVEX (EUR k)")
    a.update(h_dev=d.find("DEVEX", huf), h_sdd=d.find("Signing fee and DD", huf),
             e_fee=d.find("Participation fee", eur), e_bs=d.find("Bid security", eur),
             e_pb=f("Performance security payable", eur), e_tso=f("TSO connection-request", eur),
             e_dev=d.find("DEVEX", eur), e_sdd=d.find("Signing fee and DD", eur), sunk=f("Sunk fees and devex"))
    a.update(k_capex=d.find("Construction capex"), k_conn=f("Connection fees (TSO security"),
             k_muni=d.find("Municipal one-off"), k_dev=f("Development (participation fee"))
    brk = f("Devex breakdown")
    a.update(b1=brk + 1, b2=brk + 2, b_tot=f("Devex total", brk))
    a.update(cod_tot=f("Total capex and costs to COD (incl. developer JV buyout"))
    cb = f("RtB exit: FEP claims at the sale")
    a.update(cb_end=f("Interest accrual end", cb), cb_s20=f("Accrual factor in the end year, devex", cb),
             cb_ddays=f("Days to the first 31-Dec from devex", cb), cb_dfac=f("Accrual factor from devex", cb),
             cb_claims=[f("Claim: ", cb) + i for i in range(7)], cb_tot=d.find("Total FEP claims", cb),
             cb_head=cb, cb_wf=f("RtB exit: waterfall at the sale", cb))
    a.update(cb_c12=f("Less FEP claims at 12 %", a["cb_wf"]), cb_cdev=f("Less FEP devex return", a["cb_wf"]),
             cb_proc=d.find("Total proceeds", a["cb_wf"]), cb_flows=f("FEP cash flows, dated (outlays", a["cb_wf"]))
    a.update(cb_fdev=d.find("Devex", a["cb_flows"]), cb_cit=f("CIT: claims taxed", a["cb_flows"]))
    ca = f("RtB exit, Case B: buyer pays the grid connection fees")
    a.update(ca_head=ca, ca_out=f("FEP cash out to RtB, Case B", ca), ca_tot=f("Total FEP claims, Case B", ca),
             ca_wf=f("Waterfall at the sale, Case B", ca))
    a.update(ca_c12=f("Less FEP claims at 12 %", a["ca_wf"]), ca_cdev=f("Less FEP devex return", a["ca_wf"]),
             ca_dsa=f("Less developer success fee (DSA)", a["ca_wf"]), ca_proc=d.find("Total proceeds", a["ca_wf"]),
             ca_flows=f("FEP cash flows, dated, Case B", ca))
    a.update(ca_fdev=d.find("Devex", a["ca_flows"]), ca_rec=f("FEP receipts at the sale", a["ca_flows"]),
             ca_cit=f("CIT: claims taxed", a["ca_flows"]), ca_memo=f("Memo: pre-tax profit vs Case A", ca))
    blocks = [(nm, f(f"{nm}: annual cash flow")) for nm in NAMES]
    a.update(yr=d.find("Year", f("Annual timeline")), cod_cit=f("CIT on the gain (nil if share deal", blocks[-1][1]),
             chk_b=f("RtB waterfall closes"), chk_a=f("Case B waterfall closes"), chk_dev=f("Devex breakdown = DEVEX"))
    tl = {}
    for c, (nm, h) in zip(P, blocks):
        tl[c] = {"h": h, "inv": d.find("Investment cash flow", h), "jvb": f("Developer JV buyout at RtB", h)}

    # Case A claims block: end, start of the end year, 3 rows per event, 6 claims, DSA.
    events = [("bid", "bid submission", "Rate_Cash"), ("pb", "performance security posting", "Rate_Cash"),
              ("tso", "TSO security payment", "Rate_Cash"), ("land", "land payment (amount-weighted)", "Rate_Devex"),
              ("dev", "development devex payment (amount-weighted)", "Rate_Devex")]
    claims_a = [("fee", "participation fee (not reimbursed)", "e_fee", "bid"), ("bs", "bid security (cash)", "e_bs", "bid"),
                ("pb", "performance security payable (cash)", "e_pb", "pb"),
                ("tso", "TSO connection-request security", "e_tso", "tso"),
                ("land", "land, MAX(20 % accrued, 1× principal)", "e_land", "land"),
                ("dev", "development devex, MAX(20 % accrued, 1× principal)", "e_dev", "dev")]
    n_a = 2 + 3 * len(events) + len(claims_a) + 1
    ops = [(a["chk_a"], 2), (a["ca_fdev"], 1), (a["ca_tot"], n_a), (a["cb_fdev"], 1), (a["cb_claims"][6], 1),
           (a["cb_ddays"], 2), (a["cod_tot"], 8), (a["k_dev"], 1), (a["e_sdd"], 1), (a["h_sdd"], 1),
           (a["dsa_key"] + 1, 1), (a["ca_out"], 6)]
    dels = [(a["b1"] + 2, 4)]                                # breakdown: keep two rows (land, development)
    seq = sorted([(p, n, "ins") for p, n in ops] + [(p, n, "del") for p, n in dels], key=lambda x: -x[0])
    for p, n, kind in seq:
        if kind == "ins":
            d.insert(p, n)
        else:
            d.s.Rows.removeByIndex(p - 1, n)
    m = Rows(ops, dels)
    R = {k: (m(v) if isinstance(v, int) else [m(x) for x in v]) for k, v in a.items()}
    for c in P:
        tl[c] = {k: m(v) for k, v in tl[c].items()}
    jv = m.block(a["dsa_key"] + 1)[0]
    h_land, e_land = m.block(a["h_sdd"])[0], m.block(a["e_sdd"])[0]
    R.update(jv=jv, h_land=h_land, e_land=e_land, k_land=m.block(a["k_dev"])[0])

    # §25.1: JV flag; claims × flag; labels.
    d.copy_format(R["dsa_key"], jv, last="H")
    d.text(f"C{jv}", "JV with a developer (1 = yes; claims and DSA apply)")
    d.text(f"D{jv}", "flag")
    for c, v in zip(P, JV):
        d.number(f"{c}{jv}", v)
    doc.NamedRanges.addNewByName("Dash_JV", f"$Dashboard.$E${jv}:$G${jv}", _addr(0, col("E"), jv), 0)
    for r in R["cb_claims"]:
        for c in P:
            d.cell(f"{c}{r}").setFormula(f"={c}${jv}*" + d.cell(f"{c}{r}").getFormula()[1:])
        d.text(f"C{r}", d.cell(f"C{r}").getString().replace("Claim: ", "Claim" + JV_TAG + ": ")
               .replace(": devex,", ": development devex,"))
    d.text(f"C{R['cb_tot']}", "Total FEP claims on its shareholder loans (JV only)")
    for r in (R["cb_c12"], R["ca_c12"]):
        d.text(f"C{r}", "Less FEP claims at 12 %" + JV_TAG)
    for r in (R["cb_cdev"], R["ca_cdev"]):
        d.text(f"C{r}", "Less FEP land and devex return" + JV_TAG)

    # §26.2: devex land / development (HUF, EUR k, dates), sunk row, capex block, breakdown.
    wdate = 'SUMPRODUCT((Devex_Class="{k}")*Devex_HUF,Devex_Date)/SUMIF(Devex_Class,"{k}",Devex_HUF)'
    for r, k, lab in ((R["h_dev"], "Devex", "DEVEX – development (devex items excl. land)"),
                      (h_land, "Land", "DEVEX – land (reservation and purchase)")):
        if r == h_land:
            d.copy_format(R["h_dev"], r)
        d.text(f"C{r}", lab)
        d.text(f"D{r}", "HUF")
        for c in P:
            d.formula(f"{c}{r}", f'=SUMIF(Devex_Class,"{k}",Devex_HUF)*Inputs!{c}${18}')
        d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
        d.formula(f"J{r}", "=" + wdate.format(k=k))
    for r, src, lab in ((R["e_dev"], R["h_dev"], "DEVEX – development (devex items excl. land)"),
                        (e_land, h_land, "DEVEX – land (reservation and purchase)")):
        if r == e_land:
            d.copy_format(R["e_dev"], r)
        d.text(f"C{r}", lab)
        d.text(f"D{r}", "EUR k")
        for c in P:
            d.formula(f"{c}{r}", f"={c}{src}/FX/Thousand")
        d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
        d.formula(f"J{r}", f"=J{src}")
    for c in P:
        d.formula(f"{c}{R['sunk']}", f"={c}{R['e_fee']}+SUM({c}{R['e_dev']}:{c}{R['e_sdd']})")
    d.text(f"C{R['sunk']}", "Sunk fees and devex (signing and DD, participation fee, land, development devex)")
    r = R["k_land"]
    d.copy_format(R["k_muni"], r)
    d.text(f"C{r}", "Land (devex land items)")
    d.text(f"D{r}", "EUR k")
    for c in P:
        d.formula(f"{c}{r}", f"={c}{e_land}")
    d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
    d.text(f"C{R['k_dev']}", "Development (devex excl. land, participation fee, signing and DD)")
    for r, src, lab in ((R["b1"], e_land, "Land"), (R["b2"], R["e_dev"], "Development devex (incl. contingency)")):
        d.text(f"C{r}", lab)
        for c in P:
            d.formula(f"{c}{r}", f"={c}{src}")
    d.text(f"C{R['b_tot']}", "Devex total (= DEVEX lines above; devex items × scale)")
    d.text(f"C{R['b_tot'] - 3}", "Devex breakdown (BORA_Devex 02-Oct-26 items × the scale; tender fee and securities above)")

    # Case B: land days / factor, devex rows on the weighted development date, land claim, devex return, land flow.
    rd, rl = m.block(a["cb_ddays"])
    for r, src, x, lab in ((rd, R["cb_ddays"], f"$J${h_land}", "land payment (amount-weighted)"),
                           (rl, R["cb_dfac"], f"$J${h_land}", "land payment (amount-weighted)")):
        d.copy_format(src, r, last="H")
    for (rdays, rfac, x, lab) in ((rd, rl, f"$J${h_land}", "land payment (amount-weighted)"),
                                  (R["cb_ddays"], R["cb_dfac"], f"$J${R['h_dev']}",
                                   "development devex payment (amount-weighted)")):
        d.text(f"C{rdays}", f"Days to the first 31-Dec from {lab}")
        d.text(f"C{rfac}", f"Accrual factor from {lab}")
        d.text(f"D{rdays}", "days")
        d.text(f"D{rfac}", "x")
        for c in P:
            d.formula(f"{c}{rdays}", f"=DATE(YEAR({x})+1,1,1)-{x}")
            d.formula(f"{c}{rfac}", f"=(1+Rate_Devex*{c}{rdays}/Days_In_Year)*(1+Rate_Devex)^(YEAR({c}{R['cb_end']})"
                                    f"-YEAR({x})-1)*{c}{R['cb_s20']}")
    rc = m.block(a["cb_claims"][6])[0]
    d.copy_format(R["cb_claims"][6], rc)
    d.text(f"C{rc}", "Claim" + JV_TAG + ": land, MAX(20 % accrued, 1× principal)")
    d.text(f"D{rc}", "EUR k")
    for c in P:
        d.formula(f"{c}{rc}", f"={c}${jv}*{c}{e_land}*MAX({c}{rl}-1,Devex_Floor-1)")
    d.formula(f"H{rc}", f"=SUM(E{rc}:G{rc})")
    p = R["cb_proc"]
    for c in P:
        d.formula(f"{c}{R['cb_cdev']}", f"=-MIN(SUM({c}{p}:{c}{p + 3}),SUM({c}{rc}:{c}{R['cb_claims'][6]}))")
    for r_line, r_dev in ((m.block(a["cb_fdev"])[0], R["cb_fdev"]), (m.block(a["ca_fdev"])[0], R["ca_fdev"])):
        d.copy_format(r_dev, r_line)
        d.text(f"C{r_line}", "Land")
        d.text(f"D{r_line}", "EUR k")
        for c in P:
            d.formula(f"{c}{r_line}", f"=-{c}{e_land}")
        d.formula(f"H{r_line}", f"=SUM(E{r_line}:G{r_line})")
        for dc in "JKL":                                      # the user's date columns per project
            d.formula(f"{dc}{r_line}", f"=$J${h_land}")
            d.formula(f"{dc}{r_dev}", f"=$J${R['h_dev']}")
        d.text(f"C{r_dev}", "Development devex")

    # §25.2: Case A devex = the items paid by sale A (later items are the buyer's), above the Case A cash out.
    q = {}
    rows_a = iter(m.block(a["ca_out"]))
    sub_a = '"<="&Sale_Date_A'
    for k, cls, lab in (("l", "Land", "land"), ("d", "Devex", "development")):
        q["h" + k], q["d" + k], q["e" + k] = next(rows_a), next(rows_a), next(rows_a)
        _row(d, q["h" + k], R["h_dev"], f"Case A devex – {lab}: items paid by sale A (devex file, unscaled)", "HUF",
             f'=SUMIFS(Devex_HUF,Devex_Class,"{cls}",Devex_Date,{sub_a})', False)
        _row(d, q["d" + k], R["cb_end"], f"Case A devex – {lab}: amount-weighted payment date", "date",
             f'=SUMPRODUCT((Devex_Class="{cls}")*(Devex_Date<=Sale_Date_A)*Devex_HUF,Devex_Date)/{{c}}{q["h" + k]}',
             False)
        _row(d, q["e" + k], R["e_dev"], f"Case A devex – {lab} (× scale)", "EUR k",
             "={c}%d/FX/Thousand*Inputs!{c}$18" % q["h" + k], True)
    for c in P:
        d.formula(f"{c}{R['ca_out']}", f"={c}{R['e_fee']}+SUM({c}{R['e_bs']}:{c}{R['e_tso']})+{c}{q['el']}"
                                       f"+{c}{q['ed']}+{c}{R['e_sdd']}")
    for r_line, src, dr in ((m.block(a["ca_fdev"])[0], q["el"], q["dl"]), (R["ca_fdev"], q["ed"], q["dd"])):
        for c, dc in zip(P, "JKL"):
            d.formula(f"{c}{r_line}", f"=-{c}{src}")
            d.formula(f"{dc}{r_line}", f"={c}{dr}")

    # §25.2: Case A claims block (robust accrual rows), totals, waterfall, receipts at sale A, labels.
    rows = iter(m.block(a["ca_tot"]))
    q.update({"end": next(rows), "ys": next(rows)})
    for ev, *_ in events:
        q["a_" + ev], q["b_" + ev], q["f_" + ev] = next(rows), next(rows), next(rows)
    for key, *_ in claims_a:
        q["c_" + key] = next(rows)
    q["dsa"] = next(rows)
    put = lambda r, tmpl, lab, unit, fml, total=False: _row(d, r, tmpl, lab, unit, fml, total)  # noqa: E731
    put(q["end"], R["cb_end"], "Case A: interest accrual end (earlier of freeze and sale A)", "date",
        "=MIN(Freeze_Date,Sale_Date_A)")
    put(q["ys"], R["cb_end"], "Case A: start of the accrual end year", "date", "=DATE(YEAR({c}%d),1,1)" % q["end"])
    xs = {"bid": "Inputs!{c}$27", "pb": "Inputs!{c}$28", "tso": "Inputs!{c}$29", "land": "{c}%d" % q["dl"],
          "dev": "{c}%d" % q["dd"]}
    for ev, text, rate in events:
        x = xs[ev]
        put(q["a_" + ev], R["cb_dfac"], f"Case A, from {text}: factor in the payment year (if before the end year)",
            "x", "=1+%s*(DATE(YEAR(%s)+1,1,1)-%s)*(%s<{c}%d)/Days_In_Year" % (rate, x, x, x, q["ys"]))
        put(q["b_" + ev], R["cb_dfac"], f"Case A, from {text}: factor in the end year (to the accrual end)", "x",
            "=1+%s*MAX(0,{c}%d-MAX(%s,{c}%d)+1)/Days_In_Year" % (rate, q["end"], x, q["ys"]))
        put(q["f_" + ev], R["cb_dfac"], f"Case A, from {text}: accrual factor", "x",
            "={c}%d*(1+%s)^MAX(0,YEAR({c}%d)-YEAR(%s)-1)*{c}%d" % (q["a_" + ev], rate, q["ys"], x, q["b_" + ev]))
    for key, text, line, ev in claims_a:
        src = {"e_land": q["el"], "e_dev": q["ed"]}.get(line) or R[line]
        fml = ("={c}$%d*{c}%d*MAX({c}%d-1,Devex_Floor-1)" if ev in ("land", "dev") else "={c}$%d*{c}%d*({c}%d-1)") % (
            jv, src, q["f_" + ev])
        put(q["c_" + key], R["cb_claims"][0], f"Case A: claim{JV_TAG}: {text}", "EUR k", fml, total=True)
    put(q["dsa"], R["cb_claims"][0], "Case A: developer success fee due (DSA to sale A, capped)", "EUR k",
        "={c}%d*MIN(DSA_Rate*YEARFRAC(Inputs!{c}$35,Sale_Date_A),DSA_Cap)*{c}%d" % (R["dsa_key"], R["wind"]),
        total=True)
    r = R["ca_tot"]
    d.text(f"C{r}", "Total FEP claims on its shareholder loans, Case A (JV only; no connection-fee claims)")
    for c in P:
        d.formula(f"{c}{r}", f"=SUM({c}{q['c_fee']}:{c}{q['c_dev']})")
    p = R["ca_proc"]
    for c in P:
        d.formula(f"{c}{R['ca_c12']}", f"=-MIN(SUM({c}{p}:{c}{p + 2}),SUM({c}{q['c_fee']}:{c}{q['c_tso']}))")
        d.formula(f"{c}{R['ca_cdev']}", f"=-MIN(SUM({c}{p}:{c}{p + 3}),SUM({c}{q['c_land']}:{c}{q['c_dev']}))")
        d.formula(f"{c}{R['ca_dsa']}", f"=-MIN(SUM({c}{p}:{c}{p + 4}),{c}{q['dsa']})")
    for dc in "JKL":
        d.formula(f"{dc}{R['ca_rec']}", "=Sale_Date_A")
    relabel = {"cb_head": ("RtB exit: ", "RtB exit, Case B: "), "cb_wf": ("RtB exit: ", "RtB exit, Case B: "),
               "cb_flows": ("FEP cash flows, dated", "FEP cash flows, dated, Case B"),
               "ca_head": ("Case B", "Case A"), "ca_out": ("Case B", "Case A"), "ca_wf": ("Case B", "Case A"),
               "ca_flows": ("Case B (no connection-fee outlays)", "Case A (no connection-fee outlays; sale A)"),
               "ca_memo": ("vs Case A", "vs Case B"), "chk_b": ("RtB waterfall closes", "Case B waterfall closes"),
               "chk_a": ("Case B", "Case A")}
    for k, (x, y) in relabel.items():
        d.text(f"C{R[k]}", d.cell(f"C{R[k]}").getString().replace(x, y, 1))
    ann = d.s.getAnnotations()
    for i in reversed(range(ann.getCount())):
        pos = ann.getByIndex(i).getPosition()
        if f"{letter(pos.Column)}{pos.Row + 1}" in (f"G{R['cb_c12']}", f"J{R['ca_flows']}"):
            ann.removeByIndex(i)

    # §26.1: exits exempt.
    for r in (R["cb_cit"], R["ca_cit"], R["cod_cit"]):
        d.text(f"C{r}", EXEMPT)
        for c in P:
            d.formula(f"{c}{r}", "=0")

    # §26.3: RtB price forgone in the timeline (replaces the developer JV buy-out); COD cost stack in the summary.
    yr = R["yr"]
    for c in P:
        r = tl[c]["jvb"]
        d.text(f"C{r}", "RtB exit price forgone (premium × MVA, at the Case B sale; replaces the JV buy-out)")
        for x in TL:
            d.formula(f"{x}{r}", f"=-${c}${R['prem']}*${c}${R['mva']}*({x}${yr}=YEAR(Sale_Date))")
    stack = m.block(a["cod_tot"])
    lines = [("Construction capex", lambda c: f"={c}{R['k_capex']}"),
             ("Connection fees (direct, indirect; the TSO security is credited against the indirect fee)",
              lambda c: f"={c}{R['k_conn']}"),
             ("Municipal one-off", lambda c: f"={c}{R['k_muni']}"),
             ("Land", lambda c: f"={c}{e_land}"),
             ("Development (devex excl. land, participation fee, signing and DD)", lambda c: f"={c}{R['k_dev']}"),
             ("Bid security: returned at COD (cash out and back; net 0)", lambda c: "=0"),
             ("Performance security: returned at COD (cash out and back; net 0)", lambda c: "=0"),
             ("RtB exit price forgone (150 k€/MVA × MVA, at the Case B sale date)",
              lambda c: f"=-$H${tl[c]['jvb']}")]
    for r, (lab, fml) in zip(stack, lines):
        d.copy_format(R["cod_tot"], r, last="H")
        d.text(f"C{r}", lab)
        d.text(f"D{r}", "EUR k")
        for c in P:
            d.formula(f"{c}{r}", fml(c))
        d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
    r = R["cod_tot"]
    d.text(f"C{r}", "Total costs to COD, incl. the RtB exit price forgone")
    for c in P:
        d.formula(f"{c}{r}", f"=SUM({c}{stack[0]}:{c}{stack[-1]})")
    d.bold(f"C{r}:H{r}")

    # Checks.
    r = R["chk_dev"]
    d.text(f"C{r}", "Devex items × scale = project devex (land + development)")
    for c in P:
        d.formula(f"{c}{r}", f"=(ABS(SUM(Devex_HUF)/FX/Thousand*Inputs!{c}$18-{c}{e_land}-{c}{R['e_dev']})"
                             f">Tolerance)*1")
    r1, r2 = m.block(a["chk_a"])
    d.copy_format(R["chk_b"], r1, last="H")
    d.copy_format(R["chk_b"], r2, last="H")
    d.text(f"C{r1}", "Sale date A before the HCSSZ date (Case A exits before the connection fees)")
    d.text(f"C{r2}", "COD cost stack = timeline flows (investment flows + RtB price forgone)")
    for r in (r1, r2):
        d.text(f"D{r}", "flag")
        d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
    for c in P:
        d.formula(f"{c}{r1}", f"=N(Sale_Date_A>=Inputs!{c}$30)")
        d.formula(f"{c}{r2}", f"=N(ABS({c}{R['cod_tot']}+$H${tl[c]['inv']}+$H${tl[c]['jvb']})>Tolerance)")
    return m


def _row(d, r, tmpl, lab, unit, fml, total):
    d.copy_format(tmpl, r, last="H")
    d.text(f"C{r}", lab)
    d.text(f"D{r}", unit)
    for c in P:
        d.formula(f"{c}{r}", fml.format(c=c))
    if total:
        d.formula(f"H{r}", f"=SUM(E{r}:G{r})")


def unhide(doc):
    for name in ("Dashboard", "Inputs"):
        s = doc.Sheets.getByName(name)
        s.clearOutline()
        rows, cols = s.getRows(), s.getColumns()
        for i in range(800):
            rows.getByIndex(i).IsVisible = True
        for i in range(60):
            cols.getByIndex(i).IsVisible = True


def build():
    with Office(SRC) as lo:
        lo.save(BASE_RC)
        mi, info = inputs_inserts(lo.doc)
        md = dashboard(lo.doc)
        mi = inputs_deletes(lo.doc, mi)
        unhide(lo.doc)
        lo.calculate()
        lo.save(OUT)
    with Office(OUT) as lo:
        lo.save(RC)
    MAPS.write_text(json.dumps({"Dashboard": [md.ins, md.dels], "Inputs": [mi.ins, mi.dels]}))
    print(f"wrote {OUT}")


# ---------------------------------------------------------------------------------------------------------------------
def changed(b_wb, a_wb):
    maps = json.loads(MAPS.read_text())
    out = []
    for name in ("Dashboard", "Inputs"):
        b, a = b_wb[name], a_wb[name]
        m = Rows(*maps[name])
        for r in range(1, b.max_row + 1):
            r1 = m(r)
            if r1 is None:
                out.append((name, r, None, b.cell(r, 3).value, "row", "deleted", ""))
                continue
            for c in range(2, min(b.max_column, 46) + 1):
                x, y = b.cell(r, c).value, a.cell(r1, c).value
                if not ((x == y) or (isinstance(x, (int, float)) and isinstance(y, (int, float))
                                     and abs(x - y) <= 1e-9 * max(1, abs(x)))):
                    out.append((name, r, r1, b.cell(r, 3).value, b.cell(r, c).column_letter, x, y))
    return out


def verify():
    a_wb = openpyxl.load_workbook(RC, data_only=True)
    errs = lint_values(RC)
    print("recalc:", "no errors, Master_Check 0" if not errs else [str(e) for e in errs[:10]])
    found = [str(x) for x in lint(OUT)]
    print("lint:", "clean" if not found else f"{len(found)} finding(s)")
    for x in found[:15]:
        print("  ", x)
    wb = openpyxl.load_workbook(OUT)
    print("hidden / grouped rows:", sum(1 for s in wb.worksheets for _, d in s.row_dimensions.items()
                                        if d.hidden or d.outlineLevel),
          "| comments:", [(s.title, c.coordinate) for s in wb.worksheets for row in s.iter_rows() for c in row
                          if c.comment])
    diffs = changed(openpyxl.load_workbook(BASE_RC, data_only=True), a_wb)
    print(f"before / after: {len(diffs)} changed cell(s) in the user's rows")
    ok = not errs and recompute(a_wb)
    build_lev.DUMP = DUMP
    build_lev.dump(a_wb)
    print("dump:", DUMP)
    return ok


def recompute(wb):
    dash, inp = wb["Dashboard"], wb["Inputs"]
    lab = lambda r, ws=dash: ws.cell(r, 3).value or ""  # noqa: E731
    find = lambda t, s=1: next(r for r in range(s, dash.max_row + 1) if lab(r).startswith(t))  # noqa: E731
    irow = lambda t: next(r for r in range(1, inp.max_row + 1) if lab(r, inp).startswith(t))  # noqa: E731
    iv = lambda t, c="E": inp[f"{c}{irow(t)}"].value  # noqa: E731
    dd = lambda x: x.date() if hasattr(x, "date") else x  # noqa: E731
    # Devex items, independently from the Inputs block.
    h = irow("Devex items (BORA_Devex")
    items = [(inp.cell(r, 5).value, dd(inp.cell(r, 6).value), inp.cell(r, 7).value) for r in range(h + 1, h + 34)]
    fx = iv("FX rate")
    tot = {k: sum(x for x, _, c in items if c == k) for k in ("Land", "Devex")}
    wdate = {k: date.fromordinal(int(sum(x * t.toordinal() for x, t, c in items if c == k) / tot[k]))
             for k in tot}                                  # XIRR counts whole days (fraction truncated)
    rc, rd, floor = iv("Claim rate, cash"), iv("Claim rate, devex"), iv("Devex minimum return")
    freeze, sale_b, sale_a = dd(iv("Interest freeze")), dd(iv("Sale date, Case B")), dd(iv("Sale date, Case A"))
    sc, dsa_rate, dsa_cap = iv("Sale costs"), iv("Developer success fee (on"), iv("Developer success fee cap")

    def factor(p, r, end):
        if p > end:
            return 1.0
        bal, y, start = 1.0, p.year, p
        while True:
            stop = min(date(y, 12, 31), end)
            full = start == date(y, 1, 1) and stop == date(y, 12, 31)
            bal *= 1 + r * (1 if full else ((stop - start).days + 1) / 365)
            if stop == end:
                return bal
            y, start = y + 1, date(y + 1, 1, 1)

    def yf360(a, b):
        return ((b.year - a.year) * 360 + (b.month - a.month) * 30 + min(b.day, 30) - min(a.day, 30)) / 360

    def inputs(col):
        k = lambda t: dash[f"{col}{find(t)}"].value  # noqa: E731
        pi = lambda t: iv(t, col)  # noqa: E731
        mva, ob, bsh, hrs, pbm = k("Net connection"), k("Overbuild"), k("BESS power"), k("BESS duration"), \
            k("Performance security multiple")
        wind, mwh = mva * (1 + ob), mva * bsh * hrs
        base = (wind * iv("Security base: benchmark investment, wind") + mwh * iv("Security base: benchmark "
                                                                                  "investment, BESS")) / fx / 1000
        scale = pi("Devex scale")
        amt = {"fee": iv("Participation fee per bid") / fx / 1000, "bs": iv("Bid security, share") * base}
        amt["pb"] = iv("Performance security, minimum") * pbm * base - amt["bs"]
        amt["tso"] = iv("TSO connection-request security") * mva / fx / 1000
        amt["dir"] = pi("Direct connection fee per") / pi("Bids sharing") / fx / 1000
        amt["ind"] = pi("Indirect connection fee") * mva / fx / 1000 - amt["tso"]
        amt["land"], amt["dev"] = tot["Land"] * scale / fx / 1000, tot["Devex"] * scale / fx / 1000
        sub = {k: sum(x for x, t, c in items if c == k and t <= sale_a) for k in tot}
        amt_a = {"land": sub["Land"] * scale / fx / 1000, "dev": sub["Devex"] * scale / fx / 1000}
        date_a = {n: date.fromordinal(int(sum(x * t.toordinal() for x, t, c in items if c == k and t <= sale_a)
                                          / sub[k])) for n, k in (("land", "Land"), ("dev", "Devex"))}
        amt["sdd"] = pi("Signing fee and DD")
        pay = {"sdd": dd(pi("Signing fee and DD paid")), "fee": dd(pi("Bid submission")),
               "bs": dd(pi("Bid submission")), "pb": dd(pi("Performance security posted")),
               "tso": dd(pi("TSO security paid")), "dir": dd(pi("Connection agreement")),
               "ind": dd(pi("Connection agreement")), "land": wdate["Land"], "dev": wdate["Devex"]}
        return dict(k=k, pi=pi, mva=mva, wind=wind, mwh=mwh, amt=amt, pay=pay, amt_a=amt_a, date_a=date_a)

    def rtb(col, which):
        x = inputs(col)
        k, amt, pay = x["k"], dict(x["amt"]), dict(x["pay"])
        if which == "A":                                 # only the devex items paid by sale A
            amt.update(x["amt_a"])
            pay.update(x["date_a"])
        jv, share, prem = k("JV with a developer"), k("FEP share of distributable"), k("RtB premium")
        sale = sale_a if which == "A" else sale_b
        end = min(freeze, sale)
        lines = ["fee", "bs", "pb", "tso", "land", "dev"] + ([] if which == "A" else ["dir", "ind"])
        claims = {}
        for n in lines:
            acc = amt[n] * (factor(pay[n], rd if n in ("land", "dev") else rc, end) - 1)
            claims[n] = jv * (max(acc, amt[n] * (floor - 1)) if n in ("land", "dev") else acc)
        reimb = sum(amt[n] for n in ["bs", "pb", "tso"] + ([] if which == "A" else ["dir", "ind"]))
        dsa = k("Developer success fee (DSA) on") * min(dsa_rate * yf360(dd(x["pi"]("DSA accrual start")), sale),
                                                         dsa_cap) * x["wind"]
        cash = prem * x["mva"] + reimb
        paid = {}
        for n, due in (("sc", sc), ("rp", reimb), ("c12", sum(v for n2, v in claims.items()
                                                             if n2 not in ("land", "dev"))),
                       ("dev", claims["land"] + claims["dev"]), ("dsa", dsa)):
            paid[n] = min(cash, due)
            cash -= paid[n]
        receipts = paid["rp"] + paid["c12"] + paid["dev"] + cash * share
        outl = [(pay["sdd"], -amt["sdd"])] + [(pay[n], -amt[n]) for n in lines]
        flows = outl + [(sale, receipts)]
        return {"cash out": -sum(v for _, v in outl), "proceeds": receipts, "profit": sum(v for _, v in flows),
                "MOIC": receipts / -sum(v for _, v in outl), "IRR": xirr(sorted(flows))}

    ok = True
    for col, nm in zip(P, NAMES):
        for which, head in (("A", "RTB exit, Case A"), ("B", "RTB exit, Case B")):
            x = rtb(col, which)
            h = find(head)
            model = {"cash out": dash[f"{col}{find('Total FEP cash out', h)}"].value,
                     "proceeds": dash[f"{col}{find('Total proceeds', h)}"].value,
                     "profit": dash[f"{col}{find('FEP profit, pre-tax', h)}"].value,
                     "MOIC": dash[f"{col}{find('MOIC', h)}"].value, "IRR": dash[f"{col}{find('IRR', h)}"].value}
            diff = max(abs(model[t] - x[t]) for t in x)
            ok &= diff < 1e-6
            print(f"   {nm:<11} Case {which}: " + " | ".join(f"{t} {model[t]:,.4f}" for t in model)
                  + f"  (python diff {diff:.1e})")
    # Iváncsa COD exit: cost stack from the inputs, flows by year, price from the timeline EBITDA (model values).
    col = "F"
    x = inputs(col)
    k, amt, pay = x["k"], x["amt"], x["pay"]
    built = x["wind"] * (1 - k("Wind built below"))
    route = k("Turbine route")
    case = dash[f"H{find('Case selected')}"].value
    cc = {"Base": "E", "Low": "F", "High": "G"}[case]
    cs = lambda t: dash[f"{cc}{find(t)}"].value  # noqa: E731
    turb = cs("Turbine supply incl. transport, EU") if route == "EU" else cs("Turbine supply incl. transport, CN")
    per_mw = (turb + sum(iv(t) for t in ("Turbine installation", "Foundations", "Balance of plant – civil",
                                        "Balance of plant – electrical", "Owner's costs", "Land purchase and"))) * built
    capex = (per_mw + iv("On-site substation") * x["mva"] / iv("Source basis")) * (1 + iv("Capex contingency")) \
        + iv("BESS capex") * x["mwh"]
    conn = amt["tso"] + amt["dir"] + amt["ind"]
    muni = k("Municipal one-off") * x["mva"]
    dev = amt["dev"] + amt["fee"] + amt["sdd"]
    forgone = k("RtB premium") * x["mva"]
    stack = capex + conn + muni + amt["land"] + dev + forgone
    blk = find("Iváncsa: annual cash flow")
    cols = range(10, 10 + len(TL))
    years = [dash.cell(find("Year"), c).value for c in cols]
    ye = [dd(dash.cell(find("Year end"), c).value) for c in cols]
    eb = [dash.cell(find("EBITDA", blk), c).value for c in cols]
    ntp, cod, rel = dd(iv("Notice to proceed", col)), dd(iv("COD", col)), dd(iv("Performance security released", col))
    nc = (cod.year - ntp.year) * 12 + cod.month - ntp.month
    disc = iv("Discount rate")
    price = sum(e * (1 + disc) ** ((cod - t).days / 365) for e, t in zip(eb, ye))

    def cons(y):                     # construction months in year y (month-ends after NTP up to COD)
        mths = [(ntp.year * 12 + ntp.month + i) for i in range(1, nc + 1)]
        return sum(1 for mm in mths if (mm - 1) // 12 == y)
    flows = []
    for y, t in zip(years, ye):
        inv = sum(-v for n, v in amt.items() if pay[n].year == y)
        inv += (amt["bs"] + amt["pb"]) * (y == rel.year) - capex * cons(y) / nc - muni * (y == cod.year)
        fl = inv * (y <= cod.year) + (price + amt["bs"] + amt["pb"]) * (y == cod.year) - forgone * (y == sale_b.year)
        flows.append((t, fl))
    h = find("COD exit")
    py = {"stack": stack, "price": price, "profit": sum(v for _, v in flows),
          "MOIC": sum(v for _, v in flows if v > 0) / -sum(v for _, v in flows if v < 0), "IRR": xirr(flows)}
    model = {"stack": dash[f"{col}{find('Total costs to COD', h)}"].value,
             "price": dash[f"{col}{find('Sale price at COD', h)}"].value,
             "profit": dash[f"{col}{find('FEP profit, pre-tax', h)}"].value,
             "MOIC": dash[f"{col}{find('MOIC', h)}"].value, "IRR": dash[f"{col}{find('IRR', h)}"].value}
    diff = max(abs(model[t] - py[t]) for t in py)
    ok &= diff < 1e-6
    print("   Iváncsa COD exit: " + " | ".join(f"{t} {model[t]:,.4f}" for t in model) + f"  (python diff {diff:.1e})")
    xa = inputs("E")
    print(f"   Case A devex per project: land {xa['amt_a']['land']:,.1f} ({xa['date_a']['land']}), development "
          f"{xa['amt_a']['dev']:,.1f} ({xa['date_a']['dev']})")
    print(f"   devex per project (scale 0.25): land {tot['Land'] * 0.25 / fx / 1000:,.1f} ({wdate['Land']}), "
          f"development {tot['Devex'] * 0.25 / fx / 1000:,.1f} ({wdate['Devex']})")
    print("independent recomputation:", "PASS" if ok else "FAIL")
    return ok


if __name__ == "__main__":
    build()
    sys.exit(0 if verify() else 1)
