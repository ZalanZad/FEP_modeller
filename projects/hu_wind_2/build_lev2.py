"""No JV → no shareholder loans; Case A exits before HCSSZ; §24 re-applied (SPEC bora_v5_summary §25).

Base: reference/HU_Wind_bids_271026_flagged.xlsx (the user's latest copy; every user edit is kept). Edited in
LibreOffice through UNO like build_lev.py (kept as history): rows inserted with Rows.insertByIndex, new rows formatted
from a neighbouring row, existing formulas only extended where the SPEC asks.

  1. Key input "JV with a developer (1 = yes; claims and DSA apply)" (Bora 1, others 0); every claim line × the flag;
     claim rows relabelled "… on FEP's shareholder loan (JV only)".
  2. Inputs "Sale date, Case A (before HCSSZ)" 30-Sep-2028; Case A gets its own claims block (accrual end
     MIN(freeze, sale A), same row pattern as the Case B block), DSA to sale A, receipts dated at sale A; detail labels
     renamed so that Case A / Case B match the summary; check "Sale date A before HCSSZ"; the user's two comments
     (G195 "not correct", J232 "different exit date") removed – both are fixed.
  3. §24: operating CIT 15 % (20-year tax depreciation, no loss carry-forward), 60 % levered case (5 %, 15-year annuity),
     DSCR, debt balance, debt-repaid check, hold summary rows, hold heading.
Then verify(): recalc, no errors, Master_Check 0; before / after vs the user's file; independent Python recomputation of
Iváncsa Case B and Bora Case A waterfalls and of the hold post-tax / equity IRRs; out/dump_bids.txt.

    python3 projects/hu_wind_2/build_lev2.py
"""
import json
import sys
from datetime import date
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[1]))
from build_lev import LAST, SUMMARY, TIMELINE, TL, add_name, col, dump, letter, xirr, xnpv  # noqa: E402
from build_lev import Sheet as _Sheet  # noqa: E402
from fepm.audit.lint import lint, lint_values  # noqa: E402
from fepm.recalc import Office  # noqa: E402

SRC = HERE.parents[1] / "reference" / "HU_Wind_bids_271026_flagged.xlsx"
OUT = HERE / "out" / "HU_Wind_bids.xlsx"
RC = HERE / "out" / "HU_Wind_bids_recalc.xlsx"
BASE_RC = HERE / "out" / "HU_Wind_bids_271026_flagged_recalc.xlsx"
INSERTS = HERE / "out" / "rows_lev2.json"                       # row inserts (old position, count) per sheet
P = ["E", "F", "G"]
NAMES = ["Bora", "Iváncsa", "Dunaújváros"]
JV = [1, 0, 0]
SALE_A = date(2028, 9, 30)
NEW_INPUTS = [("Dep_Years", "Tax depreciation period (straight line from the COD year)", "years", 20,
               "Bora model Inputs G215"),
              ("Gearing", "Gearing (debt ÷ all-in cost to COD)", "%", 0.60, "Bora model Inputs G219"),
              ("Debt_Rate", "Debt interest rate", "%", 0.05, "Bora model Inputs G220"),
              ("Debt_Tenor", "Debt tenor (annuity from the first full year)", "years", 15, "Bora model Inputs G222")]
JV_TAG = " on FEP's shareholder loan (JV only)"
CASE_A_EVENTS = [("bid", "bid submission", "Rate_Cash", 27), ("pb", "performance security posting", "Rate_Cash", 28),
                 ("tso", "TSO security payment", "Rate_Cash", 29), ("dev", "devex payment (average)", "Rate_Devex", 31)]
CASE_A_CLAIMS = [("fee", "participation fee (not reimbursed)", "fee", "bid"), ("bs", "bid security (cash)", "bs", "bid"),
                 ("pb", "performance security payable (cash)", "pb", "pb"),
                 ("tso", "TSO connection-request security", "tso", "tso"),
                 ("dev", "devex, MAX(20 % accrued, 1× principal)", "dev", "dev")]
N_CASE_A = 3 + 2 * len(CASE_A_EVENTS) + len(CASE_A_CLAIMS) + 1      # end, 2 stubs, days+factor, claims, DSA


class Sheet(_Sheet):
    def formula(self, ref, f):
        """Excel-style formula → LibreOffice API grammar: ';' between arguments, `$Inputs.E27` for `Inputs!E27`."""
        self.cell(ref).setFormula(f.replace("Inputs!", "$Inputs.").replace(",", ";"))


class Rows:
    """Old row → new row after inserts (old positions; an insert at p pushes rows ≥ p down)."""
    def __init__(self):
        self.ins = []

    def add(self, pos, n):
        self.ins.append((pos, n))

    def __call__(self, r):
        return r + sum(n for p, n in self.ins if p <= r)

    def block(self, pos):
        n = dict(self.ins)[pos]
        return list(range(self(pos) - n, self(pos)))


def edit_inputs(doc):
    inp = Sheet(doc, "Inputs")
    m = Rows()
    r_sale = inp.find("Sale date (RtB exit)")
    r_cit = inp.find("Corporate income tax on exit gains", exact=False)
    r_ex = inp.find("Share deal, participation exemption", exact=False)
    r_life = inp.find("Operating life from COD")
    inp.number(f"E{r_cit}", 0.15)
    inp.text(f"C{r_cit}", "Corporate income tax on exit gains and on operating profit")
    inp.text(f"H{r_cit}", "SPEC §24.1 / §25.3 (15 %)")
    inp.text(f"C{r_sale}", "Sale date, Case B (RtB exit after HCSSZ)")
    inp.insert(r_ex + 1, len(NEW_INPUTS))
    m.add(r_ex + 1, len(NEW_INPUTS))
    inp.insert(r_sale + 1, 1)
    m.add(r_sale + 1, 1)
    idx = [doc.Sheets.getByIndex(i).Name for i in range(doc.Sheets.Count)].index("Inputs")
    r = m.block(r_sale + 1)[0]
    inp.copy_format(m(r_sale), r, last="H")
    inp.text(f"C{r}", "Sale date, Case A (before HCSSZ)")
    inp.text(f"D{r}", "date")
    inp.number(f"E{r}", (SALE_A - date(1899, 12, 30)).days)
    inp.text(f"H{r}", "SPEC §25.2 (user, 28-Oct-2026)")
    add_name(doc, "Sale_Date_A", idx, col("E"), r)
    for (name, lab, unit, v, note), r in zip(NEW_INPUTS, m.block(r_ex + 1)):
        inp.copy_format(m(r_life) if unit == "years" else m(r_cit), r, last="H")
        inp.text(f"C{r}", lab)
        inp.text(f"D{r}", unit)
        inp.number(f"E{r}", v)
        inp.text(f"H{r}", note)
        add_name(doc, name, idx, col("E"), r)
    return m


def edit_dashboard(doc):
    d = Sheet(doc, "Dashboard")
    m = Rows()
    f = lambda t, s=1: d.find(t, start=s, exact=False)  # noqa: E731
    # Anchors in the user's layout.
    a = {"dsa_key": d.find("Developer success fee (DSA) on"), "wind": d.find("Wind committed in the bid"),
         "allin": f("All-in cost to COD"), "hold": f("Hold to maturity"), "npv": f("NPV at the discount rate"),
         "eur": f("Guarantees, securities and connection fees, DEVEX (EUR k)")}
    for k, lab in (("fee", "Participation fee"), ("bs", "Bid security"), ("pb", "Performance security payable"),
                   ("tso", "TSO connection-request security"), ("dev", "DEVEX")):
        a["e_" + k] = f(lab, a["eur"])
    cb = f("RtB exit: FEP claims at the sale")                 # Case B detail (the summary's Case B)
    a.update(cb_head=cb, cb_end=f("Interest accrual end", cb), cb_s12=f("Accrual factor in the end year, cash", cb),
             cb_days=f("Days to the first 31-Dec from bid", cb), cb_fac=f("Accrual factor from bid", cb),
             cb_claims=[f("Claim: ", cb) + i for i in range(7)], cb_tot=d.find("Total FEP claims", cb),
             cb_dsa=f("Developer success fee due", cb), cb_wf=f("RtB exit: waterfall at the sale", cb))
    a.update(cb_c12=f("Less FEP claims at 12 %", a["cb_wf"]), cb_cdev=f("Less FEP devex return", a["cb_wf"]),
             cb_flows=f("FEP cash flows, dated (outlays", a["cb_wf"]))
    ca = f("RtB exit, Case B: buyer pays the grid connection fees")   # the summary's Case A
    a.update(ca_head=ca, ca_out=f("FEP cash out to RtB, Case B", ca), ca_tot=f("Total FEP claims, Case B", ca),
             ca_wf=f("Waterfall at the sale, Case B", ca))
    a.update(ca_c12=f("Less FEP claims at 12 %", a["ca_wf"]), ca_cdev=f("Less FEP devex return", a["ca_wf"]),
             ca_dsa=f("Less developer success fee (DSA)", a["ca_wf"]), ca_flows=f("FEP cash flows, dated, Case B", ca))
    a.update(ca_rec=f("FEP receipts at the sale", a["ca_flows"]), ca_memo=f("Memo: pre-tax profit vs Case A", ca))
    blocks = [(nm, f(f"{nm}: annual cash flow")) for nm in NAMES]
    blocks = [(nm, h, d.find("Unlevered pre-tax cash flow", h)) for nm, h in blocks]
    a.update(chk_b=f("RtB waterfall closes"), chk_a=f("Case B waterfall closes"))

    # Inserts, bottom-up.
    n_tl = len(TIMELINE)
    d.insert(a["chk_a"], 2)
    m.add(a["chk_a"], 2)
    for nm, h, pre in reversed(blocks):
        d.insert(pre + 1, n_tl)
        m.add(pre + 1, n_tl)
    d.insert(a["ca_tot"], N_CASE_A)
    m.add(a["ca_tot"], N_CASE_A)
    d.insert(a["npv"] + 1, len(SUMMARY))
    m.add(a["npv"] + 1, len(SUMMARY))
    d.insert(a["dsa_key"] + 1, 1)
    m.add(a["dsa_key"] + 1, 1)
    R = {k: (m(v) if isinstance(v, int) else [m(x) for x in v]) for k, v in a.items()}
    jv = m.block(a["dsa_key"] + 1)[0]
    R["jv"] = jv

    # 1. JV flag; claims × flag; relabels.
    d.copy_format(R["dsa_key"], jv, last="H")
    d.text(f"C{jv}", "JV with a developer (1 = yes; claims and DSA apply)")
    d.text(f"D{jv}", "flag")
    for c, v in zip(P, JV):
        d.number(f"{c}{jv}", v)
    add_name(doc, "Dash_JV", 0, col("E"), jv)
    doc.NamedRanges.getByName("Dash_JV").setContent(f"$Dashboard.$E${jv}:$G${jv}")
    for r in R["cb_claims"]:
        for c in P:
            old = d.cell(f"{c}{r}").getFormula()
            d.cell(f"{c}{r}").setFormula(f"={c}${jv}*" + old[1:])
        d.text(f"C{r}", d.cell(f"C{r}").getString().replace("Claim: ", "Claim" + JV_TAG + ": "))
    d.text(f"C{R['cb_tot']}", "Total FEP claims on its shareholder loans (JV only)")
    for r in (R["cb_c12"], R["ca_c12"]):
        d.text(f"C{r}", "Less FEP claims at 12 %" + JV_TAG)
    for r in (R["cb_cdev"], R["ca_cdev"]):
        d.text(f"C{r}", "Less FEP devex return" + JV_TAG)

    # 2. Case A: labels, own claims block, DSA and receipts at sale A.
    relabel = {"cb_head": ("RtB exit: ", "RtB exit, Case B: "), "cb_wf": ("RtB exit: ", "RtB exit, Case B: "),
               "cb_flows": ("FEP cash flows, dated", "FEP cash flows, dated, Case B"),
               "ca_head": ("Case B", "Case A"), "ca_out": ("Case B", "Case A"), "ca_wf": ("Case B", "Case A"),
               "ca_flows": ("Case B (no connection-fee outlays)", "Case A (no connection-fee outlays; sale A)"),
               "ca_memo": ("vs Case A", "vs Case B"), "chk_b": ("RtB waterfall closes", "Case B waterfall closes"),
               "chk_a": ("Case B", "Case A")}
    for k, (x, y) in relabel.items():
        d.text(f"C{R[k]}", d.cell(f"C{R[k]}").getString().replace(x, y, 1))
    rows = m.block(a["ca_tot"])
    q = {}
    it = iter(rows)
    q["end"], q["s12"], q["s20"] = next(it), next(it), next(it)
    for ev, *_ in CASE_A_EVENTS:
        q["d_" + ev], q["f_" + ev] = next(it), next(it)
    for key, *_ in CASE_A_CLAIMS:
        q["c_" + key] = next(it)
    q["dsa"] = next(it)
    put = lambda r, tmpl, lab, unit, fml, total=False: _row(d, r, tmpl, lab, unit, fml, total)  # noqa: E731
    put(q["end"], R["cb_end"], "Case A: interest accrual end (earlier of freeze and sale A)", "date",
        "=MIN(Freeze_Date,Sale_Date_A)")
    for key, rate, text in (("s12", "Rate_Cash", "cash"), ("s20", "Rate_Devex", "devex")):
        put(q[key], R["cb_s12"], f"Case A: accrual factor in the end year, {text} rate", "x",
            "=1+%s*({c}%d-DATE(YEAR({c}%d),1,1)+1)/Days_In_Year" % (rate, q["end"], q["end"]))
    for ev, text, rate, irow in CASE_A_EVENTS:
        x = "Inputs!{c}$%d" % irow
        put(q["d_" + ev], R["cb_days"], f"Case A: days to the first 31-Dec from {text}", "days",
            "=DATE(YEAR(%s)+1,1,1)-%s" % (x, x))
        put(q["f_" + ev], R["cb_fac"], f"Case A: accrual factor from {text}", "x",
            "=(1+%s*{c}%d/Days_In_Year)*(1+%s)^(YEAR({c}%d)-YEAR(%s)-1)*{c}%d"
            % (rate, q["d_" + ev], rate, q["end"], x, q["s20" if ev == "dev" else "s12"]))
    for key, text, line, ev in CASE_A_CLAIMS:
        pr = R["e_" + line]
        fml = ("={c}$%d*{c}%d*MAX({c}%d-1,Devex_Floor-1)" if key == "dev" else "={c}$%d*{c}%d*({c}%d-1)") % (
            jv, pr, q["f_" + ev])
        put(q["c_" + key], R["cb_claims"][0], f"Case A: claim{JV_TAG}: {text}", "EUR k", fml, total=True)
    put(q["dsa"], R["cb_dsa"], "Case A: developer success fee due (DSA to sale A, capped)", "EUR k",
        "={c}%d*MIN(DSA_Rate*YEARFRAC(Inputs!{c}$35,Sale_Date_A),DSA_Cap)*{c}%d" % (R["dsa_key"], R["wind"]),
        total=True)
    r = R["ca_tot"]
    d.text(f"C{r}", "Total FEP claims on its shareholder loans, Case A (JV only; no connection-fee claims)")
    for c in P:
        d.formula(f"{c}{r}", f"=SUM({c}{q['c_fee']}:{c}{q['c_dev']})")
    wf = R["ca_wf"]
    for c in P:
        d.formula(f"{c}{R['ca_c12']}", f"=-MIN(SUM({c}{wf + 3}:{c}{wf + 5}),SUM({c}{q['c_fee']}:{c}{q['c_tso']}))")
        d.formula(f"{c}{R['ca_cdev']}", f"=-MIN(SUM({c}{wf + 3}:{c}{wf + 6}),{c}{q['c_dev']})")
        d.formula(f"{c}{R['ca_dsa']}", f"=-MIN(SUM({c}{wf + 3}:{c}{wf + 7}),{c}{q['dsa']})")
    for dc in "JKL":
        d.formula(f"{dc}{R['ca_rec']}", "=Sale_Date_A")
    # Remove the user's two comments (both points are fixed).
    ann = d.s.getAnnotations()
    for i in reversed(range(ann.getCount())):
        pos = ann.getByIndex(i).getPosition()
        ref = f"{letter(pos.Column)}{pos.Row + 1}"
        if ref in (f"G{R['cb_c12']}", f"J{R['ca_flows']}"):
            ann.removeByIndex(i)

    # 3. §24 on this base: timeline rows, hold summary, heading.
    ye = d.find("Year end")
    tmpl = {"money": d.find("HIPA local business tax"), "bold": d.find("Unlevered pre-tax cash flow"),
            "ratio": d.find("PPA price"), "pct_bold": f("Project IRR, pre-tax", R["hold"]), "sum_money": R["npv"],
            "sum_ratio": f("MOIC, pre-tax", R["hold"] - 10)}
    tl = {}
    for c, (nm, h, pre) in zip(P, blocks):
        h, pre = m(h), m(pre)
        ref = {"allin": R["allin"], "pre": pre, "tn": d.find("Years from the COD year", h),
               "cod": d.find("COD year flag", h), "eb": d.find("EBITDA", h)}
        for k, (key, *_rest) in enumerate(TIMELINE):
            ref[key] = pre + 1 + k
        for key, lab, unit, kind, fml, total, const in TIMELINE:
            r = ref[key]
            d.copy_format(tmpl[kind], r)
            d.text(f"C{r}", lab)
            d.text(f"D{r}", unit)
            for x in TL:
                d.formula(f"{x}{r}", fml.format(x=x, w=letter(col(x) - 1), p=c, **ref))
            if total:
                d.formula(f"H{r}", f"=SUM(J{r}:{LAST}{r})")
            if const:
                d.formula(f"{c}{r}", "=Gearing*${p}${allin}*Debt_Rate/(1-(1+Debt_Rate)^-Debt_Tenor)".format(p=c, **ref))
                d.fmt(f"{c}{r}", f"J{r}")
            if kind == "ratio":
                d.fmt(f"J{r}:{LAST}{r}", f"E{tmpl['sum_ratio']}")
        tl[c] = ref
    for k, (key, lab, unit, kind) in enumerate(SUMMARY):
        r = R["npv"] + 1 + k
        d.copy_format({"money": tmpl["sum_money"], "pct_bold": tmpl["pct_bold"], "ratio": tmpl["sum_ratio"]}[kind], r,
                      last="H")
        d.text(f"C{r}", lab)
        d.text(f"D{r}", unit)
        for c in P:
            t = tl[c]
            rng = lambda k2: f"$J${t[k2]}:${LAST}${t[k2]}"  # noqa: E731
            d.formula(f"{c}{r}", {"s_cit": f"=$H${t['cit']}", "s_debt": f"=$H${t['draw']}",
                                  "s_irr": f"=XIRR({rng('post')},$J${ye}:${LAST}${ye})",
                                  "s_eirr": f"=XIRR({rng('eq')},$J${ye}:${LAST}${ye})",
                                  "s_dscr": f"=MIN({rng('dscr')})"}[key])
        if key in ("s_cit", "s_debt"):
            d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
    d.text(f"C{R['hold']}", "Hold to maturity (EBITDA; CIT 15 %; 60 % levered case)")
    # Checks, inside the Master_Check range.
    r1, r2 = m.block(a["chk_a"])
    for r, lab, fml in ((r1, "Sale date A before the HCSSZ date (Case A exits before the connection fees)",
                         "=N(Sale_Date_A>=Inputs!{c}$30)"),
                        (r2, "Debt fully repaid at the end of the tenor (balance 0)", None)):
        d.copy_format(R["chk_b"], r, last="H")
        d.text(f"C{r}", lab)
        d.text(f"D{r}", "flag")
        for c in P:
            d.formula(f"{c}{r}", fml.format(c=c) if fml else f"=N(ABS(${LAST}${tl[c]['bal']})>Tolerance)")
        d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
    return m


def _row(d, r, tmpl, lab, unit, fml, total):
    d.copy_format(tmpl, r, last="H")
    d.text(f"C{r}", lab)
    d.text(f"D{r}", unit)
    for c in P:
        d.formula(f"{c}{r}", fml.format(c=c))
    if total:
        d.formula(f"H{r}", f"=SUM(E{r}:G{r})")


def build():
    with Office(SRC) as lo:
        lo.save(BASE_RC)
        mi = edit_inputs(lo.doc)
        m = edit_dashboard(lo.doc)
        lo.calculate()
        lo.save(OUT)
    INSERTS.write_text(json.dumps({"Dashboard": m.ins, "Inputs": mi.ins}))
    with Office(OUT) as lo:
        lo.save(RC)
    print(f"wrote {OUT}")
    return m


# ---------------------------------------------------------------------------------------------------------------------
def compare(b_wb, a_wb):
    diffs = []
    for name in ("Dashboard", "Inputs"):
        b, a = b_wb[name], a_wb[name]
        m = Rows()
        for pos, n in json.loads(INSERTS.read_text())[name]:
            m.add(pos, n)
        for r in range(1, b.max_row + 1):
            for c in range(2, min(b.max_column, 46) + 1):
                x, y = b.cell(r, c).value, a.cell(m(r), c).value
                same = (x == y) or (isinstance(x, (int, float)) and isinstance(y, (int, float))
                                    and abs(x - y) <= 1e-9 * max(1, abs(x)))
                if not same:
                    diffs.append((name, r, m(r), b.cell(r, 3).value, b.cell(r, c).column_letter, x, y))
    return diffs


def verify():
    a_wb = openpyxl.load_workbook(RC, data_only=True)
    b_wb = openpyxl.load_workbook(BASE_RC, data_only=True)
    errs = lint_values(RC)
    dash, inp = a_wb["Dashboard"], a_wb["Inputs"]
    print("recalc:", "no errors, Master_Check 0" if not errs else [str(e) for e in errs[:10]])
    found = [str(f) for f in lint(OUT)]
    print("lint:", "clean" if not found else f"{len(found)} finding(s)")
    for f_ in found[:15]:
        print("  ", f_)
    print("comments left:", [(ws.title, c.coordinate) for ws in openpyxl.load_workbook(OUT).worksheets
                             for row in ws.iter_rows() for c in row if c.comment])
    diffs = compare(b_wb, a_wb)
    print(f"before / after: {len(diffs)} changed cell(s) in the user's rows")
    for sh, r0, r1, lab, cl, x, y in diffs:
        fx_ = lambda v: f"{v:,.2f}" if isinstance(v, (int, float)) else str(v)  # noqa: E731
        print(f"   {sh}!{cl}{r0}→{r1}  {str(lab)[:60]:<60} {fx_(x):>14} → {fx_(y)}")
    ok = not errs and recompute(dash, inp)
    dump(a_wb)
    print("dump: out/dump_bids.txt")
    return ok


def recompute(dash, inp):
    """Iváncsa Case B (no JV: no claims) and Bora Case A (sale 30-Sep-2028) waterfalls; hold post-tax / equity IRRs."""
    lab = lambda r, ws=dash: ws.cell(r, 3).value or ""  # noqa: E731
    find = lambda t, s=1: next(r for r in range(s, dash.max_row + 1) if lab(r).startswith(t))  # noqa: E731
    iv = lambda t: inp.cell(next(r for r in range(1, inp.max_row + 1) if lab(r, inp).startswith(t)), 5).value  # noqa
    val = lambda r, col: dash[f"{col}{r}"].value  # noqa: E731
    d = lambda x: x.date() if hasattr(x, "date") else x  # noqa: E731
    eur = find("Guarantees, securities and connection fees, DEVEX (EUR k)")
    line = {k: find(t, eur) for k, t in (("fee", "Participation fee"), ("bs", "Bid security"),
                                         ("pb", "Performance security payable"), ("tso", "TSO connection-request"),
                                         ("dir", "Direct connection fee"), ("ind", "Indirect connection fee"),
                                         ("dev", "DEVEX"), ("sdd", "Signing fee and DD"))}
    dates = {k: d(inp.cell(r, 5).value) for k, r in (("sign", 26), ("bid", 27), ("pb", 28), ("tso", 29),
                                                      ("hc", 30), ("dev", 31), ("dsa", 35))}
    rc, rd, floor = iv("Claim rate, cash"), iv("Claim rate, devex"), iv("Devex minimum return")
    freeze, sale_b, sale_a = d(iv("Interest freeze")), d(iv("Sale date, Case B")), d(iv("Sale date, Case A"))
    cit, exempt, sc = iv("Corporate income tax"), iv("Share deal"), iv("Sale costs")
    dsa_rate, dsa_cap = iv("Developer success fee (on"), iv("Developer success fee cap")

    def factor(p, r, end):
        bal, y, start = 1.0, p.year, p
        while True:
            stop = min(date(y, 12, 31), end)
            full = start == date(y, 1, 1) and stop == date(y, 12, 31)
            bal *= 1 + r * (1 if full else ((stop - start).days + 1) / 365)
            if stop == end:
                return bal
            y, start = y + 1, date(y + 1, 1, 1)

    def yf360(a, b):
        return ((b.year - a.year) * 360 + (b.month - a.month) * 30 + min(b.day, 30) - min(a.day, 30)) / 360

    def case(col, which):
        k = lambda t: val(find(t), col)  # noqa: E731
        amt = {n: val(r, col) for n, r in line.items()}
        jv, share, mva, prem = k("JV with a developer"), k("FEP share of distributable"), k("Net connection"), \
            k("RtB premium")
        dsa_on, wind = k("Developer success fee (DSA) on"), k("Wind committed")
        sale = sale_a if which == "A" else sale_b
        end = min(freeze, sale)
        lines = ["fee", "bs", "pb", "tso", "dev"] + ([] if which == "A" else ["dir", "ind"])
        pay = {"fee": dates["bid"], "bs": dates["bid"], "pb": dates["pb"], "tso": dates["tso"], "dev": dates["dev"],
               "dir": dates["hc"], "ind": dates["hc"]}
        claims = {}
        for n in lines:
            acc = amt[n] * (factor(pay[n], rd if n == "dev" else rc, end) - 1)
            claims[n] = jv * (max(acc, amt[n] * (floor - 1)) if n == "dev" else acc)
        reimb = sum(amt[n] for n in (["bs", "pb", "tso"] + ([] if which == "A" else ["dir", "ind"])))
        dsa = dsa_on * min(dsa_rate * yf360(dates["dsa"], sale), dsa_cap) * wind
        cash = prem * mva + reimb
        paid = {}
        for n, due in (("sc", sc), ("rp", reimb), ("c12", sum(v for x, v in claims.items() if x != "dev")),
                       ("dev", claims["dev"]), ("dsa", dsa)):
            paid[n] = min(cash, due)
            cash -= paid[n]
        receipts = paid["rp"] + paid["c12"] + paid["dev"] + cash * share
        outl = [(dates["sign"], -amt["sdd"])] + [(pay[n], -amt[n]) for n in lines]
        flows = outl + [(sale, receipts)]
        profit = sum(x for _, x in flows)
        tax = (max(0, paid["c12"] + paid["dev"]) if exempt == 1 else max(0, profit)) * cit
        return {"cash out": -sum(x for _, x in outl), "proceeds": receipts, "profit": profit,
                "post-tax": profit - tax, "IRR": xirr(sorted(flows))}

    ok = True
    for col, nm, which, head in (("F", "Iváncsa", "B", "RTB exit, Case B"), ("E", "Bora", "A", "RTB exit, Case A")):
        x = case(col, which)
        h = find(head)
        model = {"cash out": val(find("Total FEP cash out", h), col), "proceeds": val(find("Total proceeds", h), col),
                 "profit": val(find("FEP profit, pre-tax", h), col), "post-tax": val(find("FEP profit, post-tax", h),
                                                                                      col),
                 "IRR": val(find("IRR, pre-tax", h), col)}
        worst = max(abs(model[k] - x[k]) for k in x)
        ok &= worst < 1e-6
        print(f"   {nm} Case {which}: " + " | ".join(f"{k} {model[k]:,.4f}" for k in model)
              + f"  (python max diff {worst:.1e})")
    # Hold: post-tax project IRR and equity IRR, all three projects, from the timeline rows.
    dep_y, gear, rate, ten = iv("Tax depreciation period"), iv("Gearing"), iv("Debt interest rate"), iv("Debt tenor")
    ye_r = find("Year end")
    allin_r = find("All-in cost to COD")
    ye = [d(dash.cell(ye_r, c).value) for c in range(10, 10 + len(TL))]
    for col, nm in zip(P, NAMES):
        h = find(f"{nm}: annual cash flow")
        row = lambda t: [dash.cell(find(t, h), c).value for c in range(10, 10 + len(TL))]  # noqa: E731
        tn, eb, pre = row("Years from the COD year"), row("EBITDA"), row("Unlevered pre-tax cash flow")
        allin = val(allin_r, col)
        tax = [-cit * max(0, e - allin / dep_y * (0 <= n < dep_y)) for e, n in zip(eb, tn)]
        post = [p + t for p, t in zip(pre, tax)]
        debt = gear * allin
        ann = debt * rate / (1 - (1 + rate) ** -ten)
        eq = [p + debt * (n == 0) - ann * (1 <= n <= ten) for p, n in zip(post, tn)]
        px = {"Project IRR, post-tax": xirr(list(zip(ye, post))),
              "Equity IRR, post-tax (60 % gearing)": xirr(list(zip(ye, eq))),
              "CIT over the life (15 %, no loss carry-forward)": sum(tax)}
        diff = max(abs(val(find(k), col) - v) for k, v in px.items())
        ok &= diff < 1e-6
        print(f"   {nm} hold: " + " | ".join(f"{k.split(',')[0]} {val(find(k), col):,.4f}" for k in px)
              + f" | debt {val(find('Debt at COD'), col):,.1f} | min DSCR {val(find('Minimum DSCR'), col):.2f}"
              + f"  (python max diff {diff:.1e})")
    print("independent recomputation:", "PASS" if ok else "FAIL")
    return ok


if __name__ == "__main__":
    build()
    sys.exit(0 if verify() else 1)
