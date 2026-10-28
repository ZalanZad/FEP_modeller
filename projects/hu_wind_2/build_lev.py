"""15 % CIT and a 60 % levered hold case in the user's HU wind bids file (SPEC bora_v5_summary §24).

Base: reference/HU_Wind_bids_271026.xlsx (the user's reworked copy; every user edit is kept). The file is edited in
LibreOffice through UNO (fepm.recalc.Office): rows are inserted with Rows.insertByIndex, so every reference, name,
group and format shifts with the sheet; new rows take their format from a neighbouring row (copyRange) and then get
their own labels and formulas. No openpyxl writing.

  Inputs     CIT 15 % ("on exit gains and on operating profit"); new rows after it: tax depreciation period,
             gearing, debt interest rate, debt tenor (blue, named).
  Dashboard  per project timeline, after "Unlevered pre-tax cash flow": tax depreciation, CIT, unlevered post-tax cash
             flow, debt drawn at COD, debt service (annuity; value column: the annual amount), debt balance, equity cash
             flow (post-tax), DSCR; hold summary after "NPV at the discount rate": CIT over the life, project IRR
             post-tax, debt at COD, equity IRR post-tax, minimum DSCR; check "debt fully repaid" in the checks block.

Then verify(): recalc before / after, no errors, Master_Check 0, every existing Dashboard figure unchanged except the
CIT-driven ones (listed), independent Python recomputation of the Bora post-tax project IRR, equity IRR and minimum
DSCR, and out/dump_bids.txt.

    python3 projects/hu_wind_2/build_lev.py
"""
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
from fepm.audit.lint import lint, lint_values  # noqa: E402
from fepm.recalc import Office  # noqa: E402

SRC = ROOT / "reference" / "HU_Wind_bids_271026.xlsx"
OUT = HERE / "out" / "HU_Wind_bids.xlsx"
RC = HERE / "out" / "HU_Wind_bids_recalc.xlsx"
BASE_RC = HERE / "out" / "HU_Wind_bids_271026_recalc.xlsx"     # the user's file recalculated, for before / after
DUMP = HERE / "out" / "dump_bids.txt"

P = ["E", "F", "G"]
NAMES = ["Bora", "Iváncsa", "Dunaújváros"]
LAST = "AS"                                                     # the user's timeline ends in AS (2061)
BOLD = 150.0                                                    # com.sun.star.awt.FontWeight.BOLD
NORMAL = 100.0

# Inputs rows inserted after the CIT row: name, label, unit, value, format template row label, source note.
NEW_INPUTS = [("Dep_Years", "Tax depreciation period (straight line from the COD year)", "years", 20,
               "Operating life from COD", "Bora model Inputs G215"),
              ("Gearing", "Gearing (debt ÷ all-in cost to COD)", "%", 0.60, "Corporate income tax on exit gains",
               "Bora model Inputs G219"),
              ("Debt_Rate", "Debt interest rate", "%", 0.05, "Corporate income tax on exit gains",
               "Bora model Inputs G220"),
              ("Debt_Tenor", "Debt tenor (annuity from the first full year)", "years", 15, "Operating life from COD",
               "Bora model Inputs G222")]

# Timeline rows inserted after "Unlevered pre-tax cash flow" in each project block.
# key, label, unit, template ("money", "bold", "ratio"), timeline formula, total in H, value-column formula.
TIMELINE = [
    ("dep", "Tax depreciation (all-in cost ÷ 20 years from the COD year)", "EUR k", "money",
     "=${p}${allin}/Dep_Years*({x}{tn}>=0)*({x}{tn}<Dep_Years)", True, None),
    ("cit", "CIT (15 % × MAX(0, EBITDA − depreciation); no loss carry-forward)", "EUR k", "money",
     "=-CIT_Rate*MAX(0,{x}{eb}-{x}{dep})", True, None),
    ("post", "Unlevered post-tax cash flow", "EUR k", "bold", "={x}{pre}+{x}{cit}", True, None),
    ("draw", "Debt drawn at COD (gearing × all-in cost; no interest during construction)", "EUR k", "money",
     "=Gearing*${p}${allin}*{x}{cod}", True, None),
    ("ds", "Debt service (annuity: interest + principal; value: annual amount)", "EUR k", "money",
     "=-${p}{ds}*({x}{tn}>=1)*({x}{tn}<=Debt_Tenor)", True, "=Gearing*${p}${allin}*Debt_Rate/(1-(1+Debt_Rate)^-Debt_Tenor)"),
    ("bal", "Debt balance at year end", "EUR k", "money", "={w}{bal}*(1+Debt_Rate)+{x}{draw}+{x}{ds}", False, None),
    ("eq", "Equity cash flow (post-tax, 60 % gearing)", "EUR k", "bold", "={x}{post}+{x}{draw}+{x}{ds}", True, None),
    ("dscr", "DSCR ((EBITDA + CIT) ÷ debt service)", "x", "ratio", '=IF({x}{ds}<0,({x}{eb}+{x}{cit})/-{x}{ds},"")',
     False, None),
]
SUMMARY = [("s_cit", "CIT over the life (15 %, no loss carry-forward)", "EUR k", "money"),
           ("s_irr", "Project IRR, post-tax", "%", "pct_bold"),
           ("s_debt", "Debt at COD (60 %)", "EUR k", "money"),
           ("s_eirr", "Equity IRR, post-tax (60 % gearing)", "%", "pct_bold"),
           ("s_dscr", "Minimum DSCR", "x", "ratio")]


def col(letter):
    n = 0
    for ch in letter:
        n = n * 26 + ord(ch) - 64
    return n - 1                                                # 0-based


def letter(i):                                                  # 0-based → letters
    s = ""
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


TL = [letter(i) for i in range(col("J"), col(LAST) + 1)]


def api(formula):
    """Excel-style formula → LibreOffice API grammar (';' between arguments; no commas inside our strings)."""
    return formula.replace(",", ";")


class Sheet:
    def __init__(self, doc, name):
        self.doc, self.s = doc, doc.Sheets.getByName(name)
        self.name = name

    def cell(self, ref):
        return self.s.getCellRangeByName(ref)

    def find(self, text, start=1, end=2000, exact=True):
        for r in range(start, end):
            v = self.cell(f"C{r}").getString()
            if (v == text) if exact else v.startswith(text):
                return r
        raise KeyError(f"{self.name}: '{text}' not found from row {start}")

    def insert(self, before_row, n):
        self.s.Rows.insertByIndex(before_row - 1, n)

    def copy_format(self, src_row, dst_row, last=LAST):
        """Copy row `src_row` (format and content) onto `dst_row`, then clear the content (format stays)."""
        src = self.cell(f"B{src_row}:{last}{src_row}").getRangeAddress()
        dst = self.cell(f"B{dst_row}").getCellAddress()
        self.s.copyRange(dst, src)
        self.cell(f"B{dst_row}:{last}{dst_row}").clearContents(1 | 2 | 4 | 16)

    def text(self, ref, value):
        self.cell(ref).setString(value)

    def formula(self, ref, f):
        self.cell(ref).setFormula(api(f))

    def number(self, ref, v):
        self.cell(ref).setValue(v)

    def bold(self, ref, on=True):
        self.cell(ref).CharWeight = BOLD if on else NORMAL

    def fmt(self, ref, src):
        """Number format of `src` (a cell of the user's file) onto `ref`."""
        self.cell(ref).NumberFormat = self.cell(src).NumberFormat


def add_name(doc, name, sheet_index, ref_col, ref_row):
    from com.sun.star.table import CellAddress
    addr = CellAddress(Sheet=sheet_index, Column=ref_col, Row=ref_row - 1)
    doc.NamedRanges.addNewByName(name, f"$Inputs.${letter(ref_col)}${ref_row}", addr, 0)


def edit_inputs(doc):
    inp = Sheet(doc, "Inputs")
    r_cit = inp.find("Corporate income tax on exit gains")
    templates = {lab: inp.find(lab) for _, _, _, _, lab, _ in NEW_INPUTS}
    inp.number(f"E{r_cit}", 0.15)
    inp.text(f"C{r_cit}", "Corporate income tax on exit gains and on operating profit")
    inp.text(f"H{r_cit}", "SPEC §24.1 (15 %)")
    inp.insert(r_cit + 1, len(NEW_INPUTS))
    idx = [doc.Sheets.getByIndex(i).Name for i in range(doc.Sheets.Count)].index("Inputs")
    for k, (name, lab, unit, v, tmpl, note) in enumerate(NEW_INPUTS):
        r = r_cit + 1 + k
        t = templates[tmpl] + (len(NEW_INPUTS) if templates[tmpl] > r_cit else 0)
        inp.copy_format(t, r, last="H")
        inp.text(f"C{r}", lab)
        inp.text(f"D{r}", unit)
        inp.number(f"E{r}", v)
        inp.text(f"H{r}", note)
        add_name(doc, name, idx, col("E"), r)


def edit_dashboard(doc):
    d = Sheet(doc, "Dashboard")
    rows = {}
    # Locate the anchors in the user's layout (before any insert).
    allin = d.find("All-in cost to COD (securities excluded: returned)")
    npv = d.find("NPV at the discount rate, at the first year end")
    tmpl = {"money": d.find("HIPA local business tax"), "bold": d.find("Unlevered pre-tax cash flow"),
            "ratio": d.find("PPA price"), "pct_bold": d.find("Project IRR, pre-tax"), "sum_money": npv,
            "sum_ratio": d.find("MOIC, pre-tax", start=120)}
    d.text(f"C{d.find('Hold to maturity (EBITDA only, no debt, no tax)')}",
           "Hold to maturity (EBITDA; CIT 15 %; 60 % levered case)")              # coordinator, 28-Oct-2026
    blocks = []
    for nm in NAMES:
        head = d.find(f"{nm}: annual cash flow (EUR k)")
        blocks.append((nm, head, d.find("Unlevered pre-tax cash flow", start=head)))
    check_last = d.find("Case B waterfall closes (proceeds = all amounts paid out)")

    # Inserts bottom-up so the rows found above stay valid until their own insert.
    n_tl = len(TIMELINE)
    d.insert(check_last, 1)                                    # inside the Master_Check SUM range
    for nm, head, pre in reversed(blocks):
        d.insert(pre + 1, n_tl)
    d.insert(npv + 1, len(SUMMARY))

    # Row numbers after the inserts.
    shift = lambda r: r + (len(SUMMARY) if r > npv else 0) + n_tl * sum(1 for _, _, p in blocks if r > p)  # noqa: E731
    year_end = d.find("Year end", start=shift(npv))
    for key in tmpl:
        tmpl[key] = shift(tmpl[key])
    # Timeline rows per project.
    for c, (nm, head, pre) in zip(P, blocks):
        head, pre = shift(head), shift(pre)
        ref = {"allin": allin, "pre": pre, "tn": d.find("Years from the COD year", start=head),
               "cod": d.find("COD year flag", start=head), "eb": d.find("EBITDA", start=head)}
        for k, (key, *_rest) in enumerate(TIMELINE):
            ref[key] = pre + 1 + k
        for key, lab, unit, kind, f, total, const in TIMELINE:
            r = ref[key]
            d.copy_format(tmpl[kind], r)
            d.text(f"C{r}", lab)
            d.text(f"D{r}", unit)
            for i, x in enumerate(TL):
                w = letter(col(x) - 1)
                d.formula(f"{x}{r}", f.format(x=x, w=w, p=c, **ref))
            if total:
                d.formula(f"H{r}", f"=SUM(J{r}:{LAST}{r})")
            if const:
                d.formula(f"{c}{r}", const.format(p=c, **ref))
                d.fmt(f"{c}{r}", f"J{r}")
            if kind == "ratio":
                d.fmt(f"J{r}:{LAST}{r}", f"E{tmpl['sum_ratio']}")
        rows[c] = ref
    # Hold summary rows.
    for k, (key, lab, unit, kind) in enumerate(SUMMARY):
        r = npv + 1 + k
        d.copy_format({"money": tmpl["sum_money"], "pct_bold": tmpl["pct_bold"], "ratio": tmpl["sum_ratio"]}[kind], r,
                      last="H")
        d.text(f"C{r}", lab)
        d.text(f"D{r}", unit)
        for c in P:
            q = rows[c]
            rng = lambda key: f"$J${q[key]}:${LAST}${q[key]}"  # noqa: E731
            f = {"s_cit": f"=$H${q['cit']}", "s_irr": f"=XIRR({rng('post')},$J${year_end}:${LAST}${year_end})",
                 "s_debt": f"=$H${q['draw']}", "s_eirr": f"=XIRR({rng('eq')},$J${year_end}:${LAST}${year_end})",
                 "s_dscr": f"=MIN({rng('dscr')})"}[key]
            d.formula(f"{c}{r}", f)
        if key in ("s_cit", "s_debt"):
            d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
        rows[key] = r
    # Check row (inserted before the last check, so the Master_Check SUM covers it).
    r = shift(check_last)
    d.copy_format(r - 1, r, last="H")
    d.text(f"C{r}", "Debt fully repaid at the end of the tenor (balance 0)")
    d.text(f"D{r}", "flag")
    for c in P:
        d.formula(f"{c}{r}", f"=N(ABS(${LAST}${rows[c]['bal']})>Tolerance)")
    d.formula(f"H{r}", f"=SUM(E{r}:G{r})")
    rows["check"] = r
    return rows


def build():
    with Office(SRC) as lo:
        lo.save(BASE_RC)                                         # the user's file, recalculated, for before / after
        edit_inputs(lo.doc)
        rows = edit_dashboard(lo.doc)
        lo.calculate()
        lo.save(OUT)
    with Office(OUT) as lo:
        lo.save(RC)
    print(f"wrote {OUT}")
    return rows


# ---------------------------------------------------------------------------------------------------------------------
def xnpv(rate, flows):
    d0 = flows[0][0]
    return sum(v / (1 + rate) ** ((t - d0).days / 365) for t, v in flows)


def xirr(flows):
    lo, hi = -0.99, 10.0
    for _ in range(300):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if xnpv(mid, flows) > 0 else (lo, mid)
    return (lo + hi) / 2


def labels(ws):
    return {r: ws.cell(r, 3).value for r in range(1, ws.max_row + 1)}


def row_map(before, after):
    """Old row → new row: the labels of the old sheet in order, matched to the new sheet skipping inserted rows."""
    out, j = {}, 1
    lb, la = labels(before), labels(after)
    for r in range(1, before.max_row + 1):
        k = next((k for k in range(j, j + 30) if la.get(k) == lb.get(r)), j)   # relabelled row: same position
        out[r] = k
        j = k + 1
    return out


def compare(b_wb, a_wb):
    diffs = []
    for name in ("Dashboard", "Inputs"):
        b, a = b_wb[name], a_wb[name]
        m = row_map(b, a)
        for r in range(1, b.max_row + 1):
            for c in range(2, b.max_column + 1):
                x, y = b.cell(r, c).value, a.cell(m[r], c).value
                same = (x == y) or (isinstance(x, (int, float)) and isinstance(y, (int, float))
                                    and abs(x - y) <= 1e-9 * max(1, abs(x)))
                if not same:
                    diffs.append((name, r, m[r], b.cell(r, 3).value, b.cell(r, c).column_letter, x, y))
    return diffs


def verify(rows=None):
    a_wb = openpyxl.load_workbook(RC, data_only=True)
    b_wb = openpyxl.load_workbook(BASE_RC, data_only=True)
    errs = lint_values(RC)
    dash, inp = a_wb["Dashboard"], a_wb["Inputs"]
    print("recalc:", "no errors, Master_Check 0" if not errs else [str(e) for e in errs[:10]])
    found = [str(f) for f in lint(OUT)]
    print("lint:", "clean" if not found else f"{len(found)} finding(s)")
    for f in found[:15]:
        print("  ", f)
    # Before / after: every cell of the user's Dashboard and Inputs.
    diffs = compare(b_wb, a_wb)
    print(f"before / after: {len(diffs)} changed cell(s) in the user's rows")
    for sh, r0, r1, lab, cl, x, y in diffs:
        fx_ = lambda v: f"{v:,.2f}" if isinstance(v, (int, float)) else str(v)  # noqa: E731
        print(f"   {sh}!{cl}{r0} → {cl}{r1}  {str(lab)[:58]:<58} {fx_(x):>14} → {fx_(y)}")
    # Independent recomputation for Bora from the timeline rows of the recalculated file.
    find = lambda t, s=1: next(r for r in range(s, dash.max_row + 1) if dash.cell(r, 3).value == t)  # noqa: E731
    v = lambda name: inp.cell(next(r for r in range(1, inp.max_row + 1)  # noqa: E731
                                   if inp.cell(r, 3).value and inp.cell(r, 3).value.startswith(name)), 5).value
    cit, dep_y, gear, rate, ten = (v("Corporate income tax"), v("Tax depreciation period"), v("Gearing"),
                                  v("Debt interest rate"), v("Debt tenor"))
    head = find("Bora: annual cash flow (EUR k)")
    ye_r, tn_r = find("Year end"), find("Years from the COD year", head)
    eb_r, pre_r = find("EBITDA", head), find("Unlevered pre-tax cash flow", head)
    allin = dash.cell(find("All-in cost to COD (securities excluded: returned)"), 5).value
    cols = range(10, 10 + len(TL))
    ye = [dash.cell(ye_r, c).value.date() for c in cols]
    tn = [dash.cell(tn_r, c).value for c in cols]
    eb = [dash.cell(eb_r, c).value for c in cols]
    pre = [dash.cell(pre_r, c).value for c in cols]
    debt = gear * allin
    ann = debt * rate / (1 - (1 + rate) ** -ten)
    tax = [-cit * max(0, e - allin / dep_y * (0 <= n < dep_y)) for e, n in zip(eb, tn)]
    post = [p + t for p, t in zip(pre, tax)]
    ds = [-ann * (1 <= n <= ten) for n in tn]
    eq = [p + debt * (n == 0) + s for p, n, s in zip(post, tn, ds)]
    dscr = [(e + t) / -s for e, t, s in zip(eb, tax, ds) if s < 0]
    py = {"CIT over the life (15 %, no loss carry-forward)": sum(tax),
          "Project IRR, post-tax": xirr(list(zip(ye, post))), "Debt at COD (60 %)": debt,
          "Equity IRR, post-tax (60 % gearing)": xirr(list(zip(ye, eq))), "Minimum DSCR": min(dscr)}
    worst = 0.0
    for lab, x in py.items():
        m = dash.cell(find(lab), 5).value
        worst = max(worst, abs(m - x))
        print(f"   Bora {lab:<48} model {m:>14,.6f}  python {x:>14,.6f}")
    print(f"independent recomputation (Bora): max abs diff {worst:.2e} ->", "PASS" if worst < 1e-6 else "FAIL")
    dump(a_wb)
    print("dump:", DUMP)
    return not errs and worst < 1e-6


def dump(wb):
    def fmt(x, unit):
        if isinstance(x, datetime):
            return x.strftime("%d-%b-%Y")
        if isinstance(x, (int, float)) and not isinstance(x, bool):
            if unit and str(unit).startswith("%"):
                return f"{x:.2%}"
            return f"{x:,.4f}" if abs(x) < 10 and x != int(x) else f"{x:,.1f}"
        return "" if x is None else str(x)
    out = []
    for sheet, cols in (("Dashboard", "DEFGH"), ("Inputs", "DEFG")):
        ws = wb[sheet]
        out.append(f"#### {sheet}: label | unit | Bora | Iváncsa | Dunaújváros" + (" | Total" if "H" in cols else ""))
        for r in range(1, ws.max_row + 1):
            lab = ws[f"C{r}"].value
            if lab is None or lab == "***END***" or (sheet == "Inputs" and str(lab).startswith("Osz")):
                continue
            cells = [ws[f"{c}{r}"].value for c in cols]
            if all(x is None for x in cells[1:]):
                out.append(f"## {lab}")
                continue
            out.append(" | ".join([str(lab), str(cells[0] or "")] + [fmt(x, cells[0]) for x in cells[1:]]))
    DUMP.write_text("\n".join(out) + "\n")


if __name__ == "__main__":
    build()
    sys.exit(0 if verify() else 1)
