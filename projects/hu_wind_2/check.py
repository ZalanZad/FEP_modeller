"""Verification of projects/hu_wind_2/out/HU_Wind_bid_deposits_fees.xlsx (SPEC bora_v5_summary §21).

1. LibreOffice recalc (fepm.recalc.Office) -> out/HU_Wind_bid_deposits_fees_recalc.xlsx; no error values in the new
   rows, Master_Check 0. (The user's E22/E31 use _xlfn.SINGLE, which LibreOffice cannot evaluate; reported.)
2. Lint (fepm.audit.lint) on the new blocks (rows from 81): workbook lint filtered to the new rows, plus the
   calc-sheet rules the sheet name does not trigger: one formula per timeline row (X:BG) and direction
   (up/left; the summary rows at the top of the HTM block are a summary panel and exempt).
3. Independent Python recomputation from the input cells only: Iváncsa RtB claims (closed form per line), waterfall,
   FEP profit / MOIC / XIRR; Iváncsa HTM project IRR, NPV, COD price, COD-exit FEP profit / IRR. Must match to 1e-6.
4. out/dump.txt: label | unit | Iváncsa | Dunaújváros for the new inputs, the RtB exit and the HTM results rows.
5. §25 bids workbook (out/HU_Wind_bids.xlsx): build_lev2.verify() – recalc, lint, before / after vs the user's file,
   Iváncsa Case B, Bora Case A and the hold post-tax / equity IRRs recomputed, out/dump_bids.txt.

    python3 projects/hu_wind_2/check.py
"""
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1]))
from fepm.audit.lint import _direction, _name_targets, lint, lint_values, to_r1c1  # noqa: E402
from fepm.recalc import Office  # noqa: E402

OUT = HERE / "out" / "HU_Wind_bid_deposits_fees.xlsx"
RC = HERE / "out" / "HU_Wind_bid_deposits_fees_recalc.xlsx"
DUMP = HERE / "out" / "dump.txt"
R = json.loads((HERE / "out" / "rows.json").read_text())
FIRST = 81
TL = range(ci("X"), ci("BG") + 1)


def recalc():
    with Office(OUT) as lo:
        lo.calculate()
        lo.save(RC)


def lint_new():
    """Findings on the new rows; pre-existing findings in the user's rows are reported separately."""
    found = lint(OUT)
    new = [f for f in found if f.where.startswith("Bid_Cash!") and _row(f.where) >= FIRST]
    old = [f for f in found if f not in new]
    wb = openpyxl.load_workbook(OUT)
    ws, names = wb["Bid_Cash"], _name_targets(wb)
    lo, hi = R["summary"]
    for row in ws.iter_rows(min_row=FIRST):
        forms = {}
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                if not lo <= c.row <= hi:
                    new += _direction("Bid_Cash", c.row, c.column, c.value, names)
                if c.column in TL and c.row >= R["t_yr"]:
                    forms[c.column] = to_r1c1(c.value, c.row, c.column)
        if forms and (len(set(forms.values())) > 1 or len(forms) != len(TL)):
            new.append(f"[c-row] Bid_Cash!row {row[0].row}: timeline row not uniform or not filled")
    return new, old


def _row(where):
    m = re.search(r"(\d+)$", where)
    return int(m.group(1)) if m else 0


# ---------------------------------------------------------------------------------------------------------------------
def d(x):
    return x.date() if isinstance(x, datetime) else x


def xnpv(rate, flows):
    d0 = flows[0][0]
    return sum(v / (1 + rate) ** ((t - d0).days / 365) for t, v in flows)


def xirr(flows):
    lo, hi = -0.99, 10.0
    for _ in range(300):
        mid = (lo + hi) / 2
        if xnpv(mid, flows) > 0:
            lo = mid
        else:
            hi = mid
    return (lo + hi) / 2


def recompute(v, col="E"):
    """Independent recomputation for one project from the input cells (values) only."""
    g = lambda r: v[f"{col}{r}"].value  # noqa: E731
    e = lambda r: v[f"E{r}"].value      # common inputs and the user's key dates (column E)  # noqa: E731
    i = {k: g(R[k]) for k in ("prem", "sc", "share", "dsa", "dr", "sf", "turb", "inst", "fdn", "civ", "elec", "own",
                              "land", "subs", "basis", "ccont", "bess", "hrs", "curt", "clip", "deg", "ppa", "esc",
                              "om", "bop", "lease", "ins", "mgmt", "ptax", "adm", "bal", "grid", "mekh", "hipa",
                              "ocont", "oesc")}
    c = {k: e(R[k]) for k in ("rtb", "sale", "freeze", "devd", "r12", "r20", "floor", "life")}
    # The user's bid inputs, in HUF -> EUR k (recomputed here, not read from Bid_Cash rows 60-78).
    mva, ob, bshare, hours, muni, comm, pbm = g(5), g(6), g(7), g(8), g(9), g(10), g(12)
    fx_, bw, bb, pfee, bsr, pbr, tso, ind, dirf, sdd, devex = (g(r) for r in (45, 46, 47, 48, 49, 50, 51, 52, 53, 55, 56))
    wind, bess_mwh = mva * (1 + ob), mva * bshare * hours
    base = wind * bw + bess_mwh * bb
    k = lambda huf: huf / fx_ / 1000  # noqa: E731
    dates = {n: d(e(r)) for n, r in (("sign", 34), ("bid", 35), ("pb", 36), ("tso", 37), ("hc", 39), ("ntp", 40),
                                     ("cod", 41), ("rel", 42))}
    lines = [("fee", k(pfee), dates["bid"], c["r12"]), ("bs", k(bsr * base), dates["bid"], c["r12"]),
             ("pb", k(pbr * pbm * base - bsr * base), dates["pb"], c["r12"]),
             ("tso", k(tso * mva), dates["tso"], c["r12"]), ("dir", k(dirf), dates["hc"], c["r12"]),
             ("ind", k(ind * mva - tso * mva), dates["hc"], c["r12"]), ("dev", devex, d(c["devd"]), c["r20"])]
    end = min(d(c["freeze"]), d(c["sale"]))

    def factor(p, r):                       # closed form: act/365 both ends, simple within a year, 31-Dec capitalised
        if p.year == end.year:
            return 1 + r * ((end - p).days + 1) / 365
        first = (date(p.year, 12, 31) - p).days + 1
        last = (end - date(end.year, 1, 1)).days + 1
        return (1 + r * first / 365) * (1 + r) ** (end.year - p.year - 1) * (1 + r * last / 365)

    claims = {}
    for n, amt, p, r in lines:
        acc = amt * (factor(p, r) - 1)
        claims[n] = max(acc, amt * (c["floor"] - 1)) if n == "dev" else acc
    out = {"claim_" + n: x for n, x in claims.items()}
    amt = {n: a for n, a, _, _ in lines}
    premium = i["prem"] * mva
    reimb = sum(amt[n] for n in ("bs", "pb", "tso", "dir", "ind"))
    cash = premium + reimb
    paid = {}
    for n, due in (("sc", i["sc"]), ("rp", reimb), ("c12", sum(claims[n] for n in claims if n != "dev")),
                   ("dev", claims["dev"]), ("dsa", i["dsa"])):
        paid[n] = min(cash, due)
        cash -= paid[n]
    dist = cash
    receipts = paid["rp"] + paid["c12"] + paid["dev"] + dist * i["share"]
    flows = [(d(dates["sign"]), -sdd)] + [(p, -a) for _, a, p, _ in lines] + [(d(c["sale"]), receipts)]
    flows.sort(key=lambda t: t[0])
    out.update(rtb_premium=premium, rtb_reimb=reimb, rtb_dist=dist, rtb_receipts=receipts,
               rtb_profit=sum(x for _, x in flows),
               rtb_moic=sum(x for _, x in flows if x > 0) / -sum(x for _, x in flows if x < 0), rtb_irr=xirr(flows))

    # HTM, annual 2026-2061, flows at 31-Dec.
    built = wind * (1 - i["sf"])
    per_mw = sum(i[n] for n in ("turb", "inst", "fdn", "civ", "elec", "own", "land")) * built
    subs = i["subs"] * mva / i["basis"]
    capex = (per_mw + subs) * (1 + i["ccont"]) + i["bess"] * bess_mwh
    years = range(dates["sign"].year, dates["sign"].year + 36)

    def months(start, n):                 # the n month-ends after `start`, by calendar year
        cnt, y, m = {}, start.year, start.month
        for _ in range(n):
            m += 1
            y, m = (y + 1, 1) if m > 12 else (y, m)
            cnt[y] = cnt.get(y, 0) + 1
        return cnt
    nc = (dates["cod"].year - dates["ntp"].year) * 12 + dates["cod"].month - dates["ntp"].month
    cons, oper = months(dates["ntp"], nc), months(dates["cod"], int(c["life"]) * 12)
    out_flows = [(p, -a) for _, a, p, _ in lines] + [(dates["sign"], -sdd)]
    cf, ebitda_y, inv_y = {}, {}, {}
    for y in years:
        inv = sum(x for t, x in out_flows if t.year == y)
        inv += (amt["bs"] + amt["pb"]) * (y == dates["rel"].year)
        inv -= capex * cons.get(y, 0) / nc
        inv -= muni * mva * (y == dates["cod"].year)
        op = oper.get(y, 0) / 12
        n = y - dates["cod"].year
        gen = built * i["hrs"] * (1 - i["curt"]) * (1 - i["clip"]) * op * (1 - i["deg"]) ** max(n - 1, 0) / 1000
        rev = gen * i["ppa"] * (1 + i["esc"]) ** n
        ix = (1 + i["oesc"]) ** n
        base_opex = (sum(i[x] for x in ("om", "bop", "lease", "ins", "mgmt", "ptax")) * built * op
                     + i["adm"] * mva / i["basis"] * op + gen * (i["bal"] + i["grid"])) * ix + rev * i["mekh"]
        opex = base_opex * (1 + i["ocont"]) + gen * comm + rev * i["hipa"]
        ebitda_y[y], inv_y[y] = rev - opex, inv
        cf[y] = inv + rev - opex
    ye = {y: date(y, 12, 31) for y in years}
    price = sum(ebitda_y[y] * (1 + i["dr"]) ** ((dates["cod"] - ye[y]).days / 365) for y in years)
    cod_flows = [(ye[y], inv_y[y] * (y <= dates["cod"].year)
                  + (price + amt["bs"] + amt["pb"]) * (y == dates["cod"].year)) for y in years]
    htm = [(ye[y], cf[y]) for y in years]
    out.update(htm_irr=xirr(htm), htm_npv=xnpv(i["dr"], htm), cod_price=price,
               cod_profit=sum(x for _, x in cod_flows), cod_irr=xirr(cod_flows),
               allin=capex + amt["tso"] + amt["dir"] + amt["ind"] + muni * mva + amt["fee"] + sdd + devex,
               ebitda1=ebitda_y[dates["cod"].year + 1])
    return out


def compare(v, x):
    s = lambda key, col="E": v[f"{col}{R[key]}"].value  # noqa: E731
    model = {f"claim_{n}": s(f"c_{n if n != 'fee' else 'pf'}") for n in ("fee", "bs", "pb", "tso", "dir", "ind", "dev")}
    model.update(rtb_premium=s("w_prem"), rtb_reimb=s("w_reimb"), rtb_dist=s("w_dist"), rtb_receipts=s("f_rec"),
                 rtb_profit=s("r_profit"), rtb_moic=s("r_moic"), rtb_irr=s("r_irr"), htm_irr=s("h_irr"),
                 htm_npv=s("h_npv"), cod_price=s("x_price"), cod_profit=s("x_profit"), cod_irr=s("x_irr"),
                 allin=s("k_allin"), ebitda1=s("h_ebitda"))
    worst = 0.0
    lines = []
    for key, m in model.items():
        diff = abs(m - x[key])
        worst = max(worst, diff)
        lines.append(f"  {key:<14} model {m:>16,.6f}  python {x[key]:>16,.6f}  diff {diff:.2e}")
    return worst, lines


def dump(v):
    lo = R["h_irr"]
    hi = R["k_allin"]
    out = ["label | unit | Iváncsa | Dunaújváros"]
    ws = v["Bid_Cash"]

    def fmt(x, unit):
        if isinstance(x, datetime):
            return x.strftime("%d-%b-%Y")
        if isinstance(x, (int, float)) and not isinstance(x, bool):
            if unit and unit.startswith("%"):
                return f"{x:.2%}"
            return f"{x:,.4f}" if abs(x) < 10 and x != int(x) else f"{x:,.1f}"
        return "" if x is None else str(x)
    for r in list(range(FIRST, R["r_mva"] + 1)) + list(range(lo, hi + 1)):
        lab = ws[f"C{r}"].value
        if lab is None:
            continue
        cells = [ws[f"{c}{r}"].value for c in "DEF"]
        if ws[f"C{r}"].fill.fgColor.rgb in ("FF00563C", "FFE6F4EC") or all(x is None for x in cells):
            out.append(f"## {lab}")
            continue
        out.append(" | ".join([lab, cells[0] or ""] + [fmt(x, cells[0]) for x in cells[1:]]))
    DUMP.write_text("\n".join(out) + "\n")

# =====================================================================================================================
# §25: out/HU_Wind_bids.xlsx is the user's flagged file plus the JV flag, Case A at sale A and §24 (CIT, 60 % levered), built and
# verified by build_lev2.py (before / after comparison, Case A / B and hold recomputation, dump).
def bids():
    from build_lev2 import verify
    print("\n§25 bids workbook (build_lev2.verify)")
    return verify()


def main():
    recalc()
    vals = openpyxl.load_workbook(RC, data_only=True)
    v = vals["Bid_Cash"]
    allerr = lint_values(RC)
    errs = [f for f in allerr if f.rule != "f-error" or _row(f.where) >= FIRST]
    print("recalc (new blocks):", "no errors, Master_Check", v[f"G{R['master']}"].value if not errs else errs[:10])
    print("recalc (user's rows, untouched; LibreOffice has no _xlfn.SINGLE, Excel computes them):",
          [f"{f.where} {f.detail}" for f in allerr if f not in errs])
    new, old = lint_new()
    print("lint (new blocks):", "clean" if not new else f"{len(new)} finding(s)")
    for f in new[:30]:
        print("  ", f)
    print("lint (user's rows, untouched):", [str(f) for f in old])
    x = recompute(v)
    worst, lines = compare(v, x)
    print("\n".join(lines))
    print(f"independent recomputation (Iváncsa): max abs diff {worst:.2e} ->", "PASS" if worst < 1e-6 else "FAIL")
    dump(vals)
    print("dump:", DUMP)
    ok = bids()
    return 0 if (not errs and not new and worst < 1e-6 and ok) else 1


if __name__ == "__main__":
    sys.exit(main())
