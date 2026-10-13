"""Bana – 50 MW onshore wind (Ács Katódgyár node): acquisition, build, sale at COD or hold. Built to SPEC.md v1.

Dashboard | Inputs | Curves | Calc_1 (European turbine) | Calc_2 (Chinese turbine) | Checks,
quarterly from 01-Oct-2026 to 30-Sep-2060. Inputs come from the old Bana engine (reference/Bana_IC_model_v3.xlsm);
debt and tax follow Project Rock / Alpha with Hungarian rules. Solver.bas closes the tax circularity,
Scenarios.bas stores the key-assumption sets.
"""
import datetime as dt
import importlib.util
from pathlib import Path

from openpyxl.utils import get_column_letter as L

from fepm.xl.builder import Model
from fepm.xl.macros import package
from fepm.xl.styles import FMT, apply

HERE = Path(__file__).resolve().parent
NAME = "Bana_Model"
D = dt.date
SETS = ["Base", "COD price −10%", "Turbine +10%", "Yield −10%", "PPA −10%", "Full PB call", "Rate +1 pp",
        "Tenor 10 years"]
TO_COD = 'SUMIF({end.all},"<="&In_COD,{self.all})'          # amount to COD
EVENT = "N(AND({start}<={self.F},{end}>={self.F}))"           # line paid on its date in F
SHARE = "=MAX(0,MIN({self.F},{end})-MAX({self.E},{start})+1)/({self.F}-{self.E}+1)"   # share of a window
PATCH = {}                                                  # (calc key, column) -> number format


# =========================================================================================== inputs
def inputs(m):
    i = m.inputs
    a = i.add_input
    a("Bid_Date", "Bid submission: fee paid, bid security posted, DD", "date", D(2026, 10, 30),
      source="Kiírás VI.5")
    a("Award_Date", "Award decision", "date", D(2026, 12, 13), source="Bana v3 Inputs E96")
    a("Decision_Date", "Award decision received", "date", D(2027, 1, 31), source="GYIK Q68")
    a("PB_Days", "PB posted, TSO security paid after the decision", "days", 15, source="Kiírás XII.5, XIII.6",
      fmt="int")
    a("Closing", "Closing of the rights transfer: price paid", "date", D(2026, 12, 31),
      source="Bana deck p2: closing after the award")
    a("NTP", "Notice to proceed", "date", D(2028, 9, 30), source="Bana v3 Inputs E154")
    a("HCSSZ_Date", "Connection agreement (HCSSZ): fees paid", "date", D(2029, 1, 8), source="Bana v3 Inputs E99")
    a("COD", "Commercial operation date", "date", D(2030, 9, 30), source="Bana v3 Inputs E155, E237")
    a("Life", "Operating life", "years", 30, source="Bana v3 Inputs E157", fmt="int")
    a("PB_Release", "Performance security released (hold route)", "date", D(2033, 3, 31),
      source="Kiírás XII.8 (COD deadline + 6 m)")

    i.section("Acquisition, exit and tax")
    a("FX", "FX rate", "HUF/EUR", 365, source="Bana v3 Inputs E11", fmt="money")
    a("Purchase_Price", "Purchase price to GEI", "EUR k", 4500, source="Bana v3 Inputs E283")
    a("DD", "Transaction and DD costs", "EUR k", 150, source="Bana v3 Inputs E25")
    a("Devex", "Devex to NTP incl. contingency", "EUR k", 500, source="Bana v3 Inputs E292")
    a("COD_Price", "Sale price at COD per MW, EU turbines", "EUR k/MW", 2450, {2: 2205}, source="Bana v3 Inputs M230")
    a("COD_Price_CN", "Sale price at COD per MW, CN turbines", "EUR k/MW", 2100, {2: 1890},
      source="Bana v3 Inputs N230")
    a("Sale_Costs", "Sale costs at COD", "EUR k", 500, source="Bana v3 Inputs E26")
    a("Discount_Rate", "FEP discount rate (NPV)", "%", 0.08, source="Bana v3 Inputs E253")
    a("CIT", "Corporate income tax", "%", 0.09, source="Bana v3 Inputs E12; Tao. 19. §")
    a("Sale_Gain_Taxed", "Gain on the COD share sale taxed (0 = exempt)", "flag", 0,
      source="Bana v3 E231; participation exemption")
    a("Dep_Years", "Tax depreciation period", "years", 20, source="Bana v3 Inputs E213 (to confirm)", fmt="int")
    a("Loss_Cap", "Tax losses usable, share of the tax base", "%", 0.5, source="Tao. 17. § (2); Bana v3 E214")
    a("Loss_Years", "Tax losses usable within", "years", 5, source="Tao. 17. § (1)", fmt="int")

    i.section("Tender and grid")
    i.add_constant("Points_Decimals", "Points rounded to decimals", "#", 2, "Kiírás X.1 b")
    a("MVA", "Net connection capacity", "MVA", 50, source="MEKH Annex 1; Bana v3 Inputs E281", fmt="money")
    a("Overbuild_Built", "Overbuild built, share of MVA (bid: Dashboard)", "%", 0.10, source="User 13-Oct-26")
    a("Bottleneck_Loss", "Energy lost on overbuilt MW", "%", 0.05, source="Bora!N18 (as Bora)")
    a("BESS_Built", "BESS built (1 = yes)", "flag", 1, source="User 13-Oct-26 (as Bora)")
    a("Min_Hours", "Minimum BESS duration for ÖSZ-4", "h", 4, source="Kiírás Annex 8; GYIK Q5", fmt="money1")
    a("Tolerance_Cap", "Build shortfall tolerated without sanction", "%", 0.10, source="GYIK 28, 30")
    a("Sanction_Regime", "Sanction regime (1 pro rata, 2 full PB call)", "flag", 1, {6: 2},
      source="GYIK 28/30 (1); Kiírás XIII.17, GYIK 72 (2)")
    a("Sanction_Lag", "Sanction paid after COD", "months", 6, source="Kiírás XIII.17", fmt="int")
    a("Bench_BESS", "Security base: benchmark investment, BESS", "mHUF/MWh", 200, source="Kiírás Annex 2",
      fmt="money")
    a("Bench_Wind", "Security base: benchmark investment, wind", "mHUF/MW", 406, source="Kiírás Annex 2",
      fmt="money")
    a("Part_Fee", "Participation fee (cash, not refunded)", "mHUF", 3, source="Kiírás VI.4", fmt="money1")
    a("Bid_Rate", "Bid security (guarantee), share of base", "%", 0.005, source="Kiírás VIII.2", fmt="0.00%")
    a("PB_Rate", "Performance security (guarantee), minimum share", "%", 0.025, source="Kiírás XII.3",
      fmt="0.00%")
    a("TSO_Security", "TSO security (cash, credited to the indirect fee)", "mHUF/MVA", 4.5,
      source="Bana v3 Inputs E83", fmt="money1")
    a("Direct_Fee", "Direct connection fee per connection point (cash)", "mHUF", 460,
      source="MEKH Annex 1; Bana v3 Inputs E289", fmt="money1")
    a("Indirect_Fee", "Indirect connection fee (cash)", "mHUF/MVA", 54.830209,
      source="MEKH Annex 1; Bana v3 Inputs E288", fmt="money2")
    a("Guarantee_Fee", "Bank guarantee fee on guarantees outstanding", "%", 0.06, source="Bana v3 Inputs E113")

    i.section("Capex")
    a("Turbine", "Turbine supply, EU", "EUR k/MW", 600, {3: 660}, source="User 12-Oct-26 (Bora SPEC v2.1)")
    a("Turbine_CN", "Turbine supply, CN", "EUR k/MW", 500, {3: 550}, source="Bana v3 Inputs I162")
    a("Other_Wind", "Other wind capex (install, foundations, BoP, owner, land)", "EUR k/MW", 355,
      source="Bana v3 Inputs E163:E166, E169:E170")
    a("Substation", "Substation and 132 kV bay, EU", "EUR k", 12000, source="Bana v3 Inputs E167")
    a("Substation_CN", "Substation and 132 kV bay, CN transformer", "EUR k", 10800, source="Bana v3 Inputs I167")
    a("Capex_Contingency", "Contingency on wind and substation capex", "%", 0.05, source="Bana v3 Inputs E171")
    a("BESS_Capex", "BESS capex, all-in (no BESS revenue or opex)", "EUR k/MW", 240, source="User 12-Oct-26 (Bora)")

    i.section("Operations")
    a("P50_Hours", "Net P50 hours", "h/yr", 2460, {4: 2214}, source="Bana v3 Inputs E179", fmt="money")
    a("Curtailment", "Curtailment", "%", 0.02, source="Bana v3 Inputs E181")
    a("Degradation", "Degradation per year", "%", 0.003, source="Bana v3 Inputs E182", fmt="0.00%")
    a("PPA_Price", "PPA price at COD, whole life, 100 % of output", "EUR/MWh", 90, {5: 81},
      source="User 13-Oct-26", fmt="money1")
    a("PPA_Esc", "PPA escalation", "%", 0.01, source="Bana v3 Inputs E192")
    a("Var_Opex", "Variable opex (balancing 4, grid 1)", "EUR/MWh", 5, source="Bana v3 Inputs E194:E195",
      fmt="money1")
    a("Fixed_Opex", "Fixed opex (O&M, BoP, insurance, AM, property tax)", "EUR k/MW", 22,
      source="Bana v3 Inputs E201:E202, E204:E206")
    a("Lease", "Land lease from NTP", "EUR k/MW", 3, source="Bana v3 Inputs E203")
    a("Admin", "Administration", "EUR k", 150, source="Bana v3 Inputs E207")
    a("Opex_Contingency", "Opex contingency", "%", 0.05, source="Bana v3 Inputs E208")
    a("Opex_Esc", "Opex escalation from the COD year", "%", 0.025, source="Bana v3 Inputs E209")
    a("MEKH", "MEKH supervisory fee", "%", 0.001, source="Bana v3 Inputs E196", fmt="0.00%")
    a("HIPA", "Local business tax (HIPA), on revenue", "%", 0.02, source="Bana v3 Inputs E210")
    a("Innovation", "Innovation contribution, on the HIPA base", "%", 0.003, source="Act LXXVI of 2014, 15. §",
      fmt="0.00%")

    i.section("Senior debt")
    a("Gearing", "Debt share cap of construction cost, EU", "%", 0.80, source="Rock Assumptions G547")
    a("Gearing_CN", "Debt share cap of construction cost, CN", "%", 0.70, source="Rock G547 less 10 pp (Bana v3 CN)")
    a("DSCR", "Target DSCR for sizing and sculpting", "x", 1.20, source="Rock Assumptions G549",
      fmt="mult")
    a("Debt_Tenor", "Tenor from COD", "years", 18, {8: 10}, source="Rock G540; set 8: Bana v3 E220", fmt="int")
    a("Debt_Rate", "All-in fixed interest rate", "%", 0.05, {7: 0.06}, source="Bana v3 Inputs E218")
    a("Upfront_Fee", "Upfront fee on the facility", "%", 0.0175, source="Bana v3 Inputs E219", fmt="0.00%")
    a("Commitment_Fee", "Commitment fee on undrawn debt", "%", 0.009, source="Bana v3 Inputs E222", fmt="0.00%")
    a("DSRA_Months", "DSRA: months of debt service, funded at COD", "months", 6, source="Bana v3 Inputs E268",
      fmt="int")


# =========================================================================================== curves
TENDER = {  # lower bound, lower bound included, inner points (first entry = below the first band), weight, range
    1: ("ÖSZ-1 municipal one-off", "EUR k/MW", [10, 15, 20, 25, 30, 35, 40, 45, 50], [0] * 8 + [1],
        [0, 10, 20, 30, 40, 55, 70, 80, 90, 100], 0.20, (10, 50)),
    2: ("ÖSZ-2 community contribution", "EUR/MWh", [2, 2.1, 2.2, 2.3, 2.4, 2.5, 2.6, 2.7, 2.8, 2.9, 3],
        [0] * 10 + [1], [0, 5, 10, 20, 30, 40, 50, 60, 70, 80, 90, 100], 0.15, (2, 3)),
    4: ("ÖSZ-4 BESS ratio", "%", [x / 100 for x in (2, 5, 10, 15, 20, 25, 30, 35, 40, 50)], [0] + [1] * 8 + [0],
        list(range(0, 101, 10)), 0.15, (0.05, None)),
    3: ("ÖSZ-3 overbuild", "%", [x / 100 for x in (2, 5, 10, 15, 20, 25, 30, 35, 40, 50)], [0] + [1] * 8 + [0],
        list(range(0, 101, 10)), 0.10, None),
    6: ("ÖSZ-6 performance security multiple", "x", [1.0, 1.2, 1.3, 1.4, 1.5, 1.6, 1.7, 1.8, 1.9, 2.0],
        [0] + [1] * 9, list(range(0, 101, 10)), 0.05, None),
}
YES_NO = [("5", "ÖSZ-5 partial award", 100, 0.05), ("8", "ÖSZ-8 shared connection", 100, 0.05),
          ("9a", "ÖSZ-9 EU turbines", 80, 0.15), ("9b", "ÖSZ-9 EU transformers", 20, None)]


def curves(m):
    cv, ws = m.curves, m.curves.ws

    def line(label, unit, values, name=None, fmt=None, e=None, f=None):
        r = cv.r
        ws.cell(r, 3, label)
        apply(ws.cell(r, 4, unit), "unit")
        for k, v in enumerate(values):
            apply(ws.cell(r, 10 + k, v), "input", fmt)
        if name:
            m.define(name, f"Curves!$J${r}:${L(9 + len(values))}${r}")
        for col, item in ((5, e), (6, f)):
            if item:
                apply(ws.cell(r, col, item[1]), "input", item[2] if len(item) > 2 else fmt)
                m.define(item[0], f"Curves!${L(col)}${r}")
        cv.r += 1

    cv.r += 1
    for c in range(1, m.last + 1):
        apply(ws.cell(cv.r, c), "section")
    ws.cell(cv.r, 1, "x")
    ws.cell(cv.r, 3, "Tender scoring bands")
    ws.cell(cv.r, 5, "Kiírás Annex 3-4")
    cv.r += 1
    for c, h in {5: "Weight / min", 6: "Max", 10: "Bands →"}.items():
        apply(ws.cell(cv.r, c, h), "header")
    cv.r += 1
    for k, (label, unit, lb, inc, pts, w, rng) in TENDER.items():
        mn = (f"Osz{k}_Min", rng[0]) if rng else None
        mx = (f"Osz{k}_Max", rng[1]) if rng and rng[1] is not None else None
        line(f"{label} – lower bound", unit, lb, f"Osz{k}_LB", "general", e=mn, f=mx)
        line(f"{label} – lower bound included", "flag", inc, f"Osz{k}_Incl", "flag")
        line(f"{label} – inner points (below, bands)", "points", pts, f"Osz{k}_Pts", "int",
             e=(f"Osz{k}_W", w, "pct"))
    for k, label, pts, w in YES_NO:
        line(f"{label} – inner points", "points", [pts], f"Osz{k}_Pts", "int",
             e=(f"Osz{k[0]}_W", w, "pct") if w else None)


# =========================================================================================== calc
def fmt(key, col, f):
    PATCH[(key, col)] = f


def calc(m):
    c = m.calc
    lev = lambda k: c.lever(f"Dash_{k}")  # noqa: E731

    c.section("Timing and flags")
    c.row("constr", "Construction share (NTP → COD)", "factor", SHARE, e="In_NTP+1", f="In_COD", fmt="factor",
          ef_fmt="date")
    c.row("devex_w", "Devex share (closing → NTP)", "factor", SHARE, e="In_Closing+1", f="In_NTP", fmt="factor",
          ef_fmt="date")
    c.event("cod", "COD", "In_COD")
    c.window("ops", "Operations", "In_COD+1", "EDATE(In_COD,In_Life*Months_In_Year)")
    c.window("lease_w", "Land lease period (NTP → end of life)", "In_NTP+1", "EDATE(In_COD,In_Life*Months_In_Year)")
    c.window("tenor", "Debt repayment period", "In_COD+1", "EDATE(In_COD,In_Debt_Tenor*Months_In_Year)")
    c.event("tenor_end", "Last debt repayment", "EDATE(In_COD,In_Debt_Tenor*Months_In_Year)")
    c.row("dep", "Tax depreciation share", "factor", SHARE, e="In_COD+1",
          f="EDATE(In_COD,In_Dep_Years*Months_In_Year)", fmt="factor", ef_fmt="date")
    c.row("opyr", "Operating year", "#", "=ROUNDUP(YEARFRAC(In_COD,{end}),0)*({end}>In_COD)", total="none",
          fmt="int")
    c.row("idx", "Debt compounding index (E: at COD, F: rate per period)", "factor", "=(1+{self.F})^{period}",
          e='SUMIF({end.all},In_COD,{self.all})', f="In_Debt_Rate*Months_Per_Period/Months_In_Year",
          total="none", ef_fmt="factor")
    fmt("idx", "F", "0.000%")

    c.section("Capacity (E: committed in the bid, F: built)")
    c.row("wind", "Wind", "MW", e=f"In_MVA*(1+{lev('Overbuild')})", f="In_MVA*(1+In_Overbuild_Built)", ef_fmt="money1")
    c.row("bess", "BESS power", "MW", e=f"{lev('BESS')}*In_MVA", f="{self.E}*In_BESS_Built", ef_fmt="money1")
    c.row("mwh", "BESS energy", "MWh", e="{bess.E}*" + lev("BESS_Hours"), ef_fmt="money1")
    c.row("loss", "Grid-bottleneck loss on built wind", "%", e="({wind.F}-In_MVA)/{wind.F}*In_Bottleneck_Loss",
          ef_fmt="0.000%")
    c.row("short", "Build shortfall vs commitment (E: wind, F: BESS)", "%", e="MAX(0,1-{wind.F}/{wind.E})",
          f="(1-In_BESS_Built)*({bess.E}>0)", ef_fmt="pct")

    c.section("Tender points (E: commitment, F: weighted points)")
    band = lambda k: (f"INDEX(Osz{k}_Pts,1+SUMPRODUCT((Osz{k}_LB<{{self.E}})+(Osz{k}_LB={{self.E}})"  # noqa: E731
                      f"*Osz{k}_Incl))*Osz{k}_W")
    yes = lambda k, w: f"{{self.E}}*Osz{k}_Pts*Osz{w}_W"  # noqa: E731
    c.row("osz1", "ÖSZ-1 municipal one-off", "EUR k/MW", e=lev("Municipal"), f=band(1), ef_fmt="money2")
    c.row("osz2", "ÖSZ-2 community contribution", "EUR/MWh", e=lev("Community"), f=band(2), ef_fmt="money2")
    c.row("osz3", "ÖSZ-3 overbuild", "%", e=lev("Overbuild"), f=band(3), ef_fmt="money2")
    c.row("osz4", "ÖSZ-4 BESS power, share of MVA", "%", e=lev("BESS"), f=band(4) + "*({self.E}>=Osz4_Min)",
          ef_fmt="money2")
    c.row("osz5", "ÖSZ-5 partial award", "flag", e=lev("Partial"), f=yes(5, 5), ef_fmt="money2")
    c.row("osz6", "ÖSZ-6 performance security multiple", "x", e=lev("PB_Mult"), f=band(6), ef_fmt="money2")
    c.row("osz8", "ÖSZ-8 shared connection", "flag", e=lev("Shared"), f=yes(8, 8), ef_fmt="money2")
    c.row("osz9a", "ÖSZ-9 EU turbines (1 = European turbine route)", "flag", e=lev("EU_Turbine"), f=yes("9a", 9),
          ef_fmt="money2")
    c.row("osz9b", "ÖSZ-9 EU transformers", "flag", e=lev("EU_Transformer"), f=yes("9b", 9), ef_fmt="money2")
    for k, f in (("osz3", "pct"), ("osz4", "pct"), ("osz5", "flag"), ("osz6", "mult"), ("osz8", "flag"), ("osz9a", "flag"), ("osz9b", "flag")):
        fmt(k, "E", f)
    c.row("pts", "Points total", "points", e="ROUND(SUM({osz1.F}:{osz9b.F}),Points_Decimals)", ef_fmt="money2",
          bold=True)
    c.row("valid", "Failed tests (E: ÖSZ-1 and -2 range, F: ÖSZ-6, BESS ≤ wind, 4 h)", "#",
          e="({osz1.E}<>0)*({osz1.E}<Osz1_Min)+({osz1.E}>Osz1_Max)+({osz2.E}<>0)*({osz2.E}<Osz2_Min)"
            "+({osz2.E}>Osz2_Max)", f="N({osz6.E}<1)+N({bess.E}>{wind.E})+N({mwh.E}<{bess.E}*In_Min_Hours)",
          ef_fmt="int")

    c.section("Development, securities and fees (E: amount, F: date)")
    line = "=-{self.E}*" + EVENT
    c.row("buy", "Purchase price", "EUR k", line, e="In_Purchase_Price", f="In_Closing")
    c.row("dd", "Transaction and DD costs", "EUR k", line, e="In_DD", f="In_Bid_Date")
    c.row("devex", "Devex", "EUR k", "=-In_Devex*{devex_w}")
    c.row("pb_base", "Security base", "mHUF", e="{wind.E}*In_Bench_Wind+{mwh.E}*In_Bench_BESS", ef_fmt="money")
    c.row("fee", "Participation fee (cash)", "EUR k", line, e="In_Part_Fee*Thousand/In_FX", f="In_Bid_Date")
    c.row("bid", "Bid security outstanding (guarantee)", "EUR k",
          "={self.E}*MAX(0,(MIN(In_Decision_Date+In_PB_Days-1,{end})-MAX({self.F},{start})+1)/{days})",
          e="In_Bid_Rate*{pb_base.E}*Thousand/In_FX", f="In_Bid_Date", total="none")
    c.row("pb", "Performance security outstanding (guarantee)", "EUR k",
          "={self.E}*MAX(0,(MIN(In_PB_Release,{end})-MAX({self.F},{start})+1)/{days})",
          e="In_PB_Rate*{osz6.E}*{pb_base.E}*Thousand/In_FX", f="In_Decision_Date+In_PB_Days", total="none")
    c.row("tso", "TSO connection-request security (cash)", "EUR k", line,
          e="In_TSO_Security*In_MVA*Thousand/In_FX", f="In_Decision_Date+In_PB_Days")
    c.row("direct", "Direct connection fee (cash)", "EUR k", line, e="In_Direct_Fee*Thousand/In_FX", f="In_HCSSZ_Date")
    c.row("indirect", "Indirect connection fee, net of TSO security", "EUR k", line,
          e="(In_Indirect_Fee-In_TSO_Security)*In_MVA*Thousand/In_FX", f="In_HCSSZ_Date")
    for k in ("buy", "dd", "fee", "bid", "pb", "tso", "direct", "indirect"):
        fmt(k, "F", "date")
    c.row("conn", "Connection fees", "EUR k", "={direct}+{indirect}")
    c.row("gfee", "Guarantee fees (E: to COD)", "EUR k", "=-({bid}+{pb})*In_Guarantee_Fee*{days}/Days_In_Year",
          e=TO_COD)
    c.row("pbp", "Performance security by technology (E: wind, F: BESS)", "EUR k",
          e="{pb.E}*{wind.E}*In_Bench_Wind/{pb_base.E}", f="{pb.E}-{self.E}")
    c.row("sanc", "PB call (E: pro rata above the tolerance, F: full call)", "EUR k",
          e="({short.E}>In_Tolerance_Cap)*{short.E}*{pbp.E}+({short.F}>In_Tolerance_Cap)*{short.F}*{pbp.F}",
          f="({short.E}+{short.F}>0)*{pb.E}")
    c.row("sanction", "PB call paid, selected regime (E: date)", "EUR k",
          "=-CHOOSE(In_Sanction_Regime,{sanc.E},{sanc.F})*N(AND({start}<={self.E},{end}>={self.E}))",
          e="EDATE(In_COD,In_Sanction_Lag)")
    fmt("sanction", "E", "date")
    c.row("dev_cf", "Development cash flow", "EUR k", "={buy}+{dd}+{devex}+{fee}+{tso}+{gfee}", bold=True)

    c.section("Construction")
    c.row("tur", "Turbine supply, this route", "EUR k/MW", e="IF({osz9a.E}=1,In_Turbine,In_Turbine_CN)")
    c.row("subst", "Substation and 132 kV bay, this route", "EUR k",
          e="IF({osz9a.E}=1,In_Substation,In_Substation_CN)")
    c.row("capex", "Capex (E: total)", "EUR k", "=-{self.E}*{constr}",
          e="(({tur.E}+In_Other_Wind)*{wind.F}+{subst.E})*(1+In_Capex_Contingency)+In_BESS_Capex*{bess.F}")
    c.row("municipal", "Municipal one-off at COD", "EUR k", "=-{self.E}*{cod}", e="{osz1.E}*In_MVA")
    c.row("cost", "Construction cost (E: compounded to COD at the debt rate)", "EUR k",
          "={capex}+{conn}+{municipal}", e="SUMPRODUCT({self.all}/{idx.all})*{idx.E}", bold=True)

    c.section("Operations")
    c.row("esc", "Opex escalation index", "factor", "=(1+In_Opex_Esc)^({year}-YEAR(In_COD))", total="none")
    c.row("price", "PPA price", "EUR/MWh", "=In_PPA_Price*(1+In_PPA_Esc)^({year}-YEAR(In_COD))", total="none")
    c.row("gen", "Net generation (E: first full year)", "MWh",
          "={wind.F}*In_P50_Hours*(1-In_Curtailment)*(1-{loss.E})*(1-In_Degradation)^({opyr}-1)*{ops}*{yf}",
          e="SUMIF({year.all},YEAR(In_COD)+1,{self.all})")
    c.row("rev", "Revenue (E: first full year)", "EUR k", "={gen}*{price}/Thousand",
          e="SUMIF({year.all},YEAR(In_COD)+1,{self.all})")
    c.row("opex_var", "Variable opex (balancing, grid)", "EUR k", "=-In_Var_Opex*{gen}*{esc}/Thousand")
    c.row("opex_fix", "Fixed opex and administration", "EUR k",
          "=-(In_Fixed_Opex*{wind.F}+In_Admin)*{esc}*{ops}*{yf}")
    c.row("lease", "Land lease", "EUR k", "=-In_Lease*{wind.F}*{esc}*{lease_w}*{yf}")
    c.row("levies", "MEKH fee, HIPA and innovation contribution", "EUR k", "=-{rev}*(In_MEKH+In_HIPA+In_Innovation)")
    c.row("opex_cont", "Opex contingency", "EUR k", "=In_Opex_Contingency*({opex_var}+{opex_fix}+{lease}-{rev}*In_MEKH)")
    c.row("community", "Community contribution (ÖSZ-2)", "EUR k", "=-{osz2.E}*{gen}/Thousand")
    c.row("ebitda", "EBITDA (E: first full year)", "EUR k", "=SUM({rev}:{community})",
          e="SUMIF({year.all},YEAR(In_COD)+1,{self.all})", bold=True)

    c.section("Debt sizing on CFADS")
    c.paste("Tax", "Corporate income tax", "EUR k")
    c.row("cfads", "CFADS (EBITDA, guarantee fees, tax; after COD)", "EUR k",
          "=({ebitda}+{gfee}+{tax_paste})*({end}>In_COD)")
    c.row("ds_cap", "Debt service at target DSCR (E: PV at COD = capacity)", "EUR k", "={cfads}/In_DSCR*{tenor}",
          e="SUMPRODUCT({self.all}/{idx.all})*{idx.E}")
    c.row("gcap", "Debt share cap, this route", "%", e="IF({osz9a.E}=1,In_Gearing,In_Gearing_CN)", ef_fmt="pct")
    c.row("k", "Debt share of construction cost (E), share of capacity used (F)", "%",
          e="MIN({gcap.E},{ds_cap.E}/-{cost.E})", f="{self.E}*-{cost.E}/{ds_cap.E}", ef_fmt="pct", bold=True)

    c.section("Senior debt")
    c.row("d_int", "Interest paid", "EUR k", "=-{d_close@prev}*{idx.F}*({end}>In_COD)")
    c.row("d_ds", "Debt service (sculpted)", "EUR k", "=-{ds_cap}*{k.F}")
    c.row("d_open", "Opening balance", "EUR k", "={d_close@prev}", total="none")
    c.row("d_draw", "Drawdown", "EUR k", "=-{k.E}*{cost}")
    c.row("d_idc", "Interest capitalised", "EUR k", "={d_open}*{idx.F}*({end}<=In_COD)")
    c.row("d_rep", "Repayment", "EUR k", "={d_ds}-{d_int}")
    c.row("d_close", "Closing balance (E: at COD)", "EUR k", "=SUM({d_open}:{d_rep})",
          e="SUMIF({end.all},In_COD,{self.all})", total="seed", bold=True)
    c.row("d_up", "Upfront fee", "EUR k", "=-In_Upfront_Fee*{d_close.E}*({d_open}=0)*({d_draw}>0)")
    c.row("d_cf", "Commitment fee", "EUR k",
          "=-In_Commitment_Fee*({d_draw.I}-SUM({d_draw.todate})+{d_draw})*{yf}*({constr}>0)")
    c.row("dsra", "DSRA funded at COD, released at the end (E: balance)", "EUR k",
          "={self.E}*({tenor_end}-{cod})",
          e='-SUMIFS({d_ds.all},{end.all},">"&In_COD,{end.all},"<="&EDATE(In_COD,In_DSRA_Months))')
    c.row("dscr", "DSCR (E: minimum)", "x", '=IF({d_ds}<0,{cfads}/-{d_ds},"")', e="MIN({self.all})",
          total="none", fmt="mult", ef_fmt="mult")

    c.section("Tax")
    c.row("tax_dep", "Tax depreciation (E: base)", "EUR k", "=-{self.E}*{dep}",
          e="{d_idc.I}-({capex.I}+{conn.I}+{municipal.I}+{tso.I}+{buy.I}+{dd.I}+{devex.I})")
    c.row("taxable", "Tax base before losses", "EUR k", "={ebitda}+{tax_dep}+{d_int}+{d_up}+{d_cf}+{fee}+{gfee}")
    c.row("lcf_open", "Tax losses: opening", "EUR k", "={lcf_close@prev}", total="none")
    c.row("lcf_add", "Losses added", "EUR k", "=MAX(0,-{taxable})")
    c.row("lcf_use", "Losses used (capped)", "EUR k", "=-MIN({lcf_open},In_Loss_Cap*MAX(0,{taxable}))")
    c.row("lcf_lapse", "Losses lapsed (E: date, the NTP year + carry-forward period)", "EUR k",
          "=-({lcf_open}+{lcf_add}+{lcf_use})*N(AND({start}<={self.E},{end}>={self.E}))",
          e="EDATE(In_NTP,In_Loss_Years*Months_In_Year)")
    fmt("lcf_lapse", "E", "date")
    c.row("lcf_close", "Tax losses: closing", "EUR k", "=SUM({lcf_open}:{lcf_lapse})", total="seed")
    c.copy("Tax", "Corporate income tax", "EUR k", "=-In_CIT*(MAX(0,{taxable})+{lcf_use})")

    c.section("Funding to COD (E: uses and sources)")
    c.row("f_dev", "Purchase price, DD and devex", "EUR k", e="-({buy.I}+{dd.I}+{devex.I})")
    c.row("f_sec", "Participation fee, TSO security, guarantee fees", "EUR k", e="-({fee.I}+{tso.I}+{gfee.E})")
    c.row("f_capex", "Construction capex (wind, substation, BESS)", "EUR k", e="-{capex.I}")
    c.row("f_conn", "Connection fees, net of TSO security", "EUR k", e="-{conn.I}")
    c.row("f_mun", "Municipal one-off", "EUR k", e="-{municipal.I}")
    c.row("f_fin", "Capitalised interest and debt fees", "EUR k", e="{d_idc.I}-{d_up.I}-{d_cf.I}")
    c.row("f_lease", "Land lease during construction", "EUR k", e='-SUMIF({end.all},"<="&In_COD,{ebitda.all})')
    c.row("f_uses", "Total uses", "EUR k", e="SUM({f_dev.E}:{f_lease.E})", bold=True)
    c.row("f_equity", "FEP equity (E), peak including the DSRA (F)", "EUR k", e="{f_uses.E}-{d_close.E}",
          f="{self.E}+{dsra.E}")
    c.row("f_gear", "Debt share of total uses", "%", e="{d_close.E}/{f_uses.E}", ef_fmt="pct")

    c.section("Sale at COD")
    c.row("cod_contrib", "FEP contributions to COD", "EUR k",
          "=({dev_cf}+{cost}+{d_draw}+{d_up}+{d_cf}+{ebitda})*({end}<=In_COD)")
    c.row("v_price", "Sale price (built MW, after bottleneck loss)", "EUR k",
          e="IF({osz9a.E}=1,In_COD_Price,In_COD_Price_CN)*{wind.F}*(1-{loss.E})")
    c.row("v_net", "Net proceeds to FEP (after debt, sale costs, any PB call)", "EUR k",
          e="{v_price.E}+{sanction.I}-{d_close.E}-In_Sale_Costs")
    c.row("v_tax", "Tax on the gain", "EUR k", e="-In_CIT*In_Sale_Gain_Taxed*MAX(0,{v_net.E}+{cod_contrib.I})")
    c.row("cod_cf", "FEP cash flow – sale at COD", "EUR k", "={cod_contrib}+({v_net.E}+{v_tax.E})*{cod}", bold=True)
    c.xirr("cod_irr", "FEP IRR – sale at COD", "cod_cf")
    c.row("cod_moic", "FEP MOIC – sale at COD", "x", e="({v_net.E}+{v_tax.E})/-{cod_contrib.I}", ef_fmt="mult",
          bold=True)
    c.xnpv("cod_npv", "FEP NPV – sale at COD", "In_Discount_Rate", "cod_cf")

    c.section("Hold to maturity")
    c.row("proj_cf", "Project cash flow before financing, with actual tax", "EUR k",
          "={dev_cf}+{cost}+{ebitda}+{sanction}+{tax_copy}")
    c.xirr("proj_irr", "Project IRR, actual tax, before financing", "proj_cf")
    c.row("htm_cf", "FEP equity cash flow – hold to maturity", "EUR k",
          "={dev_cf}+{cost}+{d_draw}+{d_up}+{d_cf}+{ebitda}+{d_ds}+{dsra}+{sanction}+{tax_copy}", bold=True)
    c.xirr("htm_irr", "FEP IRR – hold to maturity", "htm_cf")
    c.moic("htm_moic", "FEP MOIC – hold to maturity", "htm_cf")
    c.xnpv("htm_npv", "FEP NPV – hold to maturity", "In_Discount_Rate", "htm_cf")


# =========================================================================================== checks
def checks(m):
    k = m.checks
    k.add_check("Levers within the tender rules", "=N({valid.E}+{valid.F}>0)", per_period=False)
    k.add_check("COD falls on a quarter end", "=N(COUNTIF({end.all},In_COD)=0)", per_period=False)
    k.add_check("Debt at COD within the DSCR capacity", "=N({k.F}>1+Tolerance)", per_period=False)
    k.add_check("Debt repaid by the end of the tenor", "=N(({end}>={tenor.F})*ABS({d_close})>Tolerance)")
    k.add_check("DSCR not below target", "=N({dscr.E}<In_DSCR-Tolerance)", per_period=False)
    k.add_check("Sources = uses to COD", "=N(ABS({f_uses.E}+{cod_contrib.I}-{d_close.E})>Tolerance)",
                per_period=False)
    k.add_check("XIRR returned a number on every route",
                "=N(ISERROR({cod_irr.E}))+N(ISERROR({htm_irr.E}))+N(ISERROR({proj_irr.E}))", per_period=False)


# =========================================================================================== dashboard
LEVERS = [  # key, label, unit, European / Chinese, offerable range, points row, list options, format
    ("Municipal", "Municipal one-off (ÖSZ-1)", "EUR k/MW", [50, 50], "0 or 10–50", "osz1", None, "money"),
    ("Community", "Community contribution (ÖSZ-2)", "EUR/MWh", [3, 3], "0 or 2–3", "osz2", None, "money2"),
    ("Overbuild", "Overbuild committed (ÖSZ-3)", "%", [0.20, 0.20], "≥ 0 %", "osz3", None, "pct"),
    ("BESS", "BESS power, share of MVA (ÖSZ-4)", "%", [0.20, 0.20], "≤ wind MW", "osz4", None, "pct"),
    ("BESS_Hours", "BESS duration", "h", [4, 4], "≥ 4 h", None, None, "money1"),
    ("PB_Mult", "Performance security multiple (ÖSZ-6)", "x", [1.0, 1.0], "≥ 1.0", "osz6", None, "mult"),
    ("Partial", "Partial award accepted (ÖSZ-5)", "flag", [0, 0], "0 / 1", "osz5", [0, 1], "flag"),
    ("Shared", "Shared connection (ÖSZ-8)", "flag", [0, 0], "0 / 1", "osz8", [0, 1], "flag"),
    ("EU_Turbine", "EU turbines ≥ 80 % (ÖSZ-9), sets the route", "flag", [1, 0], "0 / 1", "osz9a", [0, 1], "flag"),
    ("EU_Transformer", "EU transformers ≥ 80 % (ÖSZ-9)", "flag", [0, 0], "0 / 1", "osz9b", [0, 1], "flag"),
]


def dashboard(m):
    d, c, S = m.dashboard, m.calc, m.calc_sheets
    m2 = "money2"

    d.section("The bid: commitments and points", ["Offerable range", "Points EU", "Points CN"])
    for key, label, unit, vals, rng, pkey, opts, f in LEVERS:
        d.input(key, label, unit, vals, fmt=f, options=opts, extra=[rng] + ([f"={c.ref(pkey, 'F', s)}" for s in S] if pkey else []))
    d.link("Points total", "points", "{pts.E}", out="Points", fmt=m2, bold=True)
    d.status("Commitments within the offerable range", "{valid.E}+{valid.F}")

    d.section("Capacity: committed vs built")
    d.link("Wind committed in the bid", "MW", "{wind.E}", out="Wind_MW", fmt="money1")
    d.link("Wind built", "MW", "{wind.F}", out="Wind_Built", fmt="money1")
    d.link("BESS committed", "MW", "{bess.E}", out="BESS_MW", fmt="money1")
    d.link("BESS energy committed", "MWh", "{mwh.E}", out="BESS_MWh", fmt="money1")
    d.link("BESS built", "MW", "{bess.F}", out="BESS_Built", fmt="money1")
    d.link("Wind build shortfall vs commitment", "%", "{short.E}", out="Shortfall", fmt="pct")
    d.link("Grid-bottleneck energy loss", "%", "{loss.E}", out="Bottleneck_Loss", fmt="0.00%")
    d.link("PB call if the pro-rata rule applies (GYIK 28/30)", "EUR m", "{sanc.E}", scale="Thousand",
           out="PB_Call_ProRata", fmt=m2)
    d.link("PB call if the full-call rule applies (GYIK 72)", "EUR m", "{sanc.F}", scale="Thousand",
           out="PB_Call_Full", fmt=m2)

    d.section("Returns at a glance")
    for label, unit, t, scale, out in [
            ("FEP profit – sale at COD", "EUR m", "{cod_cf.I}", "Thousand", "COD_Profit"),
            ("FEP IRR – sale at COD", "%", "{cod_irr.E}", None, "COD_IRR"),
            ("FEP MOIC – sale at COD", "x", "{cod_moic.E}", None, "COD_MOIC"),
            ("FEP NPV – sale at COD", "EUR m", "{cod_npv.E}", "Thousand", "COD_NPV"),
            ("FEP profit – hold to maturity", "EUR m", "{htm_cf.I}", "Thousand", "HTM_Profit"),
            ("FEP IRR – hold to maturity", "%", "{htm_irr.E}", None, "HTM_IRR"),
            ("FEP MOIC – hold to maturity", "x", "{htm_moic.E}", None, "HTM_MOIC"),
            ("FEP NPV – hold to maturity", "EUR m", "{htm_npv.E}", "Thousand", "HTM_NPV"),
            ("Project IRR, actual tax, before financing", "%", "{proj_irr.E}", None, "Project_IRR")]:
        d.link(label, unit, t, scale=scale, out=out, fmt="pct" if unit == "%" else (m2 if scale else None),
               bold=out in ("COD_IRR", "HTM_IRR"))

    d.section("Funding to COD")
    for t, label, out, bold in [
            ("{f_dev.E}", "Purchase price, DD and devex", "Uses_Development", False),
            ("{f_sec.E}", "Participation fee, TSO security, guarantee fees", "Uses_Securities", False),
            ("{f_capex.E}", "Construction capex (wind, substation, BESS)", "Uses_Capex", False),
            ("{f_conn.E}", "Connection fees, net of TSO security", "Uses_Connection", False),
            ("{f_mun.E}", "Municipal one-off", "Uses_Municipal", False),
            ("{f_fin.E}", "Capitalised interest and debt fees", "Uses_Financing", False),
            ("{f_lease.E}", "Land lease during construction", "Uses_Lease", False),
            ("{f_uses.E}", "Total uses", "Uses_Total", True),
            ("{d_close.E}", "Senior debt at COD (incl. capitalised interest)", "Debt_COD", False),
            ("{f_equity.E}", "FEP equity to COD", "Equity_COD", True),
            ("{f_equity.F}", "FEP peak equity, hold (incl. DSRA)", "Equity_Peak_HTM", False)]:
        d.link(label, "EUR m", t, scale="Thousand", out=out, fmt=m2, bold=bold)
    d.link("Debt share of total uses", "%", "{f_gear.E}", out="Gearing_Uses", fmt="pct")

    d.section("Securities and fees", ["Form", "Posted", "Released"])
    pb_date = f"={c.ref('pb', 'F', S[0])}"
    for key, label, form, posted, released in [
            ("fee", "Participation fee", "cash", "=In_Bid_Date", "–"),
            ("bid", "Bid security", "guarantee", "=In_Bid_Date", pb_date),
            ("pb", "Performance security", "guarantee", pb_date, "=In_PB_Release"),
            ("tso", "TSO connection-request security", "cash", pb_date, "credited"),
            ("direct", "Direct connection fee", "cash", "=In_HCSSZ_Date", "–"),
            ("indirect", "Indirect connection fee, net", "cash", "=In_HCSSZ_Date", "–")]:
        r = d.link(label, "EUR m", "{%s.E}" % key, scale="Thousand", fmt="0.000", out=f"Sec_{key}",
                   extra=[form, posted, released])
        for col in (d.X + 1, d.X + 2):
            d.ws.cell(r, col).number_format = FMT["date"]
    d.link("Guarantee fees to COD", "EUR m", "-{gfee.E}", scale="Thousand", out="Guarantee_Fees_COD", fmt="0.000")

    d.section("Sale at COD")
    d.link("Sale price (enterprise value)", "EUR m", "{v_price.E}", scale="Thousand", out="COD_Price", fmt=m2)
    d.link("Net proceeds to FEP (after debt and sale costs)", "EUR m", "{v_net.E}", scale="Thousand",
           out="COD_Proceeds", fmt=m2)
    d.link("Tax on the gain", "EUR m", "{v_tax.E}", scale="Thousand", out="COD_Tax", fmt=m2)

    d.section("Hold to maturity: debt and tax")
    d.link("Debt capacity at target DSCR (PV at COD)", "EUR m", "{ds_cap.E}", scale="Thousand", out="Debt_Capacity",
           fmt=m2)
    d.link("Debt share cap of construction cost", "%", "{gcap.E}", out="Gearing_Cap", fmt="pct")
    d.link("Debt share of construction cost, sized", "%", "{k.E}", out="Debt_Share", fmt="pct")
    d.link("Share of the DSCR capacity used", "%", "{k.F}", out="Capacity_Used", fmt="pct")
    d.link("DSRA funded at COD", "EUR m", "{dsra.E}", scale="Thousand", out="DSRA", fmt=m2)
    d.link("Minimum DSCR", "x", "{dscr.E}", out="Min_DSCR")
    d.link("Net generation, first full year", "GWh", "{gen.E}", scale="Thousand", out="Gen_Year1", fmt="money1")
    d.link("Revenue, first full year", "EUR m", "{rev.E}", scale="Thousand", out="Revenue_Year1", fmt=m2)
    d.link("EBITDA, first full year", "EUR m", "{ebitda.E}", scale="Thousand", out="EBITDA_Year1", fmt=m2)
    d.link("Lifetime revenue", "EUR m", "{rev.I}", scale="Thousand", out="Revenue_Life", fmt=m2)
    d.link("Lifetime EBITDA", "EUR m", "{ebitda.I}", scale="Thousand", out="EBITDA_Life", fmt=m2)
    d.link("Lifetime corporate income tax", "EUR m", "SUM({tax_copy.all})", scale="Thousand", out="Tax_Life", fmt=m2)
    d.link("Tax losses lapsed", "EUR m", "-{lcf_lapse.I}", scale="Thousand", out="Tax_Losses_Lapsed", fmt=m2)
    d.link("Checks failed in this set (0 = OK)", "#", "Model_Checks", out="Checks", fmt="int")

    d.scenario_results(outs=["Equity_COD", "COD_IRR", "COD_MOIC", "HTM_IRR", "HTM_MOIC", "Debt_COD", "Min_DSCR",
                              "Checks"])

    d.panel_section("Key-assumption set")
    d.selector()
    d.panel_section("Key assumptions (live)")
    for label, unit, ref, f in [
            ("Purchase price", "EUR k", "In_Purchase_Price", None), ("COD", "date", "In_COD", None),
            ("Overbuild built", "%", "In_Overbuild_Built", None),
            ("Sanction regime (1 pro rata, 2 full)", "flag", "In_Sanction_Regime", None),
            ("COD price, EU", "EUR k/MW", "In_COD_Price", None), ("COD price, CN", "EUR k/MW", "In_COD_Price_CN", None),
            ("Turbine price, EU", "EUR k/MW", "In_Turbine", None), ("Turbine price, CN", "EUR k/MW", "In_Turbine_CN", None),
            ("Net P50 hours", "h/yr", "In_P50_Hours", "money"), ("PPA price at COD, whole life", "EUR/MWh", "In_PPA_Price",
                                                               "money1"),
            ("Debt rate", "%", "In_Debt_Rate", None), ("Target DSCR", "x", "In_DSCR", None),
            ("Tenor from COD", "years", "In_Debt_Tenor", None), ("Debt share cap, EU", "%", "In_Gearing", None),
            ("Debt share cap, CN", "%", "In_Gearing_CN", None), ("CIT", "%", "In_CIT", None),
            ("Discount rate", "%", "In_Discount_Rate", None), ("FX", "HUF/EUR", "In_FX", None)]:
        d.panel_link(label, unit, ref, fmt=f)


# =========================================================================================== build
def model():
    m = Model("Project Bana", "50 MW onshore wind, Ács Katódgyár – sale at COD vs hold to maturity", periods=136,
              periodicity="quarterly", start=D(2026, 10, 1), scenarios=SETS, version="v1", date=D(2026, 10, 13),
              curves=True, layout="dashboard", calc_sheets=["Calc_1", "Calc_2"],
              bids=["European turbine", "Chinese turbine"], extra_cols=len(SETS), source_col=False)
    inputs(m)
    curves(m)
    calc(m)
    checks(m)
    dashboard(m)
    return m


def build_model(out: Path) -> Path:
    m = model()
    xlsx = m.save(out / f"{NAME}.xlsx")
    for (key, col), f in PATCH.items():          # second format for rows whose E and F hold different units
        for ws in m.calc.sheets:
            ws[f"{col}{m.calc.rows[key].r}"].number_format = FMT.get(f, f)
    m.wb.save(xlsx)
    xlsm = package(xlsx, out / f"{NAME}.xlsm", run=["Scenarios.RunAllScenarios"], include=["Scenarios"])
    xlsx.unlink()
    return xlsm


def build_deck(xlsm: Path, out: Path) -> Path:
    """IC deck (deck.py next to this file), bound to the packaged, LibreOffice-recalculated workbook."""
    path = HERE / "deck.py"
    if not path.exists():
        return None
    spec = importlib.util.spec_from_file_location("bana_deck", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.build_deck(xlsm, out)


def build(out: Path):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    xlsm = build_model(out)
    return {"model": xlsm, "deck": build_deck(xlsm, out)}
