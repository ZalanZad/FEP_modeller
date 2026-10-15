"""Bora bid strategy model – BORA 1 / BORA 2 (2 x 100 MVA, Szombathely). Built to SPEC.md v2.

Dashboard | Inputs | Curves | Calc_1 | Calc_2 | Checks, quarterly from 01-Jul-2026 to 30-Sep-2061.
Inputs come from the meeting file via the old-engine extract data/old_engine/devex_lines.csv;
every input row cites its cell. BoraImpact.bas stores the per-commitment impact, Scenarios.bas the key-assumption sets.
"""
import csv
import datetime as dt
import importlib.util
from pathlib import Path

from openpyxl.utils import get_column_letter as L

from fepm.xl.builder import Model
from fepm.xl.macros import package
from fepm.xl.styles import FMT, apply

HERE = Path(__file__).resolve().parent
OLD = HERE / "data" / "old_engine"
NAME = "Bora_Bid_Model"
D = dt.date
START, NTP = D(2026, 7, 1), D(2029, 9, 30)                 # model start; NTP (devex years run to the NTP year)
SETS = ["Base", "Overbuild built", "Full PB call", "Freeze off", "Low premium", "ÖSZ-4 on wind", "Nothing delivered"]
EXIT = 'SUMIF({end.all},"<="&{exit.E},{self.all})'          # amount to the RtB exit
TO_COD = 'SUMIF({end.all},"<="&In_COD,{self.all})'          # amount to COD
EVENT = "N(AND({start}<={self.F},{end}>={self.F}))"           # line paid on its date in F
PATCH = {}                                                  # (calc key, column) -> number format


# =========================================================================================== inputs
def inputs(m):
    i = m.inputs
    a = i.add_input
    a("Invest_Date", "Signing fee and DD paid", "date", D(2026, 7, 31), source="Old engine Inputs E27")
    a("Signing_Date", "Signing (DSA accrual start)", "date", D(2026, 10, 26), source="Old engine Inputs E70")
    a("Contingency_Date", "Devex contingency paid", "date", D(2028, 12, 31), source="Old engine Devex_FEP row 38")
    a("NTP", "Notice to proceed", "date", NTP, source="Old engine Inputs E154")
    a("COD", "Commercial operation date", "date", D(2031, 9, 30), source="Bora!N16")
    a("Life", "Operating life", "years", 30, source="Old engine Inputs E157", fmt="int")
    a("Freeze_Start", "Interest freeze start", "date", D(2029, 6, 30), source="Old engine Inputs E64")
    a("Freeze_End", "Interest freeze end (accrual resumes)", "date", D(2030, 1, 1), source="Old engine Inputs E73")
    a("Freeze_On", "Interest freeze on (1 = yes)", "flag", 1, {4: 0}, source="User 09-Oct-26 (Inputs E64)")
    a("Sanction_Lag", "Sanction paid after COD", "months", 6, source="Kiírás XIII.17; old Tender E31", fmt="int")

    i.section("Money, tax and JV")
    a("FX", "FX rate", "HUF/EUR", 365, source="Old engine Inputs E11", fmt="money")
    a("CIT", "Corporate income tax", "%", 0.09, source="Old engine Inputs E12")
    a("Signing_Fee", "Signing fee to current owners", "EUR k", 250, source="Old engine Inputs E24")
    a("DD", "Transaction and DD costs", "EUR k", 150, source="Old engine Inputs E25")
    a("Sale_Costs", "Sale costs at exit", "EUR k", 500, source="Old engine Inputs E26")
    a("FEP_Share", "FEP share of distributable proceeds", "%", 0.70, source="Old engine Inputs E33")
    a("DSA_Rate", "Developer success fee (DSA) per year", "EUR k/MW", 2.5, source="Old engine Inputs E36", fmt="money1")
    a("DSA_Cap", "DSA cap", "EUR k/MW", 10, source="Old engine Inputs E37")

    i.section("Tender")
    i.add_constant("Points_Decimals", "Points rounded to decimals", "#", 2, "Kiírás X.1 b")
    a("MVA", "Net connection capacity per bid", "MVA", 100, source="Old engine Calc E74", fmt="money")
    a("Group_Cap", "Cap per group per node", "MVA", 100, source="Kiírás V.2", fmt="money")
    a("Tolerance_Cap", "Build shortfall tolerated without sanction", "%", 0.10, source="GYIK 28, 30")
    a("Min_Hours", "Minimum BESS duration", "h", 4, source="Kiírás Annex 8; GYIK Q5", fmt="money1")
    a("OSZ4_Basis", "ÖSZ-4 basis (0 = net MVA, 1 = committed wind)", "flag", 0, {6: 1},
      source="User 12-Oct-26; GYIK Q20 = 1")

    i.section("Deposits and fees (timeline order)")
    a("Bench_Wind", "Security base: benchmark investment, wind", "mHUF/MW", 406, source="Kiírás Annex 2", fmt="money")
    a("Bench_BESS", "Security base: benchmark investment, BESS", "mHUF/MWh", 200, source="Kiírás Annex 2", fmt="money")
    a("Bid_Date", "Bid submission: fee paid, bid security posted", "date", D(2026, 10, 30),
      source="Kiírás VI.5; Old engine Inputs E95")
    a("Part_Fee", "Participation fee per bid (cash, not refunded)", "mHUF", 3, source="Kiírás VI.4", fmt="money1")
    a("Bid_Rate", "Bid security (guarantee), share of base", "%", 0.005, source="Kiírás VIII.2", fmt="0.00%")
    a("Award_Date", "Award decision", "date", D(2026, 12, 13), source="Old engine Inputs E96")
    a("PB_Days", "Performance security posted after award (bid released)", "days", 15, source="Kiírás XII.5, GYIK 68",
      fmt="int")
    a("PB_Rate", "Performance security (guarantee), minimum share of base", "%", 0.025, source="Kiírás XII.3",
      fmt="0.00%")
    a("TSO_Date", "Connection request: TSO security posted", "date", D(2026, 12, 30), source="Old engine Inputs E97")
    a("TSO_Security", "TSO security (cash, credited to the indirect fee)", "mHUF/MVA", 4.5,
      source="Old engine Inputs E83", fmt="money1")
    a("HCSSZ_Date", "Connection agreement (HCSSZ): fees paid", "date", D(2029, 1, 8), source="Old engine Inputs E99")
    a("Direct_Fee", "Direct connection fee per connection point (cash)", "mHUF", 508.345804,
      source="BORA_Devex 02-Oct row 19", fmt="money2")
    a("Indirect_Fee", "Indirect connection fee (cash)", "mHUF/MVA", 73.071306, source="BORA_Devex 02-Oct row 18",
      fmt="money2")
    a("PB_Release", "Performance security released (hold route)", "date", D(2033, 3, 31),
      source="Kiírás XII.8 (COD deadline + 6 m)")
    a("Guarantee_Fee", "Bank guarantee fee on guarantees outstanding", "%", 0.06, source="Old engine Inputs E113")

    i.section("FEP claims and fees")
    a("Cash_Rate", "Return on cash lines", "%", 0.12, source="BORA_Devex 02-Oct V18")
    a("Guar_Credit", "Return on guarantee face value", "%", 0.06, source="BORA_Devex 02-Oct V17")
    a("AtRisk_Rate", "Return on devex and land (at risk)", "%", 0.20, source="BORA_Devex 02-Oct V19")
    a("Min_Multiple", "Minimum multiple on at-risk principal", "x", 2.0, source="BORA_Devex 02-Oct V20")
    a("Devex_Contingency", "Contingency on devex, land and fee", "%", 0.10, source="Old engine Inputs E68")

    i.section("Exit prices")
    a("RtB_Premium", "RtB price (no price effect)", "EUR k/MVA", 150, {5: 100}, source="Bora!N5; user 12-Oct-26")
    a("Discount_Rate", "FEP discount rate (NPV, price PV)", "%", 0.08, source="Old engine Inputs E253")
    a("COD_Price", "Sale price at COD per built MW, EU turbines", "EUR k/MW", 2450, source="Bora!N6")
    a("COD_Price_CN", "Sale price at COD per built MW, CN turbines", "EUR k/MW", 2100, source="Bora!N7")
    a("Ref_Community", "COD price reference: community contribution", "EUR/MWh", 2, source="Bora!E7")
    a("Pass_Through", "COD price effect passed into the price", "%", 1.0, source="Old engine Tender E30")
    a("Buyer_Rate", "Buyer discount rate", "%", 0.08, source="Old engine Inputs E224")

    i.section("Build switches")
    a("Overbuild_Built", "Overbuild built (1 = yes)", "flag", 0, {2: 1}, source="SPEC v2 §6 (set = meeting file)")
    a("BESS_Built", "BESS built (1 = yes)", "flag", 1, {7: 0}, source="SPEC v2 §6")
    a("Delivered", "Points commitments delivered (1 = yes)", "flag", 1, {7: 0}, source="User 15-Oct-26")
    a("Sanction_Regime", "Sanction regime (1 pro rata, 2 full call)", "flag", 1, {3: 2, 7: 2},
      source="GYIK 28/30 (1); Kiírás XIII.17, GYIK 72 (2)")
    a("Bottleneck_Loss", "Energy lost on overbuilt MW", "%", 0.05, source="Bora!N18")

    i.section("Capex")
    a("CN_Turbine", "Chinese turbines (0 EU, 1 CN)", "flag", 0, source="User 12-Oct-26 (SPEC v2.1)")
    a("Turbine", "Turbine supply, EU", "EUR k/MW", 600, source="User 12-Oct-26 (SPEC v2.1)")
    a("Turbine_CN", "Turbine supply, CN", "EUR k/MW", 500, source="Old engine Calc G101")
    a("Installation", "Turbine installation and erection", "EUR k/MW", 80, source="Old engine Inputs E163")
    a("Foundations", "Foundations", "EUR k/MW", 90, source="Old engine Inputs E164")
    a("BoP", "Balance of plant (civil 70, electrical 60)", "EUR k/MW", 130, source="Old engine Inputs E165:E166")
    a("Owner", "Owner's costs, engineering, construction insurance", "EUR k/MW", 40, source="Old engine Inputs E169")
    a("Easements", "Land purchase and easements", "EUR k/MW", 15, source="Old engine Inputs E170")
    a("Substation", "Substation and 132 kV bay, EU", "EUR k", 12000, source="Old engine Inputs E167")
    a("Substation_CN", "Substation and 132 kV bay, CN", "EUR k", 10800, source="Old engine Tender E29")
    a("Capex_Contingency", "Contingency on wind and substation capex", "%", 0.05, source="Old engine Inputs E171")
    a("BESS_Capex", "BESS capex", "EUR k/MWh", 100, source="User 15-Oct-26")

    i.section("Operations")
    a("P50_Hours", "Net P50 hours", "h/yr", 2460, source="Old engine Inputs E179", fmt="money")
    a("Curtailment", "Curtailment", "%", 0.02, source="Old engine Inputs E181")
    a("Degradation", "Degradation per year", "%", 0.003, source="Old engine Inputs E182", fmt="0.00%")
    a("PPA_Price", "PPA price at COD (whole life)", "EUR/MWh", 90, source="User 12-Oct-26 (SPEC v2.1)", fmt="money1")
    a("PPA_Esc", "PPA escalation", "%", 0.01, source="Old engine Inputs E192")
    a("Var_Opex", "Variable opex (balancing 4, grid 1)", "EUR/MWh", 5, source="Old engine Inputs E194:E195",
      fmt="money1")
    a("Fixed_Opex", "Fixed opex (O&M, BoP, insurance, AM, property tax)", "EUR k/MW", 22,
      source="Old engine Inputs E201:E202, E204:E206")
    a("Lease", "Land lease from NTP", "EUR k/MW", 3, source="Old engine Inputs E203")
    a("Admin", "Administration", "EUR k", 150, source="Old engine Inputs E207")
    a("Opex_Contingency", "Opex contingency", "%", 0.05, source="Old engine Inputs E208")
    a("Opex_Esc", "Opex escalation from the COD year", "%", 0.025, source="Old engine Inputs E209")
    a("MEKH", "MEKH fee", "%", 0.001, source="Old engine Inputs E196", fmt="0.00%")
    a("HIPA", "Local business tax (HIPA)", "%", 0.02, source="Old engine Inputs E210")

    i.section("Debt and tax")
    a("Gearing", "Gearing on capex, connection, municipal, EU", "%", 0.60, source="Old engine Inputs E217")
    a("Gearing_CN", "Gearing on capex, connection, municipal, CN", "%", 0.50, source="Old engine case 3")
    a("Debt_Rate", "Debt interest rate", "%", 0.05, source="Old engine Inputs E218")
    a("Upfront_Fee", "Upfront fee", "%", 0.0175, source="Old engine Inputs E219", fmt="0.00%")
    a("Commitment_Fee", "Commitment fee on undrawn debt", "%", 0.009, source="Old engine Inputs E222", fmt="0.00%")
    a("Debt_Tenor", "Annuity tenor from COD", "years", 10, source="Old engine Inputs E220", fmt="int")
    a("Dep_Years", "Tax depreciation period", "years", 20, source="Old engine Inputs E213", fmt="int")


# =========================================================================================== curves
TENDER = {  # lower bound, lower bound included, inner points (first entry = below the first band)
    1: ("ÖSZ-1 municipal one-off", "EUR k/MW", [10, 15, 20, 25, 30, 35, 40, 45, 50], [0] * 8 + [1],
        [0, 10, 20, 30, 40, 55, 70, 80, 90, 100], 0.20, (10, 50)),
    2: ("ÖSZ-2 community contribution", "EUR/MWh", [2, 2.1, 2.2, 2.3, 2.4, 2.5, 2.6, 2.7, 2.8, 2.9, 3],
        [0] * 10 + [1], [0, 5, 10, 20, 30, 40, 50, 60, 70, 80, 90, 100], 0.15, (2, 3)),
    3: ("ÖSZ-3 overbuild", "%", [x / 100 for x in (2, 5, 10, 15, 20, 25, 30, 35, 40, 50)], [0] + [1] * 8 + [0],
        list(range(0, 101, 10)), 0.10, None),
    4: ("ÖSZ-4 BESS ratio", "%", [x / 100 for x in (2, 5, 10, 15, 20, 25, 30, 35, 40, 50)], [0] + [1] * 8 + [0],
        list(range(0, 101, 10)), 0.15, (0.05, None)),
    6: ("ÖSZ-6 performance security multiple", "x", [1.0, 1.2, 1.3, 1.4, 1.5, 1.6, 1.7, 1.8, 1.9, 2.0],
        [0] + [1] * 9, list(range(0, 101, 10)), 0.05, None),
}
YES_NO = [("5", "ÖSZ-5 partial award", 100, 0.05), ("7", "ÖSZ-7 brownfield", 100, 0.10),
          ("8", "ÖSZ-8 shared connection", 100, 0.05), ("9a", "ÖSZ-9 EU turbines", 80, 0.15),
          ("9b", "ÖSZ-9 EU transformers", 20, None)]


def curves(m):
    cv, ws = m.curves, m.curves.ws

    def bar(title, source=""):
        cv.r += 1
        for c in range(1, m.last + 1):
            apply(ws.cell(cv.r, c), "section")
        ws.cell(cv.r, 1, "x")
        ws.cell(cv.r, 3, title)
        ws.cell(cv.r, 5, source)
        cv.r += 1

    def header(cells):
        for c, h in cells.items():
            apply(ws.cell(cv.r, c, h), "header")
        cv.r += 1

    def line(label, unit, values, name=None, fmt=None, e=None, f=None):
        r = cv.r
        ws.cell(r, 3, label)
        apply(ws.cell(r, 4, unit), "unit")
        for k, v in enumerate(values):
            apply(ws.cell(r, 10 + k, v), "input", fmt)
        if name:
            m.define(name, f"Curves!$J${r}:${L(9 + len(values))}${r}")
        for col, item in ((5, e), (6, f)):
            if item:
                apply(ws.cell(r, col, item[1]), "input", item[2] if len(item) > 2 else fmt)
                m.define(item[0], f"Curves!${L(col)}${r}")
        cv.r += 1

    bar("Tender scoring bands", "Kiírás Annex 3-4; old engine Tender D93:F161")
    header({5: "Weight / min", 6: "Max", 10: "Bands →"})
    for k, (label, unit, lb, inc, pts, w, rng) in TENDER.items():
        mn = (f"Osz{k}_Min", rng[0]) if rng and rng[0] is not None else None
        mx = (f"Osz{k}_Max", rng[1]) if rng and rng[1] is not None else None
        line(f"{label} – lower bound", unit, lb, f"Osz{k}_LB", "general", e=mn, f=mx)
        line(f"{label} – lower bound included", "flag", inc, f"Osz{k}_Incl", "flag")
        line(f"{label} – inner points (below, bands)", "points", pts, f"Osz{k}_Pts", "int",
             e=(f"Osz{k}_W", w, "pct"))
    for k, label, pts, w in YES_NO:
        line(f"{label} – inner points", "points", [pts], f"Osz{k}_Pts", "int",
             e=(f"Osz{k[0]}_W", w, "pct") if w else None)

    bar("Devex schedule", "BORA_Devex 25-Sep (meeting file Devex_FEP); excl. Other costs")
    header({3: "Line (BORA_Devex row)", 5: "HUF m", 6: "EUR k", 9: "EUR k at FX", 10: "Pay date", 11: "Land"})
    first = cv.r
    lines = [r for r in csv.DictReader(open(OLD / "devex_lines.csv", newline=""))
             if r["source"] == "engine_pasted_20260925" and r["category"] in ("Devex", "Land") and r["native"] != "MODEL"]
    for d in lines + [None, None]:
        r = cv.r
        if d:
            task, item = d["task"].strip(), d["item"].strip()
            ws.cell(r, 3, f"r{d['devex_row']} " + (f"{task} – " if task and task != item else "") + item)
        vals = (float(d["huf_native"] or 0) / 1e6, float(d["eur_native"] or 0) / 1e3,
                D.fromisoformat(d["pay_date"]), int(d["category"] == "Land")) if d else (None,) * 4
        for c, v, fmt in zip((5, 6, 10, 11), vals, ("money2", "money2", "date", "flag")):
            apply(ws.cell(r, c, v), "input", fmt)
        apply(ws.cell(r, 9, f"=E{r}*Thousand/In_FX+F{r}"), "calc", "money2")
        cv.r += 1
    for name, col in (("Devex_Keur", "I"), ("Devex_Date", "J"), ("Devex_Land", "K")):
        m.define(name, f"Curves!${col}${first}:${col}${cv.r - 1}")


# =========================================================================================== calc
def fmt(key, col, f):
    PATCH[(key, col)] = f


def calc(m):
    c = m.calc
    lev = lambda k: c.lever(f"Dash_{k}")  # noqa: E731

    c.section("Timing and flags")
    c.event("exit", "RtB exit", lev("Exit_Date"))
    share = "=MAX(0,MIN({self.F},{end})-MAX({self.E},{start})+1)/({self.F}-{self.E}+1)"
    c.row("constr", "Construction share (NTP → COD)", "factor", share, e="In_NTP+1", f="In_COD", fmt="factor",
          ef_fmt="date")
    c.event("cod", "COD", "In_COD")
    c.window("ops", "Operations", "In_COD+1", "EDATE(In_COD,In_Life*Months_In_Year)")
    c.window("lease_w", "Land lease period (NTP → end of life)", "In_NTP+1", "EDATE(In_COD,In_Life*Months_In_Year)")
    c.window("tenor", "Debt repayment period", "In_COD+1", "EDATE(In_COD,In_Debt_Tenor*Months_In_Year)")
    c.row("dep", "Tax depreciation share", "factor", share, e="In_COD+1",
          f="EDATE(In_COD,In_Dep_Years*Months_In_Year)", fmt="factor", ef_fmt="date")
    c.window("freeze", "Interest freeze (no accrual)", "In_Freeze_Start+1",
             "In_Freeze_Start+In_Freeze_On*(In_Freeze_End-In_Freeze_Start)-1")
    c.row("q4", "31-Dec capitalisation", "flag", "=N(MONTH({end})=Months_In_Year)", total="none", fmt="flag")
    c.row("df_cod", "Buyer discount factor to COD", "factor",
          "=1/(1+In_Buyer_Rate)^(MAX(0,{end}-In_COD)/Days_In_Year)", total="none")

    c.section("Capacity (E: committed, F: built)")
    c.row("wind", "Wind", "MW", e=f"In_MVA*(1+{lev('Overbuild')})", f="In_MVA+({self.E}-In_MVA)*In_Overbuild_Built",
          ef_fmt="money1")
    c.row("bess", "BESS power", "MW", e=f"{lev('BESS')}*In_MVA", f="{self.E}*In_BESS_Built", ef_fmt="money1")
    c.row("mwh", "BESS energy", "MWh", e="{bess.E}*" + lev("BESS_Hours"), f="{self.E}*In_BESS_Built",
          ef_fmt="money1")
    c.row("loss", "Grid-bottleneck loss on built wind", "%", e="({wind.F}-In_MVA)/{wind.F}*In_Bottleneck_Loss",
          ef_fmt="0.000%")

    c.section("Tender points (E: commitment, F: weighted points)")
    band = lambda k: (f"INDEX(Osz{k}_Pts,1+SUMPRODUCT((Osz{k}_LB<{{self.E}})+(Osz{k}_LB={{self.E}})"  # noqa: E731
                      f"*Osz{k}_Incl))*Osz{k}_W")
    yes = lambda k, w: f"{{self.E}}*Osz{k}_Pts*Osz{w}_W"  # noqa: E731
    c.row("osz1", "ÖSZ-1 municipal one-off", "EUR k/MW", e=lev("Municipal"), f=band(1), ef_fmt="money2")
    c.row("osz2", "ÖSZ-2 community contribution", "EUR/MWh", e=lev("Community"), f=band(2), ef_fmt="money2")
    c.row("osz3", "ÖSZ-3 overbuild", "%", e=lev("Overbuild"), f=band(3), ef_fmt="money2")
    c.row("osz4", "ÖSZ-4 BESS ratio", "%", e=lev("BESS") + "/IF(In_OSZ4_Basis=1,{wind.E}/In_MVA,1)",
          f=band(4) + "*({self.E}>=Osz4_Min)", ef_fmt="money2")
    c.row("osz5", "ÖSZ-5 partial award", "flag", e=lev("Partial"), f=yes(5, 5), ef_fmt="money2")
    c.row("osz6", "ÖSZ-6 performance security multiple", "x", e=lev("PB_Mult"), f=band(6), ef_fmt="money2")
    c.row("osz7", "ÖSZ-7 brownfield", "flag", e=lev("Brownfield"), f=yes(7, 7), ef_fmt="money2")
    c.row("osz8", "ÖSZ-8 shared connection", "flag", e=lev("Shared"), f=yes(8, 8), ef_fmt="money2")
    c.row("osz9a", "ÖSZ-9 EU turbines", "flag", e=lev("EU_Turbine"), f=yes("9a", 9), ef_fmt="money2")
    c.row("osz9b", "ÖSZ-9 EU transformers", "flag", e=lev("EU_Transformer"), f=yes("9b", 9), ef_fmt="money2")
    for k, f in (("osz3", "pct"), ("osz4", "pct"), ("osz5", "flag"), ("osz6", "mult"), ("osz7", "flag"),
                 ("osz8", "flag"), ("osz9a", "flag"), ("osz9b", "flag")):
        fmt(k, "E", f)
    c.row("pts", "Points total (E: EU route, F: CN route)", "points",
          e="ROUND(SUM({osz1.F}:{osz9b.F}),Points_Decimals)", f="ROUND(SUM({osz1.F}:{osz8.F}),Points_Decimals)",
          ef_fmt="money2", bold=True)
    c.row("pts_live", "Points, live turbine route", "points", e="IF(In_CN_Turbine=1,{pts.F},{pts.E})", ef_fmt="money2")
    c.row("valid", "Failed tests (E: capacity, F: offerable range)", "#",
          e="N(In_MVA>In_Group_Cap)+N({wind.E}<In_MVA)+N({bess.E}>{wind.E})+N({mwh.E}<{bess.E}*In_Min_Hours)",
          f="({osz1.E}<>0)*({osz1.E}<Osz1_Min)+({osz1.E}>Osz1_Max)+({osz2.E}<>0)*({osz2.E}<Osz2_Min)"
            "+({osz2.E}>Osz2_Max)", ef_fmt="int")

    c.section("Securities, fees and development costs (E: amount, F: date)")
    c.row("pb_base", "Security base", "mHUF", e="{wind.E}*In_Bench_Wind+{mwh.E}*In_Bench_BESS", ef_fmt="money")
    line = "=-{self.E}*" + EVENT
    c.row("inv", "Signing fee and DD", "EUR k", line, e="In_Signing_Fee+In_DD", f="In_Invest_Date")
    c.row("fee", "Participation fee (cash)", "EUR k", line, e="In_Part_Fee*Thousand/In_FX", f="In_Bid_Date")
    c.row("bid", "Bid security outstanding (guarantee)", "EUR k",
          "={self.E}*MAX(0,(MIN(In_Award_Date+In_PB_Days-1,{end})-MAX({self.F},{start})+1)/{days})",
          e="In_Bid_Rate*{pb_base.E}*Thousand/In_FX", f="In_Bid_Date", total="none")
    c.row("pb", "Performance security outstanding (guarantee)", "EUR k",
          "={self.E}*MAX(0,(MIN(In_PB_Release,{end})-MAX({self.F},{start})+1)/{days})",
          e="In_PB_Rate*{osz6.E}*{pb_base.E}*Thousand/In_FX", f="In_Award_Date+In_PB_Days", total="none")
    c.row("tso", "TSO connection-request security (cash)", "EUR k", line,
          e="In_TSO_Security*In_MVA*Thousand/In_FX", f="In_TSO_Date")
    c.row("direct", "Direct connection fee (cash)", "EUR k", line, e="In_Direct_Fee*Thousand/In_FX", f="In_HCSSZ_Date")
    c.row("indirect", "Indirect connection fee, net of TSO security", "EUR k", line,
          e="(In_Indirect_Fee-In_TSO_Security)*In_MVA*Thousand/In_FX", f="In_HCSSZ_Date")
    for k in ("inv", "fee", "bid", "pb", "tso", "direct", "indirect"):
        fmt(k, "F", "date")
    c.row("conn", "Connection fees", "EUR k", "={direct}+{indirect}")
    dvx = '=-SUMIFS(Devex_Keur,Devex_Date,">="&{start},Devex_Date,"<="&{end},Devex_Land,%d)'
    c.row("devex", "Devex (FEP-borne)", "EUR k", dvx % 0)
    c.row("land", "Land (E: paid to the exit)", "EUR k", dvx % 1, e="-" + EXIT)
    c.row("cont", "Devex contingency", "EUR k", line, e="In_Devex_Contingency*(SUM(Devex_Keur)+{fee.E})",
          f="In_Contingency_Date")
    fmt("cont", "F", "date")
    c.row("dev_cost", "Devex, land and contingency (E: to the exit, F: to COD)", "EUR k", "={devex}+{land}+{cont}",
          e="-" + EXIT, f="-" + TO_COD)
    c.row("dev_year", "Devex, land and contingency by calendar year (at Q4)", "EUR k",
          "={q4}*SUMIF({year.todate},{year},{dev_cost.todate})")
    c.row("gfee", "Guarantee fees (E: to the exit, F: to COD)", "EUR k",
          "=-({bid}+{pb})*In_Guarantee_Fee*{days}/Days_In_Year", e=EXIT, f=TO_COD)
    c.row("dev_cf", "Development cash flow (excl. connection fees)", "EUR k", "={inv}+{fee}+{tso}+{dev_cost}+{gfee}",
          bold=True)
    c.row("cash_cum", "Cash posted incl. land, cumulative (E: at the exit)", "EUR k",
          "={self@prev}-{fee}-{tso}-{conn}-{land}", e="SUMIF({end.all},{exit.E},{self.all})", total="seed")

    c.section("FEP claims, devex-file method, accrued to COD (E: to the exit, F: to COD)")
    acc = ("=In_{r}/Days_In_Year*(1-{{freeze}})*({{{p}_open}}*{{days}}+({{end}}+1)*{{{p}_new}}-{{{p}_dated}})"
           "*({{end}}<=In_COD)")
    cap = "={{q4}}*SUMIF({{year.todate}},{{year}},{{{p}_ret.todate}})"
    c.subheader("Cash 12 %: participation fee, TSO security, connection fees")
    c.row("a_new", "New principal (E: recoverable, to the exit)", "EUR k", "=-({fee}+{tso}+{conn})",
          e=EXIT + "+{fee.I}")
    c.row("a_dated", "Date-weighted new principal", "#",
          "=-({fee}*{fee.F}+{tso}*{tso.F}+{direct}*{direct.F}+{indirect}*{indirect.F})", total="none")
    c.row("a_open", "Opening balance", "EUR k", "={a_close@prev}", total="none")
    c.row("a_ret", "Return accrued", "EUR k", acc.format(r="Cash_Rate", p="a"), e=EXIT, f=TO_COD)
    c.row("a_cap", "Capitalised at 31-Dec", "EUR k", cap.format(p="a"))
    c.row("a_close", "Closing balance", "EUR k", "={a_open}+{a_new}+{a_cap}", total="seed")
    c.subheader("Guarantee 6 %: bid security and performance security face value")
    c.row("b_open", "Opening capitalised return", "EUR k", "={b_close@prev}", total="none")
    c.row("b_ret", "Return accrued", "EUR k",
          "=In_Guar_Credit*(1-{freeze})*({b_open}+{bid}+{pb})*{days}/Days_In_Year*({end}<=In_COD)", e=EXIT, f=TO_COD)
    c.row("b_cap", "Capitalised at 31-Dec", "EUR k", cap.format(p="b"))
    c.row("b_close", "Closing capitalised return", "EUR k", "={b_open}+{b_cap}", total="seed")
    c.subheader("At risk 20 % with 2x floor: devex, land, contingency")
    c.row("c_dated", "Date-weighted new principal", "#",
          "=SUMPRODUCT(Devex_Keur*Devex_Date*(Devex_Date>={start})*(Devex_Date<={end}))-{cont}*{cont.F}",
          total="none")
    c.row("c_open", "Opening balance", "EUR k", "={c_close@prev}", total="none")
    c.row("c_ret", "Return accrued (E, F: with the floor)", "EUR k",
          "=In_AtRisk_Rate/Days_In_Year*(1-{freeze})*({c_open}*{days}-({end}+1)*{dev_cost}-{c_dated})*({end}<=In_COD)",
          e="MAX(" + EXIT + ",(In_Min_Multiple-1)*{dev_cost.E})",
          f="MAX(" + TO_COD + ",(In_Min_Multiple-1)*{dev_cost.F})")
    c.row("c_cap", "Capitalised at 31-Dec", "EUR k", cap.format(p="c"))
    c.row("c_close", "Closing balance", "EUR k", "={c_open}-{dev_cost}+{c_cap}", total="seed")

    c.section("Overbuild and BESS sanction (E: wind / pro rata, F: BESS / full call)")
    c.row("short", "Build shortfall share", "%", e="1-{wind.F}/{wind.E}", f="(1-In_BESS_Built)*({bess.E}>0)",
          ef_fmt="pct")
    c.row("pbp", "Performance security by technology", "EUR k", e="{pb.E}*{wind.E}*In_Bench_Wind/{pb_base.E}",
          f="{pb.E}-{self.E}")
    c.row("sanc", "Sanction amount", "EUR k",
          e="({short.E}>In_Tolerance_Cap)*{short.E}*{pbp.E}+({short.F}>In_Tolerance_Cap)*{short.F}*{pbp.F}",
          f="({short.E}+{short.F}>0)*{pb.E}")
    c.row("sanction", "Sanction paid, selected regime (E: date)", "EUR k",
          "=-CHOOSE(In_Sanction_Regime,{sanc.E},{sanc.F})*N(AND({start}<={self.E},{end}>={self.E}))",
          e="EDATE(In_COD,In_Sanction_Lag)")
    fmt("sanction", "E", "date")
    c.row("mva_risk", "Connection capacity lost pro rata", "MVA", e="({short.E}>In_Tolerance_Cap)*{short.E}*In_MVA",
          ef_fmt="money1")
    c.row("excl", "3-year group exclusion possible (if the PB loss counts as a call)", "flag", e="N({sanction.I}<0)", ef_fmt="flag")

    c.section("Build and debt")
    c.row("cn", "Chinese equipment bought (1 = yes)", "flag", e="MAX(In_CN_Turbine,1-In_Delivered)", ef_fmt="flag")
    c.row("tur", "Wind capex per built MW (E: turbine bought; F: other)", "EUR k/MW",
          e="IF({cn.E}=1,In_Turbine_CN,In_Turbine)",
          f="In_Installation+In_Foundations+In_BoP+In_Owner+In_Easements")
    c.row("subst", "Substation and 132 kV bay, as bought", "EUR k",
          e="IF({cn.E}=1,In_Substation_CN,In_Substation)")
    c.row("capex", "Capex on built MW (E: total)", "EUR k", "=-{self.E}*{constr}",
          e="(({tur.E}+{tur.F})*{wind.F}+{subst.E})*(1+In_Capex_Contingency)+In_BESS_Capex*{mwh.F}")
    c.row("municipal", "Municipal one-off at COD", "EUR k", "=-{self.E}*{cod}", e="{osz1.E}*In_MVA*In_Delivered")
    c.subheader("Senior debt")
    c.row("d_open", "Opening balance", "EUR k", "={d_close@prev}", total="none")
    c.row("d_draw", "Drawdown (E: debt size)", "EUR k", "={self.E}*{constr}",
          e="-IF(In_CN_Turbine=1,In_Gearing_CN,In_Gearing)*({capex.I}+{conn.I}+{municipal.I})")
    c.row("d_idc", "Interest capitalised (F: rate per period)", "EUR k", "={d_open}*{self.F}*({end}<=In_COD)",
          f="In_Debt_Rate*Months_Per_Period/Months_In_Year")
    fmt("d_idc", "F", "0.000%")
    c.row("d_rep", "Repayment (level annuity)", "EUR k",
          "=PPMT({d_idc.F},1,In_Debt_Tenor*Months_In_Year/Months_Per_Period-SUM({tenor.todate})+1,{d_open})*{tenor}")
    c.row("d_close", "Closing balance", "EUR k", "=SUM({d_open}:{d_rep})", total="seed", bold=True)
    c.row("d_int", "Interest paid", "EUR k", "=-{d_open}*{d_idc.F}*({end}>In_COD)")
    c.row("d_up", "Upfront fee", "EUR k", "=-In_Upfront_Fee*{d_draw.E}*({d_open}=0)*({d_draw}>0)")
    c.row("d_cf", "Commitment fee", "EUR k",
          "=-In_Commitment_Fee*({d_draw.E}-SUM({d_draw.todate})+{d_draw})*{yf}*({constr}>0)")
    c.row("build", "Build cost (capex, connection, municipal, debt fees)", "EUR k",
          "={capex}+{conn}+{municipal}+{d_up}+{d_cf}", bold=True)

    c.section("Operations")
    c.row("opyr", "Operating year", "#", "=ROUNDUP(YEARFRAC(In_COD,{end}),0)*({end}>In_COD)", total="none",
          fmt="int")
    c.row("esc", "Opex escalation index", "factor", "=(1+In_Opex_Esc)^({year}-YEAR(In_COD))", total="none")
    c.row("price", "PPA price", "EUR/MWh", "=In_PPA_Price*(1+In_PPA_Esc)^({year}-YEAR(In_COD))", total="none")
    c.row("gen", "Net generation", "MWh",
          "={wind.F}*In_P50_Hours*(1-In_Curtailment)*(1-{loss.E})*(1-In_Degradation)^({opyr}-1)*{ops}*{yf}")
    c.row("rev", "Revenue", "EUR k", "={gen}*{price}/Thousand")
    c.row("opex_var", "Variable opex (balancing, grid)", "EUR k", "=-In_Var_Opex*{gen}*{esc}/Thousand")
    c.row("opex_fix", "Fixed opex and administration", "EUR k", "=-(In_Fixed_Opex*{wind.F}+In_Admin)*{esc}*{ops}*{yf}")
    c.row("lease", "Land lease", "EUR k", "=-In_Lease*{wind.F}*{esc}*{lease_w}*{yf}")
    c.row("levies", "MEKH fee and HIPA", "EUR k", "=-{rev}*(In_MEKH+In_HIPA)")
    c.row("opex_cont", "Opex contingency", "EUR k", "=In_Opex_Contingency*({opex_var}+{opex_fix}+{lease}-{rev}*In_MEKH)")
    c.row("community", "Community contribution (ÖSZ-2)", "EUR k", "=-{osz2.E}*In_Delivered*{gen}/Thousand")
    c.row("ebitda", "EBITDA", "EUR k", "=SUM({rev}:{community})", bold=True)

    c.section("Tax")
    c.row("tax_dep", "Tax depreciation (E: base)", "EUR k", "=-{self.E}*{dep}",
          e="{d_idc.I}-({capex.I}+{conn.I}+{municipal.I}+{tso.I})")
    c.row("taxable", "Taxable income", "EUR k", "={ebitda}+{tax_dep}+{d_int}+{d_up}+{d_cf}+{dev_cost}")
    c.row("lcf_open", "Tax losses: opening", "EUR k", "={lcf_close@prev}", total="none")
    c.row("lcf_add", "Losses added", "EUR k", "=MAX(0,-{taxable})")
    c.row("lcf_use", "Losses used", "EUR k", "=-MIN({lcf_open},MAX(0,{taxable}))")
    c.row("lcf_close", "Tax losses: closing", "EUR k", "=SUM({lcf_open}:{lcf_use})", total="seed")
    c.row("tax", "Corporate income tax", "EUR k", "=-In_CIT*(MAX(0,{taxable})+{lcf_use})", bold=True)

    c.section("RtB exit")
    c.row("eff", "RtB price", "EUR k", e="In_RtB_Premium*In_MVA", bold=True)
    c.subheader("Exit waterfall (E: due, F: paid)")
    c.row("w_avail", "Cash available (E: buyer reimbursement)", "EUR k", e="{a_new.E}+{land.E}",
          f="MAX(0,{eff.E}+{self.E})")
    waterfall(c, "w", [("cost", "Sale costs", "In_Sale_Costs"), ("prin", "Recoverable principal", "{a_new.E}"),
                       ("r12", "12 % return", "{a_ret.E}"), ("land", "Land principal", "{land.E}"),
                       ("r20", "At-risk return", "{c_ret.E}"), ("r6", "6 % guarantee return", "{b_ret.E}"),
                       ("dsa", "Developer success fee (DSA)",
                        "MIN(In_DSA_Cap,In_DSA_Rate*YEARFRAC(In_Signing_Date,{exit.E},0))*{wind.E}")])
    c.row("w_fep", "FEP receipts at the exit", "EUR k", f="SUM({w_prin.F}:{w_r6.F})+In_FEP_Share*{w_dist.F}",
          bold=True)
    c.row("w_dev", "Developer: DSA and share", "EUR k", f="{w_dsa.F}+(1-In_FEP_Share)*{w_dist.F}")
    c.subheader("Exit split, devex-file view (paid amounts)")
    c.row("ex_fep", "FEP (E: principal returned, F: 70 % of distributable)", "EUR k",
          e="{w_prin.F}+{w_land.F}", f="In_FEP_Share*{w_dist.F}")
    c.row("ex_dev", "FEP-borne devex, contingency and participation fee", "EUR k",
          e="{land.E}-{dev_cost.E}-{fee.E}")
    c.row("ex_claims", "FEP interest claims (E) and FEP net (F: claims + 70 % - devex)", "EUR k",
          e="{w_r12.F}+{w_r20.F}+{w_r6.F}", f="{self.E}+{ex_fep.F}+{ex_dev.E}", bold=True)
    c.row("ex_ares", "A.RES 30 % of distributable", "EUR k", e="(1-In_FEP_Share)*{w_dist.F}")
    c.row("ex_mw", "Per installed wind MW (E: FEP net, F: A.RES total)", "EUR k/MW", e="{ex_claims.F}/{wind.E}",
          f="{w_dev.F}/{wind.E}", ef_fmt="money1")
    c.subheader("FEP returns – RtB route")
    c.row("rtb_pre", "FEP pre-tax cash flow (E: tax at the exit)", "EUR k",
          "=({dev_cf}+{conn})*({end}<={exit.E})+{w_fep.F}*{exit}",
          e="-In_CIT*MAX(0,{self.I}-In_FEP_Share*{w_dist.F})")
    c.row("rtb_cf", "FEP cash flow – RtB route", "EUR k", "={rtb_pre}+{rtb_pre.E}*{exit}", bold=True)
    c.row("ex_route", "FEP route: guarantee fees, signing and DD, tax", "EUR k", e="{gfee.E}+{inv.I}+{rtb_pre.E}")
    c.xirr("rtb_irr", "FEP IRR – RtB route", "rtb_cf", iferror=True)
    c.row("rtb_moic", "FEP MOIC – RtB route (receipts after tax ÷ outlays)", "x",
          e="({w_fep.F}+{rtb_pre.E})/({w_fep.F}-{rtb_pre.I})", ef_fmt="mult", bold=True)
    c.xnpv("rtb_npv", "FEP NPV – RtB route", "In_Discount_Rate", "rtb_cf")

    c.section("Sale at COD")
    c.row("cod_contrib", "FEP contributions to COD", "EUR k", "=({dev_cf}+{build}+{d_draw}+{ebitda})*({end}<=In_COD)")
    c.row("cod_com", "Buyer: community contribution vs reference (E: PV at COD)", "EUR k",
          "=-({osz2.E}*In_Delivered-In_Ref_Community)*{gen}*(1-In_CIT)/Thousand", e="SUMPRODUCT({self.all},{df_cod.all})")
    c.row("v_price", "COD price (E: gross, F: price effect)", "EUR k",
          e="IF(In_CN_Turbine=1,In_COD_Price_CN,In_COD_Price)*{wind.F}*(1-{loss.E})",
          f="In_Pass_Through*({cod_com.E}+SUMPRODUCT({sanction.all},{df_cod.all}))")
    c.subheader("JV waterfall at COD (E: due, F: paid)")
    c.row("v_avail", "Cash available (E: debt repaid)", "EUR k", e="{d_draw.I}+{d_idc.I}",
          f="MAX(0,{v_price.E}+{v_price.F}-{self.E})")
    waterfall(c, "v", [("cost", "Sale costs", "In_Sale_Costs"),
                       ("fund", "FEP funding returned", "-{cod_contrib.I}+{inv.I}+{devex.I}+{cont.I}+{fee.I}+{gfee.F}"),
                       ("claims", "FEP claims to COD", "{a_ret.F}+{b_ret.F}+{c_ret.F}"),
                       ("dsa", "Developer success fee (DSA)",
                        "MIN(In_DSA_Cap,In_DSA_Rate*YEARFRAC(In_Signing_Date,In_COD,0))*{wind.E}")])
    c.row("v_fep", "FEP proceeds at COD", "EUR k", f="{v_fund.F}+{v_claims.F}+In_FEP_Share*{v_dist.F}", bold=True)
    c.row("cod_cf", "FEP cash flow – sale at COD", "EUR k", "={cod_contrib}+{v_fep.F}*{cod}", bold=True)
    c.xirr("cod_irr", "FEP IRR – sale at COD", "cod_cf")
    c.row("cod_moic", "FEP MOIC – sale at COD (proceeds ÷ contributions)", "x", e="{v_fep.F}/-{cod_contrib.I}",
          ef_fmt="mult", bold=True)
    c.xnpv("cod_npv", "FEP NPV – sale at COD", "In_Discount_Rate", "cod_cf")
    c.subheader("Capex breakdown to COD (two items per row)")
    c.row("cx_a", "Turbines (E), installation (F)", "EUR k", e="{tur.E}*{wind.F}", f="In_Installation*{wind.F}")
    c.row("cx_b", "Foundations (E), balance of plant (F)", "EUR k", e="In_Foundations*{wind.F}", f="In_BoP*{wind.F}")
    c.row("cx_c", "Owner's costs (E), land, easements and pre-COD lease (F)", "EUR k", e="In_Owner*{wind.F}",
          f='In_Easements*{wind.F}-SUMIF({end.all},"<="&In_COD,{ebitda.all})')
    c.row("cx_d", "Contingency (E), substation and 132 kV bay (F)", "EUR k",
          e="In_Capex_Contingency*(({tur.E}+{tur.F})*{wind.F}+{subst.E})", f="{subst.E}")
    c.row("cx_e", "BESS (E), connection fees (F)", "EUR k", e="In_BESS_Capex*{mwh.F}", f="-{conn.I}")
    c.row("cx_f", "Municipal one-off (E), development: devex, land, signing, securities (F)", "EUR k",
          e="-{municipal.I}", f="-({dev_cost.I}+{inv.I}+{fee.I}+{tso.I})")
    c.row("cx_g", "Financing: IDC, debt fees, guarantee fees", "EUR k", e="{d_idc.I}-{d_up.I}-{d_cf.I}-{gfee.F}")
    c.row("c2c", "All-in cost to COD (E: uses, F: sources)", "EUR k", e="SUM({cx_a.E}:{cx_g.F})",
          f="{d_draw.I}+{d_idc.I}-{cod_contrib.I}", bold=True)
    c.row("fn_con", "Funding: construction equity (E: capex, lease less debt), debt fees (F)", "EUR k",
          e='-({capex.I}+{d_draw.I})-SUMIF({end.all},"<="&In_COD,{ebitda.all})', f="-({d_up.I}+{d_cf.I})")
    c.subheader("Commitment")
    c.row("ov_sec", "Securities: cash posted and guarantee outstanding", "EUR k", e="{pb.E}-{fee.I}-{tso.I}")
    c.row("ov_tot", "Total commitment: all-in cost to COD and guarantee", "EUR k", e="{c2c.E}+{pb.E}", bold=True)
    c.row("ov_pen", "Penalty if overbuild not built, a call on the PB (F: full call)", "EUR k", e="-{sanction.I}",
          f="{sanc.F}")

    c.section("Hold to maturity")
    c.row("htm_cf", "FEP equity cash flow – hold to maturity", "EUR k",
          "={dev_cf}+{build}+{d_draw}+{ebitda}+{d_int}+{d_rep}+{tax}+{sanction}-{w_dev.F}*{exit}", bold=True)
    c.xirr("htm_irr", "FEP IRR – hold to maturity", "htm_cf")
    c.moic("htm_moic", "FEP MOIC – hold to maturity", "htm_cf")
    c.xnpv("htm_npv", "FEP NPV – hold to maturity", "In_Discount_Rate", "htm_cf")
    c.row("dscr", "DSCR (E: minimum)", "x", '=IF({d_int}+{d_rep}<0,({ebitda}+{tax})/-({d_int}+{d_rep}),"")',
          e="MIN({self.all})", total="none", fmt="mult", ef_fmt="mult")


def waterfall(c, p, steps):
    """One row per step: E = amount due, F = MIN(due, cash still available); distributable = cash above all dues."""
    first = f"{p}_{steps[0][0]}"
    prev = None
    for k, label, due in steps:
        rest = "{%s_avail.F}" % p + ("" if prev is None else "-SUM({%s.F}:{%s.F})" % (first, prev))
        c.row(f"{p}_{k}", label, "EUR k", e=due, f="MIN({self.E}," + rest + ")")
        prev = f"{p}_{k}"
    c.row(f"{p}_dist", "Distributable", "EUR k",
          f="MAX(0,{%s_avail.F}-SUM({%s.E}:{%s.E}))" % (p, first, prev), bold=True)


# =========================================================================================== checks
def checks(m):
    k = m.checks
    k.add_check("Levers within the tender rules", "=N({valid.E}+{valid.F}>0)", per_period=False)
    k.add_check("Exit and COD fall on a quarter end",
                "=N(COUNTIF({end.all},{exit.E})=0)+N(COUNTIF({end.all},In_COD)=0)", per_period=False)
    k.add_check("Cash posted to the exit = buyer reimbursement + participation fee",
                "=N(ABS({cash_cum.E}-{w_avail.E}-{fee.E})>Tolerance)", per_period=False)
    k.add_check("RtB waterfall: paid + distributable = cash available",
                "=N(ABS({w_fep.F}+{w_dev.F}+{w_cost.F}-{w_avail.F})>Tolerance)", per_period=False)
    k.add_check("FEP net (devex-file view) + route costs = FEP RtB profit",
                "=N(ABS({ex_claims.F}+{ex_route.E}-{rtb_cf.I})>Tolerance)", per_period=False)
    k.add_check("RtB waterfall: no step paid above its amount due",
                "=N(SUMPRODUCT(--({w_cost.F}:{w_dsa.F}>{w_cost.E}:{w_dsa.E}+Tolerance))>0)", per_period=False)
    k.add_check("COD waterfall: distributable not negative, paid within cash available",
                "=N({v_dist.F}<-Tolerance)+N(SUM({v_cost.F}:{v_dsa.F})>{v_avail.F}+Tolerance)", per_period=False)
    k.add_check("Sources = uses to COD (equity + debt vs all-in cost lines)", "=N(ABS({c2c.E}-{c2c.F})>Tolerance)",
                per_period=False)
    k.add_check("Funding table: cumulative cash = FEP contributions to COD (both bids)",
                "=N(ABS(Funding_Cum_Last-Out_Funding_Equity_To_COD)>Tolerance)", per_period=False)
    k.add_check("Capex breakdown sums to the all-in cost to COD (both bids)",
                "=N(SUM(Capex_Items)-Out_Capex_Total>Tolerance)+N(Out_Capex_Total-SUM(Capex_Items)>Tolerance)", per_period=False)
    k.add_check("Debt repaid by the end of the tenor", "=N(({end}>={tenor.F})*ABS({d_close})>Tolerance)")
    k.add_check("XIRR returned a number on every route",
                "=N(ISTEXT({rtb_irr.E}))+N(ISERROR({htm_irr.E}))+N(ISERROR({cod_irr.E}))", per_period=False)
    for b in (1, 2):
        k.add_check(f"Stored zero bid + Δ = live, base set – bid {b}",
                    f"=(Scenario_Selected=1)*(SUMPRODUCT(ABS(ZB_Live_{b}-ZB_Zero_Stored-ZB_Delta_{b}_Stored))>Tolerance)",
                    per_period=False)


# =========================================================================================== dashboard
LEVERS = [  # key, label, unit, BORA 1 / BORA 2, offerable range, points row, list options, format
    # BORA 2 = BORA 1 strategy (User 13-Oct-26: BORA 1 strategy, SPEC §0 item 21)
    ("Municipal", "Municipal one-off (ÖSZ-1)", "EUR k/MW", [50, 50], "0 or 10–50", "osz1", None, "money"),
    ("Community", "Community contribution (ÖSZ-2)", "EUR/MWh", [3, 3], "0 or 2–3", "osz2", None, "money2"),
    ("Overbuild", "Overbuild (ÖSZ-3)", "%", [0.51, 0.51], "≥ 0 %", "osz3", None, "pct"),
    ("BESS", "BESS power, share of MVA (ÖSZ-4)", "%", [0.51, 0.51], "≤ wind MW", "osz4", None, "pct"),
    ("BESS_Hours", "BESS duration", "h", [4, 4], "≥ 4 h", None, None, "money1"),
    ("PB_Mult", "Performance security multiple (ÖSZ-6)", "x", [1.0, 1.0], "≥ 1.0", "osz6", None, "mult"),
    ("Partial", "Partial award accepted (ÖSZ-5)", "flag", [1, 1], "0 / 1", "osz5", [0, 1], "flag"),
    ("Brownfield", "Brownfield site (ÖSZ-7)", "flag", [0, 0], "0 / 1", "osz7", [0, 1], "flag"),
    ("Shared", "Shared connection (ÖSZ-8)", "flag", [1, 1], "0 / 1", "osz8", [0, 1], "flag"),
    ("EU_Turbine", "EU turbines ≥ 80 % (ÖSZ-9)", "flag", [1, 1], "0 / 1", "osz9a", [0, 1], "flag"),
    ("EU_Transformer", "EU transformers ≥ 80 % (ÖSZ-9)", "flag", [1, 1], "0 / 1", "osz9b", [0, 1], "flag"),
]
IMPACT = [("Δ RtB profit €m", "money2"), ("Δ COD IRR", "0.00%"), ("Δ HTM IRR", "0.00%"),
          ("Δ securities €m", "money2")]   # macro columns, in BoraImpact.Metric order
GLANCE = [  # label, unit, calc template, scale, out name
    ("Points, live turbine route", "points", "{pts_live.E}", None, "Points"),
    ("FEP profit – RtB route", "EUR m", "{rtb_cf.I}", "Thousand", "RtB_Profit"),
    ("FEP IRR – RtB route", "%", "{rtb_irr.E}", None, "RtB_IRR"),
    ("FEP IRR – sale at COD", "%", "{cod_irr.E}", None, "COD_IRR"),        # check covers this row and below
    ("FEP NPV – sale at COD", "EUR m", "{cod_npv.E}", "Thousand", "COD_NPV"),
    ("FEP IRR – hold to maturity", "%", "{htm_irr.E}", None, "HTM_IRR"),
    ("FEP NPV – hold to maturity", "EUR m", "{htm_npv.E}", "Thousand", "HTM_NPV"),
    ("Total commitment", "EUR m", "{ov_tot.E}", "Thousand", "Commitment_Total"),
]


def dashboard(m):
    d, c, S = m.dashboard, m.calc, m.calc_sheets
    nb = len(S)
    bids = [f"B{n}" for n in range(1, nb + 1)]
    m2 = "money2"

    def grid(r1, r2, c1, n=1):
        return f"Dashboard!${L(d.X + c1 - 1)}${r1}:${L(d.X + c1 + n - 2)}${r2}"

    def col(r1, r2, b):
        return f"Dashboard!${L(d.V + b - 1)}${r1}:${L(d.V + b - 1)}${r2}"

    # 1. the bid: levers, points and the impact of each commitment against the zero bid
    head = ["Offerable range"] + [f"Points {b}" for b in bids] + [f"{h} {b}" for h, _ in IMPACT for b in bids]
    d.section("The bid: commitments, points and impact vs the zero bid (macro)", head)
    for key, label, unit, vals, rng, pkey, opts, f in LEVERS:
        pts = [f"={c.ref(pkey, 'F', s)}" for s in S] if pkey else []
        d.input(key, label, unit, vals, fmt=f, options=opts, extra=[rng] + pts)
    d.stored_block("Impact", LEVERS[0][0], LEVERS[-1][0], range(nb + 2, len(head) + 1))
    for i, (_, f) in enumerate(IMPACT):
        for b in range(nb):
            for r in range(d.rows[LEVERS[0][0]], d.rows[LEVERS[-1][0]] + 1):
                d.ws.cell(r, d.X + nb + 1 + i * nb + b).number_format = FMT.get(f, f)

    # 2. returns at a glance: the bids next to the zero bid (macro, computed on Calc_1)
    d.section("Returns at a glance", ["Zero bid"] + [f"Δ {b}" for b in bids])
    rows = []
    for label, unit, t, scale, out in GLANCE:
        f = "pct" if unit == "%" else m2
        r = d.link(label, unit, t, scale=scale, fmt=f, out=out, bold=out == "Commitment_Total")
        rows.append(r)
        for k in range(1 if out == "Points" else nb + 1):       # no Δ for points (trivial vs a zero bid)
            apply(d.ws.cell(r, d.X + k), "macro", FMT[f])
    r1, rc, r2 = rows[0], rows[3], rows[-1]
    m.define("ZeroBid_Live", f"Dashboard!${L(d.V)}${r1}:${L(d.V + nb - 1)}${r2}")
    m.define("ZeroBid_Stored", grid(r1, r2, 1))
    m.define("ZeroBid_Delta_Stored", grid(rows[1], r2, 2, nb))
    m.define("ZB_Zero_Stored", grid(rc, r2, 1))
    for b in range(1, nb + 1):
        m.define(f"ZB_Live_{b}", col(rc, r2, b))
        m.define(f"ZB_Delta_{b}_Stored", grid(rc, r2, 1 + b))

    # 3. funding need and capex breakdown to COD, both bids together (MODELLING_STANDARDS section 9)
    funding(m, d, S)
    capex_breakdown(m, d)

    d.section("Points and validity")
    d.link("Wind committed", "MW", "{wind.E}", out="Wind_MW")
    d.link("Wind built", "MW", "{wind.F}", out="Wind_Built_MW")
    d.link("BESS committed", "MW", "{bess.E}", out="BESS_MW")
    d.status("Capacity tests (group cap, overbuild, BESS ≤ wind, 4 h)", "{valid.E}")
    d.status("Commitments within the offerable range", "{valid.F}")

    d.section("Securities and fees", ["Form", "Posted", "Released"])
    pb_date = f"={c.ref('pb', 'F', S[0])}"
    for key, label, form, posted, released in [
            ("fee", "Participation fee", "cash", "=In_Bid_Date", "–"),
            ("bid", "Bid security", "guarantee", "=In_Bid_Date", pb_date),
            ("pb", "Performance security", "guarantee", pb_date, "exit / PB release"),
            ("tso", "TSO connection-request security", "cash", "=In_TSO_Date", "credited"),
            ("direct", "Direct connection fee", "cash", "=In_HCSSZ_Date", "–"),
            ("indirect", "Indirect connection fee, net", "cash", "=In_HCSSZ_Date", "–")]:
        r = d.link(label, "EUR m", "{%s.E}" % key, scale="Thousand", fmt="0.000", extra=[form, posted, released])
        for col in (d.X + 1, d.X + 2):
            d.ws.cell(r, col).number_format = FMT["date"]
    d.link("Securities: cash posted and guarantee outstanding", "EUR m", "{ov_sec.E}", scale="Thousand",
           out="Commitment_Securities", fmt=m2)
    d.link("PB call if overbuild not built (selected regime)", "EUR m", "{ov_pen.E}", scale="Thousand",
           out="Commitment_Penalty", fmt=m2)
    d.link("PB call if overbuild not built (full call)", "EUR m", "{ov_pen.F}", scale="Thousand", out="Sanction_Full",
           fmt=m2)
    d.link("Connection capacity at risk", "MVA", "{mva_risk.E}", out="MVA_At_Risk")
    d.link("3-year group exclusion possible (if the PB loss counts as a call)", "flag", "{excl.E}", out="Exclusion")

    d.section("RtB exit, devex-file waterfall (paid amounts)", ["Total", "k€/MW B1", "k€/MW B2"])
    m.define("Period_Ends", f"{S[0]}!$J$4:${L(m.last)}$4")
    d.input("Exit_Date", "RtB exit date (a quarter end)", "date", [D(2029, 12, 31)] * nb, options="Period_Ends")
    tot = lambda r: d.ws.cell(r, d.X, f"=E{r}+F{r}")  # noqa: E731  presentation total of the two bids
    for t, label, out, bold in [
            ("{ex_fep.E}", "Memo: buyer reimbursement (deposits, fees, land) to FEP", "Exit_Reimbursement", False),
            ("{eff.E}", "Exit premium (RtB price)", "RtB_Price", True),
            ("-{w_cost.F}", "Less sale costs", "WF_Sale_Costs", False),
            ("-{w_r12.F}", "Less FEP claim 12 % on deposits and connection fees", "WF_Return_12", False),
            ("-{w_r20.F}", "Less FEP claim 20 % / 2x on devex and land", "WF_Return_AtRisk", False),
            ("-{w_r6.F}", "Less FEP claim 6 % on guarantees", "WF_Return_6", False),
            ("-{ex_claims.E}", "Less FEP interest claims, total", "Exit_Claims", False),
            ("-{w_dsa.F}", "Less DSA", "WF_DSA", False),
            ("{w_dist.F}", "Distributable", "Distributable", True),
            ("{ex_fep.F}", "FEP 70 %", "Exit_FEP_70", False),
            ("{ex_dev.E}", "FEP-borne devex principal (deducted from FEP)", "Exit_FEP_Devex", False),
            ("{ex_claims.F}", "FEP net = interest claims + 70 % - devex", "Exit_FEP_Net", True),
            ("{ex_ares.E}", "A.RES 30 %", "Exit_ARES_30", False),
            ("{w_dev.F}", "A.RES total = 30 % + DSA", "Exit_ARES_Total", True),
            ("{ex_route.E}", "FEP route: guarantee fees, signing and DD, tax", "RtB_Route_Costs", False),
            ("{rtb_cf.I}", "FEP profit – RtB route", "RtB_Profit_Buildup", True)]:
        r = d.link(label, "EUR m", t, scale="Thousand", out=out, fmt=m2, bold=bold)
        apply(tot(r), "total" if bold else "calc", FMT[m2])
        if out in ("Exit_FEP_Net", "Exit_ARES_Total"):
            for b, s_ in enumerate(S):
                key = "{ex_mw.E}" if out == "Exit_FEP_Net" else "{ex_mw.F}"
                cell = d.ws.cell(r, d.X + 1 + b, "=" + m.render(key, None, None, "Dashboard", s_))
                apply(cell, "link", FMT["money1"])
                m.define(f"Out_{out}_MW_{b + 1}", f"Dashboard!${L(d.X + 1 + b)}${r}")
    d.link("FEP MOIC – RtB route", "x", "{rtb_moic.E}", out="RtB_MOIC")
    d.link("FEP NPV – RtB route", "EUR m", "{rtb_npv.E}", scale="Thousand", out="RtB_NPV", fmt=m2)

    d.section("Sale at COD")
    d.link("All-in cost to COD", "EUR m", "{c2c.E}", scale="Thousand", out="Cost_To_COD", fmt=m2, bold=True)
    d.link("Sale price (built MW, after bottleneck loss)", "EUR m", "{v_price.E}", scale="Thousand", out="COD_Price",
           fmt=m2)
    d.link("COD price effect (community, sanction)", "EUR m", "{v_price.F}", scale="Thousand",
           out="COD_Price_Effect", fmt=m2)
    d.link("Debt repaid at COD", "EUR m", "{v_avail.E}", scale="Thousand", out="COD_Debt", fmt=m2)
    d.link("FEP proceeds at COD", "EUR m", "{v_fep.F}", scale="Thousand", out="COD_Proceeds", fmt=m2)
    d.link("FEP profit – sale at COD", "EUR m", "{cod_cf.I}", scale="Thousand", out="COD_Profit", fmt=m2)
    d.link("FEP MOIC – sale at COD", "x", "{cod_moic.E}", out="COD_MOIC")

    d.section("Hold to maturity")
    d.link("FEP MOIC – hold to maturity", "x", "{htm_moic.E}", out="HTM_MOIC")
    d.link("Minimum DSCR (memo)", "x", "{dscr.E}", out="Min_DSCR")

    d.scenario_results(outs=["Points", "Distributable", "RtB_Profit", "Cost_To_COD", "COD_IRR", "HTM_IRR",
                                 "HTM_NPV", "Commitment_Penalty"])

    d.panel_section("Key-assumption set")
    d.selector()
    d.panel_section("Key assumptions (live)")
    for label, unit, ref, f in [
            ("Chinese turbines (0 EU, 1 CN)", "flag", "In_CN_Turbine", None),
            ("RtB price", "EUR k/MVA", "In_RtB_Premium", None), ("COD price, EU", "EUR k/MW", "In_COD_Price", None),
            ("Turbine price, EU", "EUR k/MW", "In_Turbine", None), ("PPA price at COD", "EUR/MWh", "In_PPA_Price", "money1"),
            ("Gearing, EU", "%", "In_Gearing", None), ("Debt rate", "%", "In_Debt_Rate", None),
            ("Discount rate", "%", "In_Discount_Rate", None), ("COD", "date", "In_COD", None), ("FX", "HUF/EUR", "In_FX", None),
            ("Interest freeze on", "flag", "In_Freeze_On", None), ("Overbuild built", "flag", "In_Overbuild_Built", None),
            ("BESS built", "flag", "In_BESS_Built", None), ("Commitments delivered", "flag", "In_Delivered", None),
            ("Sanction regime (1 pro rata, 2 full)", "flag", "In_Sanction_Regime", None),
            ("ÖSZ-4 basis (0 MVA, 1 wind)", "flag", "In_OSZ4_Basis", None)]:
        d.panel_link(label, unit, ref, fmt=f)


def funding(m, d, S):
    """Funding need to COD for all bids: date | cash | guarantee | cumulative cash equity, one line per milestone."""
    d.section("Funding need to COD – both bids", ["Guarantee €m", "Cum. cash €m"], values=["Date", "Cash €m"])
    ws, dev = m.dashboard.ws, m.calc.rows["dev_year"].r
    q4 = lambda y: L(10 + ((y - START.year) * 12 + 12 - START.month + 1) // 3 - 1)  # noqa: E731  Q4 column of year y
    devex = [(f"Devex, land and contingency {y}", f"={S[0]}!{q4(y)}$4",
              "=-(" + "+".join(f"{s}!{q4(y)}{dev}" for s in S) + ")/Thousand", None) for y in range(START.year,
                                                                                               NTP.year + 1)]
    cash = lambda t, sign="": f"={sign}({d.across(t)})/Thousand"  # noqa: E731
    lines = [("Signing fee and DD", "=In_Invest_Date", cash("{inv.I}", "-"), None),
             ("Bid: participation fee (cash), bid guarantee", "=In_Bid_Date", cash("{fee.E}"), cash("{bid.E}")),
             ("Award: performance guarantee (bid guarantee released)", f"={m.calc.ref('pb', 'F', S[0])}", None,
              cash("{pb.E}")),
             ("Connection request: TSO security", "=In_TSO_Date", cash("{tso.E}"), None)] + devex[:-1] + [
             ("HCSSZ: connection fees", "=In_HCSSZ_Date", cash("{conn.I}", "-"), None)] + devex[-1:] + [
             ("NTP → COD: construction equity (capex, lease less debt)", "=In_COD", cash("{fn_con.E}"), None),
             ("NTP → COD: debt fees (IDC is debt-funded)", "=In_COD", cash("{fn_con.F}"), None),
             ("Guarantee fees to COD", "=In_COD", cash("{gfee.F}", "-"), None),
             ("COD: municipal one-off", "=In_COD", cash("{municipal.I}", "-"), None)]
    prev = None
    for label, date, c, g in lines:
        r = d._row(None, label, "EUR m")
        apply(ws.cell(r, d.V, date), "link", FMT["date"])
        for col, f in ((d.V + 1, c), (d.X, g)):
            if f:
                apply(ws.cell(r, col, f), "calc", FMT["money2"])
        apply(ws.cell(r, d.X + 1, f"={L(d.V + 1)}{r}" + (f"+{L(d.X + 1)}{prev}" if prev else "")), "calc", FMT["money2"])
        prev = r
    r = d._row(None, "Total cash equity to COD; peak guarantee", "EUR m", bold=True)
    apply(ws.cell(r, d.V + 1, cash("{cod_contrib.I}", "-")), "total", FMT["money2"])
    apply(ws.cell(r, d.X, cash("{pb.E}")), "total", FMT["money2"])
    m.define("Out_Funding_Equity_To_COD", f"Dashboard!${L(d.V + 1)}${r}")
    m.define("Out_Funding_Peak_Guarantee", f"Dashboard!${L(d.X)}${r}")
    m.define("Funding_Cum_Last", f"Dashboard!${L(d.X + 1)}${prev}")


def capex_breakdown(m, d):
    """Capex breakdown to COD for all bids, in EUR m and EUR k per built MW, summing to the all-in cost to COD."""
    d.section("Capex breakdown to COD – both bids", values=["€ m", "k€/MW"])
    ws, mw = d.ws, d.across("{wind.F}")
    rows = []
    for t, label in [("{cx_a.E}", "Turbines"), ("{cx_a.F}", "Installation"), ("{cx_b.E}", "Foundations"),
                     ("{cx_b.F}", "Balance of plant"), ("{cx_c.E}", "Owner's costs"),
                     ("{cx_c.F}", "Land, easements, pre-COD lease"), ("{cx_d.E}", "Contingency"),
                     ("{cx_d.F}", "Substation and 132 kV bay"), ("{cx_e.E}", "BESS"), ("{cx_e.F}", "Connection fees"),
                     ("{cx_f.E}", "Municipal one-off"), ("{cx_f.F}", "Development: devex, land, signing, securities"),
                     ("{cx_g.E}", "Financing: IDC, debt fees, guarantee fees"), ("{c2c.E}", "All-in cost to COD")]:
        total = t == "{c2c.E}"
        r = d._row(None, label, "EUR m", bold=total)
        apply(ws.cell(r, d.V, f"=({d.across(t)})/Thousand"), "total" if total else "calc", FMT["money1"])
        apply(ws.cell(r, d.V + 1, f"=({d.across(t)})/({mw})"), "total" if total else "calc", FMT["money"])
        rows.append(r)
    m.define("Capex_Items", f"Dashboard!${L(d.V)}${rows[0]}:${L(d.V)}${rows[-2]}")
    m.define("Out_Capex_Total", f"Dashboard!${L(d.V)}${rows[-1]}")
    m.define("Out_Capex_Total_MW", f"Dashboard!${L(d.V + 1)}${rows[-1]}")


# =========================================================================================== build
def model():
    m = Model("Project Bora", "Bid strategy – BORA 1 / BORA 2 (2 x 100 MVA, Szombathely)", periods=141,
              periodicity="quarterly", start=START, scenarios=SETS, version="v2", date=D(2026, 10, 12),
              curves=True, layout="dashboard", calc_sheets=["Calc_1", "Calc_2"], bids=["BORA 1", "BORA 2"],
              extra_cols=1 + 2 * (1 + len(IMPACT)), source_col=False)
    inputs(m)
    curves(m)
    calc(m)
    checks(m)
    dashboard(m)
    return m


def build_model(out: Path) -> Path:
    m = model()
    xlsx = m.save(out / f"{NAME}.xlsx")
    for (key, col), f in PATCH.items():          # second format for rows whose E and F hold different units
        for ws in m.calc.sheets:
            ws[f"{col}{m.calc.rows[key].r}"].number_format = FMT.get(f, f)
    m.wb.save(xlsx)
    xlsm = package(xlsx, out / f"{NAME}.xlsm", run=["BoraImpact.RefreshAll"])
    xlsx.unlink()
    return xlsm


def build_deck(xlsm: Path, out: Path) -> Path:
    """IC deck (deck.py next to this file), bound to the packaged, LibreOffice-recalculated workbook."""
    spec = importlib.util.spec_from_file_location("bora_deck", HERE / "deck.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.build_deck(xlsm, out)


def build(out: Path):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    xlsm = build_model(out)
    return {"model": xlsm, "deck": build_deck(xlsm, out)}
