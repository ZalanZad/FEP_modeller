"""Compact text dump of the recalculated Dashboard (one line per row) for the model reviewer.

    python3 projects/bora_bidding/dump_dashboard.py   ->   projects/bora_bidding/out/dashboard_dump.txt
"""
import datetime as dt
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent / "out"


def fmt(v, number_format):
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime("%d-%b-%y")
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v)
    if "%" in number_format:
        return f"{v * 100:.2f}%"
    return f"{v:,.3f}".rstrip("0").rstrip(".") if abs(v) < 100 else f"{v:,.1f}"


def dump(recalc=OUT / "Bora_Bid_Model_recalc.xlsx", out=OUT / "dashboard_dump.txt"):
    ws = openpyxl.load_workbook(recalc, data_only=True)["Dashboard"]
    header = {c.column: c.value for c in ws[9] if c.value}
    panel = next(c for c, h in header.items() if h == "Key assumptions")
    bid1 = header[5]
    cols, lines = {}, [f"{ws['C3'].value} – {ws['C4'].value} | set: {ws['E7'].value} | master check: {ws['E8'].value}"]
    for r in range(11, ws.max_row + 1):
        label, unit = ws.cell(r, 3).value, ws.cell(r, 4).value
        cells = {c: ws.cell(r, c) for c in range(5, panel - 1) if ws.cell(r, c).value is not None}
        if label and ws.cell(r, 5).value == bid1 or (label and not unit and cells and
                                                     all(isinstance(x.value, str) for x in cells.values())):
            cols = {c: x.value for c, x in cells.items()}
            lines.append(f"\n== {label} ==  [{' | '.join(str(v) for v in cols.values())}]")
        elif label and label != "***END***":
            vals = [f"{cols.get(c, c)}={fmt(x.value, x.number_format)}" for c, x in cells.items()]
            lines.append(f"{label}" + (f" ({unit})" if unit else "") + ": " + "; ".join(vals))
    lines.append("\n== Key assumptions (live) ==")
    for r in range(10, ws.max_row + 1):
        if ws.cell(r, panel).value and ws.cell(r, panel + 2).value is not None:
            x = ws.cell(r, panel + 2)
            lines.append(f"{ws.cell(r, panel).value} ({ws.cell(r, panel + 1).value}): {fmt(x.value, x.number_format)}")
    out.write_text("\n".join(lines) + "\n")
    return out


if __name__ == "__main__":
    print(dump())
