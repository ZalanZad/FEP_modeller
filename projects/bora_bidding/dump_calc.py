"""Compact dump of Calc_1 for a logic review: row, label, unit, E / F formulas, first timeline formula (J) and the
recalculated lifetime total (I).

    python3 projects/bora_bidding/dump_calc.py   ->   projects/bora_bidding/out/calc_dump.txt
"""
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent / "out"


def dump(model=OUT / "Bora_Bid_Model.xlsm", recalc=OUT / "Bora_Bid_Model_recalc.xlsx", out=OUT / "calc_dump.txt"):
    f = openpyxl.load_workbook(model)["Calc_1"]
    v = openpyxl.load_workbook(recalc, data_only=True)["Calc_1"]
    lines = ["row | label | unit | E | F | J (first period) | I (total)"]
    for r in range(3, f.max_row + 1):
        label = f.cell(r, 3).value
        if not label:
            continue
        total = v.cell(r, 9).value
        total = f"{total:,.1f}" if isinstance(total, (int, float)) else str(total or "")
        cells = [f.cell(r, c).value for c in (5, 6, 10)]
        lines.append(" | ".join([str(r), label, f.cell(r, 4).value or ""] + [str(x) if x is not None else "" for x in cells]
                                + [total]))
    out.write_text("\n".join(lines) + "\n")
    return out


if __name__ == "__main__":
    print(dump())
