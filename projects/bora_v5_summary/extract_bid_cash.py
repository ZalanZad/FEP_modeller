"""Standalone extract in the layout of the user's bid sheet (screenshot of 27-Oct-2026): Bora's two bids beside
Iváncsa and Dunaújváros. Five blocks: the bid and its points, key dates, rates and fees, guarantees / securities /
connection fees in HUF, deposits and fees to RtB in EUR k.

Inputs (blue on cream): Bora from the current model (out/Bora_model_RTB_COD_HTM_recalc.xlsx, Bora selected; the bid
commits EU equipment), Iváncsa and Dunaújváros as in the user's sheet. Points use the model's band tables (Inputs block U),
copied into a hidden area to the right. Everything below the inputs is a formula. The model is not changed.

    PYTHONPATH=. python3 projects/bora_v5_summary/extract_bid_cash.py
"""
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1]))
from fepm.recalc import recalc  # noqa: E402

MODEL = HERE / "out" / "Bora_model_RTB_COD_HTM_recalc.xlsx"
OUT = HERE / "out" / "Bora_bid_cash_extract.xlsx"

FONT = "Aptos Narrow"
GREEN, CREAM, GREY, PURPLE, BLUE = "00563C", "FFF2CC", "EDEDED", "7030A0", "0000FF"
YELLOW, ORANGE = "FFFF00", "FFC000"
NUM, NUM1, PCT, MULT, DATE = '#,##0;(#,##0);"–"', '#,##0.0;(#,##0.0);"–"', '0.0%;(0.0%);"–"', '0.00"x"', "dd-mmm-yy"
PTS = '#,##0.0;(#,##0.0);"–"'
BIDS = "EFHI"                                                     # BORA 1, BORA 2, Iváncsa, Dunaújváros
thin = Side(style="thin", color="808080")
BOX = Border(thin, thin, thin, thin)

# (1) the bid: label, unit, format, BORA 1 = BORA 2 (None = from the model), Iváncsa = Dunaújváros
LEVERS = [("mva", "Net connection capacity", "MVA", NUM, None, 50),
          ("ob", "Overbuild (ÖSZ-3)", "%", PCT, None, 0.4),
          ("bess", "BESS power, share of MVA (ÖSZ-4)", "%", PCT, None, 0.41),
          ("hours", "BESS duration", "h", NUM, None, 4),
          ("muni", "Municipal one-off (ÖSZ-1)", "EUR k/MW", NUM, None, 50),
          ("comm", "Community contribution (ÖSZ-2)", "EUR/MWh", NUM, None, 3),
          ("part", "Partial award accepted (ÖSZ-5)", "flag", NUM, None, 0),
          ("pbm", "Performance security multiple (ÖSZ-6)", "x", MULT, None, 1),
          ("brown", "Brownfield site (ÖSZ-7)", "flag", NUM, None, 0),
          ("shared", "Shared connection (ÖSZ-8)", "flag", NUM, None, 0),
          ("euturb", "EU turbines ≥ 80 % (ÖSZ-9)", "flag", NUM, 1, 1),
          ("eutrafo", "EU transformers ≥ 80 % (ÖSZ-9)", "flag", NUM, 1, 0)]
MODEL_LEVERS = {"mva": "Dash_MVA", "ob": "Dash_Overbuild", "bess": "Dash_BESS_Share", "muni": "Dash_Municipal",
                "comm": "Dash_Community", "part": "Dash_Partial", "pbm": "Dash_PB_Mult", "brown": "Dash_Brownfield",
                "shared": "Dash_Shared"}
POINTS = [("Points – municipal one-off", "=INDEX(Osz1_Pts,1+SUMPRODUCT((Osz1_LB<{c}{muni})+(Osz1_LB={c}{muni})*Osz1_Incl))*Osz1_W"),
          ("Points – community contribution",
           "=INDEX(Osz2_Pts,1+SUMPRODUCT((Osz2_LB<{c}{comm})+(Osz2_LB={c}{comm})*Osz2_Incl))*Osz2_W"),
          ("Points – overbuild", "=INDEX(Osz3_Pts,1+SUMPRODUCT((Osz3_LB<{c}{ob})+(Osz3_LB={c}{ob})*Osz3_Incl))*Osz3_W"),
          ("Points – BESS", "=INDEX(Osz3_Pts,1+SUMPRODUCT((Osz3_LB<{c}{bess})+(Osz3_LB={c}{bess})*Osz3_Incl))*Osz4_W"),
          ("Points – partial award", "={c}{part}*Osz5_Pts*Osz5_W"),
          ("Points – performance security",
           "=INDEX(Osz6_Pts,1+SUMPRODUCT((Osz6_LB<{c}{pbm})+(Osz6_LB={c}{pbm})*Osz6_Incl))*Osz6_W"),
          ("Points – brownfield", "={c}{brown}*Osz7_Pts*Osz7_W"),
          ("Points – shared connection", "={c}{shared}*Osz8_Pts*Osz8_W"),
          ("Points – EU turbines", "={c}{euturb}*Osz9a_Pts*Osz9_W"),
          ("Points – EU transformers", "={c}{eutrafo}*Osz9b_Pts*Osz9_W")]
BANDS = ["Osz1_LB", "Osz1_Incl", "Osz1_Pts", "Osz2_LB", "Osz2_Incl", "Osz2_Pts", "Osz3_LB", "Osz3_Incl", "Osz3_Pts",
         "Osz6_LB", "Osz6_Incl", "Osz6_Pts", "Osz1_W", "Osz2_W", "Osz3_W", "Osz4_W", "Osz5_W", "Osz6_W", "Osz7_W",
         "Osz8_W", "Osz9_W", "Osz5_Pts", "Osz7_Pts", "Osz8_Pts", "Osz9a_Pts", "Osz9b_Pts"]
PB_NOTE = "model has no PB release date; Kiírás XII.13, estimate"


def model_values():
    wb = openpyxl.load_workbook(MODEL, data_only=True)

    def name(n):
        sheet, ref = next(wb.defined_names[n].destinations)
        rng = wb[sheet][ref.replace("$", "")]
        return [c.value for c in rng[0]] if isinstance(rng, tuple) else rng.value

    i, dep, dev = wb["Inputs"], wb["Deposits"], wb["Devex_FEP"]
    bora = {k: name(n) for k, n in MODEL_LEVERS.items()}
    bora["mva"] = name("Dash_MVA") / i["E82"].value
    bora["hours"] = i["E21"].value
    v = lambda r: i[f"E{r}"].value  # noqa: E731
    dates = [("Signing fee and DD paid", v(27)), ("Bid submission: fee paid, bid security posted", v(95)),
             ("Performance security posted (bid security released)", dep["K10"].value),
             ("TSO security paid", v(97)), ("Devex contingency paid", dev["F38"].value),
             ("Connection agreement (HCSSZ): fees paid", v(99)), ("Notice to proceed", v(154)), ("COD", v(155)),
             ("Performance security released", datetime(2033, 3, 31))]
    # (3) rates: key, label, unit, format, Bora, Iváncsa = Dunaújváros
    rates = [("fx", "FX rate", "HUF/EUR", NUM, v(11), 365),
             ("bw", "Security base: benchmark investment, wind", "HUF/MW", NUM, v(75) * 1e6, 406e6),
             ("bb", "Security base: benchmark investment, BESS", "HUF/MWh", NUM, v(89) * 1e6, 200e6),
             ("pf", "Participation fee per bid", "HUF", NUM, v(81) * 1e6, 3e6),
             ("bsr", "Bid security, share of base", "%", PCT, v(76), 0.005),
             ("pbr", "Performance security, minimum share of base", "%", PCT, v(77), 0.025),
             ("tso", "TSO security", "HUF/MVA", NUM, v(83) * 1e6, 4.5e6),
             ("ind", "Indirect connection fee (közvetett)", "HUF/MVA", NUM, v(297), 19826680),
             ("dir", "Direct connection fee per bid (model: per MVA of the shared point)", "HUF", NUM,
              v(84) * 1e6 * bora["mva"], 460e6),
             ("gf", "Bank guarantee fee", "%", PCT, v(113), 0.06),
             ("sdd", "Signing fee and DD", "EUR k", NUM, v(24) + v(25), 150),
             ("devex", "Devex incl. contingency", "EUR k", NUM, None, 1000),
             ("cont", "Devex contingency", "%", PCT, v(68), None)]
    bands = {n: name(n) for n in BANDS}
    check = {"bs": dep["M8"].value, "pb": dep["M8"].value + dep["M10"].value, "tso": dep["M11"].value,
             "dir": dep["M12"].value, "indn": dep["M18"].value}
    return bora, dates, rates, bands, check


def style(c, fmt=None, *, font=None, fill=None, border=None, bold=False, italic=False, align=None):
    c.font = Font(name=FONT, size=10, color=font or "000000", bold=bold, italic=italic)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = border
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def build():
    bora, dates, rates, bands, check = model_values()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bid_Cash"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJK", (2, 2, 52, 9, 12, 12, 13, 13, 13, 13, 2)):
        ws.column_dimensions[col].width = w
    R = {}

    def bar(r, title, heads):
        for c in range(2, 11):
            style(ws.cell(r, c), font="FFFFFF", fill=GREEN, bold=True, align="right" if c > 4 else None)
        ws.cell(r, 3).value = title
        for col, h in heads.items():
            ws[f"{col}{r}"].value = h

    def label(r, text, unit, bold=False, fill=None):
        style(ws[f"C{r}"], bold=bold, fill=fill)
        style(ws[f"D{r}"], font="808080", italic=True, fill=fill)
        ws[f"C{r}"].value, ws[f"D{r}"].value = text, unit

    # Hidden constants and band tables (columns M onwards).
    ws["M2"].value = "Band tables (model Inputs block U) and constants"
    r = 3
    for n, vals in bands.items():
        vals = vals if isinstance(vals, list) else [vals]
        ws.cell(r, 13).value = n
        for j, x in enumerate(vals):
            ws.cell(r, 14 + j).value = x
        ref = f"Bid_Cash!${L(14)}${r}:${L(13 + len(vals))}${r}" if len(vals) > 1 else f"Bid_Cash!$N${r}"
        wb.defined_names[n] = DefinedName(n, attr_text=ref)
        r += 1
    for n, x in (("Thousand", 1000), ("Days_In_Year", 365)):
        ws.cell(r, 13).value, ws.cell(r, 14).value = n, x
        wb.defined_names[n] = DefinedName(n, attr_text=f"Bid_Cash!$N${r}")
        r += 1
    for col in range(13, 30):
        ws.column_dimensions[L(col)].hidden = True

    # Title.
    style(ws["C1"], font=BLUE, fill=CREAM, bold=True, border=BOX)
    ws["C1"].font = Font(name=FONT, size=14, color=BLUE, bold=True)
    ws["C1"].value = "Bora, Iváncsa, Dunaújváros"
    style(ws["C2"])

    # (1) The bid: commitments and points.
    bar(4, "The bid: commitments and points", {"E": "BORA 1", "F": "BORA 2", "H": "Iváncsa", "I": "Dunaújváros"})
    r = 5
    for key, text, unit, fmt, bora_v, other in LEVERS:
        R[key] = r
        label(r, text, unit)
        bv = bora[key] if bora_v is None else bora_v
        for col, x in zip("EFGHI", (bv, bv, None, other, other)):
            style(ws[f"{col}{r}"], fmt, font=BLUE, fill=CREAM, border=BOX if col != "G" else None)
            ws[f"{col}{r}"].value = x
        r += 1
    ws["C2"].value = f'="Bid strategy – "&TEXT(E{R["ob"]},"0 %")&" overbuild (Bora)"'
    for key, text, unit, f, fmt in (("wind", "Wind committed", "MW", "={c}{mva}*(1+{c}{ob})", NUM1),
                                    ("bmw", "BESS power", "MW", "={c}{mva}*{c}{bess}", NUM1),
                                    ("mwh", "BESS energy", "MWh", "={c}{bmw}*{c}{hours}", NUM1)):
        R[key] = r
        label(r, text, unit)
        for col in BIDS:
            style(ws[f"{col}{r}"], fmt, font=PURPLE, fill=GREY, border=BOX)
            ws[f"{col}{r}"].value = f.format(c=col, **R)
        r += 1
    r += 1
    first = r
    for text, f in POINTS:
        label(r, text, "points")
        for col in BIDS:
            style(ws[f"{col}{r}"], PTS, font=PURPLE, fill=GREY, border=BOX)
            ws[f"{col}{r}"].value = f.format(c=col, **R)
        r += 1
    R["pts"] = r
    label(r, "Points total, live case", "points", bold=True)
    for col in BIDS:
        style(ws[f"{col}{r}"], "#,##0.00", bold=True)
        ws[f"{col}{r}"].value = f"=SUM({col}{first}:{col}{r - 1})"

    # (2) Key dates.
    r += 2
    bar(r, "Key dates", {"E": "Date"})
    r += 1
    for text, d in dates:
        R[text] = r
        label(r, text, "date")
        style(ws[f"E{r}"], DATE, font=BLUE, fill=CREAM, border=BOX)
        ws[f"E{r}"].value = d
        r += 1
    style(ws[f"F{r - 1}"], font="808080", italic=True)
    ws[f"F{r - 1}"].value = PB_NOTE
    D = {k: f"$E${R[t]}" for k, t in (("bid", dates[1][0]), ("pb", dates[2][0]), ("cod", "COD"),
                                       ("rel", "Performance security released"))}

    # (3) Rates and fees.
    r += 1
    bar(r, "Rates and fees", {"E": "BORA 1", "F": "BORA 2", "H": "Iváncsa", "I": "Dunaújváros"})
    r += 1
    for key, text, unit, fmt, bv, other in rates:
        R[key] = r
        label(r, text, unit)
        for col, x in zip("EFGHI", (bv, bv, None, other, other)):
            boxed = x is not None and col != "G"
            style(ws[f"{col}{r}"], fmt, font=BLUE, fill=CREAM if (x is not None or col == "G") else None,
                  border=BOX if boxed else None)
            ws[f"{col}{r}"].value = x
        r += 1

    # (4) Guarantees, securities and connection fees (HUF).
    r += 1
    bar(r, "Guarantees, securities and connection fees (HUF)",
        {"E": "BORA 1", "F": "BORA 2", "G": "BORA", "H": "Iváncsa", "I": "Dunaújváros", "J": "Total"})
    r += 1
    base = "({c}${wind}*{c}${bw}+{c}${mwh}*{c}${bb})"
    huf = [("h_pf", "Participation fee", "={c}${pf}", None, False),
           ("h_bs", "Bid security (guarantee)", "={c}${bsr}*" + base, YELLOW, False),
           ("h_pb", "Performance security (guarantee)", "={c}${pbr}*{c}${pbm}*" + base, ORANGE, False),
           ("h_tso", "TSO connection-request security (cash)", "={c}${tso}*{c}${mva}", None, False),
           ("h_dir", "Direct connection fee", "={c}${dir}", None, False),
           ("h_indg", "Indirect connection fee, gross", "={c}${ind}*{c}${mva}", None, False),
           ("h_tsoc", "Less TSO security credited", "=-{c}{h_tso}", None, False),
           ("h_indp", "Indirect connection fee, payable", "={c}{h_indg}+{c}{h_tsoc}", None, False),
           ("h_peak", "Peak guarantee (performance security)", "=MAX({c}{h_bs},{c}{h_pb})", None, True),
           ("h_cash", "Cash: fee, TSO security, connection fees", "={c}{h_pf}+{c}{h_tso}+{c}{h_dir}+{c}{h_indp}",
            None, True)]
    for key, text, f, fill, bold in huf:
        R[key] = r
        label(r, text, "HUF", bold=bold, fill=fill)
        for col in "EFGHIJ":
            top = Border(top=Side(style="thin", color="000000")) if key == "h_peak" else None
            style(ws[f"{col}{r}"], NUM, fill=fill, bold=bold, border=top)
        for col in BIDS:
            ws[f"{col}{r}"].value = f.format(c=col, **R)
        ws[f"G{r}"].value = f"=E{r}+F{r}"
        ws[f"J{r}"].value = f"=E{r}+F{r}+H{r}+I{r}"
        r += 1

    # (5) Deposits and fees to RtB (EUR k).
    r += 1
    bar(r, "Deposits and fees to RTB (EUR k)", {"E": "BORA 1", "F": "BORA 2", "H": "Iváncsa", "I": "Dunaújváros",
                                                 "J": "Total"})
    r += 1
    eur = "={c}{%s}/{c}${fx}/Thousand"
    gfee = ("={c}${gf}*((" + D["pb"] + "-" + D["bid"] + ")*{c}{e_bs}+MAX(0,MIN(" + D["cod"] + "," + D["rel"] + ")-"
            + D["pb"] + "+1)*{c}{e_pb})/Days_In_Year")
    rows5 = [("e_sdd", "Signing fee and DD", "={c}${sdd}", False),
             ("e_pf", "Participation fee", eur % "h_pf", False),
             ("e_bs", "Bid security (guarantee)", eur % "h_bs", False),
             ("e_pb", "Performance security (guarantee)", eur % "h_pb", False),
             ("e_tso", "TSO connection-request security", eur % "h_tso", False),
             ("e_dir", "Direct connection fee", eur % "h_dir", False),
             ("e_ind", "Indirect connection fee, net of TSO security", eur % "h_indp", False),
             ("e_gf", "Guarantee fees to COD (6 % p.a., bid security to PB posting, PB to COD)", gfee, False),
             ("e_dep", "Deposits and fees (PB incl. the credited bid security)",
              "={c}{e_pf}+{c}{e_pb}+{c}{e_tso}+{c}{e_dir}+{c}{e_ind}+{c}{e_gf}", True),
             ("e_tot", "Total", "={c}{e_dep}+{c}{e_sdd}", True)]
    for key, text, f, bold in rows5:
        R[key] = r
        label(r, text, "EUR k", bold=bold)
        for col in "EFHIJ":
            top = Border(top=Side(style="thin", color="000000")) if key == "e_dep" else None
            style(ws[f"{col}{r}"], NUM1 if key == "e_pf" else NUM, bold=bold, border=top)
        for col in BIDS:
            ws[f"{col}{r}"].value = f.format(c=col, **R)
        ws[f"J{r}"].value = f"=E{r}+F{r}+H{r}+I{r}"
        r += 1
    ws.freeze_panes = "E5"
    wb.save(OUT)

    # Verification: Bora per bid against the model's Deposits k€ values (cash reading), others against the image.
    rc = openpyxl.load_workbook(recalc(OUT, OUT.with_name("Bora_bid_cash_extract_recalc.xlsx")), data_only=True)
    x = rc["Bid_Cash"]
    g = lambda key, col="E": x[f"{col}{R[key]}"].value  # noqa: E731
    ties = {"bid security": (g("e_bs"), check["bs"]), "PB gross": (g("e_pb"), check["pb"]),
            "TSO": (g("e_tso"), check["tso"]), "direct": (g("e_dir"), check["dir"]),
            "indirect net": (g("e_ind"), check["indn"])}
    image = {"Iváncsa points": (g("pts", "H"), 69.5), "Bora points": (g("pts"), 80.5),
             "Iváncsa bid security HUF": (g("h_bs", "H"), 224.1e6), "Iváncsa cash HUF": (g("h_cash", "H"), 1454.334e6),
             "Bora cash HUF": (g("h_cash"), 7564303502), "Iváncsa guarantee fees": (round(g("e_gf", "H")), 836),
             "Bora guarantee fees": (round(g("e_gf")), 1591)}
    ok = all(abs(a - b) < 1e-6 * max(1, abs(b)) for a, b in {**ties, **image}.values())
    for k, (a, b) in {**ties, **image}.items():
        print(f"{k:<26} extract {a:>16,.1f}  reference {b:>16,.1f}")
    print("EUR k deposits and fees:", {c: round(g("e_dep", c), 1) for c in "EFHIJ"})
    print("EUR k total:", {c: round(g("e_tot", c), 1) for c in "EFHIJ"})
    errors = [c.coordinate for row in x.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("#")]
    print("ties:", "PASS" if ok else "FAIL", "| error cells:", errors[:5], "| rows:", r - 1)


if __name__ == "__main__":
    build()
