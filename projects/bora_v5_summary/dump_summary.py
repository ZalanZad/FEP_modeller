"""Compact dump of the Dashboard and Cash_Flow sheets (SPEC §7 step 4, §9.4) and the hand checks of SPEC §7 step 3.

Reads out/Bora_model_RTB_COD_HTM_recalc.xlsx (LibreOffice values), writes out/summary_dump.txt with one line per Dashboard
and Cash_Flow row (label | unit | E | F | G, plus the timeline ends on Cash_Flow), then recomputes capex, the PB,
the clipping interpolation, first-year generation and the points in Python from the Dashboard levers and the Inputs values and compares them with the Dashboard results.

    python3 projects/bora_v5_summary/dump_summary.py
"""
import datetime as dt
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent / "out"
TOL = 1e-6


def fmt(v):
    if isinstance(v, float):
        return f"{v:.6g}"
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    return "" if v is None else str(v)


def dump(dash, cf, path, htm=None, bid=None):
    """Dashboard: label | unit | E … K. Cash_Flow: the same plus the first three and the last timeline values.
    HTM block 17 (sale at COD): label | unit | G | timeline ends. Bid_Points summary to COD: label | unit |
    BORA 1 | BORA 2 | BORA S | Iváncsa | Dunaújváros."""
    lines = ["== Dashboard =="]
    for r in range(3, dash.max_row + 1):
        cells = [dash.cell(r, c).value for c in range(3, 12)]
        while len(cells) > 5 and cells[-1] is None:                    # columns H:K only in section 8
            cells.pop()
        if any(v is not None for v in cells):
            lines.append(" | ".join(fmt(v) for v in cells))
    lines.append("== Cash_Flow (timeline: first three years … last year) ==")
    last = max(c.column for c in cf[8] if c.value is not None)            # row 8 = year header
    for r in range(4, cf.max_row + 1):
        cells = [cf.cell(r, c).value for c in (3, 4, 5, 6, 7, 9, 10, 11)] + ["…", cf.cell(r, last).value]
        if any(v is not None for v in cells[:8]):
            lines.append(" | ".join(fmt(v) for v in cells))
    if htm is not None:
        lines.append("== HTM block 17, sale at COD (G result; timeline: first three years … last year) ==")
        for r in range(332, htm.max_row + 1):
            cells = [htm.cell(r, c).value for c in (3, 4, 7, 8, 9, 10)] + ["…", htm.cell(r, 43).value]
            if any(v is not None for v in cells[:6]):
                lines.append(" | ".join(fmt(v) for v in cells))
    if bid is not None:
        lines.append("== Bid_Points, summary to COD (BORA 1 | BORA 2 | BORA S | Iváncsa | Dunaújváros) ==")
        top = next(r for r in range(1, bid.max_row + 1) if bid.cell(r, 3).value == "Summary to COD")
        for r in range(top, top + 11):
            cells = [bid.cell(r, c).value for c in range(3, 10)]
            if any(v is not None for v in cells):
                lines.append(" | ".join(fmt(v) for v in cells))
    path.write_text("\n".join(lines) + "\n")
    return len(lines)


def band(x, lb, incl, pts):
    """Inner points: band 1 below the first bound, then one band per lower bound passed (inclusive if flagged)."""
    return pts[sum(1 for b, i in zip(lb, incl) if b < x or (b == x and i))]


def row_values(ws, r, first, n):
    return [ws.cell(r, c).value for c in range(first, first + n) if ws.cell(r, c).value is not None]


def hand_checks(wb):
    inp, s = wb["Inputs"], wb["Dashboard"]
    rows = {s.cell(r, 3).value: r for r in range(1, s.max_row + 1) if s.cell(r, 3).value}
    S = lambda label, col="E": s[f"{col}{rows[label]}"].value  # noqa: E731
    u = {inp.cell(r, 3).value: inp.cell(r, 5).value for r in range(300, inp.max_row + 1) if inp.cell(r, 3).value}
    N = lambda n: s[next(wb.defined_names[n].destinations)[1].replace("$", "")].value  # noqa: E731  Dashboard levers
    case = inp["E6"].value                               # live case column G/H
    v = lambda r: inp.cell(r, 6 + case).value if inp.cell(r, 6 + case).value is not None else inp[f"G{r}"].value  # noqa

    mva = N("Dash_MVA")
    inst = mva * (1 + N("Dash_Overbuild"))                # committed (bid): PB and points base
    built = inst * (1 - N("Dash_Build_Shortfall"))       # built: capex and generation base
    bess_mw = N("Dash_BESS_Share") * mva
    mwh = bess_mw * v(21)
    fx = v(11)
    out = []

    # Capex breakdown (€ m) from the Inputs rates; IDC and debt fees are taken from HTM (schedule-driven).
    lines = (163, 164, 165, 166, 169, 170)                # scaled to the Dashboard row 13 scenario total (AG edit)
    scale = s[f"H{label_row(s, 'Construction capex excl. turbine')}"].value / sum(v(r) for r in lines)
    wind = (v(162) + scale * sum(v(r) for r in lines)) * built / 1000
    subst = v(167) * mva / u["Source basis of the lump sums"] / 1000   # lump sum quoted for the 200 MVA source
    cont = (wind + subst) * v(173)
    bess = N("Dash_BESS_Capex_MWh") * v(21) * bess_mw / 1000
    ind = (v(85) - v(83)) * mva / fx                    # indirect fee net of the TSO security credit, mHUF → € m
    direct = v(84) * mva / fx
    muni = N("Dash_Municipal") * mva / 1000
    idc, fees = wb["HTM"]["G79"].value / 1000, -wb["HTM"]["G85"].value / 1000
    allin = wind + subst + cont + bess + ind + direct + muni + idc + fees
    out += [("All-in cost to COD, € m", allin, S("All-in cost to COD")),
            ("All-in cost, k€/MW built", allin * 1000 / built, S("All-in cost to COD", "F")),
            ("All-in cost, k€/MVA", allin * 1000 / mva, S("All-in cost to COD", "G")),
            ("BESS capex, € m", bess, S("BESS (at cost; no BESS revenue or opex modelled)")),
            ("Contingency, € m", cont, S("Contingency"))]

    # Performance security: 2.5 % × (committed MW × 406 + MWh × 200) mHUF × multiple.
    pb = v(77) * (inst * v(75) + mwh * v(89)) * N("Dash_PB_Mult")
    perf = S("Performance security (gross; bid security credited)")
    out += [("Performance security, mHUF", pb, perf), ("Performance security, € m at FX", pb / fx, perf / fx)]

    # Clipping: OVIT memo Table 7, linear between the points (last segment extended), not below 0.
    curve = {k: row_values(inp, r, 17, 12) for k, r in
             (("x", label_row(inp, "Clipping curve: installed")), ("y", label_row(inp, "Clipping curve: clipped")))}
    x = built / mva
    i = min(max(sum(1 for p in curve["x"] if p <= x), 1), len(curve["x"]) - 1) - 1
    x0, x1, y0, y1 = curve["x"][i], curve["x"][i + 1], curve["y"][i], curve["y"][i + 1]
    clip = max(0.0, y0 + (y1 - y0) / (x1 - x0) * (x - x0))
    out.append(("Clipped share of gross energy", clip, S("Clipped share of gross energy (OVIT curve)")))

    # First full year generation: built MW × P50 h × (1 − haircut) × (1 − curtailment) × yield factor × (1 − clip).
    gen = built * v(181) * (1 - v(182)) * (1 - v(183)) * v(185) * (1 - clip) / 1000
    out += [("Net generation, first full year, GWh", gen, S("Net generation, first full year")),
            ("Equivalent hours per MW built", gen * 1000 / built, S("Equivalent full-load hours, first full year", "F")),
            ("Equivalent hours per MVA", gen * 1000 / mva, S("Equivalent full-load hours, first full year", "G"))]

    # Points: band lookups from the Inputs block U tables.
    t = {}
    for r in range(300, inp.max_row + 1):
        label = inp.cell(r, 3).value or ""
        for key in ("1", "2", "3", "6"):
            if label.startswith(f"ÖSZ-{key} ") and " – " in label:
                t.setdefault(key, {})[label.split(" – ")[-1]] = row_values(inp, r, 17, 12)
    pts = lambda k, x: band(x, t[k]["lower bound"], t[k]["lower bound included"], t[k]["inner points"])  # noqa: E731
    w = {k: u[f"ÖSZ-{k} weight"] for k in "123456789"}
    ob, bs = N("Dash_Overbuild"), N("Dash_BESS_Share")
    ratio = bs / (1 + ob * N("Dash_OSZ4_Basis"))
    total = (pts("1", N("Dash_Municipal")) * w["1"] + pts("2", N("Dash_Community")) * w["2"] + pts("3", ob) * w["3"] + pts("3", ratio) * w["4"]
             + pts("6", N("Dash_PB_Mult")) * w["6"]
             + 100 * (N("Dash_Partial") * w["5"] + N("Dash_Brownfield") * w["7"]
                      + N("Dash_Shared") * w["8"])
             + (80 * N("Dash_EU_Turbine") + 20 * N("Dash_EU_Transformer")) * w["9"])
    out.append(("Total points", round(total, 2), S("Total points", "F")))
    out.append(("Bid_Points total, BORA 1 and BORA 2", wb["Bid_Points"]["E34"].value + wb["Bid_Points"]["F34"].value,
                2 * S("Total points", "F")))
    return out


def label_row(ws, text):
    return next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 3).value or "").startswith(text))


def main():
    wb = openpyxl.load_workbook(OUT / "Bora_model_RTB_COD_HTM_recalc.xlsx", data_only=True)
    n = dump(wb["Dashboard"], wb["Cash_Flow"], OUT / "summary_dump.txt", wb["HTM"], wb["Bid_Points"])
    print(f"{OUT / 'summary_dump.txt'}: {n} lines")
    ok = True
    for label, py, xl in hand_checks(wb):
        good = abs(py - xl) <= TOL * max(1, abs(py))
        ok &= good
        print(f"{label:<40} python {py:>14,.4f}  workbook {xl:>14,.4f}  {'OK' if good else 'DIFF'}")
    print("hand checks:", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
