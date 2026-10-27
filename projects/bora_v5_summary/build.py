"""Bora model RtB / COD / HTM – SPEC §20: the RtB sale for Iváncsa and Dunaújváros beside Bora.

Source (SPEC §20): reference/Bora_model_RTB_COD_HTM_271026.xlsm, the user's edit of our §19 output (scenario table, live
case EU, Inputs G99, Bid_Points "Wind built" row and capex block). Every user edit stays; the §20 edits are:
- Inputs block M: a three-column project table (Bora | Iváncsa | Dunaújváros) – seven rows added, the other
  projects' values from reference/Bora_model_RTB_COD_HTM_261026_other_projects.xlsm (Bid_Points H:I);
- Dashboard selector Dash_Project → Inputs E9; the live cells the RtB route reads pick the project from block M
  (Bora's column holds Bora's previous formula, so Bora is unchanged); Deposits capacity bases and MOIC premium read
  the project MVA; HTM / COD and Bid_Points stay Bora's;
- new sheet RtB_Projects (devex-file layout, EUR k): live column + stored block written by vba/RtBProjects.bas;
- Dashboard: selector row, 5b pointer line, key-assumption note, check row; Change_Log line.
Rows are found by names and labels. Then Scenarios.bas and RtBProjects.bas are embedded and RefreshAll runs.
build_v1.py … build_v4.py are the history (build_v4.py = §18/§19); ag_tieout.py ties Bora out to the source file.
"""
import re
import shutil
import tempfile
import zipfile
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

from fepm.recalc import recalc
from fepm.xl.macros import VBA_DIR, package

HERE = Path(__file__).resolve().parent
REF = HERE.parents[1] / "reference"
SRC = REF / "Bora_model_RTB_COD_HTM_271026.xlsm"
NAME = "Bora_model_RTB_COD_HTM"
GREEN = "FF00B050"
PROJECTS = ["Bora", "Iváncsa", "Dunaújváros"]
OTHER_SRC = "other-projects file, Bid_Points H:I"

# §20.1 block M rows added above "Gross wind capacity": label, unit, Bora (G), Iváncsa (H), Dunaújváros (I), note.
# "{old:r}" = the Bora formula that the live cell Inputs E{r} held before (case picker), moved here.
NEW_ROWS = [("mva", "Connection capacity", "MVA", "=Dash_MVA", 50, 50, OTHER_SRC),
            ("ob", "Overbuild (ÖSZ-3)", "%", "=Dash_Overbuild", 0.4, 0.4, OTHER_SRC),
            ("bess", "BESS power, share of MVA (ÖSZ-4)", "%", "=Dash_BESS_Share", 0.41, 0.41, OTHER_SRC),
            ("hours", "BESS duration", "h", "=$E$21", 4, 4, OTHER_SRC),
            ("dd", "Transaction / DD costs to FEP", "k€", "{old:25}", 0, 0, "Others: signing + DD in one row"),
            ("shared", "Direct connection fee shared by two bids (1 = yes)", "flag", "=Dash_Shared", 0, 0, OTHER_SRC),
            ("prem", "RtB premium", "k€/MVA", "{old:141}", 150, 150, "Others: placeholder (user to set)")]
# Existing block M rows (labels as in the source) and the other projects' values (H = Iváncsa, I = Dunaújváros).
OTHER_VALUES = {"Project name": ("Iváncsa", "Dunaújváros"),
                "Grid node": ("Iváncsa", "Dunaújváros Dél"),
                "Gross wind capacity": ("={c}{mva}*(1+{c}{ob})",) * 2,
                "BESS connection capacity": ("={c}{mva}*{c}{bess}",) * 2,
                "Signing fee to current owners": (150, 150),
                "JV structure live": (1, 1),
                "FEP share of net proceeds if JV": (1, 1),
                "Developer success fee live": (0, 0),
                "Number of separate bids": (1, 1),
                "Indirect connection fee (fajlagos": (19826680, 19826680),
                "Direct connection investment": (460000000, 460000000),
                "Indirect connection fee": ("={c}{ind_ft}/Million",) * 2,
                "Direct connection fee": ("={c}{dir_ft}/Million/{c}{mva}/(1+{c}{shared})",) * 2,
                "Devex target": (500, 500),
                "Devex scaling factor": ("=({c}{target}-$E$68*Deposits!$O$7)/($G${target}-$E$68*Deposits!$O$7)",)
                * 2}
NOTES_M = {"Signing fee to current owners": "Bora: case value × In_Scale; others: signing + DD",
           "FEP share of net proceeds if JV": "Others: placeholder 100 % (user to set)",
           "Developer success fee live": "Others: placeholder 0 (user to set)",
           "Devex target": "Others: devex incl. contingency (other-projects file)",
           "Devex scaling factor": "Bora: In_Scale; others: Devex_FEP E39 = the target"}

PROJECT_LABEL = "Project for the RtB route (HTM and COD meaningful for Bora only)"
PROJECT_NOTE = ("HTM and COD figures are meaningful for Bora only; with another project selected they recompute "
                "on that capacity and are not meaningful")
LINE_5B = "Other projects (Iváncsa, Dunaújváros): see RtB_Projects"

# §20.3 RtB_Projects rows: (key, label, unit, live formula, total rule) – None label = section bar.
# Total rule: "sum" = F+G+H, "mva" = total profit ÷ total MVA, None = blank.
RTB_ROWS = [
    (None, "Capacity and terms", None, None, None),
    ("mva", "Connection capacity", "MVA", "=In_Project_MVA", "sum"),
    ("wind", "Wind committed", "MW", "=Inputs!$E$16", "sum"),
    ("bmw", "BESS power", "MW", "=Inputs!$E$20", "sum"),
    ("bmwh", "BESS energy", "MWh", "=Inputs!$E$22", "sum"),
    ("share", "FEP share of distributable proceeds", "%", "=Inputs!$E$33", None),
    ("prem", "RtB premium", "k€/MVA", "=Inputs!$E$141", None),
    ("rtb", "RtB date", "date", "=Inputs!$E$56", None),
    ("sale", "Sale date", "date", "=Inputs!$E$58", None),
    ("freeze", "Interest frozen from (accrual stops, to the sale)", "date", "=Inputs!$E$64", None),
    (None, "Principal posted by class", None, None, None),
    ("p_cash", "Cash, 12 % recoverable loan (TSO security, connection fees)", "EUR k",
     "=MOIC!$G$110+MOIC!$G$117", "sum"),
    ("p_fee", "Cash, 20 % at-risk loan (participation fee, sunk)", "EUR k", "=MOIC!$G$124+MOIC!$G$134", "sum"),
    ("p_gtee", "Guarantees, 6 % (bid security, performance security face)", "EUR k", "=MOIC!$G$76", "sum"),
    ("p_dev", "Devex and land, 20 % / 2× floor", "EUR k", "=MOIC!$G$125+MOIC!$G$135", "sum"),
    ("p_tot", "Total posted", "EUR k", "=E{p_cash}+E{p_fee}+E{p_gtee}+E{p_dev}", "sum"),
    (None, "Interest claims to the sale (frozen from the date above)", None, None, None),
    ("i_cash", "12 % interest on the recoverable loan", "EUR k", "=Deposits!$Y$23", "sum"),
    ("i_gtee", "6 % interest on the guarantees", "EUR k", "=Deposits!$Y$22", "sum"),
    ("i_risk", "At-risk return (devex and land 2× floor, participation fee)", "EUR k",
     "=MOIC!$G$131+Deposits!$Y$24", "sum"),
    ("claims", "Total FEP claims (cash and at-risk principal + interest)", "EUR k",
     "=E{p_cash}+E{p_fee}+E{p_dev}+E{i_cash}+E{i_gtee}+E{i_risk}", "sum"),
    (None, "Waterfall at the sale", None, None, None),
    ("w_prem", "RtB premium", "EUR k", "=MOIC!$G$143", "sum"),
    ("w_reimb", "Buyer reimbursement of principal (securities, connection fees, land)", "EUR k", "=MOIC!$G$144",
     "sum"),
    ("w_tot", "Total proceeds", "EUR k", "=E{w_prem}+E{w_reimb}", "sum"),
    ("w_cost", "Less sale costs", "EUR k", "=MOIC!$G$146", "sum"),
    ("w_prin", "Less recoverable principal repaid", "EUR k", "=-E{p_cash}", "sum"),
    ("w_i12", "Less 12 % interest", "EUR k", "=MOIC!$G$147-E{w_prin}", "sum"),
    ("w_risk_p", "Less at-risk principal repaid (land reimbursed)", "EUR k", "=MOIC!$G$148", "sum"),
    ("w_risk", "Less at-risk return", "EUR k", "=MOIC!$G$149", "sum"),
    ("w_i6", "Less 6 % interest on the guarantees", "EUR k", "=MOIC!$G$150", "sum"),
    ("w_dsa", "Less developer success fee (DSA)", "EUR k", "=MOIC!$G$151", "sum"),
    ("dist", "Distributable proceeds", "EUR k", "=MOIC!$G$152", "sum"),
    ("fep", "FEP share", "EUR k", "=MOIC!$G$153", "sum"),
    ("dev", "Developer share", "EUR k", "=MOIC!$G$154", "sum"),
    (None, "FEP result", None, None, None),
    ("rec", "FEP receipts at the sale", "EUR k", "=MOIC!$G$171", "sum"),
    ("out", "FEP outflows (signing, DD, loans, devex, guarantee fees)", "EUR k", "=MOIC!$G$164", "sum"),
    ("devex", "of which FEP-borne devex not reimbursed", "EUR k", "=-MOIC!$G$18", "sum"),
    ("pre", "FEP net profit pre-tax", "EUR k", "=MOIC!$G$172", "sum"),
    ("post", "FEP net profit post-tax", "EUR k", "=MOIC!$G$175", "sum"),
    ("irr", "FEP IRR post-tax", "%", "=MOIC!$G$187", None),
    ("moic", "FEP MOIC", "x", "=MOIC!$G$186", None),
    ("npv", "FEP NPV at 8 %", "EUR k", "=MOIC!$G$192", "sum"),
    ("per", "FEP net profit post-tax per MVA", "k€/MVA", "=E{post}/E{mva}", "mva"),
]
FORMATS = {"MVA": "#,##0.0", "MW": "#,##0.0", "MWh": "#,##0.0", "%": "0.0%", "k€/MVA": "#,##0.0", "date": "dd-mmm-yy",
           "EUR k": "#,##0.0;(#,##0.0);\"–\"", "x": "0.00\"x\""}
PURPLE, MACRO_FILL = "FF7030A0", "FFEDEDED"


# =========================================================================================== helpers
_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?")


_ADDR = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def define(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def like(cell, model, value=None, fmt=None):
    """Write `value` with the style of `model` (and an optional number format)."""
    cell.value = value
    cell._style = copy(model._style)
    if fmt:
        cell.number_format = fmt
    return cell


def name_cell(wb, name):
    sheet, ref = next(wb.defined_names[name].destinations)
    return wb[sheet][ref.replace("$", "").split(":")[0]]


def label_row(ws, text, col=3, start=1):
    """First row at or below `start` whose label in `col` starts with `text`."""
    rows = sorted(c.row for c in ws._cells.values()
                  if c.column == col and c.row >= start and isinstance(c.value, str) and c.value.startswith(text))
    assert rows, (ws.title, text)
    return rows[0]


def _shift_addr(addr, at, n):
    return _ADDR.sub(lambda m: f"{m[1]}{m[2]}{m[3]}{int(m[4]) + n if int(m[4]) >= at else m[4]}", addr)


def _shift_formula(f, host, sheet, at, n):
    """Shift every cell reference to `sheet` rows ≥ `at` by `n` in formula `f` written on sheet `host`."""
    out = []
    for t in Tokenizer(f).items:
        v = t.value
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            ref_sheet, addr = (v.rsplit("!", 1) if "!" in v else (None, v))
            if (ref_sheet.strip("'") if ref_sheet else host) == sheet and _REF.fullmatch(addr):
                v = (ref_sheet + "!" if ref_sheet else "") + _shift_addr(addr, at, n)
        out.append(v)
    return "=" + "".join(out)


def _refs_rows(f, host, sheet):
    """Rows of `sheet` referred to as single cells or range ends in formula `f` (for the deletion guard)."""
    rows = []
    for t in Tokenizer(f).items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            ref_sheet, addr = (t.value.rsplit("!", 1) if "!" in t.value else (None, t.value))
            if (ref_sheet.strip("'") if ref_sheet else host) == sheet and _REF.fullmatch(addr):
                rows += [int(m[4]) for m in _ADDR.finditer(addr)]
    return rows


def shift_rows(wb, sheet, at, n):
    """Insert `n` rows before row `at` (n > 0) or delete −n rows from row `at` (n < 0) on `sheet`, re-pointing every
    formula, name, merged range, conditional format, validation and array-formula range (openpyxl moves cells only)."""
    ws = wb[sheet]
    first = at if n > 0 else at - n                                    # first row that moves
    gone = range(at, at - n) if n < 0 else range(0)
    for w in wb.worksheets:
        for c in list(w._cells.values()):
            v = c.value
            text = v.text if isinstance(v, ArrayFormula) else (v if c.data_type == "f" else None)
            if not text:
                continue
            assert not set(_refs_rows(text, w.title, sheet)) & set(gone), (w.title, c.coordinate, text)
            new = _shift_formula(text, w.title, sheet, first, n)
            if isinstance(v, ArrayFormula):
                v.text = new
            elif new != v:                                             # text that starts with "=" stays text
                c.value = new
    for name, dn in list(wb.defined_names.items()):
        if dn.attr_text and f"{sheet}!" in dn.attr_text:
            dn.attr_text = _shift_formula("=" + dn.attr_text, "", sheet, first, n)[1:]
    merged = [str(m) for m in ws.merged_cells.ranges]
    for m in merged:
        ws.unmerge_cells(m)
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height and r >= at}
    if n > 0:
        ws.insert_rows(at, n)
    else:
        ws.delete_rows(at, -n)
    for r in heights:
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        if r >= first:
            ws.row_dimensions[r + n].height = h
    for m in merged:
        ws.merge_cells(_shift_addr(m, first, n))
    for c in ws._cells.values():                                     # array formulas carry their own range
        if isinstance(c.value, ArrayFormula) and c.value.ref:
            c.value.ref = c.coordinate if ":" not in c.value.ref else _shift_addr(c.value.ref, first, n)
    cf = ws.conditional_formatting
    ws.conditional_formatting = ConditionalFormattingList()
    for rng in cf:
        for rule in rng.rules:
            ws.conditional_formatting.add(_shift_addr(str(rng.sqref), first, n), rule)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(_shift_addr(str(dv.sqref), first, n))
    if ws.freeze_panes:
        ws.freeze_panes = _shift_addr(ws.freeze_panes, first, n)


def _copy_cell(src, dst):
    v = src.value
    if isinstance(v, ArrayFormula):
        v = ArrayFormula(dst.coordinate, v.text)
    dst.value = v
    if src.has_style:
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.protection, dst.number_format = copy(src.alignment), copy(src.protection), src.number_format


def _link(cell, formula):
    """A link from another sheet: green font, no input fill (MODELLING_STANDARDS §5)."""
    cell.value = formula
    f = copy(cell.font)
    cell.font = Font(name=f.name, size=f.sz, bold=f.b, italic=f.i, color=GREEN)
    cell.fill = PatternFill(fill_type=None)


def _pure_link(f):
    toks = [t for t in Tokenizer(f).items if t.type != Token.WSPACE] if isinstance(f, str) and f.startswith("=") else []
    return toks[0].value if len(toks) == 1 and toks[0].type == Token.OPERAND and toks[0].subtype == Token.RANGE \
        else None


def _target(wb, host, op):
    """(sheet, coordinate) of a single-cell reference or single-cell name, else None."""
    if op in wb.defined_names:
        sheet, ref = next(wb.defined_names[op].destinations)
    else:
        sheet, ref = op.rsplit("!", 1) if "!" in op else (host, op)
    ref = ref.replace("$", "")
    return (sheet.strip("'"), ref) if re.fullmatch(r"[A-Z]{1,3}\d+", ref) else None


def collapse_chains(wb, sheet):
    """Every pure link on `sheet` that points at another pure link is re-pointed at the source (no chains)."""
    ws = wb[sheet]
    for c in ws._cells.values():
        op = _pure_link(c.value)
        while op:
            tgt = _target(wb, sheet, op)
            if not tgt:
                break
            s, coord = tgt
            inner = _pure_link(wb[s][coord].value)
            if not inner:
                break
            op = inner if ("!" in inner or inner in wb.defined_names or s == sheet) else f"{s}!{inner}"
            c.value = "=" + op


def _clear(ws, rows, cols):
    for r in rows:
        for col in cols:
            ws[f"{col}{r}"].value = None


def _copy_style(src, dst):
    dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
    dst.alignment, dst.number_format = copy(src.alignment), src.number_format


def general(path, sheet, cells, like):
    """SPEC §8.6: give `cells` the style index of `like` (a General cell) inside the saved file; LibreOffice infers
    a 0.00 % format for these sums on save (HTM G10 would show 3000.00 %)."""
    path = Path(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    wbx, rels = parts["xl/workbook.xml"].decode(), parts["xl/_rels/workbook.xml.rels"].decode()
    rid = re.search(rf'<sheet name="{sheet}"[^>]*r:id="(\w+)"', wbx).group(1)
    target = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="{rid}"', rels)
    xml_name = "xl/" + (target.group(1) or target.group(2)).lstrip("/").removeprefix("xl/")
    xml = parts[xml_name].decode()
    style = re.search(rf'<c r="{like}" s="(\d+)"', xml).group(1)
    for c in cells:
        xml, n = re.subn(rf'<c r="{c}" s="\d+"', f'<c r="{c}" s="{style}"', xml)
        assert n == 1, c
    parts[xml_name] = xml.encode()
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)
    tmp.replace(path)


# =========================================================================================== §20.1 block M
def block_m(wb):
    """Seven rows inserted above "Gross wind capacity"; the other projects' columns H:I filled; E9 = the selector."""
    ws = wb["Inputs"]
    top = label_row(ws, "M. Bora inputs")
    at = label_row(ws, "Gross wind capacity", start=top)
    old = {r: ws[f"E{r}"].value for r in (24, 25, 141)}
    old = {r: (v.text if isinstance(v, ArrayFormula) else v) for r, v in old.items()}
    shift_rows(wb, "Inputs", at, len(NEW_ROWS))
    model = at + len(NEW_ROWS)                                         # "Gross wind capacity", now moved down
    rows = {}
    for r, (key, label, unit, bora, h, i, note) in enumerate(NEW_ROWS, at):
        rows[key] = r
        for col in "CDEFGHI":
            ws[f"{col}{r}"]._style = copy(ws[f"{col}{model}"]._style)
        ws[f"C{r}"].value, ws[f"D{r}"].value, ws[f"F{r}"].value = label, unit, note
        ws[f"G{r}"].value = old[int(bora[5:-1])] if bora.startswith("{old:") else bora
        ws[f"H{r}"].value, ws[f"I{r}"].value = h, i
        ws[f"E{r}"].value = f"=INDEX($G{r}:$I{r},$E$9)"
        for col in "EGHI":
            ws[f"{col}{r}"].number_format = {"%": "0.0%", "flag": "0"}.get(unit, "#,##0.0")
    end = label_row(ws, "Grid fee reconciliation", start=top)
    labels = {ws[f"C{r}"].value: r for r in range(top, end + 1) if ws[f"C{r}"].value}
    find = lambda text: next(r for lab, r in labels.items() if lab.startswith(text))  # noqa: E731
    rows.update(ind_ft=find("Indirect connection fee (fajlagos"), dir_ft=find("Direct connection investment"),
                target=find("Devex target"))
    for text, (h, i) in OTHER_VALUES.items():
        r = find(text) if text not in ("Indirect connection fee", "Direct connection fee") else next(
            r for lab, r in labels.items() if lab == text)
        for col, v in (("H", h), ("I", i)):
            ws[f"{col}{r}"]._style = copy(ws[f"G{r}"]._style)
            ws[f"{col}{r}"].value = v.format(c=col, **rows) if isinstance(v, str) and v.startswith("=") else v
        if text in NOTES_M:
            ws[f"F{r}"].value = NOTES_M[text]
    signing, bids, scale = find("Signing fee to current owners"), find("Number of separate bids"), find("Devex scaling")
    blue = ws[f"G{signing}"]._style                                    # §20 audit: H:I inputs blue on input fill
    for r in [*range(at, at + len(NEW_ROWS)), rows["target"]]:
        for col in "HI":
            fmt = ws[f"{col}{r}"].number_format
            ws[f"{col}{r}"]._style = copy(blue)
            ws[f"{col}{r}"].number_format = fmt
    ws[f"G{signing}"].value = old[24]                                 # Bora: the previous live formula
    ws[f"G{bids}"].value = "=In_Bids"
    ws[f"G{scale}"].value = "=In_Scale"
    ws[f"E{scale}"].value = f"=INDEX($G{scale}:$I{scale},$E$9)"
    ws[f"C{top}"].value = "M. Project inputs – Bora | Iváncsa | Dunaújváros (RtB route; HTM is Bora's)"
    for col, v in zip("EGHI", ["Selected", *PROJECTS]):
        ws[f"{col}{top + 1}"].value = v
    define(wb, "In_Projects", f"Inputs!$G${top + 2}:$I${top + 2}")
    define(wb, "In_Project_MVA", f"Inputs!$E${rows['mva']}")
    # The live cells the RtB route reads pick the project (Bora's column = their previous formula).
    ws["E9"].value = "=MATCH(Dash_Project,In_Projects,0)"
    ws["C9"].value, ws["F9"].value = "Project for the RtB route (1 Bora, 2 Iváncsa, 3 Dunaújváros)", "Dash_Project"
    ws.data_validations.dataValidation = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref) != "E9"]
    ws["E16"].value = f"=$E${rows['mva']}*(1+$E${rows['ob']})"
    ws["E20"].value = f"=$E${rows['bess']}*$E${rows['mva']}"
    ws["E22"].value = f"=E20*$E${rows['hours']}"
    for live, r in ((24, signing), (25, rows["dd"]), (82, bids), (141, rows["prem"])):
        ws[f"E{live}"].value = f"=INDEX($G${r}:$I${r},$E$9)"            # the live cell picks; block M shows it
        ws[f"E{r}"].value = f"=$E${live}"
    return {**rows, "signing": signing}


def engine_links(wb, rows):
    """Deposits capacity bases, MOIC premium and bids read the project; Bid_Points keeps Bora's own rates."""
    dep, moic, bp, inp = wb["Deposits"], wb["MOIC"], wb["Bid_Points"], wb["Inputs"]
    dep["G7"].value = "=Inputs!$E$82"
    for r in (11, 12, 18):
        assert dep[f"G{r}"].value == "=Dash_MVA*Inputs!$E$17", dep[f"G{r}"].value
        dep[f"G{r}"].value = "=In_Project_MVA*Inputs!$E$17"
    assert moic["G143"].value == "=Inputs!$E$141*Dash_MVA"
    moic["G143"].value = "=Inputs!$E$141*In_Project_MVA"
    for col in range(8, 15):
        c = L(col)
        assert moic[f"{c}226"].value == f"={c}225*Dash_MVA"
        moic[f"{c}226"].value = f"={c}225*In_Project_MVA"
    ind = next(r for r in range(rows["mva"], rows["mva"] + 30) if str(inp[f"C{r}"].value).startswith("Indirect connection fee (f"))
    swaps = {"Inputs!$E$24+Inputs!$E$25": f"Inputs!$G${rows['signing']}+Inputs!$G${rows['dd']}",
             f"Inputs!$E${ind}": f"Inputs!$G${ind}", f"Inputs!$E${ind + 1}": f"Inputs!$G${ind + 1}"}
    n = 0
    for c in bp._cells.values():
        if c.data_type == "f" and isinstance(c.value, str):
            new = c.value
            for a, b in swaps.items():
                new = new.replace(a, b)
            n += new != c.value
            c.value = new
    assert n == 3, n


# =========================================================================================== §20.2 Dashboard
def dashboard(wb):
    d = wb["Dashboard"]
    at = label_row(d, "Turbine route (case)") + 1
    shift_rows(wb, "Dashboard", at, 1)
    for col in range(3, 9):
        d.cell(at, col)._style = copy(d.cell(at - 1, col)._style)
    d.cell(at, 3).value, d.cell(at, 4).value, d.cell(at, 5).value = PROJECT_LABEL, "project", PROJECTS[0]
    define(wb, "Dash_Project", f"Dashboard!$E${at}")
    dv = DataValidation(type="list", formula1="=In_Projects", allow_blank=False, showErrorMessage=True,
                        errorTitle="Project", error="Pick Bora, Iváncsa or Dunaújváros")
    dv.add(f"E{at}")
    d.add_data_validation(dv)
    # 5b: one pointer line after its last row; key assumptions: the note.
    last5b = label_row(d, "A.RES buy-out paid")
    shift_rows(wb, "Dashboard", last5b + 1, 1)
    d.cell(last5b + 1, 3)._style = copy(d.cell(label_row(d, "Wind built = bid"), 3)._style)
    d.cell(last5b + 1, 3).value = LINE_5B
    note = label_row(d, "Wind built = bid")
    shift_rows(wb, "Dashboard", note, 1)
    d.cell(note, 3)._style = copy(d.cell(note + 1, 3)._style)
    d.cell(note, 3).value = PROJECT_NOTE
    # Check row: stored = live for the selected project (inside Master_Check: before the last check row).
    cap = label_row(d, "Capacity: BESS MW")
    shift_rows(wb, "Dashboard", cap, 1)
    for col in range(3, 7):
        d.cell(cap, col)._style = copy(d.cell(cap + 1, col)._style)
    d.cell(cap, 3).value, d.cell(cap, 4).value = "RtB_Projects stored = live for the selected project", "#"
    d.cell(cap, 5).value = "=SUM(RtB_Fail)"
    d.cell(cap, 6).value = f'=IF(E{cap}=0,"OK","CHECK")'
    master = name_cell(wb, "Master_Check")
    assert re.search(rf"E\d+:E{cap + 1}\)", master.value), master.value


# =========================================================================================== §20.3 RtB_Projects
def rtb_projects(wb):
    bp = wb["Bid_Points"]
    ws = wb.create_sheet("RtB_Projects", wb.sheetnames.index("Bid_Points") + 1)
    for col, width in zip("ABCDEFGHIJK", (3, 2, 62, 9, 14, 13, 13, 13, 13, 12, 8)):
        ws.column_dimensions[col].width = width
    ws.sheet_view.showGridLines = False
    _copy_style(bp["B2"], ws["B2"])
    ws["B2"].value = "RtB sale by project (EUR k) – live project and stored results (RtBProjects.RefreshAll)"
    bar, label, unit, calc, total = bp["C7"], bp["C8"], bp["D8"], bp["E20"], bp["E34"]
    for col in "ABCDEFGHIJK":
        _copy_style(bar, ws[f"{col}4"])
    for col, v in zip("CDEFGHIJK", ("Item", "Unit", '="Live: "&Inputs!$E$281', *PROJECTS, "Total", "Stored − live",
                                    "Fail (1)")):
        ws[f"{col}4"].value = v
    stored_font = Font(name=calc.font.name, size=calc.font.sz, color=PURPLE)
    stored_fill, thin = PatternFill("solid", fgColor=MACRO_FILL), Side(style="thin", color="FFBFBFBF")
    rows, r = {}, 5
    for key, text, un, f, tot in RTB_ROWS:
        if key is None:
            r += 1 if r > 5 else 0
            for col in "ABCDEFGHIJK":
                _copy_style(bar, ws[f"{col}{r}"])
            ws[f"C{r}"].value = text
            r += 1
            continue
        rows[key] = r
        _copy_style(label, ws[f"C{r}"])
        _copy_style(unit, ws[f"D{r}"])
        ws[f"C{r}"].value, ws[f"D{r}"].value = text, un
        bold = key in ("p_tot", "claims", "w_tot", "dist", "post")
        for col in "EFGHIJK":
            _copy_style(total if bold else calc, ws[f"{col}{r}"])
            ws[f"{col}{r}"].number_format = FORMATS[un] if col not in "JK" else ("#,##0.000" if col == "J" else "0")
        if "!" in f or f.startswith("=In_"):
            _link(ws[f"E{r}"], f.format(**rows))
        else:
            ws[f"E{r}"].value = f.format(**rows)
        for col in "FGH":
            ws[f"{col}{r}"].font, ws[f"{col}{r}"].fill = copy(stored_font), copy(stored_fill)
            ws[f"{col}{r}"].border = Border(thin, thin, thin, thin)
        if tot == "sum":
            ws[f"I{r}"].value = f"=F{r}+G{r}+H{r}"
        elif tot == "mva":
            ws[f"I{r}"].value = f"=I{rows['post']}/I{rows['mva']}"
        ws[f"J{r}"].value = f"=INDEX(F{r}:H{r},Inputs!$E$9)-E{r}"
        ws[f"K{r}"].value = f"=IF(ISERROR(J{r}),1,--(ABS(J{r})>Tolerance))"
        r += 1
    first, last = min(rows.values()), max(rows.values())
    define(wb, "RtB_Live", f"RtB_Projects!$E${first}:$E${last}")
    define(wb, "RtB_Stored", f"RtB_Projects!$F${first}:$H${last}")
    define(wb, "RtB_Fail", f"RtB_Projects!$K${first}:$K${last}")
    _copy_style(unit, ws[f"C{last + 2}"])
    ws[f"C{last + 2}"].value = ("Purple: written by RtBProjects.RunProjects (RefreshAll); live column = the project "
                                "selected on the Dashboard")
    ws[f"C{last + 3}"].value = "***END***"
    _copy_style(bp["C3"], ws[f"C{last + 3}"])
    ws.freeze_panes = "E5"
    return rows


def change_log(wb):
    ws = wb["Change_Log"]
    r = max(c.row for c in ws._cells.values() if c.column == 2 and c.value is not None) + 1
    for col in "BCDEFGHIJ":
        ws[f"{col}{r}"]._style = copy(ws[f"{col}{r - 1}"]._style)
    for col, v in zip("BCDEFGHIJ", (
            f"=B{r - 1}+1", date(2026, 10, 27), "RtB sale for Iváncsa and Dunaújváros",
            "Inputs; Deposits; MOIC; Dashboard; RtB_Projects; Bid_Points",
            "Inputs E9, E16, E20, E22, E24, E25, E82, E141, block M (7 rows, columns H:I); Deposits G7, G11, G12, "
            "G18; MOIC G143, H226:N226; Bid_Points three Inputs links; Dashboard selector, 5b line, note, check",
            "Project selector Dash_Project re-enables the colleague's project switch for the RtB route; block M "
            "holds Bora | Iváncsa | Dunaújváros; RtB_Projects shows the waterfall per project, stored by "
            "vba/RtBProjects.bas (RefreshAll).",
            "Bora only (E9 fixed at 1)", "Three projects on the RtB route; HTM stays Bora's",
            "User 27-Oct-2026; SPEC §20")):
        ws[f"{col}{r}"] = v




def edit():
    wb = openpyxl.load_workbook(SRC, keep_vba=False)
    rows = block_m(wb)
    engine_links(wb, rows)
    dashboard(wb)
    rtb_projects(wb)
    change_log(wb)
    return wb


def build(out=HERE / "out", name=NAME):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    xlsx, xlsm = out / f"{name}.xlsx", out / f"{name}.xlsm"
    edit().save(xlsx)
    with tempfile.TemporaryDirectory() as vba:                         # Scenarios.bas + RtBProjects.bas
        for module in ("Scenarios.bas", "RtBProjects.bas"):
            shutil.copy(VBA_DIR / module, vba)
        package(xlsx, xlsm, run=["RtBProjects.RefreshAll"], vba_dir=Path(vba))
    general(xlsm, "HTM", ("G7", "G10", "G14"), like="G8")
    rc = recalc(xlsm, out / f"{name}_recalc.xlsx")
    return {"xlsx": xlsx, "xlsm": xlsm, "recalc": rc}


if __name__ == "__main__":
    print(build())
