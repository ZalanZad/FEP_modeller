"""Bora model RtB / COD / HTM – SPEC §18 (the user's case made proper, Bora only, euros) and §19 (Bid_Points).

Source (SPEC §18): reference/Bora_model_RTB_COD_HTM_261026_1.xlsm, the user's edit of our §17 output (scenario-table
tweaks, Inputs G162 typed 800 over the link, Inputs E322 typed 4 % over the interpolation, Bid_Points columns H:I deleted
by hand, stored results re-run). Every other cell stays as it is; the §18 edits are:
- §18.1 turbine: scenario-table Base turbine 800 and Inputs G162 back to the link; clipping: Dashboard lever
  Dash_Clip_Override (4 %, blank = memo curve) read by In_Bottleneck_Factor, E322 back to the interpolation;
- §18.2 Bora only: Bid_Points rebuilt from the §17.3 copy of the Wind HU sheet with a per-bid and a total column only;
  Inputs E9 / block M wording, Open_Items and Change_Log texts without other projects;
- §18.3 euros: Dashboard section 6 in € m (k€ where small); Bid_Points amounts in k€ (HUF only for the tender unit
  rates, one FX line);
- Change_Log: one line.
Rows are found by names and labels. Then vba/Scenarios.bas is re-embedded and RunAllScenarios stores the three cases.
build_v1.py … build_v3.py are the history (build_v3.py = §17); ag_tieout.py ties the cases out to the source file.
"""
import re
import shutil
import tempfile
import zipfile
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

from fepm.recalc import recalc
from fepm.xl.macros import VBA_DIR, package

HERE = Path(__file__).resolve().parent
REF = HERE.parents[1] / "reference"
SRC = REF / "Bora_model_RTB_COD_HTM_261026_1.xlsm"
WIND_HU = REF / "Wind_HU_points_deposits_devex_20261026.xlsx"
WIND_HU_SHEET = "Bid_Deposits_Devex"
NAME = "Bora_model_RTB_COD_HTM"

# §18.1 the user's case: EU turbine 800 k€/MW in the Base scenario, clipping 4 % instead of the memo curve.
BASE_TURBINE = 800
CLIP_OVERRIDE = 0.04
CLIP_NOTE = "Clipping override (user, 26-Oct-26); blank = OVIT memo curve for a 36 % CF park"
LOW_TURBINE = 950                                                  # coordinator: 500 is not a pessimistic Low

# §19 Bid_Points: a two-bid calculator in the Wind HU layout. Rows 1–35 stay where the Wind HU sheet has them (the bid,
# its points, the band tables further down); the sections below the points are written by bid_points() in this order.
BIDS = "EF"                                                         # BORA 1, BORA 2; column G = BORA 1 + 2
DASH_BID = {8: "=Dash_MVA/In_Bids", 9: "=Dash_Overbuild", 10: "=Dash_BESS_Share", 11: "=Inputs!$E$21",
            12: "=Dash_Municipal", 13: "=Dash_Community", 14: "=Dash_Partial", 15: "=Dash_PB_Mult",
            16: "=Dash_Brownfield", 17: "=Dash_Shared", 18: "=Dash_EU_Turbine", 19: "=Dash_EU_Transformer"}
DATES = [("d_sign", "Signing fee and DD paid", "=Inputs!$E$27"),
         ("d_bid", "Bid submission: fee paid, bid security posted", "=Inputs!$E$95"),
         ("d_award", "Award", "=Inputs!$E$96"),
         ("d_pb", "PB top-up posted (award + 17 days), bid security released", "=Deposits!$K$10"),
         ("d_tso", "TSO security paid (connection request)", "=Inputs!$E$97"),
         ("d_cont", "Devex contingency paid", "=Devex_FEP!$F$38"),
         ("d_hcssz", "Connection agreement (HCSSZ): connection fees paid", "=Inputs!$E$99"),
         ("d_ntp", "Notice to proceed", "=Inputs!$E$154"),
         ("d_cod", "COD", "=Inputs!$E$155"),
         ("d_rel", "Performance security released", None)]               # the sheet's own input (no model date)
RATES = [("fx", "FX rate", "HUF/EUR", "=Inputs!$E$11", "#,##0"),
         ("bw", "Security base: benchmark investment, wind", "HUF/MW", "=Inputs!$E$75*Million", "#,##0"),
         ("bb", "Security base: benchmark investment, BESS", "HUF/MWh", "=Inputs!$E$89*Million", "#,##0"),
         ("pf", "Participation fee per bid", "HUF", "=Inputs!$E$81*Million", "#,##0"),
         ("bsr", "Bid security, share of base", "%", "=Inputs!$E$76", "0.0%"),
         ("pbr", "Performance security, minimum share of base", "%", "=Inputs!$E$77", "0.0%"),
         ("tso", "TSO security", "HUF/MVA", "=Inputs!$E$83*Million", "#,##0"),
         ("ind", "Indirect connection fee (közvetett)", "HUF/MVA", "=Inputs!$E$290", "#,##0"),
         ("dir", "Direct connection fee per connection point (közvetlen)", "HUF", "=Inputs!$E$291", "#,##0"),
         ("gf", "Bank guarantee fee", "%", "=Inputs!$E$113", "0.0%"),
         ("sdd", "Signing fee and DD, 200 MVA project", "EUR k", "=(Inputs!$E$24+Inputs!$E$25)/In_Scale", "#,##0"),
         ("dc", "Devex contingency", "%", "=Inputs!$E$68", "0.0%"),
         ("src", "Project capacity in the devex file and the signing fee", "MVA", "=In_Source_MVA", "#,##0")]
# Securities and fees: key, label, formula ({c} = bid column, {k} = row of key k), form, date key, kind
FEES = [("share", "Share of the 200 MVA project (bid MVA ÷ 200)", "x", "={c}$8/$E${src}", None, None, "calc"),
        ("base", "Security base (wind MW × benchmark + BESS MWh × benchmark)", "EUR k",
         "=({c}$20*$E${bw}+{c}$22*$E${bb})/$E${fx}/Thousand", None, None, "calc"),
        ("sign", "Signing fee and DD", "EUR k", "=$E${sdd}*{c}{share}", "cash", "d_sign", "calc"),
        ("part", "Participation fee (sunk)", "EUR k", "=$E${pf}/$E${fx}/Thousand", "cash", "d_bid", "calc"),
        ("bsg", "Bid security, gross", "EUR k", "=$E${bsr}*{c}{base}", "guarantee", "d_bid", "calc"),
        ("pbg", "Performance security, gross (rate × multiple × base)", "EUR k", "=$E${pbr}*{c}$15*{c}{base}",
         "guarantee", "d_pb", "calc"),
        ("bsc", "Less bid security credited (XII.11; its guarantee released)", "EUR k", "=-{c}{bsg}", None, "d_pb",
         "calc"),
        ("pbt", "Performance security top-up posted at award", "EUR k", "={c}{pbg}+{c}{bsc}", "guarantee", "d_pb",
         "calc"),
        ("tsor", "TSO connection-request security", "EUR k", "=$E${tso}*{c}$8/$E${fx}/Thousand", "cash", "d_tso",
         "calc"),
        ("indg", "Indirect connection fee, gross", "EUR k", "=$E${ind}*{c}$8/$E${fx}/Thousand", None, "d_hcssz",
         "calc"),
        ("tsoc", "Less TSO security credited (XIII.6, Annex 1)", "EUR k", "=-{c}{tsor}", None, "d_hcssz", "calc"),
        ("indp", "Indirect connection fee payable", "EUR k", "={c}{indg}+{c}{tsoc}", "cash", "d_hcssz", "calc"),
        ("dirp", "Direct connection fee (per point ÷ bids sharing it, ÖSZ-8)", "EUR k",
         "=$E${dir}/$E${fx}/Thousand/MAX(1,{c}$17*SUM($E$17:$F$17))", "cash", "d_hcssz", "calc"),
        ("sub1", "Cash payable excl. connection fees", "EUR k", "={c}{part}+{c}{tsor}", None, None, "total"),
        ("sub2", "Connection fees payable", "EUR k", "={c}{indp}+{c}{dirp}", None, None, "total"),
        ("sub3", "Cash payable incl. connection fees", "EUR k", "={c}{sub1}+{c}{sub2}", None, None, "total"),
        ("peak", "Guarantees outstanding at peak (PB gross)", "EUR k", "=MAX({c}{bsg},{c}{pbg})", None, None, "total"),
        ("gfee", "Guarantee fees to COD (on the guarantee outstanding)", "EUR k",
         "=$E${gf}*(($E${d_pb}-$E${d_bid})*{c}{bsg}+MAX(0,MIN($E${d_cod},$E${d_rel})-$E${d_pb}+1)*{c}{pbg})"
         "/Days_In_Year", None, None, "total")]
SUMMARY = [("s_pts", "Points", "points", "={c}$34"),
           ("s_sign", "Signing fee and DD", "EUR k", "={c}{sign}"),
           ("s_dev", "Devex, land and contingency", "EUR k", "={c}{dtot}"),
           ("s_cash", "Cash deposits and fees excl. connection fees", "EUR k", "={c}{sub1}"),
           ("s_conn", "Connection fees", "EUR k", "={c}{sub2}"),
           ("s_ex", "Total cash to COD excl. connection fees", "EUR k", "={c}{s_sign}+{c}{s_dev}+{c}{s_cash}"),
           ("s_in", "Total cash to COD incl. connection fees", "EUR k", "={c}{s_ex}+{c}{s_conn}"),
           ("s_gfee", "Guarantee fees to COD", "EUR k", "={c}{gfee}"),
           ("s_peak", "Peak guarantee", "EUR k", "={c}{peak}"),
           ("s_cod", "COD", "date", "=$E${d_cod}")]
DEVEX_LINES = range(5, 38)                                          # Devex_FEP rows (200 MVA amounts in column L)
DEVEX_OTHER = 36                                                   # other costs: outside the contingency base
QUARTERS = 21
PB_RELEASE = date(2033, 3, 31)
EU_NOTE = "the bid commits EU equipment; the CN case is a build switch on the Dashboard"
PB_RELEASE_NOTE = "model has no PB release date; Kiírás XII.13, estimate"
CONSTANTS = {"R": ("Million", "#", "Unit conversion"), "S": ("Days in a year", "days", "Guarantee-fee day count"),
             "T": ("Months in a quarter", "months", "Calendar")}

# §18.2 texts that named other projects (Change_Log history condensed to Bora; Open_Items E43 wording).
TEXTS = {
    ("Inputs", "C9"): "Bora", ("Inputs", "F9"): "Fixed at 1",
    ("Inputs", "C279"): "M. Bora inputs – grid node fees from MEKH/MAVIR Kapacitás pályázati kiírás 2026, Annex 1 "
                        "(1. melléklet)",
    ("Inputs", "E280"): "Bora",
    ("Open_Items", "E43"): "Single asset with one COD versus several SPVs with staggered RTB / COD. Determines whether "
                           "Stage 2 / 3 run per SPV or on a blended asset.",
    ("Change_Log", "F99"): "Inputs E9, rows 16/20/24/31/32/38/82/84–88, block M C277:J296; Devex_FEP E5:E37, L4:L39; "
                           "Dashboard G14",
    ("Change_Log", "G99"): "Selector added at Inputs E9 (fixed to Bora since #127). Grid fees from Kapacitás pályázati "
                           "kiírás 2026 Annex 1. Restored Dashboard G14 formula (had been overwritten with 0).",
    ("Change_Log", "I99"): "Bora unchanged; RTB FEP IRR 27.5%",
    ("Change_Log", "E100"): "Dashboard; Inputs; Run_Map",
    ("Change_Log", "G100"): "Added Dashboard scenario control panel: case selector and case inputs (Base / European / "
                            "Chinese: EYA, capex, PPA, Aurora, opex, gearing, rate …). Inputs now link to the panel "
                            "(green). Run_Map holds live result formulas; RunAllScenarios VBA (text in Run_Map B100+) "
                            "rewrites Dashboard section 11. Button added – macro must be pasted and assigned by the user.",
    ("Change_Log", "E103"): "Dashboard; Run_Map",
    ("Change_Log", "F103"): "Dashboard E61/E63, C61/C63, K76, K90, H90; Run_Map C22:C23",
    ("Change_Log", "G107"): "Stale cell-address labels updated after row inserts; 400 kV → 132 kV notes; E220 tenor "
                            "label; Dashboard F8 stale-snapshot warning; Open_Items D6/E7/F9/E10/F12 refreshed",
    ("Change_Log", "E109"): "MOIC; Dashboard; Run_Map",
    ("Change_Log", "F109"): "MOIC rows 270, 272–278 deleted; Dashboard row 150 deleted; Dashboard C65:E66; Run_Map "
                            "B36:C37",
    ("Change_Log", "E111"): "Inputs; Run_Map",
    ("Change_Log", "F111"): "Inputs C9, J278, J279, F283, F284, F292; Run_Map C8",
    ("Change_Log", "G111"): "Name spelling aligned across Inputs and Run_Map on user instruction (superseded by #127).",
    ("Change_Log", "H111"): "n/a", ("Change_Log", "I111"): "n/a",
    ("Change_Log", "G113"): "Base turbine price 500 → 950 k€/MW (European) to score the ≥80% EU-equipment tender "
                            "points.",
    ("Change_Log", "E123"): "Dashboard; Run_Map",
    ("Change_Log", "F123"): "Dashboard E7:E8; Dashboard section 11 E/F; Run_Map C3; Control_Panel E8",
    ("Change_Log", "E125"): "Inputs; Dashboard; Control_Panel; Run_Map",
    ("Change_Log", "F125"): "Inputs H7:I276 ↔ K7:L276; Dashboard F7:G8, E241:F241; macro",
    ("Change_Log", "F127"): "Control_Panel E6, F6, F11:H22; Inputs E9, H278:J296; Dashboard E211:F240; Run_Map C6:C8, "
                            "VBA text",
    ("Change_Log", "G127"): "Model reduced to Bora only: selector fixed to 1, inputs and notes not for Bora removed. "
                            "Dashboard section 11 re-populated with Bora case 2 / case 3 results.",
    ("Change_Log", "H127"): "Selector over several sites",
}
OTHER_PROJECTS = re.compile(r"Iváncsa|Dunaújváros|\bBana\b|Bána|Ócsa|Comparison sheet|project #|project sheet", re.I)

# §18.3 Dashboard section 6 (label → unit, formula): € m at the Inputs FX (mHUF ÷ HUF/EUR), k€ where small.
SECTION6 = {"Participation fee": ("k€", "=Deposits!$M$7", "#,##0.0"),
            "Bid security": ("€ m", "=Deposits!$H$8/Inputs!$E$11", None),
            "Performance security (gross": ("€ m", "=Deposits!$H$10/Inputs!$E$11", None),
            "TSO connection-request security": ("€ m", "=Deposits!$H$11/Inputs!$E$11", None),
            "Connection fee, indirect (before TSO credit)": ("€ m", "=Deposits!$H$18/Inputs!$E$11", None),
            "Connection fee, direct": ("€ m", "=Deposits!$H$12/Inputs!$E$11", None),
            "Total gross": ("€ m", "=Deposits!$H$19/Inputs!$E$11", None),
            "Total payable after netting": ("€ m", "=Deposits!$J$19/Inputs!$E$11", None)}

GREEN = "FF00B050"


# =========================================================================================== helpers
def define(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def like(cell, model, value=None, fmt=None):
    """Write `value` with the style of `model` (and an optional number format)."""
    cell.value = value
    cell._style = copy(model._style)
    if fmt:
        cell.number_format = fmt
    return cell


def name_cell(wb, name):
    sheet, ref = next(wb.defined_names[name].destinations)
    return wb[sheet][ref.replace("$", "").split(":")[0]]


def label_row(ws, text, col=3, start=1):
    """First row at or below `start` whose label in `col` starts with `text`."""
    rows = sorted(c.row for c in ws._cells.values()
                  if c.column == col and c.row >= start and isinstance(c.value, str) and c.value.startswith(text))
    assert rows, (ws.title, text)
    return rows[0]


# =========================================================================================== row insertion / deletion
_ADDR = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?")


def _shift_addr(addr, at, n):
    return _ADDR.sub(lambda m: f"{m[1]}{m[2]}{m[3]}{int(m[4]) + n if int(m[4]) >= at else m[4]}", addr)


def _shift_formula(f, host, sheet, at, n):
    """Shift every cell reference to `sheet` rows ≥ `at` by `n` in formula `f` written on sheet `host`."""
    out = []
    for t in Tokenizer(f).items:
        v = t.value
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            ref_sheet, addr = (v.rsplit("!", 1) if "!" in v else (None, v))
            if (ref_sheet.strip("'") if ref_sheet else host) == sheet and _REF.fullmatch(addr):
                v = (ref_sheet + "!" if ref_sheet else "") + _shift_addr(addr, at, n)
        out.append(v)
    return "=" + "".join(out)


def _refs_rows(f, host, sheet):
    """Rows of `sheet` referred to as single cells or range ends in formula `f` (for the deletion guard)."""
    rows = []
    for t in Tokenizer(f).items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            ref_sheet, addr = (t.value.rsplit("!", 1) if "!" in t.value else (None, t.value))
            if (ref_sheet.strip("'") if ref_sheet else host) == sheet and _REF.fullmatch(addr):
                rows += [int(m[4]) for m in _ADDR.finditer(addr)]
    return rows


def shift_rows(wb, sheet, at, n):
    """Insert `n` rows before row `at` (n > 0) or delete −n rows from row `at` (n < 0) on `sheet`, re-pointing every
    formula, name, merged range, conditional format, validation and array-formula range (openpyxl moves cells only)."""
    ws = wb[sheet]
    first = at if n > 0 else at - n                                    # first row that moves
    gone = range(at, at - n) if n < 0 else range(0)
    for w in wb.worksheets:
        for c in list(w._cells.values()):
            v = c.value
            text = v.text if isinstance(v, ArrayFormula) else (v if c.data_type == "f" else None)
            if not text:
                continue
            assert not set(_refs_rows(text, w.title, sheet)) & set(gone), (w.title, c.coordinate, text)
            new = _shift_formula(text, w.title, sheet, first, n)
            if isinstance(v, ArrayFormula):
                v.text = new
            elif new != v:                                             # text that starts with "=" stays text
                c.value = new
    for name, dn in list(wb.defined_names.items()):
        if dn.attr_text and f"{sheet}!" in dn.attr_text:
            dn.attr_text = _shift_formula("=" + dn.attr_text, "", sheet, first, n)[1:]
    merged = [str(m) for m in ws.merged_cells.ranges]
    for m in merged:
        ws.unmerge_cells(m)
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height and r >= at}
    if n > 0:
        ws.insert_rows(at, n)
    else:
        ws.delete_rows(at, -n)
    for r in heights:
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        if r >= first:
            ws.row_dimensions[r + n].height = h
    for m in merged:
        ws.merge_cells(_shift_addr(m, first, n))
    for c in ws._cells.values():                                     # array formulas carry their own range
        if isinstance(c.value, ArrayFormula) and c.value.ref:
            c.value.ref = c.coordinate if ":" not in c.value.ref else _shift_addr(c.value.ref, first, n)
    cf = ws.conditional_formatting
    ws.conditional_formatting = ConditionalFormattingList()
    for rng in cf:
        for rule in rng.rules:
            ws.conditional_formatting.add(_shift_addr(str(rng.sqref), first, n), rule)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(_shift_addr(str(dv.sqref), first, n))
    if ws.freeze_panes:
        ws.freeze_panes = _shift_addr(ws.freeze_panes, first, n)

# =========================================================================================== §17.3 Bid_Points
def _copy_cell(src, dst):
    v = src.value
    if isinstance(v, ArrayFormula):
        v = ArrayFormula(dst.coordinate, v.text)
    dst.value = v
    if src.has_style:
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.protection, dst.number_format = copy(src.alignment), copy(src.protection), src.number_format


def _link(cell, formula):
    """A link from another sheet: green font, no input fill (MODELLING_STANDARDS §5)."""
    cell.value = formula
    f = copy(cell.font)
    cell.font = Font(name=f.name, size=f.sz, bold=f.b, italic=f.i, color=GREEN)
    cell.fill = PatternFill(fill_type=None)


def _pure_link(f):
    toks = [t for t in Tokenizer(f).items if t.type != Token.WSPACE] if isinstance(f, str) and f.startswith("=") else []
    return toks[0].value if len(toks) == 1 and toks[0].type == Token.OPERAND and toks[0].subtype == Token.RANGE \
        else None


def _target(wb, host, op):
    """(sheet, coordinate) of a single-cell reference or single-cell name, else None."""
    if op in wb.defined_names:
        sheet, ref = next(wb.defined_names[op].destinations)
    else:
        sheet, ref = op.rsplit("!", 1) if "!" in op else (host, op)
    ref = ref.replace("$", "")
    return (sheet.strip("'"), ref) if re.fullmatch(r"[A-Z]{1,3}\d+", ref) else None


def collapse_chains(wb, sheet):
    """Every pure link on `sheet` that points at another pure link is re-pointed at the source (no chains)."""
    ws = wb[sheet]
    for c in ws._cells.values():
        op = _pure_link(c.value)
        while op:
            tgt = _target(wb, sheet, op)
            if not tgt:
                break
            s, coord = tgt
            inner = _pure_link(wb[s][coord].value)
            if not inner:
                break
            op = inner if ("!" in inner or inner in wb.defined_names or s == sheet) else f"{s}!{inner}"
            c.value = "=" + op


# =========================================================================================== §19 Bid_Points
def _clear(ws, rows, cols):
    for r in rows:
        for col in cols:
            ws[f"{col}{r}"].value = None


class _Writer:
    """Writes the Bid_Points rows with the Wind HU styles (header bar, label, unit, input, calc, total, date)."""

    def __init__(self, ws, src, first):
        self.ws, self.src, self.r, self.rows = ws, src, first, {}

    def blank(self):
        self.r += 1

    def header(self, title, heads):
        for col in range(1, 17):
            _copy_style(self.src.cell(47, col), self.ws.cell(self.r, col))
        self.ws[f"A{self.r}"].value, self.ws[f"C{self.r}"].value = "x", title
        for col, v in heads.items():
            self.ws[f"{col}{self.r}"].value = v
        self.r += 1

    def row(self, key, label, unit, values, kind="calc", fmt=None):
        r, s = self.r, self.src
        total = kind == "total"
        _copy_style(s["C72"] if total else s["C48"], self.ws[f"C{r}"])
        _copy_style(s["D48"], self.ws[f"D{r}"])
        self.ws[f"C{r}"].value, self.ws[f"D{r}"].value = label, unit
        for col, v in values.items():
            if col in "EFG":
                model = {"input": s["E48"], "date_in": s["E37"], "date": s["L63"], "total": s["E72"]}.get(kind, s["E63"])
            else:
                model = {"H": s["D48"], "L": s["L63"]}.get(col, s["E63"])
            _copy_style(model, self.ws[f"{col}{r}"])
            self.ws[f"{col}{r}"].value = v
            if fmt and col in "EFG":
                self.ws[f"{col}{r}"].number_format = fmt
        if key:
            self.rows[key] = r
        self.r += 1
        return r


def _copy_style(src, dst):
    dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
    dst.alignment, dst.number_format = copy(src.alignment), src.number_format


def _fmt(f, rows, c="E"):
    return f.format(c=c, **{k: v for k, v in rows.items()})


def bid_points(wb):
    """The Wind HU sheet copied in (§17.3), then §19: a self-contained two-bid calculator, Bora only, EUR k."""
    at = wb.sheetnames.index("Bid_Points")
    del wb["Bid_Points"]
    src = openpyxl.load_workbook(WIND_HU)[WIND_HU_SHEET]
    ws = wb.create_sheet("Bid_Points", at)
    for row in src.iter_rows():
        for c in row:
            if c.value is not None or c.has_style:
                _copy_cell(c, ws.cell(c.row, c.column))
    for k, dim in src.column_dimensions.items():
        ws.column_dimensions[k].width, ws.column_dimensions[k].hidden = dim.width, dim.hidden
    ws.column_dimensions["H"].width = 11
    for r, dim in src.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[r].height = dim.height
    ws.freeze_panes = src.freeze_panes
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines
    cf_rules = [rule for rng in src.conditional_formatting for rule in rng.rules]

    # Title, constants and the bid (rows 1–35): two typed bids pre-filled with the model's bid, the Dashboard bid
    # alongside for comparison; the other sites' columns go.
    _clear(ws, range(1, 36), "GHIJ")
    ws["C3"].value, ws["C4"].value = "Bora", "Bora – bid, deposits and devex to COD"
    ws["Q4"].value = None
    _copy_cell(src["S4"], ws["T4"])
    ws["T4"].value = 3
    for col, (label, unit, source) in CONSTANTS.items():
        _copy_style(src["E48"], ws[f"{col}4"])
        ws[f"{col}4"].number_format = "#,##0"
        for r, v in ((3, label), (5, unit), (6, source)):
            _copy_style(src["D48"], ws[f"{col}{r}"])
            ws[f"{col}{r}"].value = v
    define(wb, "Million", "Bid_Points!$R$4")
    define(wb, "Days_In_Year", "Bid_Points!$S$4")
    define(wb, "Qtr_Months", "Bid_Points!$T$4")
    ws["E7"].value, ws["F7"].value, ws["G7"].value, ws["H7"].value = "BORA 1", "BORA 2", "BORA 1 + 2", "Dashboard bid"
    _copy_style(src["G7"], ws["H7"])
    bid = model_bid(wb)
    for r, f in DASH_BID.items():
        for col in BIDS:
            ws[f"{col}{r}"].value = bid[r]
        _copy_style(src["E20"], ws[f"H{r}"])
        ws[f"H{r}"].number_format = ws[f"E{r}"].number_format
        _link(ws[f"H{r}"], f)
    ws["D12"].value = "EUR k/MVA"
    _copy_style(src["D48"], ws["I18"])
    ws["I18"].value = EU_NOTE
    for r in (8, 20, 21, 22):
        _copy_style(src["E20"], ws[f"G{r}"])
        ws[f"G{r}"].value = f"=E{r}+F{r}"
    ws["C34"].value = "Points total"
    _copy_style(src["C24"], ws["C35"])
    _copy_style(src["D24"], ws["D35"])
    ws["C35"].value, ws["D35"].value = "Same as the Dashboard bid (1 = yes)", "flag"
    for col in BIDS:
        _copy_style(src["E20"], ws[f"{col}35"])
        ws[f"{col}35"].value = f"=N(SUMPRODUCT(--({col}8:{col}19<>$H$8:$H$19))=0)"
        ws[f"{col}35"].number_format = "0"
    ws["F222"].value = 1                                              # ÖSZ-6: x ≥ 2.0 scores 100 (GYIK Q34)

    # Everything between the points and the scoring tables is rewritten: clear it, close the gap, open the rows.
    ws.conditional_formatting = ConditionalFormattingList()
    _clear(ws, range(36, 157), [L(c) for c in range(1, 21)])
    shift_rows(wb, "Bid_Points", 36, -121)
    n = 3 + len(DATES) + 2 + len(RATES) + 2 + len(FEES) + 2 + len(DEVEX_LINES) + 3 + len(SUMMARY) + 2 + QUARTERS + 3
    shift_rows(wb, "Bid_Points", 36, n)
    w = _Writer(ws, src, 37)
    heads = {"E": "BORA 1", "F": "BORA 2", "G": "BORA 1 + 2"}

    w.header("Key dates", {"E": "Date"})
    for key, label, f in DATES:
        if f:
            w.row(key, label, "date", {"E": f}, "date")
            _link(ws[f"E{w.rows[key]}"], f)
        else:
            w.row(key, label, "date", {"E": PB_RELEASE, "F": PB_RELEASE_NOTE}, "date_in")
            _copy_style(src["D48"], ws[f"F{w.rows[key]}"])
    w.blank()
    w.header("Rates and fees (tender unit rates in HUF, converted at the FX rate)", {"E": "Rate"})
    for key, label, unit, f, fmt in RATES:
        r = w.row(key, label, unit, {"E": f}, "calc", fmt)
        _link(ws[f"E{r}"], f)
    w.blank()
    w.header("Securities and fees to COD (EUR k)", {**heads, "H": "Form", "L": "Date"})
    for key, label, unit, f, form, dkey, kind in FEES:
        w.rows[key] = w.r                                              # rows may refer to their own key
        vals = {col: _fmt(f, w.rows, col) for col in BIDS}
        vals["G"] = None if key == "share" else f"=E{w.r}+F{w.r}"
        if form:
            vals["H"] = form
        if dkey:
            vals["L"] = f"=$E${w.rows[dkey]}"
        w.row(key, label, unit, {k: v for k, v in vals.items() if v is not None}, kind,
              "0.00" if unit == "x" else "#,##0.0")
    w.blank()
    w.header("Devex, land and contingency to COD (EUR k)", {**heads, "H": "Form", "L": "Date"})
    dev = wb["Devex_FEP"]
    first = w.r
    for d in DEVEX_LINES:
        task, item = dev[f"C{d}"].value, dev[f"D{d}"].value
        label = f"{task.strip()} – {item}" if task and task.strip() != item.strip() else item
        r = w.row(f"dev{d}", label, "EUR k", {**{col: f"=Devex_FEP!$L${d}*{col}${w.rows['share']}" for col in BIDS},
                                               "G": f"=E{w.r}+F{w.r}", "H": "cash"}, "calc", "#,##0.0")
        _link(ws[f"L{r}"], f"=Devex_FEP!$F${d}")
        _copy_style(src["L63"], ws[f"L{r}"])
        ws[f"L{r}"].font = Font(name=ws[f"L{r}"].font.name, size=ws[f"L{r}"].font.sz, color=GREEN)
    last, other = w.r - 1, w.rows[f"dev{DEVEX_OTHER}"]
    w.row("cont", "Contingency on devex (excl. other costs) and participation fee", "EUR k",
          {**{col: f"=$E${w.rows['dc']}*(SUM({col}{first}:{col}{last})-{col}{other}+{col}{w.rows['part']})"
              for col in BIDS}, "G": f"=E{w.r}+F{w.r}", "H": "cash", "L": f"=$E${w.rows['d_cont']}"}, "calc", "#,##0.0")
    w.row("dtot", "Devex, land and contingency", "EUR k",
          {**{col: f"=SUM({col}{first}:{col}{w.r - 1})" for col in BIDS}, "G": f"=E{w.r}+F{w.r}"}, "total", "#,##0.0")
    items = (w.rows["sign"], w.rows["cont"])
    w.blank()
    w.header("Summary to COD (EUR k)", heads)
    for key, label, unit, f in SUMMARY:
        w.rows[key] = w.r
        vals = {col: _fmt(f, w.rows, col) for col in (BIDS if key in ("s_pts", "s_cod") else "EFG")}
        if key == "s_cod":
            vals = {"E": vals["E"]}
        w.row(key, label, unit, vals, "total" if key in ("s_ex", "s_in") else ("date" if key == "s_cod" else "calc"),
              None if key == "s_cod" else ("#,##0.0" if unit == "EUR k" else "#,##0.00"))
    w.blank()

    # Quarterly cash: the dated cash items (Form = cash) per quarter; cumulative and guarantees on the total column.
    define(wb, "Bid_Dates", f"Bid_Points!$L${items[0]}:$L${items[1]}")
    define(wb, "Bid_Form", f"Bid_Points!$H${items[0]}:$H${items[1]}")
    w.header("Cash to COD by quarter (EUR k)", {**heads, "L": "Cumulative", "M": "Guarantees"})
    q0 = w.r
    R = w.rows
    for i in range(QUARTERS):
        r = w.r
        span = f"$C{r}"
        vals = {col: (f"=SUMPRODUCT((Bid_Dates>EOMONTH({span},-Qtr_Months))*(Bid_Dates<={span})"
                      f'*(Bid_Form="cash"),{col}${items[0]}:{col}${items[1]})') for col in BIDS}
        vals["G"] = f"=E{r}+F{r}"
        vals["L"] = f"=G{r}" if i == 0 else f"=L{r - 1}+G{r}"
        vals["M"] = (f"=({span}<=$E${R['d_cod']})*(($E${R['d_bid']}<={span})*($E${R['d_pb']}>{span})*$G${R['bsg']}"
                     f"+($E${R['d_pb']}<={span})*($E${R['d_rel']}>{span})*$G${R['pbg']})")
        w.row(None, None, "EUR k", vals, "calc", "#,##0.0")
        ws[f"L{r}"].number_format = ws[f"M{r}"].number_format = "#,##0.0"
        _copy_style(src["C134"] if i == 0 else src["L63"], ws[f"C{r}"])
        ws[f"C{r}"].value = date(2026, 9, 30) if i == 0 else f"=EOMONTH(C{r - 1},Qtr_Months)"
    q1 = w.r - 1
    w.row("qtot", "Total", "EUR k", {col: f"=SUM({col}{q0}:{col}{q1})" for col in "EFG"}, "total", "#,##0.0")
    chk = w.row("qchk", "Check: quarters vs summary", "flag",
                {col: f'=IF(ABS({col}{w.rows["qtot"]}-{col}{R["s_in"]})<Tolerance,"OK","CHECK")' for col in "EFG"},
                "calc")
    for rule in cf_rules:
        ws.conditional_formatting.add(f"E{chk}:G{chk}", rule)
    assert w.r == 36 + n - 1, (w.r, n)                                  # one blank row above the band tables
    mx = label_row(ws, "Maximum achievable points")
    top, bottom = label_row(ws, "ÖSZ-1 – Municipal"), label_row(ws, "Generator transformers")
    ws[f"I{mx}"].value = ws[f"I{mx}"].value.replace(")*100", f")*MAX(H{top}:H{bottom})")
    collapse_chains(wb, "Bid_Points")
    return ws


def model_bid(wb):
    """The model's bid per bid (typed into BORA 1 and BORA 2): Dashboard levers; the bid commits EU equipment (the CN
    case is a build switch on the Dashboard), so ÖSZ-9 is 1 for both items."""
    v = lambda n: name_cell(wb, n).value  # noqa: E731
    eu = 1
    bids = -(-v("Dash_MVA") // name_cell(wb, "In_MVA_Per_Bid").value)
    return {8: v("Dash_MVA") / bids, 9: v("Dash_Overbuild"), 10: v("Dash_BESS_Share"), 11: wb["Inputs"]["G21"].value,
            12: v("Dash_Municipal"), 13: v("Dash_Community"), 14: v("Dash_Partial"), 15: v("Dash_PB_Mult"),
            16: v("Dash_Brownfield"), 17: v("Dash_Shared"), 18: eu, 19: eu}



# =========================================================================================== §18.1 the user's case
def turbine(wb):
    """Base turbine 800 in the scenario table; Inputs G162 links to it again (the user typed 800 over the link).
    Low turbine 950 (coordinator, for the user to confirm: 500 is not a pessimistic value)."""
    first = name_cell(wb, "Dash_Key_Inputs")
    first.value = BASE_TURBINE
    first.offset(column=1).value = LOW_TURBINE
    inp = wb["Inputs"]
    inp[f"G{label_row(inp, 'Turbine supply incl. transport')}"].value = "=Dashboard!$H$8"


def clip_override(wb):
    """Dashboard lever Dash_Clip_Override (blank = memo curve); E322 back to the interpolation (typed 4 % over it)."""
    inp, d = wb["Inputs"], wb["Dashboard"]
    clip = label_row(inp, "Clipped share at the built ratio")
    inp[f"E{clip}"].value = f"=MAX(0,E{clip - 3}+E{clip - 1}*(E{clip - 8}-E{clip - 5}))"
    inp[f"G{clip}"].value = None                                      # the user's stray copy of the formula
    at = label_row(d, "Turbine route (case)")
    shift_rows(wb, "Dashboard", at, 1)
    for col in range(3, 9):
        like(d.cell(at, col), d.cell(at - 1, col))
    d.cell(at, 3).value, d.cell(at, 4).value = "Clipped share of gross energy (blank = OVIT memo curve)", "%"
    d.cell(at, 5).value = CLIP_OVERRIDE
    d.cell(at, 5).number_format = "0.0%"
    define(wb, "Dash_Clip_Override", f"Dashboard!$E${at}")
    factor = name_cell(wb, "In_Bottleneck_Factor")
    factor.value = f'=1-IF(Dash_Clip_Override="",E{clip},Dash_Clip_Override)'
    inp.cell(factor.row, 3).value = "Generation kept after clipping (1 − clipped share applied)"
    inp.cell(factor.row, 6).value = "Dash_Clip_Override, else the curve"
    d.cell(label_row(d, "Clipped share of gross energy (OVIT"), 3).value = "Clipped share of gross energy applied"
    d.cell(label_row(d, "Clipping per OVIT memo"), 3).value = CLIP_NOTE


# =========================================================================================== §18.2 Bora only
def bora_only(wb):
    for (sheet, ref), text in TEXTS.items():
        wb[sheet][ref].value = text
    left = [(ws.title, c.coordinate) for ws in wb.worksheets if ws.title != "BORA_Devex" for c in ws._cells.values()
            if isinstance(c.value, str) and OTHER_PROJECTS.search(c.value)]
    assert not left, left


# =========================================================================================== §18.3 euros first
def euros(wb):
    d = wb["Dashboard"]
    top = label_row(d, "6. Securities and fees")
    d.cell(top, 3).value = "6. Securities and fees (Deposits; € m at the Inputs FX)"
    model = label_row(d, "Of which cash", start=top)
    for label, (unit, f, fmt) in SECTION6.items():
        r = label_row(d, label, start=top)
        d.cell(r, 4).value = unit
        like(d.cell(r, 5), d.cell(model, 5), f, fmt)


def checks(wb):
    """Dashboard section 8 (§19.6): the Bid_Points points check leaves Master_Check (Bid_Points shows "Same as the
    Dashboard bid" instead); the Bid_Points quarterly cash check (all three columns) goes in."""
    d, bp = wb["Dashboard"], wb["Bid_Points"]
    old = label_row(d, "Bid_Points total for BORA 1")
    _clear(d, [old], [L(c) for c in range(1, 12)])
    shift_rows(wb, "Dashboard", old, -1)
    chk = label_row(bp, "Check: quarters vs summary")
    at = label_row(d, "Capacity: BESS MW")                              # before the last check row: inside the SUM
    shift_rows(wb, "Dashboard", at, 1)
    for col in range(3, 7):
        like(d.cell(at, col), d.cell(at + 1, col))
    d.cell(at, 3).value, d.cell(at, 4).value = "Bid_Points quarterly cash = summary to COD", "#"
    d.cell(at, 5).value = f'=COUNTIF(Bid_Points!$E${chk}:$G${chk},"CHECK")'
    d.cell(at, 6).value = f'=IF(E{at}=0,"OK","CHECK")'
    master = name_cell(wb, "Master_Check")
    assert re.search(rf"E\d+:E{at + 1}\)", master.value), master.value


def change_log(wb):
    ws = wb["Change_Log"]
    r = max(c.row for c in ws._cells.values() if c.column == 2 and c.value is not None) + 1
    for col in "BCDEFGHIJ":
        ws[f"{col}{r}"]._style = copy(ws[f"{col}{r - 1}"]._style)
    for col, v in zip("BCDEFGHIJ", (
            f"=B{r - 1}+1", date(2026, 10, 26), "User case made proper; Bora only; euros",
            "Dashboard; Inputs; Bid_Points; Open_Items; Change_Log",
            "Dashboard E8, clipping lever row, section 6, key assumptions; Inputs G162, E322, In_Bottleneck_Factor, "
            "C9, F9, C279, E280; Bid_Points rebuilt",
            "Base turbine 800 via the scenario table; Dash_Clip_Override 4 % (blank = memo curve) replaces the typed "
            "E322; Bid_Points per bid and total only, amounts in k€; Dashboard section 6 in € m; texts naming other "
            "sites removed.",
            "G162 and E322 typed over formulas; Bid_Points with other sites; section 6 in mHUF",
            "Same numbers from levers and links; Bora only; € on the presentation sheets",
            "User 26-Oct-2026; SPEC §18")):
        ws[f"{col}{r}"] = v
    r += 1
    for col in "BCDEFGHIJ":
        ws[f"{col}{r}"]._style = copy(ws[f"{col}{r - 1}"]._style)
    for col, v in zip("BCDEFGHIJ", (
            f"=B{r - 1}+1", date(2026, 10, 27), "Bid_Points two-bid calculator", "Bid_Points; Dashboard",
            "Bid_Points rows 7–157; Dashboard section 8 (points check removed, quarterly check E:G)",
            "BORA 1 | BORA 2 | BORA 1 + 2 with typed bid levers, securities and fees with the deductions (bid security "
            "into the PB, TSO security against the indirect fee), devex per bid on the 200 MVA lines, summary excl. / "
            "incl. connection fees, quarterly cash.",
            "Bid_Points per bid and total, linked to the Dashboard levers",
            "Self-contained two-bid calculator in EUR k; 'Same as the Dashboard bid' line",
            "User 27-Oct-2026; SPEC §19")):
        ws[f"{col}{r}"] = v



# =========================================================================================== build
def general(path, sheet, cells, like):
    """SPEC §8.6: give `cells` the style index of `like` (a General cell) inside the saved file; LibreOffice infers
    a 0.00 % format for these sums on save (HTM G10 would show 3000.00 %)."""
    path = Path(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    wbx, rels = parts["xl/workbook.xml"].decode(), parts["xl/_rels/workbook.xml.rels"].decode()
    rid = re.search(rf'<sheet name="{sheet}"[^>]*r:id="(\w+)"', wbx).group(1)
    target = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="{rid}"', rels)
    xml_name = "xl/" + (target.group(1) or target.group(2)).lstrip("/").removeprefix("xl/")
    xml = parts[xml_name].decode()
    style = re.search(rf'<c r="{like}" s="(\d+)"', xml).group(1)
    for c in cells:
        xml, n = re.subn(rf'<c r="{c}" s="\d+"', f'<c r="{c}" s="{style}"', xml)
        assert n == 1, c
    parts[xml_name] = xml.encode()
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)
    tmp.replace(path)




def edit():
    wb = openpyxl.load_workbook(SRC, keep_vba=False)
    turbine(wb)
    clip_override(wb)
    bid_points(wb)
    bora_only(wb)
    euros(wb)
    checks(wb)
    change_log(wb)
    return wb


def build(out=HERE / "out", name=NAME):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    xlsx, xlsm = out / f"{name}.xlsx", out / f"{name}.xlsm"
    edit().save(xlsx)
    with tempfile.TemporaryDirectory() as vba:                         # Scenarios.bas only (as is)
        shutil.copy(VBA_DIR / "Scenarios.bas", vba)
        package(xlsx, xlsm, run=["Scenarios.RunAllScenarios"], vba_dir=Path(vba))
    general(xlsm, "HTM", ("G7", "G10", "G14"), like="G8")
    rc = recalc(xlsm, out / f"{name}_recalc.xlsx")
    return {"xlsx": xlsx, "xlsm": xlsm, "recalc": rc}


if __name__ == "__main__":
    print(build())
