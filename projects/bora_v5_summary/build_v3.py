"""Bora model RtB / COD / HTM – SPEC §17: the OVIT overbuild strategy, the clipping curve and the Wind HU points sheet.

Source (SPEC §17): reference/Bora_model_RTB_COD_HTM_AG_20261026.xlsm, the user's edited copy (scenario table,
construction-capex scenario row, yield-on-cost rows, CIT 15 %). Every user edit stays; the §17 edits are:
- §17.1 strategy levers on the Dashboard and two key-assumption notes;
- §17.2 Inputs block U: the OVIT clipping curve (memo Table 7) and the interpolation helpers above
  In_Bottleneck_Factor; Dash_Bottleneck_Loss (5 % of overbuilt MW) retired with its Dashboard row;
- §17.3 new sheet Bid_Points (copy of the user's Wind HU sheet "Bid_Deposits_Devex"), BORA 1 / BORA 2 linked to the
  model, ÖSZ-6 top band made inclusive, quarterly cash rows rewritten under the formula-length rule; a Dashboard check
  row (Bid_Points total = Dashboard points) inside Master_Check;
- Change_Log: one line.
Rows are found by names and labels (never by the old row numbers). Then vba/Scenarios.bas is re-embedded and
RunAllScenarios stores the three cases. build_v1.py / build_v2.py are the history; ag_tieout.py is the regression.
"""
import re
import shutil
import tempfile
import zipfile
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

from fepm.recalc import recalc
from fepm.xl.macros import VBA_DIR, package

HERE = Path(__file__).resolve().parent
REF = HERE.parents[1] / "reference"
SRC = REF / "Bora_model_RTB_COD_HTM_AG_20261026.xlsm"
WIND_HU = REF / "Wind_HU_points_deposits_devex_20261026.xlsx"
WIND_HU_SHEET = "Bid_Deposits_Devex"
NAME = "Bora_model_RTB_COD_HTM"

# §17.1 levers (Dashboard names). Built = bid × (1 − shortfall) = 151 MW × 90 % = 135.9 MW on 100 MVA (Inputs E308).
LEVERS = {"Dash_Overbuild": 0.51, "Dash_Build_Shortfall": 0.10, "Dash_BESS_Share": 0.30, "Dash_Municipal": 50,
          "Dash_Community": 3, "Dash_PB_Mult": 1.0, "Dash_Partial": 1, "Dash_Brownfield": 0, "Dash_Shared": 1}
# ag_tieout.py: the AG file's own levers and its 5 % rule (literal on purpose: regression build only).
AG_LEVERS = {"Dash_Overbuild": 0.11, "Dash_BESS_Share": 0, "Dash_Municipal": 0, "Dash_Community": 0}
AG_FACTOR = "=1-0.05*(In_Wind_Built-Dash_MVA)/In_Wind_Built"

# §17.2 OVIT optimal-overbuild memo (23-Oct-2026), Table 7: installed ÷ cap ratio → clipped share of gross energy.
CLIP = [(1.01, 0.0), (1.08, 0.010), (1.15, 0.032), (1.22, 0.056), (1.30, 0.080), (1.37, 0.103), (1.44, 0.126),
        (1.51, 0.148), (1.58, 0.169), (1.66, 0.190), (1.73, 0.209), (1.80, 0.228)]
CLIP_SRC = "OVIT memo Table 7, drawn year"
NOTES = ["Overbuild per OVIT optimal-overbuild memo (23-Oct-2026), capped at the maximum-points band",
         "Clipping per OVIT memo Table 7 (36 % P50 park); conservative at 2,500 h"]

# §17.3 Bid_Points links (Wind HU row → formula; BORA 1 = column E, BORA 2 = column F, identical bids).
BID_LEVERS = {8: "=Dash_MVA", 9: "=Dash_Overbuild", 10: "=Dash_BESS_Share", 11: "=Inputs!$E$21",
              12: "=Dash_Municipal", 13: "=Dash_Community", 14: "=Dash_Partial", 15: "=Dash_PB_Mult",
              16: "=Dash_Brownfield", 17: "=Dash_Shared", 18: "=Dash_EU_Turbine", 19: "=Dash_EU_Transformer"}
BID_RATES = {48: "=Inputs!$E$11", 49: "=Inputs!$E$75*Million", 50: "=Inputs!$E$89*Million",
             51: "=Inputs!$E$81*Million", 52: "=Inputs!$E$76", 53: "=Inputs!$E$77", 54: "=Inputs!$E$83*Million",
             55: "=Inputs!$E$85*Million", 56: "=Inputs!$E$84*Million*{c}$8", 57: "=Inputs!$E$113",
             58: "=Inputs!$E$24+Inputs!$E$25", 60: "=Inputs!$E$68"}
BID_DATES = {37: "=Inputs!$E$27", 38: "=Inputs!$E$95", 39: "=Deposits!$K$10", 40: "=Inputs!$E$97",
             41: "=Devex_FEP!$F$38", 42: "=Inputs!$E$99", 43: "=Inputs!$E$154", 44: "=Inputs!$E$155"}
# Devex lines: Wind HU row → Devex_FEP row (the model's k€ per line, scaled by In_Scale). The Wind HU sheet has no
# line for Devex_FEP row 36 (other costs): inserted as Bid_Points row 117. Contingency = Devex_FEP row 38.
DEVEX = {**{r: r - 81 for r in range(86, 117)}, 117: 37}
DEVEX_OTHER, DEVEX_CONT = 36, 38
PB_RELEASE_NOTE = "model has no PB release date; Kiírás XII.13, estimate"
CONSTANTS = {"R": ("Million", "#", "Unit conversion"), "S": ("Days in a year", "days", "Guarantee-fee day count"),
             "T": ("Months in a quarter", "months", "Calendar")}
CASH_LINES = (75, 76, 79, 80, 81, *range(86, 119))                   # the lines the quarterly cash rows add up

GREEN = "FF00B050"


# =========================================================================================== helpers
def define(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def like(cell, model, value=None, fmt=None):
    """Write `value` with the style of `model` (and an optional number format)."""
    cell.value = value
    cell._style = copy(model._style)
    if fmt:
        cell.number_format = fmt
    return cell


def name_cell(wb, name):
    sheet, ref = next(wb.defined_names[name].destinations)
    return wb[sheet][ref.replace("$", "").split(":")[0]]


def label_row(ws, text, col=3, start=1):
    """First row at or below `start` whose label in `col` starts with `text`."""
    rows = sorted(c.row for c in ws._cells.values()
                  if c.column == col and c.row >= start and isinstance(c.value, str) and c.value.startswith(text))
    assert rows, (ws.title, text)
    return rows[0]


# =========================================================================================== row insertion / deletion
_ADDR = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?")


def _shift_addr(addr, at, n):
    return _ADDR.sub(lambda m: f"{m[1]}{m[2]}{m[3]}{int(m[4]) + n if int(m[4]) >= at else m[4]}", addr)


def _shift_formula(f, host, sheet, at, n):
    """Shift every cell reference to `sheet` rows ≥ `at` by `n` in formula `f` written on sheet `host`."""
    out = []
    for t in Tokenizer(f).items:
        v = t.value
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            ref_sheet, addr = (v.rsplit("!", 1) if "!" in v else (None, v))
            if (ref_sheet.strip("'") if ref_sheet else host) == sheet and _REF.fullmatch(addr):
                v = (ref_sheet + "!" if ref_sheet else "") + _shift_addr(addr, at, n)
        out.append(v)
    return "=" + "".join(out)


def _refs_rows(f, host, sheet):
    """Rows of `sheet` referred to as single cells or range ends in formula `f` (for the deletion guard)."""
    rows = []
    for t in Tokenizer(f).items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            ref_sheet, addr = (t.value.rsplit("!", 1) if "!" in t.value else (None, t.value))
            if (ref_sheet.strip("'") if ref_sheet else host) == sheet and _REF.fullmatch(addr):
                rows += [int(m[4]) for m in _ADDR.finditer(addr)]
    return rows


def shift_rows(wb, sheet, at, n):
    """Insert `n` rows before row `at` (n > 0) or delete −n rows from row `at` (n < 0) on `sheet`, re-pointing every
    formula, name, merged range, conditional format, validation and array-formula range (openpyxl moves cells only)."""
    ws = wb[sheet]
    first = at if n > 0 else at - n                                    # first row that moves
    gone = range(at, at - n) if n < 0 else range(0)
    for w in wb.worksheets:
        for c in list(w._cells.values()):
            v = c.value
            text = v.text if isinstance(v, ArrayFormula) else (v if c.data_type == "f" else None)
            if not text:
                continue
            assert not set(_refs_rows(text, w.title, sheet)) & set(gone), (w.title, c.coordinate, text)
            new = _shift_formula(text, w.title, sheet, first, n)
            if isinstance(v, ArrayFormula):
                v.text = new
            elif new != v:                                             # text that starts with "=" stays text
                c.value = new
    for name, dn in list(wb.defined_names.items()):
        if dn.attr_text and f"{sheet}!" in dn.attr_text:
            dn.attr_text = _shift_formula("=" + dn.attr_text, "", sheet, first, n)[1:]
    merged = [str(m) for m in ws.merged_cells.ranges]
    for m in merged:
        ws.unmerge_cells(m)
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height and r >= at}
    if n > 0:
        ws.insert_rows(at, n)
    else:
        ws.delete_rows(at, -n)
    for r in heights:
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        if r >= first:
            ws.row_dimensions[r + n].height = h
    for m in merged:
        ws.merge_cells(_shift_addr(m, first, n))
    for c in ws._cells.values():                                     # array formulas carry their own range
        if isinstance(c.value, ArrayFormula) and c.value.ref:
            c.value.ref = c.coordinate if ":" not in c.value.ref else _shift_addr(c.value.ref, first, n)
    cf = ws.conditional_formatting
    ws.conditional_formatting = ConditionalFormattingList()
    for rng in cf:
        for rule in rng.rules:
            ws.conditional_formatting.add(_shift_addr(str(rng.sqref), first, n), rule)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(_shift_addr(str(dv.sqref), first, n))
    if ws.freeze_panes:
        ws.freeze_panes = _shift_addr(ws.freeze_panes, first, n)


# =========================================================================================== §17.2 clipping curve
def clipping(wb):
    """Inputs block U: Table 7 (two rows, ratio and clipped share, columns Q:AB like the points bands) and the
    interpolation helpers, inserted above In_Bottleneck_Factor, which becomes 1 − the interpolated clipped share."""
    ws = wb["Inputs"]
    at = name_cell(wb, "In_Bottleneck_Factor").row
    model = at - 1                                                     # a block U calc row (label, unit, E, note)
    band = label_row(ws, "ÖSZ-3 overbuild and ÖSZ-4 BESS ratio – lower bound")
    rows = [  # label, unit, E formula ({r} = row of the helper, {x} = built ratio row …), note, number format
        ("Clipping curve: installed wind ÷ connection capacity (Q:AB)", "x", None, CLIP_SRC, "0.00"),
        ("Clipping curve: clipped share of gross energy (Q:AB)", "%", None, CLIP_SRC, "0.0%"),
        ("Built wind ÷ connection capacity", "x", "=In_Wind_Built/Dash_MVA", "Inputs E308", "0.000"),
        ("Curve point at or below the built ratio (1 below the curve)", "#",
         "=IFERROR(MATCH({x},In_Clip_Ratios,1),1)", "Interpolation", "0"),
        ("Curve segment used (the last one above the curve)", "#", "=MIN({i0},COUNT(In_Clip_Ratios)-1)",
         "Interpolation", "0"),
        ("Segment start: ratio", "x", "=INDEX(In_Clip_Ratios,{i})", "Interpolation", "0.00"),
        ("Segment end: ratio", "x", "=INDEX(In_Clip_Ratios,{i}+1)", "Interpolation", "0.00"),
        ("Segment start: clipped share", "%", "=INDEX(In_Clip_Shares,{i})", "Interpolation", "0.0%"),
        ("Segment end: clipped share", "%", "=INDEX(In_Clip_Shares,{i}+1)", "Interpolation", "0.0%"),
        ("Segment slope", "% per x", "=({y1}-{y0})/({x1}-{x0})", "Interpolation", "0.000"),
        ("Clipped share at the built ratio (linear, not below 0)", "%", "=MAX(0,{y0}+{s}*({x}-{x0}))",
         "Interpolation", "0.00%"),
    ]
    shift_rows(wb, "Inputs", at, len(rows))
    keys = dict(zip(("t1", "t2", "x", "i0", "i", "x0", "x1", "y0", "y1", "s", "clip"), range(at, at + len(rows))))
    ref = {k: f"E{r}" for k, r in keys.items()}
    for r, (label, unit, f, note, fmt) in enumerate(rows, at):
        like(ws.cell(r, 3), ws.cell(model, 3), label)
        like(ws.cell(r, 4), ws.cell(model, 4), unit)
        like(ws.cell(r, 6), ws.cell(model, 6), note)
        if f:
            like(ws.cell(r, 5), ws.cell(model, 5), f.format(**ref), fmt)
    for i, (x, y) in enumerate(CLIP):
        like(ws.cell(keys["t1"], 17 + i), ws.cell(band, 17), x, "0.00")
        like(ws.cell(keys["t2"], 17 + i), ws.cell(band, 17), y, "0.0%")
    define(wb, "In_Clip_Ratios", f"Inputs!$Q${keys['t1']}:${L(16 + len(CLIP))}${keys['t1']}")
    define(wb, "In_Clip_Shares", f"Inputs!$Q${keys['t2']}:${L(16 + len(CLIP))}${keys['t2']}")
    factor = name_cell(wb, "In_Bottleneck_Factor")
    factor.value = f"=1-{ref['clip']}"
    ws.cell(factor.row, 3).value = "Generation kept after clipping (1 − clipped share)"
    ws.cell(factor.row, 6).value = "Used in HTM E36"


def retire_bottleneck(wb):
    """Dash_Bottleneck_Loss is no longer read (In_Bottleneck_Factor now uses the curve): name and row removed."""
    cell = name_cell(wb, "Dash_Bottleneck_Loss")
    for w in wb.worksheets:
        for c in w._cells.values():
            v = c.value.text if isinstance(c.value, ArrayFormula) else c.value
            assert not (c.data_type == "f" and "Dash_Bottleneck_Loss" in str(v)), (w.title, c.coordinate)
    del wb.defined_names["Dash_Bottleneck_Loss"]
    shift_rows(wb, "Dashboard", cell.row, -1)
    d = wb["Dashboard"]
    d.cell(label_row(d, "Bottleneck loss applied"), 3).value = "Clipped share of gross energy (OVIT curve)"
    d.cell(label_row(d, "Energy lost to the bottleneck"), 3).value = "Energy clipped, first full year"


# =========================================================================================== §17.1 levers and notes
def set_levers(wb, levers):
    for name, v in levers.items():
        name_cell(wb, name).value = v


def key_notes(wb):
    """Two key-assumption notes above the closing tolerance note."""
    d = wb["Dashboard"]
    note = label_row(d, "Wind built = bid")
    shift_rows(wb, "Dashboard", note, len(NOTES))
    for r, text in enumerate(NOTES, note):
        like(d.cell(r, 3), d.cell(note + len(NOTES), 3), text)


# =========================================================================================== §17.3 Bid_Points
def _copy_cell(src, dst):
    v = src.value
    if isinstance(v, ArrayFormula):
        v = ArrayFormula(dst.coordinate, v.text)
    dst.value = v
    if src.has_style:
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.protection, dst.number_format = copy(src.alignment), copy(src.protection), src.number_format


def _link(cell, formula):
    """A link from another sheet: green font, no input fill (MODELLING_STANDARDS §5)."""
    cell.value = formula
    f = copy(cell.font)
    cell.font = Font(name=f.name, size=f.sz, bold=f.b, italic=f.i, color=GREEN)
    cell.fill = PatternFill(fill_type=None)


def _pure_link(f):
    toks = [t for t in Tokenizer(f).items if t.type != Token.WSPACE] if isinstance(f, str) and f.startswith("=") else []
    return toks[0].value if len(toks) == 1 and toks[0].type == Token.OPERAND and toks[0].subtype == Token.RANGE \
        else None


def _target(wb, host, op):
    """(sheet, coordinate) of a single-cell reference or single-cell name, else None."""
    if op in wb.defined_names:
        sheet, ref = next(wb.defined_names[op].destinations)
    else:
        sheet, ref = op.rsplit("!", 1) if "!" in op else (host, op)
    ref = ref.replace("$", "")
    return (sheet.strip("'"), ref) if re.fullmatch(r"[A-Z]{1,3}\d+", ref) else None


def collapse_chains(wb, sheet):
    """Every pure link on `sheet` that points at another pure link is re-pointed at the source (no chains)."""
    ws = wb[sheet]
    for c in ws._cells.values():
        op = _pure_link(c.value)
        while op:
            tgt = _target(wb, sheet, op)
            if not tgt:
                break
            s, coord = tgt
            inner = _pure_link(wb[s][coord].value)
            if not inner:
                break
            op = inner if ("!" in inner or inner in wb.defined_names or s == sheet) else f"{s}!{inner}"
            c.value = "=" + op


def bid_points(wb):
    src = openpyxl.load_workbook(WIND_HU)[WIND_HU_SHEET]
    ws = wb.create_sheet("Bid_Points", wb.sheetnames.index("Dashboard") + 1)
    for row in src.iter_rows():
        for c in row:
            if c.value is not None or c.has_style:
                _copy_cell(c, ws.cell(c.row, c.column))
    for k, dim in src.column_dimensions.items():
        ws.column_dimensions[k].width, ws.column_dimensions[k].hidden = dim.width, dim.hidden
    for r, dim in src.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[r].height = dim.height
    ws.freeze_panes = src.freeze_panes
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines
    for rng in src.conditional_formatting:
        for rule in rng.rules:
            ws.conditional_formatting.add(str(rng.sqref), rule)
    # Constants: the workbook's Thousand (Inputs block U) replaces Q4; Million, Days_In_Year and the quarter length
    # are the sheet's own constants.
    ws["Q4"].value = None
    _copy_cell(src["S4"], ws["T4"])
    ws["T4"].value = 3
    define(wb, "Million", "Bid_Points!$R$4")
    define(wb, "Days_In_Year", "Bid_Points!$S$4")
    define(wb, "Qtr_Months", "Bid_Points!$T$4")
    ws["C4"].value = '="Bid strategy – "&TEXT($E$9,"0 %")&" overbuild"'
    for col, (label, unit, source) in CONSTANTS.items():                # blue inputs with unit and source
        like(ws[f"{col}4"], ws["H8"], ws[f"{col}4"].value, "#,##0")
        like(ws[f"{col}3"], ws["C8"], label)
        like(ws[f"{col}5"], ws["D8"], unit)
        like(ws[f"{col}6"], ws["D8"], source)
    for r in (34, 122):
        ws[f"C{r}"].value = "Points total, live case"
    ws["E63"].value, ws["F63"].value = "=$E$51*In_Bids", "=$F$51*In_Bids"
    like(ws["F45"], ws["D45"], PB_RELEASE_NOTE)

    for col in "EF":                                                  # BORA 1 and BORA 2 = the model's bid
        for r, f in {**BID_LEVERS, **BID_RATES}.items():
            _link(ws[f"{col}{r}"], f.format(c=col))
    ws["C56"].value = "Direct connection fee per bid (model: per MVA of the shared point)"
    for r, f in BID_DATES.items():
        _link(ws[f"E{r}"], f)
    for r, d in DEVEX.items():                                       # devex per line = the model's (scaled) k€
        for col in "EF":
            _link(ws[f"{col}{r}"], f"=Devex_FEP!$E${d}")
        _link(ws[f"L{r}"], f"=Devex_FEP!$F${d}")
    for col in "EF":
        _link(ws[f"{col}118"], f"=Devex_FEP!$E${DEVEX_CONT}")
    for r in range(85, 119):                                          # HUF m / EUR k source split: unscaled, dropped
        for col in "NO":
            ws[f"{col}{r}"].value = None
            ws[f"{col}{r}"]._style = copy(ws[f"K{r}"]._style)
    ws["F222"].value = 1                                              # ÖSZ-6: x ≥ 2.0 scores 100 (GYIK Q34)

    # Quarterly cash: one SUMPRODUCT per cell over the cash lines (flag column M) instead of the 40-term sums.
    like(ws["M74"], ws["L74"], "In cash")
    for r in range(75, 119):
        if ws[f"E{r}"].value is not None and r != 85:
            like(ws[f"M{r}"], ws["F222"], 1 if r in CASH_LINES else 0, "0")
    define(wb, "Bid_Dates", "Bid_Points!$L$75:$L$118")
    define(wb, "Bid_Cash", "Bid_Points!$M$75:$M$118")
    first, last = 134, label_row(ws, "Total", start=134) - 1
    for r in range(first, last + 1):
        if r > first:
            ws[f"C{r}"].value = f"=EOMONTH(C{r - 1},Qtr_Months)"
        for col in "EFHI":
            ws[f"{col}{r}"].value = (f"=SUMPRODUCT((Bid_Dates>EOMONTH($C{r},-Qtr_Months))*(Bid_Dates<=$C{r})"
                                     f"*Bid_Cash,{col}$75:{col}$118)")
        ws[f"M{r}"].value = (f"=($C{r}<=$E$44)*(($E$38<=$C{r})*($E$39>$C{r})*$J$77"
                             f"+($E$39<=$C{r})*($E$45>$C{r})*$J$78)")
    chk = label_row(ws, "Check: quarters vs summary")
    for col in "EFHI":
        ws[f"{col}{chk}"].value = f'=IF(ABS({col}{chk - 1}-{col}126)<Tolerance,"OK","CHECK")'
    mx = label_row(ws, "Maximum achievable points")
    ws[f"I{mx}"].value = "=(E159+E171+E185+E198+E211+E225+E226+E227+E229)*MAX(H160:H231)"
    other(wb, ws)
    collapse_chains(wb, "Bid_Points")
    return ws


def other(wb, ws):
    """Devex_FEP row 36 (other costs) has no Wind HU line: a row is inserted above the RtB land purchase."""
    at = 117
    shift_rows(wb, "Bid_Points", at, 1)
    for col in range(2, 14):
        like(ws.cell(at, col), ws.cell(at + 1, col))
    ws[f"C{at}"].value = "Other costs (HR, car etc.)"
    ws[f"D{at}"].value = ws[f"D{at + 1}"].value
    for col in "EF":
        _link(ws[f"{col}{at}"], f"=Devex_FEP!$E${DEVEX_OTHER}")
    for col in "HI":
        below = Translator(ws[f"{col}{at + 1}"].value, f"{col}{at + 1}")
        ws[f"{col}{at}"].value = below.translate_formula(f"{col}{at}")
    ws[f"J{at}"].value = f"=SUM(E{at}:I{at})"
    _link(ws[f"L{at}"], f"=Devex_FEP!$F${DEVEX_OTHER}")
    ws[f"M{at}"].value = 1


def dashboard_check(wb):
    """Check row inside Master_Check: Bid_Points total for BORA 1 = Dashboard total points."""
    d = wb["Dashboard"]
    at = label_row(d, "Capacity: BESS MW")                              # before the last check row: inside the SUM
    total = label_row(d, "Total points")
    shift_rows(wb, "Dashboard", at, 1)
    for col in range(3, 7):
        like(d.cell(at, col), d.cell(at + 1, col))
    d.cell(at, 3).value, d.cell(at, 4).value = "Bid_Points total for BORA 1 = Dashboard total points", "#"
    d.cell(at, 5).value = f"=N(ABS(Bid_Points!$E$34-$F${total})>Tolerance)"
    d.cell(at, 6).value = f'=IF(E{at}=0,"OK","CHECK")'
    master = name_cell(wb, "Master_Check")
    assert f"E{at}" in master.value or re.search(rf"E\d+:E{at + 1}\)", master.value), master.value


def change_log(wb):
    ws = wb["Change_Log"]
    r = max(c.row for c in ws._cells.values() if c.column == 2 and c.value is not None) + 1
    for col in "BCDEFGHIJ":
        ws[f"{col}{r}"]._style = copy(ws[f"{col}{r - 1}"]._style)
    for col, v in zip("BCDEFGHIJ", (
            f"=B{r - 1}+1", date(2026, 10, 26), "OVIT overbuild, clipping curve, Bid_Points",
            "Dashboard; Inputs; Bid_Points",
            "Dashboard levers, key assumptions, check row; Inputs block U clipping curve and helpers; new sheet "
            "Bid_Points; Dash_Bottleneck_Loss removed",
            "Overbuild 51 % (max-points band), shortfall 10 %, BESS 30 % × 4 h, municipal 50, community 3; clipping "
            "per OVIT memo Table 7 replaces the 5 % rule; Wind HU points sheet linked to the model (ÖSZ-6 top band "
            "inclusive).",
            "Overbuild 11 %, BESS 0, municipal 0, community 0, 5 % of overbuilt MW lost",
            "Overbuild 51 %, BESS 30 MW / 120 MWh, municipal 50, community 3, OVIT clipping curve",
            "User 26-Oct-2026; SPEC §17")):
        ws[f"{col}{r}"] = v


# =========================================================================================== build
def general(path, sheet, cells, like):
    """SPEC §8.6: give `cells` the style index of `like` (a General cell) inside the saved file; LibreOffice infers
    a 0.00 % format for these sums on save (HTM G10 would show 3000.00 %)."""
    path = Path(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    wbx, rels = parts["xl/workbook.xml"].decode(), parts["xl/_rels/workbook.xml.rels"].decode()
    rid = re.search(rf'<sheet name="{sheet}"[^>]*r:id="(\w+)"', wbx).group(1)
    target = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="{rid}"', rels)
    xml_name = "xl/" + (target.group(1) or target.group(2)).lstrip("/").removeprefix("xl/")
    xml = parts[xml_name].decode()
    style = re.search(rf'<c r="{like}" s="(\d+)"', xml).group(1)
    for c in cells:
        xml, n = re.subn(rf'<c r="{c}" s="\d+"', f'<c r="{c}" s="{style}"', xml)
        assert n == 1, c
    parts[xml_name] = xml.encode()
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)
    tmp.replace(path)


def edit(ag=False):
    wb = openpyxl.load_workbook(SRC, keep_vba=False)
    clipping(wb)
    retire_bottleneck(wb)
    key_notes(wb)
    bid_points(wb)
    dashboard_check(wb)
    change_log(wb)
    set_levers(wb, LEVERS)
    if ag:                                                             # ag_tieout.py: the AG levers and 5 % rule
        set_levers(wb, AG_LEVERS)
        name_cell(wb, "In_Bottleneck_Factor").value = AG_FACTOR
    return wb


def build(out=HERE / "out", ag=False, name=NAME):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    xlsx, xlsm = out / f"{name}.xlsx", out / f"{name}.xlsm"
    edit(ag).save(xlsx)
    with tempfile.TemporaryDirectory() as vba:                         # Scenarios.bas only (as is)
        shutil.copy(VBA_DIR / "Scenarios.bas", vba)
        package(xlsx, xlsm, run=["Scenarios.RunAllScenarios"], vba_dir=Path(vba))
    general(xlsm, "HTM", ("G7", "G10", "G14"), like="G8")
    rc = recalc(xlsm, out / f"{name}_recalc.xlsx")
    return {"xlsx": xlsx, "xlsm": xlsm, "recalc": rc}


if __name__ == "__main__":
    print(build())
