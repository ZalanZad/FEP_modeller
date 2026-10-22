"""SPEC §7 step 1: with the levers at the source values the edited engine reproduces the v5 file.

Builds the workbook with build.SOURCE into out/regression/ (case 1 = v5 case 2 EU, case 2 = v5 case 3 CN), compares
the stored case results (and HTM capex per MW, read live per case) with the colleague's own snapshot on the source Bora
sheet (tolerance 1e-6 relative), and reconciles the Deposits schedule to the source totals (Inputs G87 / G88, the old
Deposits E34 test, 0.01 mHUF).

    python3 projects/bora_v5_summary/regression.py
"""
import importlib.util
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1]))
from fepm.recalc import Office  # noqa: E402

spec = importlib.util.spec_from_file_location("bora_v5_summary_build", HERE / "build.py")
build = importlib.util.module_from_spec(spec)
spec.loader.exec_module(build)

TOL = 1e-6
# metric, source cell on the v5 Bora sheet (snapshot of 22-Oct-26 16:48), Dashboard name, case column, scale
METRICS = [("RtB FEP net profit post-tax, k€", "D23", "Out_rtb_profit", 1, 1000),
           ("RtB FEP IRR post-tax", "D24", "Out_rtb_irr", 1, 1),
           ("RtB FEP MOIC", "D27", "Out_rtb_moic", 1, 1),
           ("HTM EU FEP equity IRR", "E24", "Out_eq_irr", 1, 1),
           ("HTM EU project IRR post-tax", "E49", "Out_proj_irr", 1, 1),
           ("HTM EU capex per MW, k€/MW", "E42", "HTM!G150", 1, 1),
           ("HTM CN FEP equity IRR", "F24", "Out_eq_irr", 2, 1)]


def row(wb, name):
    return int(next(wb.defined_names[name].destinations)[1].split(":")[0].split("$")[-1])


def stored(wb, name, case):
    """Stored_Results value of the Live_Results row `name` (an Out_ name) for case 1-2."""
    r = row(wb, "Stored_Results") + row(wb, name) - row(wb, "Live_Results")
    return wb["Dashboard"].cell(r, 4 + case).value


def main():
    paths = build.build(HERE / "out" / "regression", levers=build.SOURCE, name="Bora_v5_Summary_regression")
    src = openpyxl.load_workbook(build.SRC, data_only=True)["Bora"]
    wb = openpyxl.load_workbook(paths["recalc"], data_only=True)
    names = [c.value for c in wb["Inputs"]["G7":"H7"][0]]
    live = {}
    with Office(paths["xlsm"]) as lo:                       # HTM capex per MW is not a stored result: read per case
        for case, n in enumerate(names, 1):
            lo.set("Scenario_Name", n)
            lo.calculate()
            live[case] = lo.get("HTM!G150")
        lo.set("Scenario_Name", names[0])
        lo.calculate()
        master = lo.get("Master_Check")
    ok = True
    print(f"{'metric':<34} {'source':>16} {'rebuilt':>16} {'rel diff':>10}")
    for label, cell, name, case, scale in METRICS:
        a = src[cell].value
        b = live[case] if "!" in name else stored(wb, name, case) * scale
        rel = abs(b - a) / abs(a)
        ok &= rel <= TOL
        print(f"{label:<34} {a:>16.10g} {b:>16.10g} {rel:>10.1e}{'' if rel <= TOL else '  FAIL'}")
    src_in = openpyxl.load_workbook(build.SRC, data_only=True)["Inputs"]
    for label, ours, theirs in (("Deposits gross, FEP basis (S19) vs Inputs G87", "S19", "G87"),
                                ("Deposits payable, FEP basis (T19) vs Inputs G88", "T19", "G88")):
        a, b = src_in[theirs].value, wb["Deposits"][ours].value
        ok &= abs(a - b) < 0.01
        print(f"{label:<48} {a:>12,.2f} {b:>12,.3f} {'' if abs(a - b) < 0.01 else '  FAIL'}")
    print(f"Master_Check (regression build, case 1): {master}")
    print("regression:", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
